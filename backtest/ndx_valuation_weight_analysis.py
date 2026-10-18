"""Paired analysis for the v9 point-in-time FCF Yield sensitivity study."""
from __future__ import annotations

import hashlib
import json
import os
import re
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

from backtest.ndx_experiment_analysis import _read_verified_bundle


ANALYSIS_POLICY_ID = "ndx_fcf_value_weight_analysis_v1"
PROFILE_PATTERN = re.compile(
    r"^(?P<composition>.+)__(?:val|guard)(?P<weight>00|05|10|15)$"
)
DECLARED_WEIGHTS = (0.0, 0.05, 0.10, 0.15)
MIN_CANDIDATE_IMPROVEMENT_RATE = 0.75
MIN_WINDOW_XIRR_WIN_RATE = 0.60
MAX_WINDOW_DRAWDOWN_P90_DETERIORATION = 0.005


class NdxValuationWeightAnalysisError(ValueError):
    """Raised when v9 evidence is incomplete or violates its frozen contract."""


def _content_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _atomic_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        frame.to_csv(temporary, index=False, encoding="utf-8-sig")
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def _atomic_json(payload: dict[str, Any], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        temporary.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def _profile_parts(profile_id: object) -> tuple[str, float]:
    match = PROFILE_PATTERN.fullmatch(str(profile_id))
    if not match:
        raise NdxValuationWeightAnalysisError(
            f"unexpected v9 profile id: {profile_id}"
        )
    return match.group("composition"), int(match.group("weight")) / 100.0


def _attach_profile_parts(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    parts = result["profile_id"].map(_profile_parts)
    result["composition_id"] = parts.map(lambda value: value[0])
    result["declared_valuation_weight"] = parts.map(lambda value: value[1])
    observed = pd.to_numeric(result["valuation_score_weight"], errors="coerce").fillna(0.0)
    if not np.allclose(observed, result["declared_valuation_weight"]):
        raise NdxValuationWeightAnalysisError(
            "profile valuation weights disagree with profile ids"
        )
    return result


def paired_candidate_effects(
    candidates: pd.DataFrame,
    *,
    declared_weights: tuple[float, ...] = DECLARED_WEIGHTS,
) -> pd.DataFrame:
    """Pair every non-zero valuation weight with its composition-specific val00."""
    data = _attach_profile_parts(candidates)
    keys = ["composition_id", "lookback_years", "trade_cost_bps"]
    expected = set(declared_weights)
    observed_groups = data.groupby(keys, dropna=False)[
        "declared_valuation_weight"
    ].agg(lambda values: set(float(value) for value in values))
    if len(observed_groups) != 36 or any(values != expected for values in observed_groups):
        raise NdxValuationWeightAnalysisError(
            "v9 requires 36 complete composition/lookback/cost groups"
        )
    metrics = [
        "strategy_xirr_p50",
        "strategy_xirr_p10",
        "strategy_maxdd_p50",
        "strategy_maxdd_p90",
        "incumbent_xirr_p50_margin",
        "incumbent_xirr_p10_margin",
        "turnover_x_p50",
        "cost_x_p50",
        "selection_churn_p50",
    ]
    baseline = data[data["declared_valuation_weight"].eq(0.0)][keys + metrics].copy()
    baseline = baseline.rename(columns={metric: f"baseline_{metric}" for metric in metrics})
    challenger = data[data["declared_valuation_weight"].gt(0.0)].copy()
    paired = challenger.merge(baseline, on=keys, how="left", validate="many_to_one")
    if paired[[f"baseline_{metric}" for metric in metrics]].isna().any().any():
        raise NdxValuationWeightAnalysisError("one or more candidate baselines are missing")
    for metric in metrics:
        paired[f"delta_{metric}"] = (
            pd.to_numeric(paired[metric], errors="raise")
            - pd.to_numeric(paired[f"baseline_{metric}"], errors="raise")
        )
    paired["drawdown_p50_relief"] = -paired["delta_strategy_maxdd_p50"]
    paired["drawdown_p90_relief"] = -paired["delta_strategy_maxdd_p90"]
    return paired


def paired_window_effects(
    windows: pd.DataFrame,
    *,
    expected_pairs: int = 1944,
) -> pd.DataFrame:
    data = _attach_profile_parts(windows)
    keys = [
        "composition_id",
        "lookback_years",
        "trade_cost_bps",
        "start",
        "end",
    ]
    metrics = [
        "strategy_xirr",
        "strategy_maxdd",
        "incumbent_xirr_margin",
        "incumbent_dd_relief",
        "turnover_x",
        "cost_x",
        "selection_churn_p50",
    ]
    baseline = data[data["declared_valuation_weight"].eq(0.0)][keys + metrics].copy()
    baseline = baseline.rename(columns={metric: f"baseline_{metric}" for metric in metrics})
    challenger = data[data["declared_valuation_weight"].gt(0.0)].copy()
    paired = challenger.merge(baseline, on=keys, how="left", validate="many_to_one")
    if len(paired) != int(expected_pairs):
        raise NdxValuationWeightAnalysisError(
            f"expected {int(expected_pairs)} paired windows, got {len(paired)}"
        )
    if paired[[f"baseline_{metric}" for metric in metrics]].isna().any().any():
        raise NdxValuationWeightAnalysisError("one or more window baselines are missing")
    for metric in metrics:
        paired[f"delta_{metric}"] = (
            pd.to_numeric(paired[metric], errors="raise")
            - pd.to_numeric(paired[f"baseline_{metric}"], errors="raise")
        )
    paired["drawdown_relief"] = -paired["delta_strategy_maxdd"]
    return paired


def summarize_effects(
    candidate_effects: pd.DataFrame,
    window_effects: pd.DataFrame,
    *,
    declared_weights: tuple[float, ...] = DECLARED_WEIGHTS,
) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for weight in declared_weights[1:]:
        candidates = candidate_effects[
            candidate_effects["declared_valuation_weight"].eq(weight)
        ]
        windows = window_effects[
            window_effects["declared_valuation_weight"].eq(weight)
        ]
        candidate_p50_rate = float((candidates["delta_strategy_xirr_p50"] > 0).mean())
        candidate_p10_rate = float((candidates["delta_strategy_xirr_p10"] > 0).mean())
        window_xirr_rate = float((windows["delta_strategy_xirr"] > 0).mean())
        dd_p90_delta = float(windows["delta_strategy_maxdd"].quantile(0.90))
        shortlisted = bool(
            candidate_p50_rate >= MIN_CANDIDATE_IMPROVEMENT_RATE
            and candidate_p10_rate >= MIN_CANDIDATE_IMPROVEMENT_RATE
            and window_xirr_rate >= MIN_WINDOW_XIRR_WIN_RATE
            and float(candidates["delta_strategy_xirr_p50"].median()) > 0
            and float(candidates["delta_strategy_xirr_p10"].median()) > 0
            and dd_p90_delta <= MAX_WINDOW_DRAWDOWN_P90_DETERIORATION
        )
        rows.append({
            "valuation_weight": weight,
            "candidate_pairs": len(candidates),
            "candidate_p50_improvement_rate": candidate_p50_rate,
            "candidate_p10_improvement_rate": candidate_p10_rate,
            "candidate_xirr_p50_delta_median": float(
                candidates["delta_strategy_xirr_p50"].median()
            ),
            "candidate_xirr_p10_delta_median": float(
                candidates["delta_strategy_xirr_p10"].median()
            ),
            "window_pairs": len(windows),
            "window_xirr_win_rate": window_xirr_rate,
            "window_xirr_delta_p10": float(windows["delta_strategy_xirr"].quantile(0.10)),
            "window_xirr_delta_p50": float(windows["delta_strategy_xirr"].median()),
            "window_drawdown_delta_p50": float(
                windows["delta_strategy_maxdd"].median()
            ),
            "window_drawdown_delta_p90": dd_p90_delta,
            "turnover_delta_p50": float(windows["delta_turnover_x"].median()),
            "cost_delta_p50": float(windows["delta_cost_x"].median()),
            "shortlist_passed": shortlisted,
        })
    return pd.DataFrame(rows)


def effect_breakdowns(
    candidate_effects: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    aggregations = {
        "candidate_pairs": ("profile_id", "size"),
        "candidate_p50_improvement_rate": (
            "delta_strategy_xirr_p50", lambda values: (values > 0).mean()
        ),
        "candidate_p10_improvement_rate": (
            "delta_strategy_xirr_p10", lambda values: (values > 0).mean()
        ),
        "candidate_xirr_p50_delta_median": (
            "delta_strategy_xirr_p50", "median"
        ),
        "candidate_xirr_p10_delta_median": (
            "delta_strategy_xirr_p10", "median"
        ),
        "candidate_maxdd_p90_delta_median": (
            "delta_strategy_maxdd_p90", "median"
        ),
        "turnover_delta_p50": ("delta_turnover_x_p50", "median"),
    }
    by_composition = candidate_effects.groupby(
        ["composition_id", "declared_valuation_weight"], as_index=False
    ).agg(**aggregations)
    by_lookback = candidate_effects.groupby(
        ["lookback_years", "declared_valuation_weight"], as_index=False
    ).agg(**aggregations)
    return by_composition, by_lookback


def absolute_gate_summary(candidates: pd.DataFrame) -> pd.DataFrame:
    data = _attach_profile_parts(candidates)
    return data.groupby("declared_valuation_weight", as_index=False).agg(
        candidate_rows=("profile_id", "size"),
        target_gate_passed_rows=("target_gate_passed", "sum"),
        incumbent_gate_passed_rows=("incumbent_gate_passed", "sum"),
        admission_passed_rows=("admission_passed", "sum"),
        best_incumbent_xirr_p50_margin=("incumbent_xirr_p50_margin", "max"),
        best_incumbent_xirr_p10_margin=("incumbent_xirr_p10_margin", "max"),
        best_incumbent_dd_relief_p50=("incumbent_dd_relief_p50", "max"),
        best_incumbent_dd_relief_p90=("incumbent_dd_relief_p90", "max"),
    )


def valuation_evidence_summary(selections: pd.DataFrame) -> pd.DataFrame:
    data = _attach_profile_parts(selections)
    data = data[data["declared_valuation_weight"].gt(0.0)].copy()
    data["selected_tickers"] = data["tickers"].fillna("").map(
        lambda value: len({ticker for ticker in str(value).split(";") if ticker})
    )
    return data.groupby(
        ["composition_id", "declared_valuation_weight"], as_index=False
    ).agg(
        selection_rows=("selection_date", "size"),
        valuation_score_coverage_min=("valuation_score_coverage", "min"),
        valuation_observed_coverage_min=("valuation_observed_coverage", "min"),
        valuation_missing_neutral_p50=("valuation_missing_neutral_count", "median"),
        valuation_missing_neutral_max=("valuation_missing_neutral_count", "max"),
        selected_tickers_p50=("selected_tickers", "median"),
    )


def selection_change_attribution(
    selections: pd.DataFrame,
    *,
    expected_pairs: int = 7740,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Describe valuation-driven membership changes without causal claims."""
    data = _attach_profile_parts(selections)
    keys = [
        "composition_id",
        "lookback_years",
        "trade_cost_bps",
        "window_start",
        "selection_date",
    ]
    baseline = data[data["declared_valuation_weight"].eq(0.0)][keys + ["tickers"]]
    baseline = baseline.rename(columns={"tickers": "baseline_tickers"})
    paired = data[data["declared_valuation_weight"].gt(0.0)].merge(
        baseline, on=keys, how="left", validate="many_to_one"
    )
    if len(paired) != int(expected_pairs) or paired["baseline_tickers"].isna().any():
        raise NdxValuationWeightAnalysisError(
            f"selection attribution requires {int(expected_pairs)} complete paired states"
        )

    def ticker_set(value: object) -> set[str]:
        return {ticker for ticker in str(value or "").split(";") if ticker}

    paired["selected_set"] = paired["tickers"].map(ticker_set)
    paired["baseline_set"] = paired["baseline_tickers"].map(ticker_set)
    paired["added_tickers"] = paired.apply(
        lambda row: ";".join(sorted(row["selected_set"] - row["baseline_set"])),
        axis=1,
    )
    paired["removed_tickers"] = paired.apply(
        lambda row: ";".join(sorted(row["baseline_set"] - row["selected_set"])),
        axis=1,
    )
    paired["replacement_count"] = paired.apply(
        lambda row: len(row["selected_set"] - row["baseline_set"]), axis=1
    )
    paired["selection_changed"] = paired["replacement_count"].gt(0)
    paired["overlap_ratio"] = paired.apply(
        lambda row: len(row["selected_set"] & row["baseline_set"])
        / max(1, len(row["baseline_set"])),
        axis=1,
    )
    summary = paired.groupby(
        ["composition_id", "declared_valuation_weight"], as_index=False
    ).agg(
        paired_selection_rows=("selection_date", "size"),
        changed_selection_rate=("selection_changed", "mean"),
        replacement_count_p50=("replacement_count", "median"),
        replacement_count_p90=("replacement_count", lambda values: values.quantile(0.90)),
        replacement_count_max=("replacement_count", "max"),
        overlap_ratio_p50=("overlap_ratio", "median"),
    )
    ticker_rows: list[dict[str, Any]] = []
    for weight, group in paired.groupby("declared_valuation_weight", sort=True):
        added: dict[str, int] = {}
        removed: dict[str, int] = {}
        for value in group["added_tickers"]:
            for ticker in ticker_set(value):
                added[ticker] = added.get(ticker, 0) + 1
        for value in group["removed_tickers"]:
            for ticker in ticker_set(value):
                removed[ticker] = removed.get(ticker, 0) + 1
        for ticker in sorted(set(added) | set(removed)):
            ticker_rows.append({
                "declared_valuation_weight": float(weight),
                "ticker": ticker,
                "added_count": added.get(ticker, 0),
                "removed_count": removed.get(ticker, 0),
                "net_entry_count": added.get(ticker, 0) - removed.get(ticker, 0),
                "attribution_scope": "descriptive_only_no_causal_claim",
            })
    ticker_summary = pd.DataFrame(ticker_rows).sort_values(
        ["declared_valuation_weight", "net_entry_count", "ticker"],
        ascending=[True, False, True],
        kind="stable",
    )
    return summary, ticker_summary


CHINESE_COLUMNS = {
    "composition_id": "组合构成",
    "profile_id": "方案编号",
    "declared_valuation_weight": "估值权重",
    "valuation_weight": "估值权重",
    "lookback_years": "回看年数",
    "trade_cost_bps": "交易成本基点",
    "start": "窗口开始",
    "end": "窗口结束",
    "candidate_pairs": "候选配对数",
    "candidate_p50_improvement_rate": "候选收益中位改善率",
    "candidate_p10_improvement_rate": "候选收益尾部改善率",
    "candidate_xirr_p50_delta_median": "年化收益中位差中位数",
    "candidate_xirr_p10_delta_median": "年化收益尾部差中位数",
    "window_pairs": "窗口配对数",
    "window_xirr_win_rate": "逐窗口收益改善率",
    "window_xirr_delta_p10": "逐窗口收益差P10",
    "window_xirr_delta_p50": "逐窗口收益差P50",
    "window_drawdown_delta_p50": "逐窗口回撤变化P50",
    "window_drawdown_delta_p90": "逐窗口回撤恶化P90",
    "candidate_maxdd_p90_delta_median": "候选最大回撤P90变化中位数",
    "turnover_delta_p50": "换手变化P50",
    "cost_delta_p50": "成本变化P50",
    "shortlist_passed": "训练候选门通过",
    "candidate_rows": "候选行数",
    "target_gate_passed_rows": "012922参照门通过行数",
    "incumbent_gate_passed_rows": "当前8只守门通过行数",
    "admission_passed_rows": "最终准入行数",
    "best_incumbent_xirr_p50_margin": "相对当前8只最佳收益P50差",
    "best_incumbent_xirr_p10_margin": "相对当前8只最佳收益P10差",
    "best_incumbent_dd_relief_p50": "相对当前8只最佳回撤改善P50",
    "best_incumbent_dd_relief_p90": "相对当前8只最佳回撤改善P90",
    "selection_rows": "重选记录数",
    "valuation_score_coverage_min": "估值评分覆盖率下限",
    "valuation_observed_coverage_min": "估值实值覆盖率下限",
    "valuation_missing_neutral_p50": "中性缺失数量P50",
    "valuation_missing_neutral_max": "中性缺失数量最大值",
    "selected_tickers_p50": "入选股票数P50",
    "paired_selection_rows": "配对重选记录数",
    "changed_selection_rate": "重选名单变化率",
    "replacement_count_p50": "替换股票数P50",
    "replacement_count_p90": "替换股票数P90",
    "replacement_count_max": "替换股票数最大值",
    "overlap_ratio_p50": "名单重合率P50",
    "ticker": "股票代码",
    "added_count": "新增入选次数",
    "removed_count": "退出入选次数",
    "net_entry_count": "净进入次数",
    "attribution_scope": "归因权限",
    "strategy_xirr_p50": "方案年化收益P50",
    "strategy_xirr_p10": "方案年化收益P10",
    "baseline_strategy_xirr_p50": "零估值权重年化收益P50",
    "baseline_strategy_xirr_p10": "零估值权重年化收益P10",
    "delta_strategy_xirr_p50": "年化收益P50差",
    "delta_strategy_xirr_p10": "年化收益P10差",
    "strategy_maxdd_p50": "方案最大回撤P50",
    "strategy_maxdd_p90": "方案最大回撤P90",
    "baseline_strategy_maxdd_p50": "零估值权重最大回撤P50",
    "baseline_strategy_maxdd_p90": "零估值权重最大回撤P90",
    "delta_strategy_maxdd_p50": "最大回撤P50变化",
    "delta_strategy_maxdd_p90": "最大回撤P90变化",
    "drawdown_p50_relief": "最大回撤P50改善",
    "drawdown_p90_relief": "最大回撤P90改善",
    "incumbent_xirr_p50_margin": "相对当前8只收益P50差",
    "incumbent_xirr_p10_margin": "相对当前8只收益P10差",
    "baseline_incumbent_xirr_p50_margin": "零估值权重相对当前8只收益P50差",
    "baseline_incumbent_xirr_p10_margin": "零估值权重相对当前8只收益P10差",
    "delta_incumbent_xirr_p50_margin": "相对当前8只收益P50边际变化",
    "delta_incumbent_xirr_p10_margin": "相对当前8只收益P10边际变化",
    "turnover_x_p50": "方案换手倍数P50",
    "baseline_turnover_x_p50": "零估值权重换手倍数P50",
    "delta_turnover_x_p50": "换手倍数P50变化",
    "cost_x_p50": "方案成本倍数P50",
    "baseline_cost_x_p50": "零估值权重成本倍数P50",
    "delta_cost_x_p50": "成本倍数P50变化",
    "selection_churn_p50": "方案重选变动率P50",
    "baseline_selection_churn_p50": "零估值权重重选变动率P50",
    "delta_selection_churn_p50": "重选变动率P50变化",
    "strategy_xirr": "方案年化收益",
    "baseline_strategy_xirr": "零估值权重年化收益",
    "delta_strategy_xirr": "年化收益差",
    "strategy_maxdd": "方案最大回撤",
    "baseline_strategy_maxdd": "零估值权重最大回撤",
    "delta_strategy_maxdd": "最大回撤变化",
    "drawdown_relief": "最大回撤改善",
    "incumbent_xirr_margin": "相对当前8只年化收益差",
    "baseline_incumbent_xirr_margin": "零估值权重相对当前8只年化收益差",
    "delta_incumbent_xirr_margin": "相对当前8只收益边际变化",
    "incumbent_dd_relief": "相对当前8只回撤改善",
    "baseline_incumbent_dd_relief": "零估值权重相对当前8只回撤改善",
    "delta_incumbent_dd_relief": "相对当前8只回撤边际变化",
    "turnover_x": "方案换手倍数",
    "baseline_turnover_x": "零估值权重换手倍数",
    "delta_turnover_x": "换手倍数变化",
    "cost_x": "方案成本倍数",
    "baseline_cost_x": "零估值权重成本倍数",
    "delta_cost_x": "成本倍数变化",
    "baseline_selection_churn_p50": "零估值权重重选变动率P50",
    "delta_selection_churn_p50": "重选变动率P50变化",
}


def _localized(frame: pd.DataFrame) -> pd.DataFrame:
    missing = sorted(set(frame.columns) - set(CHINESE_COLUMNS))
    if missing:
        raise NdxValuationWeightAnalysisError(
            f"Chinese report headers are missing: {missing}"
        )
    return frame.rename(
        columns={column: CHINESE_COLUMNS.get(column, column) for column in frame.columns}
    )


def _format_workbook(path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    percentage_markers = ("权重", "改善率", "覆盖率", "收益", "回撤", "换手", "成本")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            header = str(cell.value or "")
            values = [header] + [
                str(sheet.cell(row=row, column=cell.column).value or "")
                for row in range(2, min(sheet.max_row, 200) + 1)
            ]
            sheet.column_dimensions[cell.column_letter].width = min(
                42, max(10, max(len(value) for value in values) + 2)
            )
            if "基点" not in header and any(
                marker in header for marker in percentage_markers
            ):
                for row in range(2, sheet.max_row + 1):
                    sheet.cell(row=row, column=cell.column).number_format = "0.00%"
    workbook.save(path)


def build_valuation_weight_analysis_bundle(
    experiment_dir: str | Path,
    *,
    output_dir: str | Path | None = None,
    strategy_config_path: str | Path | None = None,
    expected_stage: str = "neutral_missing_valuation_weight",
    analysis_directory_name: str = "valuation_weight",
    workbook_filename: str = "v9_fcf_valuation_weight_analysis.xlsx",
) -> dict[str, Any]:
    root = Path(experiment_dir)
    manifest, frames = _read_verified_bundle(
        root, strategy_config_path=strategy_config_path
    )
    contract = manifest.get("analysis_contract") or {}
    if contract.get("research_stage") != expected_stage:
        raise NdxValuationWeightAnalysisError("experiment research stage changed")
    declared_weights = tuple(
        float(value) for value in contract.get("valuation_weights") or ()
    )
    if (
        len(declared_weights) < 2
        or declared_weights[0] != 0.0
        or tuple(sorted(set(declared_weights))) != declared_weights
    ):
        raise NdxValuationWeightAnalysisError("declared valuation weights are invalid")
    if contract.get("locked_validation_allowed") is not False:
        raise NdxValuationWeightAnalysisError("v9 must remain training-only")
    versions = set(frames["candidates"]["valuation_score_version_id"].dropna().astype(str))
    if versions != {str(contract.get("valuation_score_version_id"))}:
        raise NdxValuationWeightAnalysisError("valuation evidence version changed")

    expected_window_pairs = 36 * (len(declared_weights) - 1) * 18
    expected_selection_pairs = int(
        len(frames["selections"]) / len(declared_weights)
        * (len(declared_weights) - 1)
    )
    candidate_effects = paired_candidate_effects(
        frames["candidates"], declared_weights=declared_weights
    )
    window_effects = paired_window_effects(
        frames["windows"], expected_pairs=expected_window_pairs
    )
    summary = summarize_effects(
        candidate_effects,
        window_effects,
        declared_weights=declared_weights,
    )
    by_composition, by_lookback = effect_breakdowns(candidate_effects)
    absolute_gates = absolute_gate_summary(frames["candidates"])
    evidence = valuation_evidence_summary(frames["selections"])
    selection_changes, ticker_attribution = selection_change_attribution(
        frames["selections"], expected_pairs=expected_selection_pairs
    )
    shortlist = summary[summary["shortlist_passed"]]
    incumbent_passed = int(absolute_gates["incumbent_gate_passed_rows"].sum())
    if shortlist.empty:
        decision = "NO_VALUATION_WEIGHT_PASSED_ROBUSTNESS_GATE"
    elif incumbent_passed == 0:
        decision = "COMPONENT_SHORTLIST_INCUMBENT_GATE_FAILED_FORWARD_WATCH_ONLY"
    else:
        decision = "COMPONENT_SHORTLIST_REQUIRES_FRESH_VALIDATION"

    output = (
        Path(output_dir)
        if output_dir
        else root / "analysis" / analysis_directory_name
    )
    output.mkdir(parents=True, exist_ok=True)
    paths = {
        "summary": output / "valuation_weight_summary.csv",
        "candidate_effects": output / "paired_candidate_effects.csv",
        "window_effects": output / "paired_window_effects.csv",
        "evidence": output / "valuation_evidence_summary.csv",
        "by_composition": output / "effect_by_composition.csv",
        "by_lookback": output / "effect_by_lookback.csv",
        "absolute_gates": output / "absolute_gate_summary.csv",
        "selection_changes": output / "selection_change_summary.csv",
        "ticker_attribution": output / "ticker_entry_exit_attribution.csv",
        "workbook": output / workbook_filename,
        "manifest": output / "manifest.json",
    }
    candidate_columns = [
        "composition_id", "profile_id", "declared_valuation_weight",
        "lookback_years", "trade_cost_bps",
        "strategy_xirr_p50", "baseline_strategy_xirr_p50",
        "delta_strategy_xirr_p50", "strategy_xirr_p10",
        "baseline_strategy_xirr_p10", "delta_strategy_xirr_p10",
        "strategy_maxdd_p50", "baseline_strategy_maxdd_p50",
        "delta_strategy_maxdd_p50", "drawdown_p50_relief",
        "strategy_maxdd_p90", "baseline_strategy_maxdd_p90",
        "delta_strategy_maxdd_p90", "drawdown_p90_relief",
        "incumbent_xirr_p50_margin", "baseline_incumbent_xirr_p50_margin",
        "delta_incumbent_xirr_p50_margin", "incumbent_xirr_p10_margin",
        "baseline_incumbent_xirr_p10_margin",
        "delta_incumbent_xirr_p10_margin", "turnover_x_p50",
        "baseline_turnover_x_p50", "delta_turnover_x_p50", "cost_x_p50",
        "baseline_cost_x_p50", "delta_cost_x_p50", "selection_churn_p50",
        "baseline_selection_churn_p50", "delta_selection_churn_p50",
    ]
    window_columns = [
        "composition_id", "profile_id", "declared_valuation_weight",
        "lookback_years", "trade_cost_bps", "start", "end",
        "strategy_xirr", "baseline_strategy_xirr", "delta_strategy_xirr",
        "strategy_maxdd", "baseline_strategy_maxdd", "delta_strategy_maxdd",
        "drawdown_relief", "incumbent_xirr_margin",
        "baseline_incumbent_xirr_margin", "delta_incumbent_xirr_margin",
        "incumbent_dd_relief", "baseline_incumbent_dd_relief",
        "delta_incumbent_dd_relief", "turnover_x", "baseline_turnover_x",
        "delta_turnover_x", "cost_x", "baseline_cost_x", "delta_cost_x",
        "selection_churn_p50", "baseline_selection_churn_p50",
        "delta_selection_churn_p50",
    ]
    candidate_report = candidate_effects[candidate_columns]
    window_report = window_effects[window_columns]
    localized_summary = _localized(summary)
    localized_candidates = _localized(candidate_report)
    localized_windows = _localized(window_report)
    localized_evidence = _localized(evidence)
    localized_by_composition = _localized(by_composition)
    localized_by_lookback = _localized(by_lookback)
    localized_absolute_gates = _localized(absolute_gates)
    localized_selection_changes = _localized(selection_changes)
    localized_ticker_attribution = _localized(ticker_attribution)
    _atomic_csv(localized_summary, paths["summary"])
    _atomic_csv(localized_candidates, paths["candidate_effects"])
    _atomic_csv(localized_windows, paths["window_effects"])
    _atomic_csv(localized_evidence, paths["evidence"])
    _atomic_csv(localized_by_composition, paths["by_composition"])
    _atomic_csv(localized_by_lookback, paths["by_lookback"])
    _atomic_csv(localized_absolute_gates, paths["absolute_gates"])
    _atomic_csv(localized_selection_changes, paths["selection_changes"])
    _atomic_csv(localized_ticker_attribution, paths["ticker_attribution"])
    with pd.ExcelWriter(paths["workbook"], engine="openpyxl") as writer:
        localized_summary.to_excel(writer, sheet_name="权重结论", index=False)
        localized_candidates.to_excel(writer, sheet_name="候选配对", index=False)
        localized_windows.to_excel(writer, sheet_name="逐窗口配对", index=False)
        localized_evidence.to_excel(writer, sheet_name="估值证据覆盖", index=False)
        localized_by_composition.to_excel(writer, sheet_name="按组成稳定性", index=False)
        localized_by_lookback.to_excel(writer, sheet_name="按回看期稳定性", index=False)
        localized_absolute_gates.to_excel(writer, sheet_name="绝对准入门", index=False)
        localized_selection_changes.to_excel(writer, sheet_name="重选变化", index=False)
        localized_ticker_attribution.to_excel(writer, sheet_name="股票进出频次", index=False)
        pd.DataFrame([
            {"规则": "分析策略", "设定": ANALYSIS_POLICY_ID},
            {"规则": "候选P50和P10改善率下限", "设定": MIN_CANDIDATE_IMPROVEMENT_RATE},
            {"规则": "逐窗口收益改善率下限", "设定": MIN_WINDOW_XIRR_WIN_RATE},
            {"规则": "逐窗口回撤P90最大恶化", "设定": MAX_WINDOW_DRAWDOWN_P90_DETERIORATION},
            {"规则": "训练结果权限", "设定": "只可形成锁定验证候选，不可修改实盘"},
        ]).to_excel(writer, sheet_name="预注册门槛", index=False)
    _format_workbook(paths["workbook"])

    payload = {
        "schema_version": 1,
        "analysis_policy_id": ANALYSIS_POLICY_ID,
        "experiment_id": manifest.get("experiment_id"),
        "experiment_manifest_sha256": _content_hash(root / "manifest.json"),
        "candidate_pairs": len(candidate_effects),
        "window_pairs": len(window_effects),
        "paired_selection_rows": int(selection_changes["paired_selection_rows"].sum()),
        "shortlisted_weights": shortlist["valuation_weight"].tolist(),
        "incumbent_gate_passed_rows": incumbent_passed,
        "decision": decision,
        "research_only": True,
        "live_change_allowed": False,
        "files": {
            key: {"filename": path.name, "sha256": _content_hash(path)}
            for key, path in paths.items()
            if key != "manifest"
        },
    }
    _atomic_json(payload, paths["manifest"])
    return payload
