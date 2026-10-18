"""White-box experiment orchestration for Nasdaq-100 stock competition.

The lab keeps experiment planning, checkpoint identity, diagnostics and exports
separate from the backtest engine. A failed row means only that the current
shape failed its declared gates; it never terminates the research direction.
"""
from __future__ import annotations

import copy
import hashlib
import json
import math
import os
import platform
import re
import sys
from datetime import datetime, timezone
from itertools import product
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
import numpy as np
import scipy
import yaml

from backtest.ndx_quality import QUALITY_POLICY_ID
from data.ndx_valuation_score import SUPPORTED_VALUATION_SCORE_POLICY_IDS
from core.helpers import atomic_write_text


LAB_SCHEMA_VERSION = 1
QQQ_ROLE = "benchmark_only_not_candidate"
OBJECTIVE_CONTRACT_KEYS = {
    "candidate_universe_source",
    "primary_target",
    "qqq_performance_role",
    "incumbent_role",
}
ALLOWED_PHASE_ROLES = {
    "exploratory_training",
    "sensitivity",
    "locked_validation",
}
GRID_KEYS = (
    "top_ns",
    "weight_modes",
    "safe_weights",
    "rebalance_months",
    "max_weights",
)
SWEEP_KEYS = (
    "lookback_years",
    "trade_cost_bps",
    "min_history_years",
    "safe_annual",
    "step_months",
    "max_windows",
)
CONSTRUCTION_EXPORT_KEYS = (
    "profile_id",
    "challenger_sleeve_weight",
    "selection_mode",
    "correlation_lookback_days",
    "correlation_threshold",
    "correlation_pool_size",
    "max_names_per_cluster",
    "cluster_max_weight",
    "top_weight_count",
    "top_weight_cap",
    "min_effective_n",
    "fundamental_score_weight",
    "quality_policy_id",
    "valuation_score_weight",
    "valuation_policy_id",
    "quality_selection_mode",
    "quality_challenge_pool_size",
    "max_quality_replacements",
    "max_names_per_sic4",
    "max_names_per_sic2",
    "sic4_max_weight",
    "sic2_max_weight",
    "candidate_tickers",
    "defense_mode",
    "defense_signal",
    "target_vol",
    "defense_band",
    "defense_vol_window",
    "defense_exec_band",
    "defense_drawdown_lookback_days",
    "defense_drawdown_trigger",
    "defense_drawdown_full",
    "defense_max_safe_weight",
    "defense_release_step",
)
RUN_DEFAULTS = {
    "monthly": 3000.0,
    "total_cap": 37000.0,
    "target": "QQQ",
    "min_years": 1.0,
    "lookback_years": 2.0,
    "min_history_years": 0.5,
    "trade_cost_bps": 10.0,
    "safe_annual": 0.04,
    "step_months": 1,
    "max_windows": 0,
    "local_only": True,
    "include_samples": True,
    "comparison_priority": "target_then_incumbent",
}
ENGINE_SOURCE_FILES = (
    "backtest/ndx_competition.py",
    "backtest/ndx_quality.py",
    "backtest/engine.py",
    "backtest/metrics.py",
    "backtest/preference.py",
    "backtest/preference_walkforward.py",
    "backtest/risk_envelope.py",
    "backtest/weights.py",
    "core/engine.py",
    "data/ndx_history.py",
    "data/ndx_fundamentals.py",
    "data/ndx_fundamental_normalization.py",
    "data/ndx_sic_history.py",
    "data/candidate_pool.py",
    "data/fetch.py",
    "data/funds.py",
    "data/versioning.py",
    "core/config.py",
)

EXCEL_SHEET_NAMES = {
    "summary": "实验汇总",
    "jobs": "作业状态",
    "candidates": "组合候选",
    "windows": "滚动窗口",
    "selections": "调仓明细",
    "gate_failures": "门槛失败",
    "engine_skips": "引擎跳过",
    "versions": "数据版本",
    "current_holdings": "当前持仓",
}

EXCEL_COLUMN_NAMES = {
    "field": "项目",
    "value": "内容",
    "experiment_id": "实验编号",
    "phase_id": "阶段编号",
    "phase_role": "阶段角色",
    "job_id": "作业编号",
    "job_sha256": "作业指纹",
    "engine_sha256": "引擎指纹",
    "runtime_sha256": "运行环境指纹",
    "strategy_config_sha256": "策略配置指纹",
    "status": "状态",
    "error": "错误信息",
    "profile_id": "方案编号",
    "shape_id": "形态编号",
    "basket": "组合名称",
    "candidate_assets": "候选资产范围",
    "candidate_tickers": "冻结候选股票",
    "candidate_member_count": "当期候选成分数",
    "qqq_role": "QQQ角色",
    "target_benchmark": "第二层参照",
    "research_start": "研究开始日期",
    "research_end": "研究结束日期",
    "start": "开始日期",
    "end": "结束日期",
    "window_start": "滚动窗口开始",
    "window_end": "滚动窗口结束",
    "selection_date": "调仓日期",
    "lookback_years": "回看年数",
    "top_n": "持仓数量",
    "weight_mode": "权重方法",
    "rebalance_months": "调仓月数",
    "safe_weight": "安全资产比例",
    "challenger_sleeve_weight": "挑战组合比例",
    "max_weight": "单股比例上限",
    "min_weight": "单股比例下限",
    "trade_cost_bps": "交易成本（基点）",
    "windows": "滚动窗口数",
    "incumbent_win_rate": "战胜当前8只胜率",
    "incumbent_dd_win_rate": "回撤优于当前8只胜率",
    "incumbent_xirr_p50_margin": "相对当前8只XIRR中位差",
    "incumbent_xirr_p10_margin": "相对当前8只XIRR尾部差",
    "incumbent_dd_relief_p50": "相对当前8只回撤改善中位",
    "incumbent_dd_relief_p90": "相对当前8只回撤改善压力值",
    "incumbent_excess_p50": "相对当前8只收益中位差",
    "incumbent_excess_p10": "相对当前8只收益尾部差",
    "target_win_rate": "战胜第二层参照胜率",
    "target_dd_win_rate": "回撤优于第二层参照胜率",
    "target_excess_p50": "相对第二层参照收益中位差",
    "target_excess_p10": "相对第二层参照收益尾部差",
    "target_dd_relief_p50": "相对第二层参照回撤改善中位",
    "target_dd_relief_p10": "相对第二层参照回撤改善尾部",
    "strategy_xirr_p10": "策略XIRR尾部",
    "strategy_xirr_p50": "策略XIRR中位",
    "strategy_xirr_p90": "策略XIRR高位",
    "strategy_maxdd_p50": "策略最大回撤中位",
    "strategy_maxdd_p90": "策略最大回撤压力值",
    "turnover_x_p50": "换手倍数中位",
    "cost_x_p50": "成本占投入中位",
    "selection_churn_p50": "持仓变更率中位",
    "effective_n_p10": "有效持仓数尾部",
    "top1_p90": "最大单股比例压力值",
    "top3_p90": "前三大比例压力值",
    "selection_mode": "选股方法",
    "fundamental_score_weight": "基本面评分权重",
    "quality_policy_id": "质量规则编号",
    "valuation_score_weight": "估值评分权重",
    "valuation_policy_id": "估值规则编号",
    "max_names_per_sic4": "同SIC四位最多股票数",
    "max_names_per_sic2": "同SIC二位最多股票数",
    "sic4_max_weight": "SIC四位比例上限",
    "sic2_max_weight": "SIC二位比例上限",
    "max_sic4_weight_p90": "SIC四位集中度压力值",
    "max_sic2_weight_p90": "SIC二位集中度压力值",
    "admission_passed": "原主目标准入通过",
    "incumbent_gate_passed": "当前8只守门通过",
    "target_gate_passed": "第二层参照守门通过",
    "concentration_guard_passed": "集中度守门通过",
    "failed_checks": "失败检查项",
    "shape_diagnosis": "形态诊断",
    "next_action": "下一步动作",
    "direction_status": "研究方向状态",
    "tickers_observed": "历史出现股票",
    "tickers": "当期持仓股票",
    "weights": "当期风险仓权重",
    "sic2_by_ticker": "股票SIC二位分类",
    "sic4_by_ticker": "股票SIC四位分类",
    "sic2_weights": "SIC二位权重",
    "sic4_weights": "SIC四位权重",
    "trigger": "调仓触发原因",
    "report_date": "成分报告日期",
    "available_date": "数据可用日期",
    "price_coverage": "价格覆盖率",
    "score_coverage": "评分覆盖率",
    "future_data_guard_passed": "未来数据防护通过",
    "candidate_pool_passed": "冻结候选池审计通过",
    "weight_sum_passed": "权重合计审计通过",
    "weight_cap_passed": "单股上限审计通过",
    "weight_floor_passed": "单股下限审计通过",
    "ticker_count_passed": "持仓数量审计通过",
    "selection_kind": "持仓类型",
    "ticker": "股票代码",
    "portfolio_weight": "组合实际比例",
    "risk_sleeve_weight": "风险仓内部比例",
    "source_name": "数据源名称",
    "version_id": "版本编号",
    "content_sha256": "内容指纹",
    "source": "来源",
    "source_locator": "来源位置",
    "data_start": "数据开始日期",
    "data_end": "数据结束日期",
    "row_count": "数据行数",
    "fetched_at": "获取时间",
    "cached": "是否缓存",
    "stale": "是否过期",
}

EXCEL_COLUMN_NAMES.update({
    "candidate_rows": "候选方案行数",
    "admission_passed_rows": "原主目标准入行数",
    "best_target_excess_p50": "相对第二层参照最佳收益中位差",
    "best_qqq_excess_p50": "相对QQQ最佳收益中位差",
    "best_qqq_excess_p10": "相对QQQ最佳收益尾部差",
    "best_qqq_dd_relief_p50": "相对QQQ最佳回撤改善中位",
    "direction_terminal": "研究方向是否终止",
    "expected_candidate_rows": "预计候选行数",
    "actual_candidate_rows": "实际候选行数",
    "actual_window_rows": "实际窗口行数",
    "duration_seconds": "运行秒数",
    "started_at": "开始运行时间",
    "completed_at": "完成时间",
    "target": "第二层参照代码",
    "top_ns": "测试持仓数量",
    "weight_modes": "测试权重方法",
    "safe_weights": "测试安全资产比例",
    "max_weights": "测试单股上限",
    "construction_profile_count": "构建方案数",
    "construction_profile_ids": "构建方案编号",
    "include_samples": "是否输出逐窗口样本",
    "correlation_lookback_days": "相关性回看天数",
    "correlation_threshold": "相关性阈值",
    "correlation_pool_size": "相关性候选池数量",
    "max_names_per_cluster": "同相关簇最多股票数",
    "cluster_max_weight": "相关簇比例上限",
    "top_weight_count": "前几大股票约束数量",
    "top_weight_cap": "前几大风险仓位比例上限",
    "min_effective_n": "最低有效持仓数",
    "quality_selection_mode": "质量选股方式",
    "quality_challenge_pool_size": "质量挑战池数量",
    "max_quality_replacements": "最多质量替换数",
    "defense_mode": "防御模式",
    "defense_signal": "防御信号",
    "target_vol": "目标波动率",
    "defense_band": "防御调整区间",
    "defense_vol_window": "防御波动窗口",
    "defense_exec_band": "防御执行带",
    "defense_drawdown_lookback_days": "防御回撤回看天数",
    "defense_drawdown_trigger": "防御回撤触发值",
    "defense_drawdown_full": "防御回撤满档值",
    "defense_max_safe_weight": "防御安全资产上限",
    "defense_release_step": "防御释放步长",
    "qqq_win_rate": "战胜QQQ胜率",
    "qqq_excess_p50": "相对QQQ收益中位差",
    "qqq_excess_p10": "相对QQQ收益尾部差",
    "qqq_dd_relief_p50": "相对QQQ回撤改善中位",
    "qqq_dd_relief_p10": "相对QQQ回撤改善尾部",
    "incumbent_dd_relief_direct_p50": "相对当前8只回撤直接改善中位",
    "incumbent_dd_relief_direct_p10": "相对当前8只回撤直接改善尾部",
    "avg_safe_weight_p10": "平均安全资产比例尾部",
    "avg_safe_weight_p50": "平均安全资产比例中位",
    "avg_safe_weight_p90": "平均安全资产比例高位",
    "max_pair_correlation_p90": "最大两股相关性压力值",
    "max_cluster_weight_p90": "最大相关簇比例压力值",
    "max_top_weight_p90": "前几大风险仓位比例压力值",
    "observed_min_effective_n": "实际最低有效持仓数",
    "correlation_rejected_p50": "相关性剔除数量中位",
    "min_quality_score_coverage": "最低质量评分覆盖率",
    "quality_replacements_p50": "质量替换数量中位",
    "quality_replacements_p90": "质量替换数量压力值",
    "sic_rejected_p50": "SIC约束剔除数量中位",
    "defense_signal_prehistory_seeded": "防御信号是否含起点前历史",
    "defense_signal_strict_asof_passed": "防御信号时点审计通过",
    "min_price_coverage": "最低价格覆盖率",
    "min_score_coverage": "最低评分覆盖率",
    "target_classification": "第二层参照结果分类",
    "target_return_passed": "第二层参照收益门通过",
    "target_risk_passed": "第二层参照风险门通过",
    "terminal_direction_decision": "是否形成终止方向结论",
    "input_version_bundle_id": "输入数据版本包编号",
    "input_version_ids": "输入数据版本编号",
    "history_version_id": "历史成分版本编号",
    "price_version_id": "价格版本编号",
    "target_version_id": "第二层参照版本编号",
    "quality_score_version_id": "质量评分版本编号",
    "valuation_score_version_id": "估值评分版本编号",
    "future_data_guard": "未来数据防护规则",
    "strategy_xirr": "策略XIRR",
    "strategy_maxdd": "策略最大回撤",
    "strategy_final": "策略期末价值",
    "target_xirr": "第二层参照XIRR",
    "target_maxdd": "第二层参照最大回撤",
    "qqq_xirr": "QQQ XIRR",
    "qqq_maxdd": "QQQ最大回撤",
    "incumbent_xirr": "当前8只XIRR",
    "incumbent_maxdd": "当前8只最大回撤",
    "target_excess_xirr": "相对第二层参照XIRR差",
    "target_dd_relief": "相对第二层参照回撤改善",
    "qqq_excess_xirr": "相对QQQ XIRR差",
    "qqq_dd_relief": "相对QQQ回撤改善",
    "incumbent_xirr_margin": "相对当前8只XIRR差",
    "incumbent_dd_relief": "相对当前8只回撤改善",
    "turnover_x": "换手倍数",
    "cost_x": "成本占投入比例",
    "avg_safe_weight": "平均安全资产比例",
    "min_safe_weight": "最低安全资产比例",
    "max_safe_weight": "最高安全资产比例",
    "safe_turnover_x": "安全资产换手倍数",
    "defense_signal_first_date": "防御信号首日",
    "defense_signal_last_date": "防御信号末日",
    "selection_count": "调仓次数",
    "accession_number": "SEC申报编号",
    "expected_security_count": "当期完整成分证券数",
    "expected_issuer_count": "当期完整成分发行人数",
    "price_count": "有效价格股票数",
    "scored_count": "有效评分股票数",
    "selection_mode_observed": "实际选股方式",
    "correlation_observations": "相关性观测数",
    "correlation_price_start": "相关性价格开始日期",
    "correlation_price_end": "相关性价格结束日期",
    "correlation_strict_asof_passed": "相关性时点审计通过",
    "correlation_pool_size_observed": "实际相关性候选池数量",
    "max_pair_correlation": "最大两股相关性",
    "cluster_labels": "股票相关簇编号",
    "cluster_weights": "相关簇权重",
    "observed_max_cluster_weight": "实际最大相关簇比例",
    "observed_top_weight": "实际前几大风险仓位比例",
    "observed_effective_n": "实际有效持仓数",
    "observed_max_names_per_cluster": "实际同簇最多股票数",
    "rejected_correlated_tickers": "因相关性剔除股票",
    "fundamental_score_weight_observed": "实际基本面评分权重",
    "quality_policy_id_observed": "实际质量规则编号",
    "valuation_score_weight_observed": "实际估值评分权重",
    "valuation_policy_id_observed": "实际估值规则编号",
    "valuation_score_coverage": "估值评分覆盖率",
    "valuation_observed_coverage": "估值观测覆盖率",
    "valuation_score_date": "估值评分日期",
    "valuation_future_data_guard_passed": "估值未来数据守门通过",
    "valuation_missing_neutral_count": "估值缺失中性数量",
    "valuation_evidence_passed": "估值证据审计通过",
    "valuation_score_coverage_passed": "估值评分覆盖率通过",
    "valuation_future_data_guard_audit_passed": "估值未来数据审计通过",
    "quality_score_coverage": "质量评分覆盖率",
    "quality_score_date": "质量评分日期",
    "quality_future_data_guard_observed": "质量数据时点审计结果",
    "quality_selection_mode_observed": "实际质量选股方式",
    "quality_challenge_pool_size_observed": "实际质量挑战池数量",
    "max_quality_replacements_observed": "实际质量替换上限",
    "quality_challenge_pool_tickers": "质量挑战池股票",
    "quality_replacement_pairs": "质量替换明细",
    "observed_quality_replacements": "实际质量替换数",
    "observed_max_names_per_sic4": "实际同SIC四位最多股票数",
    "observed_max_names_per_sic2": "实际同SIC二位最多股票数",
    "observed_max_sic4_weight": "实际SIC四位最大比例",
    "observed_max_sic2_weight": "实际SIC二位最大比例",
    "rejected_sic_tickers": "因SIC约束剔除股票",
    "rejected_sic_reasons": "SIC约束剔除原因",
    "observed_weight_sum": "实际权重合计",
    "observed_max_weight": "实际最大单股比例",
    "observed_min_weight": "实际最小单股比例",
    "effective_max_weight": "有效单股比例上限",
    "effective_min_weight": "有效单股比例下限",
    "churn": "持仓变更率",
    "turnover": "换手金额",
    "cost": "交易成本金额",
    "correlation_evidence_passed": "相关性证据审计通过",
    "correlation_future_data_guard_passed": "相关性未来数据防护通过",
    "correlation_pair_cap_passed": "两股相关性上限通过",
    "cluster_name_cap_passed": "相关簇股票数上限通过",
    "cluster_weight_cap_passed": "相关簇比例上限通过",
    "top_weight_cap_passed": "前几大风险仓位上限通过",
    "effective_n_floor_passed": "有效持仓数下限通过",
    "quality_evidence_passed": "质量证据审计通过",
    "quality_score_coverage_passed": "质量评分覆盖率通过",
    "quality_future_data_guard_passed": "质量未来数据防护通过",
    "quality_replacement_guard_passed": "质量替换审计通过",
    "sic4_name_cap_passed": "SIC四位股票数上限通过",
    "sic2_name_cap_passed": "SIC二位股票数上限通过",
    "sic4_weight_cap_passed": "SIC四位比例上限通过",
    "sic2_weight_cap_passed": "SIC二位比例上限通过",
    "defense_signal_guard_passed": "防御信号审计通过",
    "failed_check": "失败检查项",
    "selection_rows": "去重前调仓记录数",
    "window_conditioned_rows": "窗口条件记录数",
    "duplicate_window_rows": "重叠窗口重复记录数",
    "effective_n_p50": "有效持仓数中位",
    "top1_p50": "最大单股比例中位",
    "top3_p50": "前三大比例中位",
    "window_conditioned_churn_p50": "窗口条件持仓变更率中位",
    "concentration_effective_n_margin": "相对当前8只有效持仓数差",
    "concentration_top1_margin": "相对当前8只最大单股比例余量",
    "concentration_top3_margin": "相对当前8只前三大比例余量",
    "analysis_role": "分析角色代码",
    "analysis_role_cn": "分析角色",
    "dimension": "比较维度",
    "incumbent_win_rate_median": "战胜当前8只胜率中位",
    "incumbent_win_rate_best": "战胜当前8只最佳胜率",
    "incumbent_xirr_p50_margin_median": "相对当前8只XIRR中位差的中位",
    "incumbent_xirr_p50_margin_best": "相对当前8只XIRR中位差的最佳值",
    "incumbent_xirr_p10_margin_median": "相对当前8只XIRR尾部差的中位",
    "incumbent_xirr_p10_margin_best": "相对当前8只XIRR尾部差的最佳值",
    "incumbent_dd_relief_p50_median": "相对当前8只回撤改善中位的中位",
    "incumbent_dd_relief_p90_median": "相对当前8只回撤压力改善的中位",
    "concentration_pass_rate": "集中度合规率",
    "turnover_x_p50_median": "换手倍数中位的中位",
    "sic2": "SEC SIC二位分类",
})


class NdxExperimentLabError(ValueError):
    """Raised when an experiment spec or checkpoint is unsafe to compare."""


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def _json_hash(value: Any) -> str:
    payload = json.dumps(value, ensure_ascii=True, sort_keys=True, default=str)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _engine_source_hash() -> str:
    root = Path(__file__).resolve().parents[1]
    digest = hashlib.sha256()
    for relative in ENGINE_SOURCE_FILES:
        path = root / relative
        if not path.exists():
            raise NdxExperimentLabError(f"engine source file is missing: {relative}")
        digest.update(relative.encode("utf-8"))
        digest.update(b"\0")
        digest.update(path.read_bytes())
        digest.update(b"\0")
    return digest.hexdigest()


def _slug(value: Any) -> str:
    text = str(value).strip().lower().replace(".", "p")
    text = re.sub(r"[^a-z0-9_-]+", "-", text).strip("-")
    return text or "value"


def _as_list(value: Any, *, name: str) -> list:
    if isinstance(value, (list, tuple)):
        values = list(value)
    elif value is None:
        values = []
    else:
        values = [value]
    if not values:
        raise NdxExperimentLabError(f"{name} must contain at least one value")
    return values


def _joined(value: Iterable[Any]) -> str:
    return ",".join(str(item) for item in value)


def load_experiment_spec(path: str | Path) -> dict:
    source = Path(path)
    with source.open("r", encoding="utf-8") as handle:
        raw = yaml.safe_load(handle) or {}
    if not isinstance(raw, dict):
        raise NdxExperimentLabError("experiment spec root must be a mapping")
    if int(raw.get("schema_version") or 0) != LAB_SCHEMA_VERSION:
        raise NdxExperimentLabError(
            f"unsupported experiment schema: {raw.get('schema_version')!r}"
        )
    experiment_id = str(raw.get("experiment_id") or "").strip()
    if not experiment_id:
        raise NdxExperimentLabError("experiment_id is required")
    if raw.get("research_only") is not True:
        raise NdxExperimentLabError("research_only must be true")
    if raw.get("candidate_assets") != "nasdaq_100_individual_securities":
        raise NdxExperimentLabError(
            "candidate_assets must be nasdaq_100_individual_securities"
        )
    if raw.get("qqq_role") != QQQ_ROLE:
        raise NdxExperimentLabError(
            f"qqq_role must be {QQQ_ROLE}; QQQ can never be a candidate"
        )
    objective_contract = raw.get("objective_contract")
    if objective_contract is not None:
        if not isinstance(objective_contract, dict):
            raise NdxExperimentLabError("objective_contract must be a mapping")
        missing_contract = sorted(OBJECTIVE_CONTRACT_KEYS - set(objective_contract))
        if missing_contract:
            raise NdxExperimentLabError(
                f"objective_contract is missing fields: {missing_contract}"
            )
        primary_target = str(objective_contract.get("primary_target") or "").upper()
        if not primary_target or primary_target == "QQQ":
            raise NdxExperimentLabError(
                "objective_contract.primary_target must be a non-QQQ optimization target"
            )
        if objective_contract.get("qqq_performance_role") != "auxiliary_reference_only":
            raise NdxExperimentLabError(
                "objective_contract.qqq_performance_role must be auxiliary_reference_only"
            )
    analysis_contract = raw.get("analysis_contract")
    if analysis_contract is not None:
        if not isinstance(analysis_contract, dict):
            raise NdxExperimentLabError("analysis_contract must be a mapping")
        if not str(analysis_contract.get("component_family") or "").strip():
            raise NdxExperimentLabError(
                "analysis_contract.component_family is required"
            )
        if not str(analysis_contract.get("baseline_profile_id") or "").strip():
            raise NdxExperimentLabError(
                "analysis_contract.baseline_profile_id is required"
            )
    phases = raw.get("phases") or []
    if not isinstance(phases, list) or not phases:
        raise NdxExperimentLabError("at least one experiment phase is required")

    normalized = dict(raw)
    normalized["_spec_path"] = str(source.resolve())
    normalized["_spec_sha256"] = _json_hash(raw)
    normalized["_engine_sha256"] = _engine_source_hash()
    normalized["_runtime_versions"] = {
        "python": platform.python_version(),
        "python_implementation": platform.python_implementation(),
        "numpy": np.__version__,
        "pandas": pd.__version__,
        "scipy": scipy.__version__,
        "byteorder": sys.byteorder,
    }
    normalized["_runtime_sha256"] = _json_hash(normalized["_runtime_versions"])
    return normalized


def bind_strategy_config(spec: dict, path: str | Path) -> dict:
    config_path = Path(path).resolve()
    if not config_path.exists():
        raise NdxExperimentLabError(f"strategy config is missing: {config_path}")
    bound = dict(spec)
    bound["_strategy_config_path"] = str(config_path)
    bound["_strategy_config_sha256"] = hashlib.sha256(
        config_path.read_bytes()
    ).hexdigest()
    return bound


def _normalize_grid(grid: dict, *, prefix: str) -> dict[str, list]:
    if not isinstance(grid, dict):
        raise NdxExperimentLabError(f"{prefix}.shared_grid must be a mapping")
    normalized = {}
    for key in GRID_KEYS:
        normalized[key] = _as_list(grid.get(key), name=f"{prefix}.{key}")
    modes = [str(item).strip() for item in normalized["weight_modes"]]
    if any(not item for item in modes):
        raise NdxExperimentLabError(f"{prefix}.weight_modes contains an empty value")
    normalized["weight_modes"] = modes
    return normalized


def _profile_id_value(value: Any) -> str:
    if isinstance(value, float):
        value = f"{value:g}"
    return re.sub(r"[^a-zA-Z0-9_.-]+", "-", str(value)).strip("-") or "none"


def _expand_construction_profile_families(
    families: Any,
    *,
    prefix: str,
) -> list[dict]:
    if not isinstance(families, list) or not families:
        raise NdxExperimentLabError(
            f"{prefix}.construction_profile_families must be a non-empty list"
        )
    profiles = []
    seen_ids = set()
    for family_index, family in enumerate(families, start=1):
        if not isinstance(family, dict):
            raise NdxExperimentLabError(
                f"{prefix}.construction_profile_families[{family_index}] must be a mapping"
            )
        id_prefix = str(family.get("id_prefix") or "").strip()
        fixed = family.get("fixed") or {}
        grid = family.get("grid") or {}
        if not id_prefix or not isinstance(fixed, dict) or not isinstance(grid, dict):
            raise NdxExperimentLabError(
                f"{prefix}.construction_profile_families[{family_index}] is invalid"
            )
        if not grid:
            raise NdxExperimentLabError(
                f"{prefix}.{id_prefix}.grid must not be empty"
            )
        grid_keys = list(grid)
        dimensions = [
            _as_list(grid[key], name=f"{prefix}.{id_prefix}.grid.{key}")
            for key in grid_keys
        ]
        for combination in product(*dimensions):
            variable = dict(zip(grid_keys, combination))
            profile_id = "__".join([
                _profile_id_value(id_prefix),
                *(
                    f"{_profile_id_value(key)}-{_profile_id_value(value)}"
                    for key, value in variable.items()
                ),
            ])
            if profile_id in seen_ids:
                raise NdxExperimentLabError(
                    f"{prefix}.construction_profile_families generated duplicate id "
                    f"{profile_id!r}"
                )
            seen_ids.add(profile_id)
            profiles.append({"id": profile_id, **copy.deepcopy(fixed), **variable})
    return profiles


def _normalize_run_params(values: dict, grid: dict, *, prefix: str) -> dict:
    params = {**RUN_DEFAULTS, **values, **grid}
    if not params.get("start"):
        raise NdxExperimentLabError(f"{prefix}.start is required")
    params["start"] = str(params["start"])
    params["end"] = str(params["end"]) if params.get("end") else None
    params["target"] = str(params.get("target") or "QQQ").upper().strip()
    params["monthly"] = float(params["monthly"])
    params["total_cap"] = float(params["total_cap"])
    params["min_years"] = float(params["min_years"])
    params["lookback_years"] = float(params["lookback_years"])
    params["min_history_years"] = float(params["min_history_years"])
    params["trade_cost_bps"] = float(params["trade_cost_bps"])
    params["safe_annual"] = float(params["safe_annual"])
    params["step_months"] = int(params["step_months"])
    params["max_windows"] = int(params["max_windows"])
    params["local_only"] = bool(params["local_only"])
    params["include_samples"] = bool(params["include_samples"])
    params["comparison_priority"] = str(
        params.get("comparison_priority") or "target_then_incumbent"
    ).strip().lower()
    if params["comparison_priority"] not in {
        "target_then_incumbent",
        "incumbent_first",
    }:
        raise NdxExperimentLabError(
            f"{prefix}.comparison_priority must be target_then_incumbent or incumbent_first"
        )
    if not params["local_only"]:
        raise NdxExperimentLabError(
            f"{prefix}.local_only must be true; refresh data before starting the lab"
        )
    if params["step_months"] <= 0:
        raise NdxExperimentLabError(f"{prefix}.step_months must be positive")
    if params["min_years"] <= 0 or params["lookback_years"] <= 0:
        raise NdxExperimentLabError(f"{prefix} year parameters must be positive")

    for key in GRID_KEYS:
        params[key] = list(grid[key])
    params["top_ns"] = [int(item) for item in params["top_ns"]]
    params["safe_weights"] = [float(item) for item in params["safe_weights"]]
    params["rebalance_months"] = [int(item) for item in params["rebalance_months"]]
    params["max_weights"] = [float(item) for item in params["max_weights"]]
    construction_profiles = params.get("construction_profiles")
    profile_families = params.pop("construction_profile_families", None)
    if construction_profiles is not None and profile_families is not None:
        raise NdxExperimentLabError(
            f"{prefix} cannot declare both construction_profiles and "
            "construction_profile_families"
        )
    if profile_families is not None:
        construction_profiles = _expand_construction_profile_families(
            profile_families,
            prefix=prefix,
        )
        params["construction_profiles"] = construction_profiles
    if construction_profiles is not None:
        if not isinstance(construction_profiles, list) or not construction_profiles:
            raise NdxExperimentLabError(
                f"{prefix}.construction_profiles must be a non-empty list"
            )
        ids = []
        for index, profile in enumerate(construction_profiles, start=1):
            if not isinstance(profile, dict):
                raise NdxExperimentLabError(
                    f"{prefix}.construction_profiles[{index}] must be a mapping"
                )
            profile_id = str(profile.get("id") or profile.get("profile_id") or "").strip()
            if not profile_id:
                raise NdxExperimentLabError(
                    f"{prefix}.construction_profiles[{index}].id is required"
                )
            ids.append(profile_id)
        if len(ids) != len(set(ids)):
            raise NdxExperimentLabError(
                f"{prefix}.construction_profiles contains duplicate ids"
            )
        params["construction_profiles"] = copy.deepcopy(construction_profiles)
    return params


def _make_job(
    spec: dict,
    phase: dict,
    params: dict,
    *,
    job_hint: str,
) -> dict:
    phase_id = str(phase.get("phase_id") or "").strip()
    role = str(phase.get("role") or "").strip()
    objective_contract = dict(spec.get("objective_contract") or {})
    primary_target = str(objective_contract.get("primary_target") or "").upper()
    if primary_target and params.get("target") != primary_target:
        raise NdxExperimentLabError(
            f"{phase_id}.target={params.get('target')} violates primary target {primary_target}"
        )
    identity = {
        "experiment_id": spec["experiment_id"],
        "engine_sha256": spec["_engine_sha256"],
        "runtime_sha256": spec["_runtime_sha256"],
        "strategy_config_sha256": spec.get("_strategy_config_sha256"),
        "phase_id": phase_id,
        "phase_role": role,
        "params": params,
    }
    job_sha256 = _json_hash(identity)
    job_id = f"{_slug(phase_id)}__{_slug(job_hint)}__{job_sha256[:8]}"
    expected_rows = (
        len(params["construction_profiles"])
        if params.get("construction_profiles") is not None
        else math.prod(len(params[key]) for key in GRID_KEYS)
    )
    return {
        "schema_version": LAB_SCHEMA_VERSION,
        "experiment_id": spec["experiment_id"],
        "research_question": spec.get("research_question"),
        "objective_contract": objective_contract or None,
        "analysis_contract": copy.deepcopy(spec.get("analysis_contract")) or None,
        "direction_status": str(spec.get("direction_status") or "active_iteration"),
        "spec_sha256": spec["_spec_sha256"],
        "engine_sha256": spec["_engine_sha256"],
        "runtime_sha256": spec["_runtime_sha256"],
        "runtime_versions": spec["_runtime_versions"],
        "strategy_config_path": spec.get("_strategy_config_path"),
        "strategy_config_sha256": spec.get("_strategy_config_sha256"),
        "spec_path": spec["_spec_path"],
        "phase_id": phase_id,
        "phase_role": role,
        "phase_rationale": phase.get("rationale"),
        "job_id": job_id,
        "job_sha256": job_sha256,
        "expected_candidate_rows": int(expected_rows),
        "params": params,
    }


def expand_experiment_jobs(
    spec: dict,
    *,
    phase_ids: Iterable[str] | None = None,
) -> list[dict]:
    selected = {str(item) for item in (phase_ids or []) if str(item)}
    defaults = dict(spec.get("defaults") or {})
    jobs = []
    seen_phase_ids = set()

    for phase in spec.get("phases") or []:
        if not isinstance(phase, dict):
            raise NdxExperimentLabError("every phase must be a mapping")
        phase_id = str(phase.get("phase_id") or "").strip()
        if not phase_id or phase_id in seen_phase_ids:
            raise NdxExperimentLabError(f"invalid or duplicate phase_id: {phase_id!r}")
        seen_phase_ids.add(phase_id)
        if phase.get("enabled") is not True:
            continue
        if selected and phase_id not in selected:
            continue
        role = str(phase.get("role") or "").strip()
        if role not in ALLOWED_PHASE_ROLES:
            raise NdxExperimentLabError(
                f"{phase_id}.role must be one of {sorted(ALLOWED_PHASE_ROLES)}"
            )
        grid = _normalize_grid(
            {**(defaults.get("shared_grid") or {}), **(phase.get("shared_grid") or {})},
            prefix=phase_id,
        )
        base_values = {
            key: value
            for key, value in {**defaults, **phase}.items()
            if key not in {"shared_grid", "sweep", "locked_candidates"}
        }

        if role == "locked_validation":
            candidates = phase.get("locked_candidates") or []
            if not candidates:
                raise NdxExperimentLabError(
                    f"{phase_id}: enabled locked_validation requires explicit locked_candidates"
                )
            for index, candidate in enumerate(candidates, start=1):
                if not isinstance(candidate, dict):
                    raise NdxExperimentLabError(
                        f"{phase_id}.locked_candidates[{index}] must be a mapping"
                    )
                candidate_grid = dict(grid)
                for key in GRID_KEYS:
                    if key in candidate:
                        candidate_grid[key] = _as_list(
                            candidate[key], name=f"{phase_id}.{key}"
                        )
                    if len(candidate_grid[key]) != 1:
                        raise NdxExperimentLabError(
                            f"{phase_id}: locked candidate {key} must contain exactly one value"
                        )
                values = {**base_values, **candidate}
                params = _normalize_run_params(
                    values, candidate_grid, prefix=f"{phase_id}.locked_candidates[{index}]"
                )
                hint = str(candidate.get("id") or f"locked-{index}")
                jobs.append(_make_job(spec, phase, params, job_hint=hint))
            continue

        sweep = dict(phase.get("sweep") or {})
        unknown_sweep = sorted(set(sweep) - set(SWEEP_KEYS))
        if unknown_sweep:
            raise NdxExperimentLabError(
                f"{phase_id}.sweep contains unsupported keys: {unknown_sweep}"
            )
        sweep_keys = list(sweep)
        sweep_values = [
            _as_list(sweep[key], name=f"{phase_id}.sweep.{key}")
            for key in sweep_keys
        ]
        combinations = product(*sweep_values) if sweep_values else [()]
        for combination in combinations:
            varied = dict(zip(sweep_keys, combination))
            values = {**base_values, **varied}
            params = _normalize_run_params(values, grid, prefix=phase_id)
            hint = "__".join(
                f"{_slug(key)}-{_slug(value)}" for key, value in varied.items()
            ) or "base"
            jobs.append(_make_job(spec, phase, params, job_hint=hint))

    if selected:
        missing = selected - seen_phase_ids
        if missing:
            raise NdxExperimentLabError(f"unknown phase filters: {sorted(missing)}")
    fingerprints = [job["job_sha256"] for job in jobs]
    if len(fingerprints) != len(set(fingerprints)):
        raise NdxExperimentLabError("experiment expands to duplicate jobs")
    return jobs


def _construction_contract_error(job: dict, result: dict) -> str | None:
    expected_profiles = (job.get("params") or {}).get("construction_profiles")
    if expected_profiles is None:
        return None
    expected_by_id = {
        str(profile.get("id") or profile.get("profile_id")): profile
        for profile in expected_profiles
    }
    rows = result.get("rows") or []
    observed_by_id = {str(row.get("profile_id")): row for row in rows}
    if len(rows) != len(observed_by_id):
        return "backtest contains duplicate construction profile rows"
    if set(observed_by_id) != set(expected_by_id):
        return "backtest construction profile ids changed"
    keys = {
        "top_n",
        "weight_mode",
        "rebalance_months",
        "safe_weight",
        "max_weight",
        "min_weight",
        *CONSTRUCTION_EXPORT_KEYS[1:],
    }
    for profile_id, expected in expected_by_id.items():
        observed = observed_by_id[profile_id]
        for key in keys & set(expected):
            left = expected.get(key)
            right = observed.get(key)
            if key == "candidate_tickers":
                def normalized_tickers(value: Any) -> tuple[str, ...]:
                    if value is None:
                        return ()
                    items = (
                        value.replace(",", ";").split(";")
                        if isinstance(value, str)
                        else value
                    )
                    return tuple(sorted({
                        str(item or "").strip().upper()
                        for item in items
                        if str(item or "").strip()
                    }))

                matches = normalized_tickers(left) == normalized_tickers(right)
            elif isinstance(left, (int, float)) and not isinstance(left, bool):
                try:
                    matches = bool(np.isclose(float(left), float(right), atol=1e-12))
                except (TypeError, ValueError):
                    matches = False
            else:
                matches = left == right
            if not matches:
                return f"construction profile {profile_id}.{key} changed"
    sample_ids = {
        str(sample.get("profile_id")) for sample in result.get("samples") or []
    }
    if sample_ids and sample_ids != set(expected_by_id):
        return "rolling samples do not cover every construction profile"
    return None


def checkpoint_payload(
    job: dict,
    result: dict,
    *,
    started_at: str,
    duration_seconds: float,
) -> dict:
    summary = result.get("summary") or {}
    if summary.get("qqq_role") != QQQ_ROLE:
        raise NdxExperimentLabError("backtest result does not preserve QQQ benchmark-only role")
    expected_target = str((job.get("params") or {}).get("target") or "").upper()
    if str(summary.get("target_benchmark") or "").upper() != expected_target:
        raise NdxExperimentLabError("backtest target does not match the pre-registered job")
    objective_contract = job.get("objective_contract") or {}
    if (
        objective_contract
        and summary.get("qqq_performance_role")
        != objective_contract.get("qqq_performance_role")
    ):
        raise NdxExperimentLabError("backtest QQQ performance role changed")
    profile_error = _construction_contract_error(job, result)
    if profile_error:
        raise NdxExperimentLabError(profile_error)
    return {
        "lab_job": {
            **job,
            "started_at": started_at,
            "completed_at": utc_now_iso(),
            "duration_seconds": round(float(duration_seconds), 3),
            "actual_candidate_rows": len(result.get("rows") or []),
            "actual_window_rows": len(result.get("samples") or []),
        },
        "result": result,
    }


def validate_checkpoint(payload: dict, job: dict) -> tuple[bool, str]:
    if not isinstance(payload, dict):
        return False, "checkpoint is not a mapping"
    meta = payload.get("lab_job") or {}
    result = payload.get("result") or {}
    if meta.get("job_sha256") != job.get("job_sha256"):
        return False, "job fingerprint changed"
    if meta.get("spec_sha256") != job.get("spec_sha256"):
        return False, "experiment spec changed"
    if meta.get("engine_sha256") != job.get("engine_sha256"):
        return False, "backtest engine source changed"
    if meta.get("runtime_sha256") != job.get("runtime_sha256"):
        return False, "Python or numerical-library runtime changed"
    if meta.get("strategy_config_sha256") != job.get("strategy_config_sha256"):
        return False, "strategy config changed"
    summary = result.get("summary") or {}
    if summary.get("qqq_role") != QQQ_ROLE:
        return False, "QQQ role is not benchmark-only"
    expected_target = str((job.get("params") or {}).get("target") or "").upper()
    if str(summary.get("target_benchmark") or "").upper() != expected_target:
        return False, "backtest target changed"
    objective_contract = job.get("objective_contract") or {}
    if (
        objective_contract
        and summary.get("qqq_performance_role")
        != objective_contract.get("qqq_performance_role")
    ):
        return False, "QQQ performance role changed"
    profile_error = _construction_contract_error(job, result)
    if profile_error:
        return False, profile_error
    manifest = result.get("data_manifest") or {}
    if not ((manifest.get("ndx_history") or {}).get("version_id")):
        return False, "NDX history version is missing"
    if not ((manifest.get("price_universe") or {}).get("version_id")):
        return False, "price universe version is missing"
    expected_profiles = (job.get("params") or {}).get("construction_profiles") or []
    if any(_profile_requires_quality_evidence(profile) for profile in expected_profiles):
        quality_version = manifest.get("quality_scores") or {}
        if not quality_version.get("version_id"):
            return False, "quality score version is missing"
        if quality_version.get("quality_policy_id") != QUALITY_POLICY_ID:
            return False, "quality score policy changed"
    if any(_profile_requires_valuation_evidence(profile) for profile in expected_profiles):
        valuation_version = manifest.get("valuation_scores") or {}
        if not valuation_version.get("version_id"):
            return False, "valuation score version is missing"
        expected_policies = {
            str(profile.get("valuation_policy_id") or "")
            for profile in expected_profiles
            if _profile_requires_valuation_evidence(profile)
        }
        if (
            len(expected_policies) != 1
            or "" in expected_policies
            or not expected_policies.issubset(SUPPORTED_VALUATION_SCORE_POLICY_IDS)
            or valuation_version.get("valuation_policy_id")
            != next(iter(expected_policies))
        ):
            return False, "valuation score policy changed"
    audit = selection_constraint_audit(result)
    if audit["available"] and not audit["passed"]:
        return False, "selection constraint audit failed"
    return True, "ok"


def write_checkpoint(path: str | Path, payload: dict) -> None:
    atomic_write_text(
        Path(path),
        json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
    )


def read_checkpoint(path: str | Path) -> dict:
    with Path(path).open("r", encoding="utf-8") as handle:
        return json.load(handle)


def candidate_diagnosis(row: dict) -> dict:
    target = row.get("target_envelope") or {}
    incumbent = row.get("incumbent_guard") or {}
    target_checks = [
        f"target:{item.get('name')}"
        for item in target.get("checks") or []
        if item.get("passed") is False
    ]
    incumbent_checks = [
        f"incumbent:{item.get('name')}"
        for item in incumbent.get("checks") or []
        if item.get("passed") is False
    ]
    if target.get("passed") is False and not target_checks:
        target_checks.append("target:gate")
    if incumbent.get("passed") is False and not incumbent_checks:
        incumbent_checks.append("incumbent:gate")
    failed_checks = target_checks + incumbent_checks
    admission = bool(row.get("admission_passed"))
    qqq_near = bool(
        row.get("qqq_win_rate") is not None
        and float(row["qqq_win_rate"]) >= 0.50
        and row.get("qqq_excess_p50") is not None
        and float(row["qqq_excess_p50"]) >= 0.0
    )
    if admission:
        diagnosis = "PASSED_ALL_GATES"
        next_action = "PROMOTION_REVIEW"
    elif target.get("classification") == "INSUFFICIENT_EVIDENCE":
        diagnosis = "INSUFFICIENT_EVIDENCE"
        next_action = "EXPAND_OR_REPAIR_EVIDENCE"
    elif qqq_near:
        diagnosis = "NEAR_QQQ_BUT_GATE_FAILED"
        next_action = "ITERATE_ALPHA_OR_RISK_CONTROLS"
    elif target.get("return_passed") and (
        not target.get("risk_passed") or not incumbent.get("passed")
    ):
        diagnosis = "RETURN_PASSED_RISK_FAILED"
        next_action = "ITERATE_RISK_CONTROLS"
    else:
        diagnosis = "CURRENT_SHAPE_RETURN_FAILED"
        next_action = "ITERATE_ALPHA_FEATURES"
    return {
        "shape_diagnosis": diagnosis,
        "next_action": next_action,
        "direction_status": "PROMOTION_REVIEW" if admission else "ACTIVE_ITERATION",
        "terminal_direction_decision": False,
        "failed_checks": failed_checks,
    }


def _profile_requires_quality_evidence(profile: dict) -> bool:
    return bool(
        float(profile.get("fundamental_score_weight") or 0.0) > 0
        or profile.get("quality_policy_id") is not None
        or str(profile.get("quality_selection_mode") or "none") != "none"
        or profile.get("max_names_per_sic4") is not None
        or profile.get("max_names_per_sic2") is not None
        or profile.get("sic4_max_weight") is not None
        or profile.get("sic2_max_weight") is not None
    )


def _profile_requires_valuation_evidence(profile: dict) -> bool:
    return bool(float(profile.get("valuation_score_weight") or 0.0) > 0)


def _valuation_selection_audit(
    sample: dict,
    selection: dict,
    *,
    tolerance: float,
) -> dict:
    required = _profile_requires_valuation_evidence(sample)
    defaults = {
        "valuation_evidence_passed": True,
        "valuation_score_coverage_passed": True,
        "valuation_future_data_guard_audit_passed": True,
    }
    if not required:
        return defaults
    expected_weight = float(sample.get("valuation_score_weight") or 0.0)
    try:
        observed_weight = float(selection.get("valuation_score_weight"))
        score_coverage = float(selection.get("valuation_score_coverage"))
        observed_coverage = float(selection.get("valuation_observed_coverage"))
    except (TypeError, ValueError):
        observed_weight = math.nan
        score_coverage = math.nan
        observed_coverage = math.nan
    expected_policy = str(sample.get("valuation_policy_id") or "")
    policy_passed = bool(
        expected_policy in SUPPORTED_VALUATION_SCORE_POLICY_IDS
        and selection.get("valuation_policy_id") == expected_policy
        and np.isfinite(observed_weight)
        and np.isclose(observed_weight, expected_weight, atol=tolerance)
    )
    coverage_passed = bool(
        np.isfinite(score_coverage)
        and score_coverage >= 0.85 - tolerance
        and np.isfinite(observed_coverage)
        and observed_coverage >= 0.70 - tolerance
        and int(selection.get("valuation_missing_neutral_count") or 0) >= 0
    )
    selection_date = str(selection.get("date") or "")
    score_date = str(selection.get("valuation_score_date") or "")
    future_passed = bool(
        selection.get("valuation_future_data_guard_passed") is True
        and selection_date
        and score_date == selection_date
    )
    return {
        "valuation_evidence_passed": policy_passed,
        "valuation_score_coverage_passed": coverage_passed,
        "valuation_future_data_guard_audit_passed": future_passed,
    }


def _quality_sic_selection_audit(
    sample: dict,
    selection: dict,
    weights: dict[str, float],
    *,
    tolerance: float,
) -> dict:
    required = _profile_requires_quality_evidence(sample)
    defaults = {
        "quality_evidence_passed": True,
        "quality_score_coverage_passed": True,
        "quality_future_data_guard_passed": True,
        "quality_replacement_guard_passed": True,
        "observed_quality_replacements": None,
        "sic4_name_cap_passed": True,
        "sic2_name_cap_passed": True,
        "sic4_weight_cap_passed": True,
        "sic2_weight_cap_passed": True,
        "observed_max_names_per_sic4": None,
        "observed_max_names_per_sic2": None,
        "observed_max_sic4_weight": None,
        "observed_max_sic2_weight": None,
    }
    if not required:
        return defaults

    expected_weight = float(sample.get("fundamental_score_weight") or 0.0)
    expected_quality_mode = str(
        sample.get("quality_selection_mode") or "none"
    )
    observed_quality_mode = str(
        selection.get("quality_selection_mode") or "none"
    )
    try:
        observed_weight = float(selection.get("fundamental_score_weight"))
    except (TypeError, ValueError):
        observed_weight = math.nan
    policy_passed = bool(
        sample.get("quality_policy_id") == QUALITY_POLICY_ID
        and selection.get("quality_policy_id") == QUALITY_POLICY_ID
        and np.isfinite(observed_weight)
        and np.isclose(observed_weight, expected_weight, atol=tolerance)
        and observed_quality_mode == expected_quality_mode
    )
    selection_date = str(selection.get("date") or "")
    quality_date = str(selection.get("quality_score_date") or "")
    future_passed = bool(
        selection.get("quality_future_data_guard_passed") is True
        and selection_date
        and quality_date == selection_date
    )
    try:
        quality_coverage = float(selection.get("quality_score_coverage"))
    except (TypeError, ValueError):
        quality_coverage = math.nan
    coverage_passed = bool(
        (expected_weight <= 0 and expected_quality_mode == "none")
        or (np.isfinite(quality_coverage) and quality_coverage >= 0.85 - tolerance)
    )

    replacement_guard_passed = True
    observed_quality_replacements = None
    if expected_quality_mode == "bounded_challenge":
        top_n = int(sample.get("top_n") or 0)
        expected_pool_size = int(
            sample.get("quality_challenge_pool_size") or 0
        )
        expected_budget = int(sample.get("max_quality_replacements") or 0)
        pool = selection.get("quality_challenge_pool_tickers") or []
        pairs = selection.get("quality_replacement_pairs") or []
        observed_quality_replacements = len(pairs) if isinstance(pairs, list) else None
        replacement_guard_passed = bool(
            isinstance(pool, list)
            and isinstance(pairs, list)
            and len(pool) == expected_pool_size
            and len(pool) == len(set(pool))
            and int(selection.get("quality_challenge_pool_size") or 0)
            == expected_pool_size
            and int(selection.get("max_quality_replacements") or 0)
            == expected_budget
            and len(pairs) <= expected_budget
            and top_n > 0
            and expected_pool_size > top_n
        )
        outgoing: list[str] = []
        incoming: list[str] = []
        if replacement_guard_passed:
            rank_by_ticker = {
                str(ticker): rank
                for rank, ticker in enumerate(pool, start=1)
            }
            baseline = set(map(str, pool[:top_n]))
            challengers = set(map(str, pool[top_n:]))
            for pair in pairs:
                if not isinstance(pair, dict):
                    replacement_guard_passed = False
                    break
                out = str(pair.get("outgoing_ticker") or "")
                inc = str(pair.get("incoming_ticker") or "")
                try:
                    out_rank = int(pair.get("outgoing_price_rank"))
                    inc_rank = int(pair.get("incoming_price_rank"))
                    out_score = float(pair.get("outgoing_quality_score"))
                    inc_score = float(pair.get("incoming_quality_score"))
                    score_delta = float(pair.get("quality_score_delta"))
                except (TypeError, ValueError):
                    replacement_guard_passed = False
                    break
                pair_passed = bool(
                    out in baseline
                    and inc in challengers
                    and rank_by_ticker.get(out) == out_rank
                    and rank_by_ticker.get(inc) == inc_rank
                    and inc_score > out_score
                    and np.isclose(
                        inc_score - out_score, score_delta, atol=tolerance
                    )
                    and score_delta > 0
                )
                if not pair_passed:
                    replacement_guard_passed = False
                    break
                outgoing.append(out)
                incoming.append(inc)
            final_tickers = set(map(str, weights))
            replacement_guard_passed = bool(
                replacement_guard_passed
                and len(outgoing) == len(set(outgoing))
                and len(incoming) == len(set(incoming))
                and final_tickers
                == (baseline - set(outgoing)) | set(incoming)
            )

    raw_sic4 = selection.get("sic4_by_ticker") or {}
    raw_sic2 = selection.get("sic2_by_ticker") or {}
    try:
        sic4 = {str(ticker): int(value) for ticker, value in raw_sic4.items()}
        sic2 = {str(ticker): str(value) for ticker, value in raw_sic2.items()}
    except (AttributeError, TypeError, ValueError):
        sic4 = {}
        sic2 = {}
    mapping_passed = set(sic4) == set(weights) and set(sic2) == set(weights)

    sic4_names: dict[int, int] = {}
    sic2_names: dict[str, int] = {}
    sic4_weights: dict[int, float] = {}
    sic2_weights: dict[str, float] = {}
    if mapping_passed:
        for ticker, weight in weights.items():
            group4 = sic4[ticker]
            group2 = sic2[ticker]
            sic4_names[group4] = sic4_names.get(group4, 0) + 1
            sic2_names[group2] = sic2_names.get(group2, 0) + 1
            sic4_weights[group4] = sic4_weights.get(group4, 0.0) + float(weight)
            sic2_weights[group2] = sic2_weights.get(group2, 0.0) + float(weight)

    def reported_weights_match(raw: object, observed: dict) -> bool:
        if not isinstance(raw, dict) or set(map(str, raw)) != set(map(str, observed)):
            return False
        try:
            normalized = {str(key): float(value) for key, value in raw.items()}
        except (TypeError, ValueError):
            return False
        return all(
            np.isclose(normalized[str(key)], float(value), atol=tolerance)
            for key, value in observed.items()
        )

    reported_passed = bool(
        mapping_passed
        and reported_weights_match(selection.get("sic4_weights"), sic4_weights)
        and reported_weights_match(selection.get("sic2_weights"), sic2_weights)
    )
    evidence_passed = bool(policy_passed and mapping_passed and reported_passed)

    max_sic4_names = max(sic4_names.values()) if sic4_names else None
    max_sic2_names = max(sic2_names.values()) if sic2_names else None
    max_sic4_weight = max(sic4_weights.values()) if sic4_weights else None
    max_sic2_weight = max(sic2_weights.values()) if sic2_weights else None
    sic4_name_cap = sample.get("max_names_per_sic4")
    sic2_name_cap = sample.get("max_names_per_sic2")
    sic4_weight_cap = sample.get("sic4_max_weight")
    sic2_weight_cap = sample.get("sic2_max_weight")
    return {
        "quality_evidence_passed": evidence_passed,
        "quality_score_coverage_passed": coverage_passed,
        "quality_future_data_guard_passed": future_passed,
        "quality_replacement_guard_passed": replacement_guard_passed,
        "observed_quality_replacements": observed_quality_replacements,
        "sic4_name_cap_passed": bool(
            sic4_name_cap is None
            or (
                max_sic4_names is not None
                and max_sic4_names <= int(sic4_name_cap)
            )
        ),
        "sic2_name_cap_passed": bool(
            sic2_name_cap is None
            or (
                max_sic2_names is not None
                and max_sic2_names <= int(sic2_name_cap)
            )
        ),
        "sic4_weight_cap_passed": bool(
            sic4_weight_cap is None
            or (
                max_sic4_weight is not None
                and max_sic4_weight <= float(sic4_weight_cap) + tolerance
            )
        ),
        "sic2_weight_cap_passed": bool(
            sic2_weight_cap is None
            or (
                max_sic2_weight is not None
                and max_sic2_weight <= float(sic2_weight_cap) + tolerance
            )
        ),
        "observed_max_names_per_sic4": max_sic4_names,
        "observed_max_names_per_sic2": max_sic2_names,
        "observed_max_sic4_weight": max_sic4_weight,
        "observed_max_sic2_weight": max_sic2_weight,
    }


def selection_constraint_audit(result: dict, *, tolerance: float = 1e-8) -> dict:
    counts = {
        "selection_rows": 0,
        "parse_failures": 0,
        "weight_sum_failures": 0,
        "cap_failures": 0,
        "floor_failures": 0,
        "ticker_count_failures": 0,
        "future_data_guard_failures": 0,
        "correlation_evidence_failures": 0,
        "correlation_future_data_guard_failures": 0,
        "correlation_pair_cap_failures": 0,
        "cluster_name_cap_failures": 0,
        "cluster_weight_cap_failures": 0,
        "top_weight_cap_failures": 0,
        "effective_n_floor_failures": 0,
        "defense_signal_guard_failures": 0,
        "quality_evidence_failures": 0,
        "quality_score_coverage_failures": 0,
        "quality_future_data_guard_failures": 0,
        "quality_replacement_guard_failures": 0,
        "valuation_evidence_failures": 0,
        "valuation_score_coverage_failures": 0,
        "valuation_future_data_guard_failures": 0,
        "sic4_name_cap_failures": 0,
        "sic2_name_cap_failures": 0,
        "sic4_weight_cap_failures": 0,
        "sic2_weight_cap_failures": 0,
        "candidate_pool_failures": 0,
    }
    max_cap_excess = 0.0
    max_sum_error = 0.0
    for sample in result.get("samples") or []:
        top_n = int(sample.get("top_n") or 0)
        if top_n <= 0:
            continue
        max_weight = float(sample.get("max_weight") or 0.0)
        min_weight = float(sample.get("min_weight") or 0.02)
        effective_cap = max(max_weight, 1.0 / top_n)
        effective_floor = min(max(0.0, min_weight), 1.0 / top_n)
        selection_mode = str(sample.get("selection_mode") or "rank")
        cluster_cap = sample.get("cluster_max_weight")
        top_weight_count = sample.get("top_weight_count")
        top_weight_cap = sample.get("top_weight_cap")
        min_effective_n = sample.get("min_effective_n")
        correlation_required = selection_mode != "rank" or cluster_cap is not None
        if sample.get("defense_mode") in {"vol_target", "drawdown_guard"}:
            defense_signal = str(sample.get("defense_signal") or "")
            defense_passed = sample.get("defense_signal_strict_asof_passed") is True
            if defense_signal in {
                "selected_basket_ex_ante_vol",
                "selected_basket_ex_ante_drawdown",
            }:
                defense_passed = bool(
                    defense_passed
                    and sample.get("defense_signal_prehistory_seeded") is True
                    and sample.get("defense_signal_first_date")
                    and sample.get("start")
                    and str(sample["defense_signal_first_date"]) < str(sample["start"])
                )
            if not defense_passed:
                counts["defense_signal_guard_failures"] += 1
        for selection in sample.get("selections") or []:
            counts["selection_rows"] += 1
            raw_weights = selection.get("weights") or {}
            try:
                weights = [float(value) for value in raw_weights.values()]
            except (AttributeError, TypeError, ValueError):
                counts["parse_failures"] += 1
                continue
            if not weights or not np.isfinite(np.asarray(weights, dtype=float)).all():
                counts["parse_failures"] += 1
                continue
            sum_error = abs(sum(weights) - 1.0)
            max_sum_error = max(max_sum_error, sum_error)
            if sum_error > tolerance:
                counts["weight_sum_failures"] += 1
            cap_excess = max(weights) - effective_cap
            max_cap_excess = max(max_cap_excess, cap_excess)
            if cap_excess > tolerance:
                counts["cap_failures"] += 1
            if min(weights) < effective_floor - tolerance:
                counts["floor_failures"] += 1
            if len(weights) != top_n:
                counts["ticker_count_failures"] += 1
            if (top_weight_count is None) != (top_weight_cap is None):
                counts["parse_failures"] += 1
            elif top_weight_count is not None:
                try:
                    parsed_top_count = int(top_weight_count)
                    parsed_top_cap = float(top_weight_cap)
                    observed_top_weight = sum(
                        sorted(weights, reverse=True)[:parsed_top_count]
                    )
                    if (
                        not 1 <= parsed_top_count < top_n
                        or not np.isfinite(parsed_top_cap)
                        or observed_top_weight > parsed_top_cap + tolerance
                    ):
                        counts["top_weight_cap_failures"] += 1
                except (TypeError, ValueError):
                    counts["parse_failures"] += 1
            if min_effective_n is not None:
                try:
                    parsed_effective_floor = float(min_effective_n)
                    observed_effective_n = 1.0 / sum(
                        weight * weight for weight in weights
                    )
                    if (
                        not np.isfinite(parsed_effective_floor)
                        or observed_effective_n
                        < parsed_effective_floor - tolerance
                    ):
                        counts["effective_n_floor_failures"] += 1
                except (TypeError, ValueError, ZeroDivisionError):
                    counts["parse_failures"] += 1
            raw_candidate_tickers = sample.get("candidate_tickers")
            if isinstance(raw_candidate_tickers, str):
                candidate_tickers = {
                    item.strip().upper()
                    for item in raw_candidate_tickers.replace(",", ";").split(";")
                    if item.strip()
                }
            else:
                candidate_tickers = {
                    str(item).strip().upper()
                    for item in (raw_candidate_tickers or [])
                    if str(item).strip()
                }
            if candidate_tickers and not {
                str(ticker).upper() for ticker in raw_weights
            }.issubset(candidate_tickers):
                counts["candidate_pool_failures"] += 1
            if not (
                selection.get("available_date")
                and selection.get("date")
                and str(selection["available_date"]) <= str(selection["date"])
            ):
                counts["future_data_guard_failures"] += 1
            quality_sic = _quality_sic_selection_audit(
                sample,
                selection,
                {str(ticker): float(value) for ticker, value in raw_weights.items()},
                tolerance=tolerance,
            )
            for field, counter in (
                ("quality_evidence_passed", "quality_evidence_failures"),
                ("quality_score_coverage_passed", "quality_score_coverage_failures"),
                (
                    "quality_future_data_guard_passed",
                    "quality_future_data_guard_failures",
                ),
                (
                    "quality_replacement_guard_passed",
                    "quality_replacement_guard_failures",
                ),
                ("sic4_name_cap_passed", "sic4_name_cap_failures"),
                ("sic2_name_cap_passed", "sic2_name_cap_failures"),
                ("sic4_weight_cap_passed", "sic4_weight_cap_failures"),
                ("sic2_weight_cap_passed", "sic2_weight_cap_failures"),
            ):
                if not quality_sic[field]:
                    counts[counter] += 1
            valuation = _valuation_selection_audit(
                sample,
                selection,
                tolerance=tolerance,
            )
            for field, counter in (
                ("valuation_evidence_passed", "valuation_evidence_failures"),
                (
                    "valuation_score_coverage_passed",
                    "valuation_score_coverage_failures",
                ),
                (
                    "valuation_future_data_guard_audit_passed",
                    "valuation_future_data_guard_failures",
                ),
            ):
                if not valuation[field]:
                    counts[counter] += 1
            if not correlation_required:
                continue
            labels = selection.get("cluster_labels") or {}
            cluster_weights = selection.get("cluster_weights") or {}
            if (
                str(selection.get("selection_mode") or "") != selection_mode
                or int(selection.get("correlation_observations") or 0) < 60
                or set(labels) != set(raw_weights)
                or not cluster_weights
            ):
                counts["correlation_evidence_failures"] += 1
                continue
            if not (
                selection.get("correlation_strict_asof_passed") is True
                and selection.get("correlation_price_end")
                and selection.get("date")
                and str(selection["correlation_price_end"]) < str(selection["date"])
            ):
                counts["correlation_future_data_guard_failures"] += 1
            threshold = float(sample.get("correlation_threshold") or 0.0)
            if selection_mode == "correlation_replacement":
                max_pair = selection.get("max_pair_correlation")
                if (
                    max_pair is None
                    or not np.isfinite(float(max_pair))
                    or float(max_pair) >= threshold - tolerance
                ):
                    counts["correlation_pair_cap_failures"] += 1
            if selection_mode == "cluster_round_robin":
                name_counts: dict[int, int] = {}
                for cluster_id in labels.values():
                    parsed = int(cluster_id)
                    name_counts[parsed] = name_counts.get(parsed, 0) + 1
                if max(name_counts.values()) > int(
                    sample.get("max_names_per_cluster") or 0
                ):
                    counts["cluster_name_cap_failures"] += 1
            if cluster_cap is not None and max(
                float(value) for value in cluster_weights.values()
            ) > float(cluster_cap) + tolerance:
                counts["cluster_weight_cap_failures"] += 1
    failure_total = sum(
        value for key, value in counts.items()
        if key not in {"selection_rows"}
    )
    return {
        "available": counts["selection_rows"] > 0,
        "passed": counts["selection_rows"] > 0 and failure_total == 0,
        **counts,
        "max_cap_excess": max(0.0, float(max_cap_excess)),
        "max_weight_sum_error": float(max_sum_error),
    }


def _nested(value: dict, *keys: str) -> Any:
    current: Any = value
    for key in keys:
        if not isinstance(current, dict):
            return None
        current = current.get(key)
    return current


def _version_bundle(manifest: dict) -> tuple[str, list[str]]:
    ids = []
    for value in manifest.values():
        if isinstance(value, dict):
            version_id = value.get("version_id")
            if version_id:
                ids.append(str(version_id))
    ids = sorted(set(ids))
    return f"ndx_lab_inputs:{_json_hash(ids)[:16]}", ids


def _shape_id(meta: dict, row: dict) -> str:
    identity = {
        "job": meta.get("job_sha256"),
        "lookback_years": row.get("lookback_years"),
        "top_n": row.get("top_n"),
        "weight_mode": row.get("weight_mode"),
        "rebalance_months": row.get("rebalance_months"),
        "safe_weight": row.get("safe_weight"),
        "max_weight": row.get("max_weight"),
        "min_weight": row.get("min_weight"),
        "trade_cost_bps": row.get("trade_cost_bps"),
        **{key: row.get(key) for key in CONSTRUCTION_EXPORT_KEYS},
    }
    return f"ndx_shape:{_json_hash(identity)[:16]}"


def _candidate_row(checkpoint: dict, row: dict) -> dict:
    meta = checkpoint["lab_job"]
    result = checkpoint["result"]
    summary = result.get("summary") or {}
    target = row.get("target_envelope") or {}
    incumbent = row.get("incumbent_guard") or {}
    manifest = result.get("data_manifest") or {}
    version_bundle_id, version_ids = _version_bundle(manifest)
    diagnosis = candidate_diagnosis(row)
    distribution = row.get("distribution") or {}
    margins = incumbent.get("margins") or {}
    candidate = {
        "experiment_id": meta.get("experiment_id"),
        "phase_id": meta.get("phase_id"),
        "phase_role": meta.get("phase_role"),
        "job_id": meta.get("job_id"),
        "job_sha256": meta.get("job_sha256"),
        "shape_id": _shape_id(meta, row),
        "candidate_assets": summary.get("candidate_assets"),
        "qqq_role": summary.get("qqq_role"),
        "target_benchmark": summary.get("target_benchmark"),
        "research_start": (meta.get("params") or {}).get("start"),
        "research_end": (meta.get("params") or {}).get("end"),
        "basket": row.get("basket"),
        **{key: row.get(key) for key in CONSTRUCTION_EXPORT_KEYS},
        "lookback_years": row.get("lookback_years"),
        "top_n": row.get("top_n"),
        "weight_mode": row.get("weight_mode"),
        "rebalance_months": row.get("rebalance_months"),
        "safe_weight": row.get("safe_weight"),
        "max_weight": row.get("max_weight"),
        "min_weight": row.get("min_weight"),
        "trade_cost_bps": row.get("trade_cost_bps"),
        "windows": row.get("windows"),
        "window_start": (row.get("range") or [None, None])[0],
        "window_end": (row.get("range") or [None, None])[1],
        "target_win_rate": row.get("win_rate"),
        "target_dd_win_rate": row.get("dd_win_rate"),
        "target_excess_p50": row.get("excess_p50"),
        "target_excess_p10": row.get("excess_p10"),
        "target_dd_relief_p50": row.get("dd_relief_p50"),
        "target_dd_relief_p10": row.get("dd_relief_p10"),
        "qqq_win_rate": row.get("qqq_win_rate"),
        "qqq_excess_p50": row.get("qqq_excess_p50"),
        "qqq_excess_p10": row.get("qqq_excess_p10"),
        "qqq_dd_relief_p50": row.get("qqq_dd_relief_p50"),
        "qqq_dd_relief_p10": row.get("qqq_dd_relief_p10"),
        "incumbent_win_rate": row.get("incumbent_win_rate"),
        "incumbent_dd_win_rate": row.get("incumbent_dd_win_rate"),
        "incumbent_excess_p50": row.get("incumbent_excess_p50"),
        "incumbent_excess_p10": row.get("incumbent_excess_p10"),
        "incumbent_dd_relief_direct_p50": row.get("incumbent_dd_relief_p50"),
        "incumbent_dd_relief_direct_p10": row.get("incumbent_dd_relief_p10"),
        "strategy_xirr_p10": _nested(distribution, "strategy_xirr", "p10"),
        "strategy_xirr_p50": _nested(distribution, "strategy_xirr", "p50"),
        "strategy_xirr_p90": _nested(distribution, "strategy_xirr", "p90"),
        "strategy_maxdd_p50": _nested(distribution, "strategy_maxdd", "p50"),
        "strategy_maxdd_p90": _nested(distribution, "strategy_maxdd", "p90"),
        "turnover_x_p50": row.get("turnover_x_p50"),
        "cost_x_p50": row.get("cost_x_p50"),
        "selection_churn_p50": row.get("selection_churn_p50"),
        "avg_safe_weight_p10": row.get("avg_safe_weight_p10"),
        "avg_safe_weight_p50": row.get("avg_safe_weight_p50"),
        "avg_safe_weight_p90": row.get("avg_safe_weight_p90"),
        "max_pair_correlation_p90": row.get("max_pair_correlation_p90"),
        "max_cluster_weight_p90": row.get("max_cluster_weight_p90"),
        "max_top_weight_p90": row.get("max_top_weight_p90"),
        "observed_min_effective_n": row.get("observed_min_effective_n"),
        "correlation_rejected_p50": row.get("correlation_rejected_p50"),
        "min_quality_score_coverage": row.get("min_quality_score_coverage"),
        "quality_replacements_p50": row.get("quality_replacements_p50"),
        "quality_replacements_p90": row.get("quality_replacements_p90"),
        "max_sic4_weight_p90": row.get("max_sic4_weight_p90"),
        "max_sic2_weight_p90": row.get("max_sic2_weight_p90"),
        "sic_rejected_p50": row.get("sic_rejected_p50"),
        "defense_signal_prehistory_seeded": row.get(
            "defense_signal_prehistory_seeded"
        ),
        "defense_signal_strict_asof_passed": row.get(
            "defense_signal_strict_asof_passed"
        ),
        "min_price_coverage": row.get("min_price_coverage"),
        "min_score_coverage": row.get("min_score_coverage"),
        "target_classification": target.get("classification"),
        "target_return_passed": target.get("return_passed"),
        "target_risk_passed": target.get("risk_passed"),
        "target_gate_passed": target.get("passed"),
        "incumbent_gate_passed": incumbent.get("passed"),
        "incumbent_xirr_p50_margin": margins.get("xirr_p50_vs_incumbent"),
        "incumbent_xirr_p10_margin": margins.get("xirr_p10_vs_incumbent"),
        "incumbent_dd_relief_p50": margins.get("dd_relief_p50_vs_incumbent"),
        "incumbent_dd_relief_p90": margins.get("dd_relief_p90_vs_incumbent"),
        "admission_passed": bool(row.get("admission_passed")),
        "shape_diagnosis": diagnosis["shape_diagnosis"],
        "next_action": diagnosis["next_action"],
        "direction_status": diagnosis["direction_status"],
        "terminal_direction_decision": diagnosis["terminal_direction_decision"],
        "failed_checks": ";".join(diagnosis["failed_checks"]),
        "tickers_observed": ";".join(row.get("tickers_observed") or []),
        "input_version_bundle_id": version_bundle_id,
        "input_version_ids": ";".join(version_ids),
        "history_version_id": _nested(manifest, "ndx_history", "version_id"),
        "price_version_id": _nested(manifest, "price_universe", "version_id"),
        "target_version_id": _nested(manifest, "target_nav", "version_id"),
        "quality_score_version_id": _nested(
            manifest, "quality_scores", "version_id"
        ),
        "valuation_score_version_id": _nested(
            manifest, "valuation_scores", "version_id"
        ),
        "future_data_guard": manifest.get("future_data_guard"),
    }
    return candidate


def _window_rows(checkpoint: dict) -> list[dict]:
    meta = checkpoint["lab_job"]
    rows = []
    for sample in (checkpoint["result"].get("samples") or []):
        strategy = sample.get("strategy") or {}
        target = sample.get("benchmark") or {}
        qqq = sample.get("qqq") or {}
        incumbent = sample.get("incumbent") or {}
        rows.append({
            "experiment_id": meta.get("experiment_id"),
            "phase_id": meta.get("phase_id"),
            "job_id": meta.get("job_id"),
            "job_sha256": meta.get("job_sha256"),
            "start": sample.get("start"),
            "end": sample.get("end"),
            "basket": sample.get("basket"),
            **{key: sample.get(key) for key in CONSTRUCTION_EXPORT_KEYS},
            "lookback_years": sample.get("lookback_years"),
            "top_n": sample.get("top_n"),
            "weight_mode": sample.get("weight_mode"),
            "rebalance_months": sample.get("rebalance_months"),
            "safe_weight": sample.get("safe_weight"),
            "max_weight": sample.get("max_weight"),
            "min_weight": sample.get("min_weight"),
            "trade_cost_bps": sample.get("trade_cost_bps"),
            "strategy_xirr": strategy.get("xirr"),
            "strategy_maxdd": strategy.get("maxdd"),
            "strategy_final": strategy.get("final"),
            "target_xirr": target.get("xirr"),
            "target_maxdd": target.get("maxdd"),
            "qqq_xirr": qqq.get("xirr"),
            "qqq_maxdd": qqq.get("maxdd"),
            "incumbent_xirr": incumbent.get("xirr"),
            "incumbent_maxdd": incumbent.get("maxdd"),
            "target_excess_xirr": sample.get("excess_xirr"),
            "target_dd_relief": sample.get("dd_relief"),
            "qqq_excess_xirr": sample.get("qqq_excess_xirr"),
            "qqq_dd_relief": sample.get("qqq_dd_relief"),
            "incumbent_xirr_margin": (
                strategy.get("xirr") - incumbent.get("xirr")
                if strategy.get("xirr") is not None and incumbent.get("xirr") is not None
                else None
            ),
            "incumbent_dd_relief": (
                incumbent.get("maxdd") - strategy.get("maxdd")
                if incumbent.get("maxdd") is not None and strategy.get("maxdd") is not None
                else None
            ),
            "turnover_x": sample.get("turnover_x"),
            "cost_x": sample.get("cost_x"),
            "selection_churn_p50": sample.get("selection_churn_p50"),
            "avg_safe_weight": sample.get("avg_safe_weight"),
            "min_safe_weight": sample.get("min_safe_weight"),
            "max_safe_weight": sample.get("max_safe_weight"),
            "safe_turnover_x": sample.get("safe_turnover_x"),
            "defense_signal_prehistory_seeded": sample.get(
                "defense_signal_prehistory_seeded"
            ),
            "defense_signal_strict_asof_passed": sample.get(
                "defense_signal_strict_asof_passed"
            ),
            "defense_signal_first_date": sample.get("defense_signal_first_date"),
            "defense_signal_last_date": sample.get("defense_signal_last_date"),
            "max_pair_correlation_p90": sample.get("max_pair_correlation_p90"),
            "max_cluster_weight_p90": sample.get("max_cluster_weight_p90"),
            "max_top_weight_p90": sample.get("max_top_weight_p90"),
            "observed_min_effective_n": sample.get(
                "observed_min_effective_n"
            ),
            "correlation_rejected_p50": sample.get("correlation_rejected_p50"),
            "min_quality_score_coverage": sample.get(
                "min_quality_score_coverage"
            ),
            "quality_replacements_p50": sample.get(
                "quality_replacements_p50"
            ),
            "quality_replacements_p90": sample.get(
                "quality_replacements_p90"
            ),
            "max_sic4_weight_p90": sample.get("max_sic4_weight_p90"),
            "max_sic2_weight_p90": sample.get("max_sic2_weight_p90"),
            "sic_rejected_p50": sample.get("sic_rejected_p50"),
            "min_price_coverage": sample.get("min_price_coverage"),
            "min_score_coverage": sample.get("min_score_coverage"),
            "selection_count": len(sample.get("selections") or []),
            "tickers_observed": ";".join(sample.get("tickers_observed") or []),
        })
    return rows


def _selection_rows(checkpoint: dict) -> list[dict]:
    meta = checkpoint["lab_job"]
    rows = []
    for sample in (checkpoint["result"].get("samples") or []):
        shape = {
            "experiment_id": meta.get("experiment_id"),
            "phase_id": meta.get("phase_id"),
            "job_id": meta.get("job_id"),
            "window_start": sample.get("start"),
            "basket": sample.get("basket"),
            **{key: sample.get(key) for key in CONSTRUCTION_EXPORT_KEYS},
            "lookback_years": sample.get("lookback_years"),
            "top_n": sample.get("top_n"),
            "weight_mode": sample.get("weight_mode"),
            "rebalance_months": sample.get("rebalance_months"),
            "safe_weight": sample.get("safe_weight"),
            "max_weight": sample.get("max_weight"),
            "min_weight": sample.get("min_weight"),
            "trade_cost_bps": sample.get("trade_cost_bps"),
        }
        defense_signal_guard_passed = True
        if sample.get("defense_mode") in {"vol_target", "drawdown_guard"}:
            defense_signal_guard_passed = bool(
                sample.get("defense_signal_strict_asof_passed") is True
            )
            if sample.get("defense_signal") in {
                "selected_basket_ex_ante_vol",
                "selected_basket_ex_ante_drawdown",
            }:
                defense_signal_guard_passed = bool(
                    defense_signal_guard_passed
                    and sample.get("defense_signal_prehistory_seeded") is True
                    and sample.get("defense_signal_first_date")
                    and sample.get("start")
                    and str(sample["defense_signal_first_date"]) < str(sample["start"])
                )
        for selection in sample.get("selections") or []:
            weights = selection.get("weights") or {}
            weight_values = [float(value) for value in weights.values()]
            parsed_weights = {
                str(ticker): float(value) for ticker, value in weights.items()
            }
            raw_candidate_pool = sample.get("candidate_tickers")
            if isinstance(raw_candidate_pool, str):
                candidate_pool = {
                    item.strip().upper()
                    for item in raw_candidate_pool.replace(",", ";").split(";")
                    if item.strip()
                }
            else:
                candidate_pool = {
                    str(item).strip().upper()
                    for item in (raw_candidate_pool or [])
                    if str(item).strip()
                }
            top_n = int(sample.get("top_n") or 0)
            max_weight = float(sample.get("max_weight") or 0.0)
            min_weight = float(sample.get("min_weight") or 0.02)
            effective_cap = max(max_weight, 1.0 / top_n) if top_n > 0 else max_weight
            effective_floor = min(min_weight, 1.0 / top_n) if top_n > 0 else min_weight
            top_weight_count = sample.get("top_weight_count")
            top_weight_cap = sample.get("top_weight_cap")
            parsed_top_count = (
                int(top_weight_count) if top_weight_count is not None else None
            )
            observed_top_weight = (
                sum(sorted(weight_values, reverse=True)[:parsed_top_count])
                if weight_values and parsed_top_count is not None else None
            )
            observed_effective_n = (
                1.0 / sum(value * value for value in weight_values)
                if weight_values else None
            )
            top_weight_cap_passed = bool(
                top_weight_count is None
                and top_weight_cap is None
                or (
                    parsed_top_count is not None
                    and 1 <= parsed_top_count < top_n
                    and top_weight_cap is not None
                    and observed_top_weight is not None
                    and observed_top_weight <= float(top_weight_cap) + 1e-8
                )
            )
            effective_n_floor_passed = bool(
                sample.get("min_effective_n") is None
                or (
                    observed_effective_n is not None
                    and observed_effective_n
                    >= float(sample["min_effective_n"]) - 1e-8
                )
            )
            selection_mode = str(sample.get("selection_mode") or "rank")
            labels = selection.get("cluster_labels") or {}
            cluster_weights = selection.get("cluster_weights") or {}
            correlation_required = (
                selection_mode != "rank" or sample.get("cluster_max_weight") is not None
            )
            correlation_evidence_passed = bool(
                not correlation_required
                or (
                    str(selection.get("selection_mode") or "") == selection_mode
                    and int(selection.get("correlation_observations") or 0) >= 60
                    and set(labels) == set(weights)
                    and bool(cluster_weights)
                )
            )
            correlation_asof_passed = bool(
                not correlation_required
                or (
                    selection.get("correlation_strict_asof_passed") is True
                    and selection.get("correlation_price_end")
                    and selection.get("date")
                    and str(selection["correlation_price_end"]) < str(selection["date"])
                )
            )
            pair_cap_passed = bool(
                selection_mode != "correlation_replacement"
                or (
                    selection.get("max_pair_correlation") is not None
                    and float(selection["max_pair_correlation"])
                    < float(sample.get("correlation_threshold") or 0.0) - 1e-8
                )
            )
            label_counts = pd.Series(list(labels.values())).value_counts() if labels else pd.Series(dtype=float)
            cluster_name_cap_passed = bool(
                selection_mode != "cluster_round_robin"
                or (
                    not label_counts.empty
                    and int(label_counts.max())
                    <= int(sample.get("max_names_per_cluster") or 0)
                )
            )
            cluster_weight_cap_passed = bool(
                sample.get("cluster_max_weight") is None
                or (
                    bool(cluster_weights)
                    and max(float(value) for value in cluster_weights.values())
                    <= float(sample["cluster_max_weight"]) + 1e-8
                )
            )
            quality_sic = _quality_sic_selection_audit(
                sample,
                selection,
                parsed_weights,
                tolerance=1e-8,
            )
            valuation = _valuation_selection_audit(
                sample,
                selection,
                tolerance=1e-8,
            )
            sic4_by_ticker = selection.get("sic4_by_ticker") or {}
            sic2_by_ticker = selection.get("sic2_by_ticker") or {}
            sic4_weights = selection.get("sic4_weights") or {}
            sic2_weights = selection.get("sic2_weights") or {}
            rows.append({
                **shape,
                "selection_date": selection.get("date"),
                "trigger": selection.get("trigger"),
                "report_date": selection.get("report_date"),
                "available_date": selection.get("available_date"),
                "accession_number": selection.get("accession_number"),
                "expected_security_count": selection.get("expected_count"),
                "expected_issuer_count": selection.get("expected_issuer_count"),
                "candidate_member_count": selection.get("candidate_member_count"),
                "price_count": selection.get("price_count"),
                "scored_count": selection.get("scored_count"),
                "price_coverage": selection.get("price_coverage"),
                "score_coverage": selection.get("score_coverage"),
                "tickers": ";".join(selection.get("tickers") or []),
                "weights": ";".join(
                    f"{ticker}:{float(weight):.8f}"
                    for ticker, weight in sorted(weights.items())
                ),
                "selection_mode_observed": selection.get("selection_mode"),
                "correlation_observations": selection.get(
                    "correlation_observations"
                ),
                "correlation_price_start": selection.get("correlation_price_start"),
                "correlation_price_end": selection.get("correlation_price_end"),
                "correlation_strict_asof_passed": selection.get(
                    "correlation_strict_asof_passed"
                ),
                "correlation_pool_size_observed": selection.get(
                    "correlation_pool_size"
                ),
                "max_pair_correlation": selection.get("max_pair_correlation"),
                "cluster_labels": ";".join(
                    f"{ticker}:{int(cluster_id)}"
                    for ticker, cluster_id in sorted(labels.items())
                ),
                "cluster_weights": ";".join(
                    f"{int(cluster_id)}:{float(weight):.8f}"
                    for cluster_id, weight in sorted(
                        cluster_weights.items(),
                        key=lambda item: int(item[0]),
                    )
                ),
                "observed_max_cluster_weight": (
                    max(cluster_weights.values()) if cluster_weights else None
                ),
                "observed_top_weight": observed_top_weight,
                "observed_effective_n": observed_effective_n,
                "observed_max_names_per_cluster": (
                    int(label_counts.max()) if not label_counts.empty else None
                ),
                "rejected_correlated_tickers": ";".join(
                    selection.get("rejected_correlated_tickers") or []
                ),
                "fundamental_score_weight_observed": selection.get(
                    "fundamental_score_weight"
                ),
                "quality_policy_id_observed": selection.get("quality_policy_id"),
                "valuation_score_weight_observed": selection.get(
                    "valuation_score_weight"
                ),
                "valuation_policy_id_observed": selection.get(
                    "valuation_policy_id"
                ),
                "valuation_score_coverage": selection.get(
                    "valuation_score_coverage"
                ),
                "valuation_observed_coverage": selection.get(
                    "valuation_observed_coverage"
                ),
                "valuation_score_date": selection.get("valuation_score_date"),
                "valuation_future_data_guard_passed": selection.get(
                    "valuation_future_data_guard_passed"
                ),
                "valuation_missing_neutral_count": selection.get(
                    "valuation_missing_neutral_count"
                ),
                "quality_score_coverage": selection.get(
                    "quality_score_coverage"
                ),
                "quality_score_date": selection.get("quality_score_date"),
                "quality_future_data_guard_observed": selection.get(
                    "quality_future_data_guard_passed"
                ),
                "quality_selection_mode_observed": selection.get(
                    "quality_selection_mode"
                ),
                "quality_challenge_pool_size_observed": selection.get(
                    "quality_challenge_pool_size"
                ),
                "max_quality_replacements_observed": selection.get(
                    "max_quality_replacements"
                ),
                "quality_challenge_pool_tickers": ";".join(
                    selection.get("quality_challenge_pool_tickers") or []
                ),
                "quality_replacement_pairs": json.dumps(
                    selection.get("quality_replacement_pairs") or [],
                    ensure_ascii=True,
                    sort_keys=True,
                ),
                "observed_quality_replacements": quality_sic[
                    "observed_quality_replacements"
                ],
                "sic4_by_ticker": ";".join(
                    f"{ticker}:{int(value)}"
                    for ticker, value in sorted(sic4_by_ticker.items())
                ),
                "sic2_by_ticker": ";".join(
                    f"{ticker}:{str(value)}"
                    for ticker, value in sorted(sic2_by_ticker.items())
                ),
                "sic4_weights": ";".join(
                    f"{key}:{float(value):.8f}"
                    for key, value in sorted(
                        sic4_weights.items(), key=lambda item: int(item[0])
                    )
                ),
                "sic2_weights": ";".join(
                    f"{key}:{float(value):.8f}"
                    for key, value in sorted(
                        sic2_weights.items(), key=lambda item: str(item[0])
                    )
                ),
                "observed_max_names_per_sic4": quality_sic[
                    "observed_max_names_per_sic4"
                ],
                "observed_max_names_per_sic2": quality_sic[
                    "observed_max_names_per_sic2"
                ],
                "observed_max_sic4_weight": quality_sic[
                    "observed_max_sic4_weight"
                ],
                "observed_max_sic2_weight": quality_sic[
                    "observed_max_sic2_weight"
                ],
                "rejected_sic_tickers": ";".join(
                    selection.get("rejected_sic_tickers") or []
                ),
                "rejected_sic_reasons": ";".join(
                    f"{ticker}:{reason}"
                    for ticker, reason in sorted(
                        (selection.get("rejected_sic_reasons") or {}).items()
                    )
                ),
                "observed_weight_sum": sum(weight_values),
                "observed_max_weight": max(weight_values) if weight_values else None,
                "observed_min_weight": min(weight_values) if weight_values else None,
                "effective_max_weight": effective_cap,
                "effective_min_weight": effective_floor,
                "weight_sum_passed": bool(
                    weight_values and abs(sum(weight_values) - 1.0) <= 1e-8
                ),
                "weight_cap_passed": bool(
                    weight_values and max(weight_values) <= effective_cap + 1e-8
                ),
                "weight_floor_passed": bool(
                    weight_values and min(weight_values) >= effective_floor - 1e-8
                ),
                "ticker_count_passed": len(weight_values) == top_n,
                "candidate_pool_passed": bool(
                    not candidate_pool
                    or set(map(str.upper, parsed_weights)).issubset(candidate_pool)
                ),
                "churn": selection.get("churn"),
                "turnover": selection.get("turnover"),
                "cost": selection.get("cost"),
                "future_data_guard_passed": bool(
                    selection.get("available_date")
                    and selection.get("date")
                    and str(selection["available_date"]) <= str(selection["date"])
                ),
                "correlation_evidence_passed": correlation_evidence_passed,
                "correlation_future_data_guard_passed": correlation_asof_passed,
                "correlation_pair_cap_passed": pair_cap_passed,
                "cluster_name_cap_passed": cluster_name_cap_passed,
                "cluster_weight_cap_passed": cluster_weight_cap_passed,
                "top_weight_cap_passed": top_weight_cap_passed,
                "effective_n_floor_passed": effective_n_floor_passed,
                "quality_evidence_passed": quality_sic[
                    "quality_evidence_passed"
                ],
                "quality_score_coverage_passed": quality_sic[
                    "quality_score_coverage_passed"
                ],
                "quality_future_data_guard_passed": quality_sic[
                    "quality_future_data_guard_passed"
                ],
                "quality_replacement_guard_passed": quality_sic[
                    "quality_replacement_guard_passed"
                ],
                "valuation_evidence_passed": valuation[
                    "valuation_evidence_passed"
                ],
                "valuation_score_coverage_passed": valuation[
                    "valuation_score_coverage_passed"
                ],
                "valuation_future_data_guard_audit_passed": valuation[
                    "valuation_future_data_guard_audit_passed"
                ],
                "sic4_name_cap_passed": quality_sic["sic4_name_cap_passed"],
                "sic2_name_cap_passed": quality_sic["sic2_name_cap_passed"],
                "sic4_weight_cap_passed": quality_sic[
                    "sic4_weight_cap_passed"
                ],
                "sic2_weight_cap_passed": quality_sic[
                    "sic2_weight_cap_passed"
                ],
                "defense_signal_guard_passed": defense_signal_guard_passed,
            })
    return rows


def _version_rows(checkpoint: dict) -> list[dict]:
    meta = checkpoint["lab_job"]
    manifest = checkpoint["result"].get("data_manifest") or {}
    rows = []
    for source_name, value in manifest.items():
        if not isinstance(value, dict):
            continue
        rows.append({
            "experiment_id": meta.get("experiment_id"),
            "phase_id": meta.get("phase_id"),
            "job_id": meta.get("job_id"),
            "source_name": source_name,
            "version_id": value.get("version_id"),
            "content_sha256": value.get("content_sha256"),
            "source": value.get("source"),
            "source_locator": value.get("source_locator"),
            "data_start": value.get("data_start"),
            "data_end": value.get("data_end"),
            "row_count": value.get("row_count") or value.get("rows"),
            "fetched_at": value.get("fetched_at"),
            "cached": value.get("cached"),
            "stale": value.get("stale"),
        })
    return rows


def _current_holding_rows(checkpoint: dict) -> list[dict]:
    meta = checkpoint["lab_job"]
    result = checkpoint["result"]
    rows = []
    for selection_kind, key in (
        ("governance_best", "current_selection"),
        ("closest_qqq", "current_selection_best_qqq"),
    ):
        selection = result.get(key) or {}
        evidence = selection.get("membership_evidence") or {}
        for holding in selection.get("holdings") or []:
            rows.append({
                "experiment_id": meta.get("experiment_id"),
                "phase_id": meta.get("phase_id"),
                "job_id": meta.get("job_id"),
                "selection_kind": selection_kind,
                "selection_date": selection.get("selection_date"),
                "ticker": holding.get("ticker"),
                "portfolio_weight": holding.get("portfolio_weight"),
                "risk_sleeve_weight": holding.get("risk_sleeve_weight"),
                "safe_weight": selection.get("safe_weight"),
                "weight_mode": selection.get("weight_mode"),
                "max_weight": selection.get("max_weight"),
                "weight_report_date": evidence.get("weight_report_date"),
                "weight_available_date": evidence.get("weight_available_date"),
                "weight_is_current": evidence.get("weight_is_current"),
                "research_only": True,
            })
    return rows


def build_export_frames(
    checkpoints: Iterable[dict],
    *,
    jobs: Iterable[dict] | None = None,
    errors: dict[str, str] | None = None,
) -> dict[str, pd.DataFrame]:
    checkpoints = list(checkpoints)
    completed_by_id = {
        checkpoint["lab_job"]["job_id"]: checkpoint for checkpoint in checkpoints
    }
    error_map = errors or {}
    job_rows = []
    for job in jobs or [item["lab_job"] for item in checkpoints]:
        checkpoint = completed_by_id.get(job["job_id"])
        meta = checkpoint["lab_job"] if checkpoint else job
        status = "completed" if checkpoint else "failed" if job["job_id"] in error_map else "pending"
        construction_profiles = (
            (job.get("params") or {}).get("construction_profiles") or []
        )
        job_rows.append({
            "experiment_id": job.get("experiment_id"),
            "phase_id": job.get("phase_id"),
            "phase_role": job.get("phase_role"),
            "job_id": job.get("job_id"),
            "job_sha256": job.get("job_sha256"),
            "engine_sha256": job.get("engine_sha256"),
            "runtime_sha256": job.get("runtime_sha256"),
            "strategy_config_sha256": job.get("strategy_config_sha256"),
            "status": status,
            "error": error_map.get(job["job_id"]),
            "expected_candidate_rows": job.get("expected_candidate_rows"),
            "actual_candidate_rows": meta.get("actual_candidate_rows"),
            "actual_window_rows": meta.get("actual_window_rows"),
            "duration_seconds": meta.get("duration_seconds"),
            "started_at": meta.get("started_at"),
            "completed_at": meta.get("completed_at"),
            "start": (job.get("params") or {}).get("start"),
            "end": (job.get("params") or {}).get("end"),
            "target": (job.get("params") or {}).get("target"),
            "lookback_years": (job.get("params") or {}).get("lookback_years"),
            "trade_cost_bps": (job.get("params") or {}).get("trade_cost_bps"),
            "top_ns": _joined((job.get("params") or {}).get("top_ns") or []),
            "weight_modes": _joined((job.get("params") or {}).get("weight_modes") or []),
            "safe_weights": _joined((job.get("params") or {}).get("safe_weights") or []),
            "rebalance_months": _joined((job.get("params") or {}).get("rebalance_months") or []),
            "max_weights": _joined((job.get("params") or {}).get("max_weights") or []),
            "construction_profile_count": len(construction_profiles),
            "construction_profile_ids": _joined(
                profile.get("id") or profile.get("profile_id")
                for profile in construction_profiles
            ),
            "include_samples": (job.get("params") or {}).get("include_samples"),
        })

    candidate_rows = []
    window_rows = []
    selection_rows = []
    failure_rows = []
    version_rows = []
    holding_rows = []
    for checkpoint in checkpoints:
        meta = checkpoint["lab_job"]
        result = checkpoint["result"]
        candidate_rows.extend(
            _candidate_row(checkpoint, row) for row in result.get("rows") or []
        )
        window_rows.extend(_window_rows(checkpoint))
        selection_rows.extend(_selection_rows(checkpoint))
        version_rows.extend(_version_rows(checkpoint))
        holding_rows.extend(_current_holding_rows(checkpoint))
        for skipped in result.get("skipped") or []:
            failure_rows.append({
                "experiment_id": meta.get("experiment_id"),
                "phase_id": meta.get("phase_id"),
                "job_id": meta.get("job_id"),
                **skipped,
            })

    candidates = pd.DataFrame(candidate_rows)
    gate_failure_rows = []
    if not candidates.empty:
        for row in candidates.to_dict(orient="records"):
            for failed_check in str(row.get("failed_checks") or "").split(";"):
                if not failed_check:
                    continue
                gate_failure_rows.append({
                    "experiment_id": row.get("experiment_id"),
                    "phase_id": row.get("phase_id"),
                    "job_id": row.get("job_id"),
                    "shape_id": row.get("shape_id"),
                    "basket": row.get("basket"),
                    **{key: row.get(key) for key in CONSTRUCTION_EXPORT_KEYS},
                    "lookback_years": row.get("lookback_years"),
                    "top_n": row.get("top_n"),
                    "weight_mode": row.get("weight_mode"),
                    "rebalance_months": row.get("rebalance_months"),
                    "safe_weight": row.get("safe_weight"),
                    "max_weight": row.get("max_weight"),
                    "trade_cost_bps": row.get("trade_cost_bps"),
                    "failed_check": failed_check,
                    "shape_diagnosis": row.get("shape_diagnosis"),
                    "next_action": row.get("next_action"),
                    "direction_status": row.get("direction_status"),
                    "terminal_direction_decision": False,
                })
    summary_rows = []
    if not candidates.empty:
        grouped = candidates.groupby(
            ["phase_id", "shape_diagnosis", "next_action"], dropna=False
        )
        for (phase_id, diagnosis, next_action), frame in grouped:
            summary_rows.append({
                "phase_id": phase_id,
                "shape_diagnosis": diagnosis,
                "next_action": next_action,
                "candidate_rows": len(frame),
                "admission_passed_rows": int(frame["admission_passed"].fillna(False).sum()),
                "best_target_excess_p50": frame["target_excess_p50"].max(),
                "best_qqq_excess_p50": frame["qqq_excess_p50"].max(),
                "best_qqq_excess_p10": frame["qqq_excess_p10"].max(),
                "best_qqq_dd_relief_p50": frame["qqq_dd_relief_p50"].max(),
                "direction_terminal": False,
            })
    return {
        "summary": pd.DataFrame(summary_rows),
        "jobs": pd.DataFrame(job_rows),
        "candidates": candidates,
        "windows": pd.DataFrame(window_rows),
        "selections": pd.DataFrame(selection_rows),
        "gate_failures": pd.DataFrame(gate_failure_rows),
        "engine_skips": pd.DataFrame(failure_rows),
        "versions": pd.DataFrame(version_rows),
        "current_holdings": pd.DataFrame(holding_rows),
    }


def _atomic_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        frame.to_csv(temporary, index=False, encoding="utf-8-sig")
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def write_plan_csv(
    jobs: Iterable[dict],
    path: str | Path,
    *,
    checkpoints: Iterable[dict] | None = None,
    errors: dict[str, str] | None = None,
) -> None:
    frames = build_export_frames(
        list(checkpoints or []), jobs=list(jobs), errors=errors
    )
    _atomic_csv(frames["jobs"], Path(path))


EXCEL_VALUE_NAMES = {
    True: "是",
    False: "否",
    "completed": "已完成",
    "pending": "待运行",
    "failed": "失败",
    "exploratory_training": "探索训练",
    "sensitivity": "敏感性测试",
    "locked_validation": "锁定验证",
    "equal": "等权",
    "inverse_vol": "逆波动率",
    "risk_parity": "风险平价",
    "score_risk": "评分风险权重",
    "source_weight": "指数源权重",
    "rank": "评分排序",
    "static": "静态安全仓",
}


def _excel_column_title(column: object) -> str:
    name = str(column)
    return EXCEL_COLUMN_NAMES.get(name, f"技术字段（{name}）")


def _localized_excel_frame(frame: pd.DataFrame) -> tuple[pd.DataFrame, list[dict]]:
    localized = frame.copy()
    if not localized.empty:
        localized = localized.dropna(axis=1, how="all")
    mapping = {}
    dictionary = []
    used = set()
    for column in localized.columns:
        title = _excel_column_title(column)
        base = title
        suffix = 2
        while title in used:
            title = f"{base}（{suffix}）"
            suffix += 1
        used.add(title)
        mapping[column] = title
        dictionary.append({"程序字段": str(column), "中文表头": title})
    localized = localized.rename(columns=mapping)
    for column in localized.columns:
        if localized[column].dtype == object or pd.api.types.is_bool_dtype(
            localized[column]
        ):
            localized[column] = localized[column].map(
                lambda value: EXCEL_VALUE_NAMES.get(value, value)
            )
    return localized, dictionary


def _format_workbook(path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    percent_markers = (
        "weight", "rate", "xirr", "maxdd", "relief", "margin",
        "coverage", "turnover_x", "cost_x", "churn",
        "比例", "胜率", "回撤", "收益", "覆盖率", "换手", "成本",
    )
    for sheet in workbook.worksheets:
        if sheet.max_row >= 1 and sheet.max_column >= 1:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
        headers = {
            cell.column: str(cell.value or "") for cell in sheet[1]
        } if sheet.max_row >= 1 else {}
        for column_index, header in headers.items():
            values = [header]
            for row_index in range(2, min(sheet.max_row, 200) + 1):
                value = sheet.cell(row=row_index, column=column_index).value
                if value is not None:
                    values.append(str(value))
                if any(marker in header for marker in percent_markers):
                    sheet.cell(row=row_index, column=column_index).number_format = "0.00%"
            width = min(45, max(10, max(len(value) for value in values) + 2))
            sheet.column_dimensions[sheet.cell(1, column_index).column_letter].width = width
    workbook.save(path)


def write_export_bundle(
    spec: dict,
    jobs: Iterable[dict],
    checkpoints: Iterable[dict],
    output_dir: str | Path,
    *,
    errors: dict[str, str] | None = None,
    write_xlsx: bool = True,
) -> dict:
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    jobs = list(jobs)
    checkpoints = list(checkpoints)
    frames = build_export_frames(checkpoints, jobs=jobs, errors=errors)
    file_map = {
        "summary": "summary.csv",
        "jobs": "job_status.csv",
        "candidates": "candidates.csv",
        "windows": "windows.csv",
        "selections": "selections.csv",
        "gate_failures": "gate_failures.csv",
        "engine_skips": "engine_skips.csv",
        "versions": "versions.csv",
        "current_holdings": "current_holdings.csv",
    }
    for key, filename in file_map.items():
        _atomic_csv(frames[key], output / filename)

    evidence_files = {
        key: {
            "filename": filename,
            "rows": len(frames[key]),
            "sha256": hashlib.sha256((output / filename).read_bytes()).hexdigest(),
        }
        for key, filename in file_map.items()
    }

    phase_versions = {}
    candidates = frames["candidates"]
    if not candidates.empty:
        for phase_id, frame in candidates.groupby("phase_id"):
            bundles = sorted(set(frame["input_version_bundle_id"].dropna().astype(str)))
            phase_versions[str(phase_id)] = {
                "bundle_ids": bundles,
                "comparison_ready": len(bundles) == 1,
            }
    manifest = {
        "schema_version": LAB_SCHEMA_VERSION,
        "experiment_id": spec.get("experiment_id"),
        "objective_contract": spec.get("objective_contract"),
        "analysis_contract": spec.get("analysis_contract"),
        "spec_sha256": spec.get("_spec_sha256"),
        "engine_sha256": spec.get("_engine_sha256"),
        "runtime_sha256": spec.get("_runtime_sha256"),
        "runtime_versions": spec.get("_runtime_versions"),
        "strategy_config_path": spec.get("_strategy_config_path"),
        "strategy_config_sha256": spec.get("_strategy_config_sha256"),
        "generated_at": utc_now_iso(),
        "research_only": True,
        "qqq_role": QQQ_ROLE,
        "direction_status": "active_iteration",
        "terminal_direction_decision": False,
        "planned_jobs": len(jobs),
        "completed_jobs": len(checkpoints),
        "failed_jobs": len(errors or {}),
        "candidate_rows": len(frames["candidates"]),
        "window_rows": len(frames["windows"]),
        "selection_rows": len(frames["selections"]),
        "phase_version_consistency": phase_versions,
        "selection_constraint_audit": {
            "available": not frames["selections"].empty,
            "passed": (
                bool(
                    frames["selections"][
                        [
                            "weight_sum_passed",
                            "weight_cap_passed",
                            "weight_floor_passed",
                            "ticker_count_passed",
                            "candidate_pool_passed",
                            "future_data_guard_passed",
                            "correlation_evidence_passed",
                            "correlation_future_data_guard_passed",
                            "correlation_pair_cap_passed",
                            "cluster_name_cap_passed",
                            "cluster_weight_cap_passed",
                            "top_weight_cap_passed",
                            "effective_n_floor_passed",
                            "quality_evidence_passed",
                            "quality_score_coverage_passed",
                            "quality_future_data_guard_passed",
                            "quality_replacement_guard_passed",
                            "valuation_evidence_passed",
                            "valuation_score_coverage_passed",
                            "valuation_future_data_guard_audit_passed",
                            "sic4_name_cap_passed",
                            "sic2_name_cap_passed",
                            "sic4_weight_cap_passed",
                            "sic2_weight_cap_passed",
                            "defense_signal_guard_passed",
                        ]
                    ].all(axis=None)
                )
                if not frames["selections"].empty else None
            ),
            "weight_sum_failures": (
                int((~frames["selections"]["weight_sum_passed"]).sum())
                if not frames["selections"].empty else None
            ),
            "cap_failures": (
                int((~frames["selections"]["weight_cap_passed"]).sum())
                if not frames["selections"].empty else None
            ),
            "floor_failures": (
                int((~frames["selections"]["weight_floor_passed"]).sum())
                if not frames["selections"].empty else None
            ),
            "ticker_count_failures": (
                int((~frames["selections"]["ticker_count_passed"]).sum())
                if not frames["selections"].empty else None
            ),
            "candidate_pool_failures": (
                int((~frames["selections"]["candidate_pool_passed"]).sum())
                if not frames["selections"].empty else None
            ),
            "future_data_guard_failures": (
                int((~frames["selections"]["future_data_guard_passed"]).sum())
                if not frames["selections"].empty else None
            ),
            "correlation_evidence_failures": (
                int((~frames["selections"]["correlation_evidence_passed"]).sum())
                if not frames["selections"].empty else None
            ),
            "correlation_future_data_guard_failures": (
                int((~frames["selections"]["correlation_future_data_guard_passed"]).sum())
                if not frames["selections"].empty else None
            ),
            "correlation_pair_cap_failures": (
                int((~frames["selections"]["correlation_pair_cap_passed"]).sum())
                if not frames["selections"].empty else None
            ),
            "cluster_name_cap_failures": (
                int((~frames["selections"]["cluster_name_cap_passed"]).sum())
                if not frames["selections"].empty else None
            ),
            "cluster_weight_cap_failures": (
                int((~frames["selections"]["cluster_weight_cap_passed"]).sum())
                if not frames["selections"].empty else None
            ),
            "top_weight_cap_failures": (
                int((~frames["selections"]["top_weight_cap_passed"]).sum())
                if not frames["selections"].empty else None
            ),
            "effective_n_floor_failures": (
                int((~frames["selections"]["effective_n_floor_passed"]).sum())
                if not frames["selections"].empty else None
            ),
            "quality_evidence_failures": (
                int((~frames["selections"]["quality_evidence_passed"]).sum())
                if not frames["selections"].empty else None
            ),
            "quality_score_coverage_failures": (
                int(
                    (~frames["selections"]["quality_score_coverage_passed"]).sum()
                )
                if not frames["selections"].empty else None
            ),
            "quality_future_data_guard_failures": (
                int(
                    (~frames["selections"]["quality_future_data_guard_passed"]).sum()
                )
                if not frames["selections"].empty else None
            ),
            "quality_replacement_guard_failures": (
                int(
                    (~frames["selections"]["quality_replacement_guard_passed"]).sum()
                )
                if not frames["selections"].empty else None
            ),
            "valuation_evidence_failures": (
                int((~frames["selections"]["valuation_evidence_passed"]).sum())
                if not frames["selections"].empty else None
            ),
            "valuation_score_coverage_failures": (
                int(
                    (~frames["selections"]["valuation_score_coverage_passed"]).sum()
                )
                if not frames["selections"].empty else None
            ),
            "valuation_future_data_guard_failures": (
                int(
                    (~frames["selections"][
                        "valuation_future_data_guard_audit_passed"
                    ]).sum()
                )
                if not frames["selections"].empty else None
            ),
            "sic4_name_cap_failures": (
                int((~frames["selections"]["sic4_name_cap_passed"]).sum())
                if not frames["selections"].empty else None
            ),
            "sic2_name_cap_failures": (
                int((~frames["selections"]["sic2_name_cap_passed"]).sum())
                if not frames["selections"].empty else None
            ),
            "sic4_weight_cap_failures": (
                int((~frames["selections"]["sic4_weight_cap_passed"]).sum())
                if not frames["selections"].empty else None
            ),
            "sic2_weight_cap_failures": (
                int((~frames["selections"]["sic2_weight_cap_passed"]).sum())
                if not frames["selections"].empty else None
            ),
            "defense_signal_guard_failures": (
                int((~frames["selections"]["defense_signal_guard_passed"]).sum())
                if not frames["selections"].empty else None
            ),
        },
        "files": {key: filename for key, filename in file_map.items()},
        "evidence_files": evidence_files,
        "comparison_rule": (
            "Only compare rows within a phase when comparison_ready=true; "
            "a failed shape never terminates the research direction."
        ),
    }

    workbook_error = None
    if write_xlsx:
        workbook_path = output / "ndx_experiment_report.xlsx"
        temporary = workbook_path.with_name(workbook_path.stem + ".tmp.xlsx")
        try:
            readme = pd.DataFrame([
                {"项目": "实验编号", "内容": spec.get("experiment_id")},
                {"项目": "研究问题", "内容": spec.get("research_question")},
                {
                    "项目": "目标约束",
                    "内容": json.dumps(
                        spec.get("objective_contract"), ensure_ascii=False, sort_keys=True
                    ),
                },
                {"项目": "QQQ角色", "内容": "仅作成分代理和辅助参照"},
                {"项目": "主竞争基准", "内容": "当前8只策略配置"},
                {"项目": "研究状态", "内容": "持续迭代"},
                {"项目": "是否形成终局结论", "内容": "否"},
                {"项目": "是否需要人工复核", "内容": "是"},
                {
                    "项目": "比较规则",
                    "内容": "仅比较同阶段且数据版本一致的结果；单个方案失败不终止研究方向。",
                },
            ])
            localized_frames = {}
            field_dictionary = []
            for key, frame in frames.items():
                localized, entries = _localized_excel_frame(frame)
                sheet_name = EXCEL_SHEET_NAMES[key]
                localized_frames[sheet_name] = localized
                field_dictionary.extend(
                    {"所属工作表": sheet_name, **entry} for entry in entries
                )
            with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
                readme.to_excel(writer, sheet_name="使用说明", index=False)
                for sheet_name, frame in localized_frames.items():
                    frame.to_excel(writer, sheet_name=sheet_name, index=False)
                pd.DataFrame(field_dictionary).drop_duplicates().to_excel(
                    writer, sheet_name="字段字典", index=False
                )
            _format_workbook(temporary)
            os.replace(temporary, workbook_path)
            manifest["files"]["workbook"] = workbook_path.name
        except Exception as exc:
            workbook_error = f"{type(exc).__name__}: {exc}"
        finally:
            if temporary.exists():
                temporary.unlink()
    if workbook_error:
        manifest["workbook_error"] = workbook_error
    atomic_write_text(
        output / "manifest.json",
        json.dumps(manifest, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
    )
    return manifest
