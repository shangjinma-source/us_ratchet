"""Calculate diagnostic valuation values without ranking securities."""
from __future__ import annotations

import argparse
import json
import os
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from core.config import load_config
from core.helpers import atomic_write_text
from data.ndx_fundamentals import fundamental_cache_dir
from data.ndx_valuation_diagnostics import (
    VALUATION_DIAGNOSTIC_POLICY_ID,
    calculate_valuation_values,
    valuation_distributions,
    valuation_integrity_failures,
    valuation_tail_review,
)
from data.ndx_valuation_readiness import VALUATION_READINESS_POLICY_ID
from data.versioning import file_sha256, manifest_path, utc_now_iso, write_dataset_manifest


def _resolve(path: str | Path) -> Path:
    value = Path(path)
    return value if value.is_absolute() else ROOT / value


def _atomic_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        frame.to_csv(temporary, index=False, encoding="utf-8-sig")
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def _read_json(path: Path) -> dict:
    if not path.exists():
        raise FileNotFoundError(path)
    return json.loads(path.read_text(encoding="utf-8"))


def _verified_csv(path: Path) -> tuple[pd.DataFrame, dict]:
    sidecar = manifest_path(path)
    manifest = _read_json(sidecar)
    if manifest.get("content_sha256") != file_sha256(path):
        raise ValueError(f"versioned input hash mismatch: {path}")
    return pd.read_csv(path, low_memory=False), manifest


def _verified_readiness(directory: Path) -> tuple[pd.DataFrame, dict, dict]:
    root = _read_json(directory / "manifest.json")
    checks = {
        "policy_id": root.get("policy_id") == VALUATION_READINESS_POLICY_ID,
        "integrity_passed": bool(root.get("integrity_passed")),
        "price_basis_ready": bool(root.get("price_basis_ready")),
        "debt_concepts_ready": bool(root.get("debt_concepts_ready")),
        "valuation_scoring_allowed": not bool(root.get("valuation_scoring_allowed")),
    }
    failed = sorted(key for key, passed in checks.items() if not passed)
    if failed:
        raise ValueError(f"valuation readiness manifest failed checks: {failed}")
    values, manifest = _verified_csv(directory / "valuation_readiness.csv")
    declared = (root.get("evidence_files") or {}).get("readiness") or {}
    if declared.get("sha256") != manifest.get("content_sha256"):
        raise ValueError("valuation readiness hash differs from root manifest")
    return values, manifest, root


def _ticker_coverage(values: pd.DataFrame) -> pd.DataFrame:
    output = []
    for ticker, group in values.groupby("ticker", sort=True):
        fcf = group["fcf_yield_value_ready"].fillna(False).astype(bool)
        ev = group["ev_sales_value_ready"].fillna(False).astype(bool)
        fcf_blockers = sorted({
            item
            for text in group["fcf_yield_blockers"].fillna("").astype(str)
            for item in text.split(";")
            if item
        })
        ev_blockers = sorted({
            item
            for text in group["ev_sales_blockers"].fillna("").astype(str)
            for item in text.split(";")
            if item
        })
        output.append({
            "ticker": ticker,
            "selection_rows": len(group),
            "first_selection_date": group["selection_date"].min(),
            "last_selection_date": group["selection_date"].max(),
            "fcf_yield_ready_rows": int(fcf.sum()),
            "fcf_yield_ready_rate": float(fcf.mean()),
            "ev_sales_ready_rows": int(ev.sum()),
            "ev_sales_ready_rate": float(ev.mean()),
            "fcf_yield_blockers": ";".join(fcf_blockers),
            "ev_sales_blockers": ";".join(ev_blockers),
        })
    return pd.DataFrame(output)


def _review_rows(values: pd.DataFrame) -> pd.DataFrame:
    mask = (
        values["negative_enterprise_value"].fillna(False).astype(bool)
        | values["fcf_yield_unit_review"].fillna(False).astype(bool)
        | values["ev_sales_unit_review"].fillna(False).astype(bool)
    )
    columns = [
        "selection_date",
        "ticker",
        "cik",
        "market_cap_value",
        "free_cash_flow_value",
        "fcf_yield_value",
        "debt_value",
        "cash_value",
        "enterprise_value",
        "revenue_value",
        "ev_sales_value",
        "negative_enterprise_value",
        "fcf_yield_unit_review",
        "ev_sales_unit_review",
    ]
    return values.loc[mask, columns].sort_values(
        ["selection_date", "ticker"], kind="stable"
    ).reset_index(drop=True)


def _style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    fill = PatternFill("solid", fgColor="DCE6F1")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            width = max(len(str(cell.value or "")) for cell in column[:300]) + 2
            sheet.column_dimensions[letter].width = min(max(width, 10), 48)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)


def _write_workbook(
    path: Path,
    frames: dict[str, pd.DataFrame],
) -> None:
    translations = {
        "selection_date": "重选日",
        "ticker": "股票代码",
        "cik": "SEC CIK",
        "price": "未复权价格",
        "price_date": "价格日期",
        "shares_value": "申报股数",
        "shares_period_end": "股数期末日",
        "market_cap_value": "市值",
        "market_cap_unit": "市值单位",
        "free_cash_flow_value": "自由现金流",
        "free_cash_flow_unit": "自由现金流单位",
        "fcf_yield_value": "FCF收益率",
        "fcf_yield_value_ready": "FCF收益率就绪",
        "debt_value": "债务",
        "debt_unit": "债务单位",
        "debt_method": "债务方法",
        "cash_value": "现金",
        "cash_unit": "现金单位",
        "revenue_value": "营业收入",
        "revenue_unit": "收入单位",
        "enterprise_value": "企业价值",
        "enterprise_value_unit": "企业价值单位",
        "ev_sales_value": "EV销售额倍数",
        "ev_sales_value_ready": "EV销售额就绪",
        "negative_enterprise_value": "企业价值为负",
        "fcf_yield_unit_review": "FCF收益率单位复核",
        "ev_sales_unit_review": "EV销售额单位复核",
        "valuation_value_integrity_ready": "数值完整性通过",
        "valuation_scoring_allowed": "允许评分",
        "factor": "因子",
        "universe_count": "候选数",
        "ready_count": "就绪数",
        "ready_coverage": "就绪率",
        "negative_count": "负值数",
        "negative_rate": "负值率",
        "minimum": "最小值",
        "maximum": "最大值",
        "tail": "尾部",
        "tail_rank": "尾部序号",
        "value": "数值",
        "review_only": "仅复核",
        "selection_rows": "重选行数",
        "first_selection_date": "首个重选日",
        "last_selection_date": "末个重选日",
        "fcf_yield_ready_rows": "FCF收益率就绪行",
        "fcf_yield_ready_rate": "FCF收益率就绪率",
        "ev_sales_ready_rows": "EV销售额就绪行",
        "ev_sales_ready_rate": "EV销售额就绪率",
        "failure_type": "失败类型",
        "source_name": "来源名称",
        "version_id": "版本号",
        "content_sha256": "内容哈希",
        "rows": "行数",
        "note": "说明",
    }

    def zh(frame: pd.DataFrame) -> pd.DataFrame:
        return frame.rename(columns=translations)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet, frame in frames.items():
            zh(frame).to_excel(writer, sheet_name=sheet, index=False)
    _style_workbook(path)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Calculate valuation values for diagnostics only; never rank."
    )
    parser.add_argument("--config", default="config/config.yaml")
    parser.add_argument("--input-dir", default="")
    parser.add_argument("--output-dir", default="")
    parser.add_argument("--tail-size", type=int, default=5)
    args = parser.parse_args()

    cfg = load_config(_resolve(args.config))
    fundamental_dir = fundamental_cache_dir(cfg)
    source = (
        _resolve(args.input_dir)
        if args.input_dir
        else fundamental_dir / "valuation_readiness_v4"
    )
    output = (
        _resolve(args.output_dir)
        if args.output_dir
        else fundamental_dir / "valuation_diagnostics_v3"
    )
    output.mkdir(parents=True, exist_ok=True)
    readiness, readiness_manifest, readiness_root = _verified_readiness(source)
    values = calculate_valuation_values(readiness)
    distributions = valuation_distributions(values)
    tails = valuation_tail_review(values, tail_size=int(args.tail_size))
    ticker_coverage = _ticker_coverage(values)
    reviews = _review_rows(values)
    failures = valuation_integrity_failures(values)
    versions = pd.DataFrame([{
        "source_name": "valuation_readiness_v4",
        "version_id": readiness_manifest.get("version_id"),
        "content_sha256": readiness_manifest.get("content_sha256"),
        "rows": len(readiness),
        "note": (
            f"policy={readiness_root.get('policy_id')}；价格口径与债务v2均通过；"
            "评分关闭"
        ),
    }])
    frames = {
        "values": values,
        "distributions": distributions,
        "tail_review": tails,
        "ticker_coverage": ticker_coverage,
        "review_rows": reviews,
        "integrity_failures": failures,
        "versions": versions,
    }
    paths = {key: output / f"{key}.csv" for key in frames}
    evidence = {}
    for key, frame in frames.items():
        _atomic_csv(frame, paths[key])
        manifest = write_dataset_manifest(
            paths[key],
            dataset_type="ndx_valuation_diagnostics",
            dataset_key=f"{VALUATION_DIAGNOSTIC_POLICY_ID}:{key}",
            source="derived_from_versioned_valuation_readiness",
            fetched_at_evidence="derived_runtime_utc",
            row_count=len(frame),
            data_start=(
                str(values["selection_date"].min())
                if "selection_date" in frame.columns and not frame.empty else None
            ),
            data_end=(
                str(values["selection_date"].max())
                if "selection_date" in frame.columns and not frame.empty else None
            ),
            extra={
                "policy_id": VALUATION_DIAGNOSTIC_POLICY_ID,
                "research_only": True,
                "valuation_scoring_allowed": False,
            },
        )
        evidence[key] = {
            "filename": paths[key].name,
            "rows": len(frame),
            "sha256": manifest["content_sha256"],
            "version_id": manifest["version_id"],
        }
    fcf_values = values["fcf_yield_value"].dropna()
    ev_values = values["ev_sales_value"].dropna()
    summary = {
        "schema_version": 1,
        "generated_at": utc_now_iso(),
        "policy_id": VALUATION_DIAGNOSTIC_POLICY_ID,
        "research_only": True,
        "source_policy_id": readiness_root.get("policy_id"),
        "source_readiness_version_id": readiness_manifest.get("version_id"),
        "source_readiness_sha256": readiness_manifest.get("content_sha256"),
        "selection_rows": len(values),
        "selection_dates": int(values["selection_date"].nunique()),
        "ticker_count": int(values["ticker"].nunique()),
        "fcf_yield_value_rows": len(fcf_values),
        "ev_sales_value_rows": len(ev_values),
        "fcf_yield_minimum": float(fcf_values.min()) if len(fcf_values) else None,
        "fcf_yield_maximum": float(fcf_values.max()) if len(fcf_values) else None,
        "ev_sales_minimum": float(ev_values.min()) if len(ev_values) else None,
        "ev_sales_maximum": float(ev_values.max()) if len(ev_values) else None,
        "negative_enterprise_value_rows": int(
            values["negative_enterprise_value"].sum()
        ),
        "fcf_yield_unit_review_rows": int(values["fcf_yield_unit_review"].sum()),
        "ev_sales_unit_review_rows": int(values["ev_sales_unit_review"].sum()),
        "review_rows": len(reviews),
        "integrity_failure_rows": len(failures),
        "integrity_passed": failures.empty,
        "training_spec_allowed": False,
        "valuation_scoring_allowed": False,
        "next_required_evidence": [
            "manual_review_of_extreme_and_negative_enterprise_values",
            "frozen_missing_value_and_cross_section_normalization_policy",
            "point_in_time_fx_adr_and_multi_security_allocation",
        ],
        "files": {
            **{key: value["filename"] for key, value in evidence.items()},
            "workbook": "ndx_valuation_diagnostics.xlsx",
        },
    }
    atomic_write_text(
        output / "summary.json",
        json.dumps(summary, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
    )
    atomic_write_text(
        output / "manifest.json",
        json.dumps(
            {
                "schema_version": 1,
                "generated_at": summary["generated_at"],
                "policy_id": VALUATION_DIAGNOSTIC_POLICY_ID,
                "source_readiness_version_id": summary[
                    "source_readiness_version_id"
                ],
                "source_readiness_sha256": summary["source_readiness_sha256"],
                "evidence_files": evidence,
                "integrity_passed": failures.empty,
                "training_spec_allowed": False,
                "valuation_scoring_allowed": False,
            },
            ensure_ascii=False,
            indent=2,
            allow_nan=False,
        ) + "\n",
    )
    conclusion = pd.DataFrame([
        {"项目": "研究目的", "内容": "计算估值值域并复核单位；不排名、不打分"},
        {"项目": "点时候选行", "内容": len(values)},
        {"项目": "FCF收益率数值行", "内容": len(fcf_values)},
        {"项目": "EV销售额数值行", "内容": len(ev_values)},
        {"项目": "FCF收益率范围", "内容": f"{summary['fcf_yield_minimum']}..{summary['fcf_yield_maximum']}"},
        {"项目": "EV销售额范围", "内容": f"{summary['ev_sales_minimum']}..{summary['ev_sales_maximum']}"},
        {"项目": "负企业价值行", "内容": summary["negative_enterprise_value_rows"]},
        {"项目": "FCF单位复核行", "内容": summary["fcf_yield_unit_review_rows"]},
        {"项目": "EV销售额单位复核行", "内容": summary["ev_sales_unit_review_rows"]},
        {"项目": "完整性失败", "内容": len(failures)},
        {"项目": "允许生成训练规格", "内容": False},
        {"项目": "估值评分/实盘", "内容": "关闭；不改配置、不改今日建议"},
    ])
    workbook = output / "ndx_valuation_diagnostics.xlsx"
    _write_workbook(
        workbook,
        {
            "结论": conclusion,
            "逐期估值值": values,
            "逐期分布": distributions,
            "尾部复核": tails,
            "逐股覆盖": ticker_coverage,
            "异常复核": reviews,
            "完整性失败": failures,
            "数据版本": versions,
        },
    )
    print("NDX点时估值值域诊断（research only）")
    print(
        f"FCF Yield: {len(fcf_values)} rows | "
        f"range {summary['fcf_yield_minimum']:.4f}..{summary['fcf_yield_maximum']:.4f}"
    )
    print(
        f"EV/Sales: {len(ev_values)} rows | "
        f"range {summary['ev_sales_minimum']:.4f}..{summary['ev_sales_maximum']:.4f}"
    )
    print(
        f"Review rows: {len(reviews)} | integrity failures: {len(failures)} | "
        "training/scoring disabled"
    )
    print(f"Workbook: {workbook}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
