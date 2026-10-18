"""Build versioned unadjusted price evidence for NDX valuation research."""
from __future__ import annotations

import argparse
import json
import os
import sys
import time
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from backtest.ndx_competition import _price_universe_version
from backtest.preference_walkforward import _load_series, _research_cfg
from core.config import load_config
from core.helpers import atomic_write_text
from data.ndx_fundamentals import fundamental_cache_dir
from data.ndx_price_basis import (
    PRICE_BASIS_POLICY_ID,
    RAW_PRICE_CACHE_POLICY_ID,
    PriceBasisPolicy,
    crosscheck_share_split_events,
    fetch_unadjusted_bars,
    load_verified_unadjusted_cache,
    point_in_time_raw_price_rows,
    raw_cache_path,
    raw_split_events,
    write_unadjusted_cache,
)
from data.ndx_security_identity import (
    apply_security_lifecycle,
    load_security_lifecycle,
)
from data.versioning import file_sha256, manifest_path, utc_now_iso, write_dataset_manifest


ATTEMPT_COLUMNS = [
    "policy_id",
    "ticker",
    "requested_begin",
    "requested_end",
    "source_chain",
    "status",
    "source",
    "source_request_mode",
    "rows",
    "data_start",
    "data_end",
    "version_id",
    "content_sha256",
    "error",
    "attempted_at",
    "duration_seconds",
]


def _resolve(path: str | Path) -> Path:
    value = Path(path)
    return value if value.is_absolute() else ROOT / value


def _atomic_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        frame.to_csv(temporary, index=False, encoding="utf-8-sig")
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def _read_json(path: Path) -> dict:
    if not path.exists():
        raise FileNotFoundError(path)
    return json.loads(path.read_text(encoding="utf-8"))


def _verified_csv(path: Path) -> tuple[pd.DataFrame, dict]:
    sidecar = manifest_path(path)
    manifest = _read_json(sidecar)
    if manifest.get("content_sha256") != file_sha256(path):
        raise ValueError(f"versioned input hash mismatch: {path}")
    return pd.read_csv(path, low_memory=False), manifest


def _load_attempts(path: Path) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame(columns=ATTEMPT_COLUMNS)
    frame = pd.read_csv(path, dtype=str, keep_default_na=False)
    for column in ATTEMPT_COLUMNS:
        if column not in frame.columns:
            frame[column] = ""
    return frame[ATTEMPT_COLUMNS].drop_duplicates("ticker", keep="last")


def _matching_attempts(
    attempts: pd.DataFrame,
    *,
    begin: str,
    end: str,
    source_chain: str,
) -> pd.DataFrame:
    if attempts.empty:
        return attempts.copy()
    return attempts[
        attempts["policy_id"].eq(RAW_PRICE_CACHE_POLICY_ID)
        & attempts["requested_begin"].eq(begin)
        & attempts["requested_end"].eq(end)
        & attempts["source_chain"].eq(source_chain)
    ].copy()


def _upsert_attempt(attempts: pd.DataFrame, row: dict) -> pd.DataFrame:
    frame = attempts[~attempts["ticker"].eq(row["ticker"])].copy()
    appended = pd.DataFrame([row], columns=ATTEMPT_COLUMNS)
    frame = (
        appended
        if frame.empty
        else pd.concat([frame, appended], ignore_index=True)
    )
    return frame.sort_values("ticker", kind="stable").reset_index(drop=True)


def _cache_inventory(
    output: Path,
    tickers: list[str],
    *,
    begin: str,
    end: str,
) -> tuple[dict[str, pd.DataFrame], dict[str, dict], dict[str, str]]:
    series: dict[str, pd.DataFrame] = {}
    manifests: dict[str, dict] = {}
    errors: dict[str, str] = {}
    for ticker in tickers:
        try:
            frame, manifest = load_verified_unadjusted_cache(
                output,
                ticker,
                requested_begin=begin,
                requested_end=end,
            )
            series[ticker] = frame
            manifests[ticker] = manifest
        except Exception as exc:
            errors[ticker] = str(exc)
    return series, manifests, errors


def _coverage_by_date(rows: pd.DataFrame) -> pd.DataFrame:
    output = []
    for selection_date, group in rows.groupby("selection_date", sort=True):
        ready = group["price_ready"].fillna(False).astype(bool)
        basis = group["price_basis_ready"].fillna(False).astype(bool)
        aligned = group["adjusted_alignment_status"].eq("same_date_ratio_available")
        output.append({
            "selection_date": selection_date,
            "universe_count": len(group),
            "price_basis_ready_count": int(basis.sum()),
            "price_basis_ready_coverage": float(basis.mean()) if len(group) else None,
            "raw_price_ready_count": int(ready.sum()),
            "raw_price_ready_coverage": float(ready.mean()) if len(group) else None,
            "same_date_adjusted_overlap_count": int(aligned.sum()),
            "same_date_adjusted_overlap_coverage": (
                float(aligned.mean()) if len(group) else None
            ),
        })
    return pd.DataFrame(output)


def _source_summary(
    tickers: list[str],
    attempts: pd.DataFrame,
    manifests: dict[str, dict],
    coverage: pd.DataFrame,
) -> pd.DataFrame:
    attempt_map = {
        str(row.ticker): row._asdict()
        for row in attempts.itertuples(index=False)
    }
    output = []
    for ticker in tickers:
        manifest = manifests.get(ticker) or {}
        attempt = attempt_map.get(ticker) or {}
        source_column = (
            coverage["price_ticker"]
            if "price_ticker" in coverage.columns
            else coverage["ticker"]
        )
        group = coverage[source_column.eq(ticker)]
        output.append({
            "ticker": ticker,
            "fetch_status": "ready" if ticker in manifests else attempt.get("status", "pending"),
            "source": manifest.get("source") or attempt.get("source"),
            "source_request_mode": (
                manifest.get("source_request_mode") or attempt.get("source_request_mode")
            ),
            "source_semantics_evidence": manifest.get("source_semantics_evidence"),
            "raw_rows": manifest.get("row_count") or attempt.get("rows"),
            "data_start": manifest.get("data_start") or attempt.get("data_start"),
            "data_end": manifest.get("data_end") or attempt.get("data_end"),
            "version_id": manifest.get("version_id") or attempt.get("version_id"),
            "content_sha256": (
                manifest.get("content_sha256") or attempt.get("content_sha256")
            ),
            "selection_rows": len(group),
            "ready_selection_rows": int(
                group["price_ready"].fillna(False).astype(bool).sum()
            ),
            "error": attempt.get("error", ""),
        })
    return pd.DataFrame(output)


def _integrity_failures(
    plan: pd.DataFrame,
    coverage: pd.DataFrame,
    crosscheck: pd.DataFrame,
) -> pd.DataFrame:
    output = []
    if len(coverage) != len(plan):
        output.append({
            "scope": "coverage",
            "ticker": "",
            "selection_date": "",
            "reason": f"row_count_mismatch:{len(coverage)}!={len(plan)}",
        })
    if coverage.duplicated(["selection_date", "ticker"]).any():
        output.append({
            "scope": "coverage",
            "ticker": "",
            "selection_date": "",
            "reason": "duplicate_selection_ticker",
        })
    invalid_guard = coverage[
        coverage["price_ready"].fillna(False).astype(bool)
        & ~coverage["price_future_data_guard_passed"].fillna(False).astype(bool)
    ]
    for row in invalid_guard.itertuples(index=False):
        output.append({
            "scope": "future_guard",
            "ticker": row.ticker,
            "selection_date": row.selection_date,
            "reason": "ready_price_is_not_strictly_prior",
        })
    ratio = pd.to_numeric(coverage["adjusted_to_raw_ratio"], errors="coerce")
    invalid_ratio = coverage[ratio.notna() & (~np.isfinite(ratio) | (ratio <= 0))]
    for row in invalid_ratio.itertuples(index=False):
        output.append({
            "scope": "adjusted_raw_alignment",
            "ticker": row.ticker,
            "selection_date": row.selection_date,
            "reason": "non_positive_or_non_finite_ratio",
        })
    if not crosscheck.empty:
        failed = crosscheck[
            crosscheck["share_split_like"].fillna(False).astype(bool)
            & crosscheck["crosscheck_status"].isin(
                ["expected_raw_split_missing", "raw_split_factor_mismatch"]
            )
        ]
        for row in failed.itertuples(index=False):
            output.append({
                "scope": "share_split_crosscheck",
                "ticker": row.ticker,
                "selection_date": row.selection_date,
                "reason": row.crosscheck_status,
            })
    return pd.DataFrame(
        output,
        columns=["scope", "ticker", "selection_date", "reason"],
    )


def _style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="DCE6F1")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            width = max(len(str(cell.value or "")) for cell in column[:300]) + 2
            sheet.column_dimensions[letter].width = min(max(width, 10), 48)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)


def _write_workbook(
    path: Path,
    *,
    conclusion: pd.DataFrame,
    identity: pd.DataFrame,
    fetch_plan: pd.DataFrame,
    coverage: pd.DataFrame,
    by_date: pd.DataFrame,
    sources: pd.DataFrame,
    split_events: pd.DataFrame,
    crosscheck: pd.DataFrame,
    failures: pd.DataFrame,
    versions: pd.DataFrame,
) -> None:
    translations = {
        "selection_date": "重选日",
        "ticker": "股票代码",
        "cik": "SEC CIK",
        "price_ticker": "实际取价代码",
        "ticker_alias_applied": "已应用代码别名",
        "ticker_lifecycle_rule_id": "代码生命周期规则",
        "ticker_lifecycle_status": "代码生命周期状态",
        "security_cusip": "证券CUSIP",
        "ticker_lifecycle_evidence_url": "代码变更官方证据",
        "price_status": "价格状态",
        "price_ready": "原始价格就绪",
        "price_basis_ready": "价格口径就绪",
        "price": "未复权收盘价",
        "price_date": "价格日期",
        "price_age_days": "价格滞后天数",
        "price_future_data_guard_passed": "未来数据守门通过",
        "historical_price_basis": "历史价格口径",
        "source": "来源",
        "source_request_mode": "来源请求模式",
        "source_semantics_evidence": "未复权语义证据",
        "raw_price_version_id": "原始价格版本",
        "raw_price_sha256": "原始价格哈希",
        "adjusted_close_same_date": "同日收益回测价",
        "adjusted_to_raw_ratio": "收益价比原始价",
        "adjusted_alignment_status": "同日对齐状态",
        "adjusted_reference_ticker": "收益价参照代码",
        "universe_count": "候选数",
        "price_basis_ready_count": "口径就绪数",
        "price_basis_ready_coverage": "口径就绪率",
        "raw_price_ready_count": "原始价就绪数",
        "raw_price_ready_coverage": "原始价就绪率",
        "same_date_adjusted_overlap_count": "同日重叠数",
        "same_date_adjusted_overlap_coverage": "同日重叠率",
        "fetch_status": "抓取状态",
        "raw_rows": "原始行情行数",
        "data_start": "数据起日",
        "data_end": "数据止日",
        "version_id": "版本号",
        "content_sha256": "内容哈希",
        "selection_rows": "重选行数",
        "ready_selection_rows": "就绪重选行数",
        "error": "错误",
        "prior_date": "前一交易日",
        "event_date": "价格事件日",
        "prior_close": "事件前收盘",
        "close": "事件日收盘",
        "raw_price_ratio": "原始价格比",
        "split_factor": "拆股倍数",
        "split_kind": "拆股方向",
        "prior_selection_date": "前一重选日",
        "share_split_like": "股数变化类似拆股",
        "expected_split_factor": "股数隐含倍数",
        "matched_raw_event_date": "匹配原始价格事件日",
        "matched_raw_split_factor": "匹配原始价格倍数",
        "factor_relative_error": "倍数相对误差",
        "crosscheck_status": "交叉核验状态",
        "scope": "范围",
        "reason": "原因",
        "requested_begin": "请求起日",
        "requested_end": "请求止日",
        "first_selection_date": "首个重选日",
        "last_selection_date": "末个重选日",
        "status": "计划状态",
        "note": "说明",
    }

    def zh(frame: pd.DataFrame) -> pd.DataFrame:
        return frame.rename(columns=translations)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        conclusion.to_excel(writer, sheet_name="结论", index=False)
        zh(identity).to_excel(writer, sheet_name="证券代码生命周期", index=False)
        zh(fetch_plan).to_excel(writer, sheet_name="抓取计划", index=False)
        zh(coverage).to_excel(writer, sheet_name="逐期价格口径", index=False)
        zh(by_date).to_excel(writer, sheet_name="日期覆盖率", index=False)
        zh(sources).to_excel(writer, sheet_name="逐股来源", index=False)
        zh(split_events).to_excel(writer, sheet_name="原始价拆股事件", index=False)
        zh(crosscheck).to_excel(writer, sheet_name="股数价格交叉核验", index=False)
        zh(failures).to_excel(writer, sheet_name="完整性失败", index=False)
        zh(versions).to_excel(writer, sheet_name="数据版本", index=False)
    _style_workbook(path)


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Build unadjusted price evidence for point-in-time valuation; "
            "never overwrite the adjusted return-engine cache."
        )
    )
    parser.add_argument("--config", default="config/config.yaml")
    parser.add_argument("--plan", default="")
    parser.add_argument("--output-dir", default="")
    parser.add_argument("--raw-cache-dir", default="")
    parser.add_argument(
        "--security-lifecycle",
        default="config/ndx_security_lifecycle.yaml",
    )
    parser.add_argument("--source-chain", default="sina")
    parser.add_argument("--history-buffer-days", type=int, default=400)
    parser.add_argument("--maximum-price-age-days", type=int, default=7)
    parser.add_argument("--max-tickers", type=int, default=0)
    parser.add_argument("--plan-only", action="store_true")
    parser.add_argument("--refresh", action="store_true")
    parser.add_argument("--retry-failed", action="store_true")
    args = parser.parse_args()

    cfg = load_config(_resolve(args.config))
    fundamental_dir = fundamental_cache_dir(cfg)
    plan_path = _resolve(args.plan) if args.plan else fundamental_dir / "plan.csv"
    output = (
        _resolve(args.output_dir)
        if args.output_dir
        else fundamental_dir / "price_basis_v2"
    )
    raw_cache = (
        _resolve(args.raw_cache_dir)
        if args.raw_cache_dir
        else fundamental_dir / "price_basis_v1"
    )
    output.mkdir(parents=True, exist_ok=True)
    raw_cache.mkdir(parents=True, exist_ok=True)
    plan = pd.read_csv(plan_path, low_memory=False)
    required = {"selection_date", "ticker", "cik"}
    if not required.issubset(plan.columns):
        raise ValueError(f"NDX plan lacks {sorted(required - set(plan.columns))}")
    plan["selection_date"] = plan["selection_date"].astype(str).str[:10]
    plan["ticker"] = plan["ticker"].astype(str).str.upper().str.strip()
    if plan.duplicated(["selection_date", "ticker"]).any():
        raise ValueError("NDX plan contains duplicate selection_date/ticker")
    lifecycle_path = _resolve(args.security_lifecycle)
    lifecycle = load_security_lifecycle(lifecycle_path)
    identity_plan = apply_security_lifecycle(plan, lifecycle)
    selection_dates = pd.to_datetime(plan["selection_date"])
    begin = (
        selection_dates.min() - pd.Timedelta(days=int(args.history_buffer_days))
    ).strftime("%Y-%m-%d")
    end = selection_dates.max().strftime("%Y-%m-%d")
    logical_tickers = sorted(plan["ticker"].unique())
    tickers = sorted(identity_plan["price_ticker"].unique())
    source_chain = [
        item.strip().lower() for item in args.source_chain.split(",") if item.strip()
    ]
    if not source_chain:
        raise ValueError("source chain is empty")
    source_chain_key = ";".join(source_chain)
    policy = PriceBasisPolicy(maximum_price_age_days=int(args.maximum_price_age_days))
    policy.validate()
    attempts_path = raw_cache / "attempts.csv"
    attempts = _load_attempts(attempts_path)
    matching = _matching_attempts(
        attempts,
        begin=begin,
        end=end,
        source_chain=source_chain_key,
    )
    raw_series, raw_manifests, _cache_errors = _cache_inventory(
        raw_cache,
        tickers,
        begin=begin,
        end=end,
    )
    failed_checkpoints = set(
        matching.loc[matching["status"].eq("failed"), "ticker"].astype(str)
    )
    pending = []
    for ticker in tickers:
        if args.refresh:
            pending.append(ticker)
        elif ticker in raw_series:
            continue
        elif ticker in failed_checkpoints and not args.retry_failed:
            continue
        else:
            pending.append(ticker)
    selected = pending[: int(args.max_tickers)] if int(args.max_tickers) > 0 else pending

    print("NDX 未复权价格口径证据计划")
    print(
        f"策略代码: {len(logical_tickers)} | 实际取价代码: {len(tickers)} | "
        f"已有有效缓存: {len(raw_series)} | "
        f"失败检查点: {len(failed_checkpoints)} | 待执行: {len(pending)}"
    )
    print(
        f"代码别名行: {int(identity_plan['ticker_alias_applied'].sum())} | "
        f"生命周期: {lifecycle.mapping_id}"
    )
    print(f"区间: {begin}..{end} | 来源链: {source_chain_key}")
    print(f"输出: {output}")
    if args.plan_only:
        return 0

    durations = []
    for index, ticker in enumerate(selected, start=1):
        started = time.perf_counter()
        print(f"[{index}/{len(selected)}] {ticker}", flush=True)
        try:
            bars, source_mode, source_errors = fetch_unadjusted_bars(
                ticker,
                begin=begin,
                end=end,
                source_chain=source_chain,
            )
            manifest = write_unadjusted_cache(
                raw_cache,
                ticker,
                bars,
                source_mode=source_mode,
                requested_begin=begin,
                requested_end=end,
            )
            duration = time.perf_counter() - started
            row = {
                "policy_id": RAW_PRICE_CACHE_POLICY_ID,
                "ticker": ticker,
                "requested_begin": begin,
                "requested_end": end,
                "source_chain": source_chain_key,
                "status": "ready",
                "source": manifest.get("source"),
                "source_request_mode": source_mode,
                "rows": len(bars),
                "data_start": manifest.get("data_start"),
                "data_end": manifest.get("data_end"),
                "version_id": manifest.get("version_id"),
                "content_sha256": manifest.get("content_sha256"),
                "error": json.dumps(source_errors, ensure_ascii=False),
                "attempted_at": utc_now_iso(),
                "duration_seconds": round(duration, 3),
            }
            print(
                f"  完成: {len(bars)} 行 | {source_mode} | {duration:.1f}s",
                flush=True,
            )
        except Exception as exc:
            duration = time.perf_counter() - started
            row = {
                "policy_id": RAW_PRICE_CACHE_POLICY_ID,
                "ticker": ticker,
                "requested_begin": begin,
                "requested_end": end,
                "source_chain": source_chain_key,
                "status": "failed",
                "source": "",
                "source_request_mode": "",
                "rows": 0,
                "data_start": "",
                "data_end": "",
                "version_id": "",
                "content_sha256": "",
                "error": f"{type(exc).__name__}: {exc}",
                "attempted_at": utc_now_iso(),
                "duration_seconds": round(duration, 3),
            }
            print(f"  失败: {row['error']}", flush=True)
        attempts = _upsert_attempt(attempts, row)
        _atomic_csv(attempts, attempts_path)
        durations.append(duration)
        remaining = len(selected) - index
        eta = (sum(durations) / len(durations)) * remaining if durations else 0
        print(f"  本次剩余 ETA: {eta / 60:.1f}m", flush=True)

    raw_series, raw_manifests, cache_errors = _cache_inventory(
        raw_cache,
        tickers,
        begin=begin,
        end=end,
    )
    matching = _matching_attempts(
        attempts,
        begin=begin,
        end=end,
        source_chain=source_chain_key,
    )
    attempted_tickers = set(raw_series).union(
        matching.loc[matching["status"].eq("failed"), "ticker"].astype(str)
    )
    audit_complete = set(tickers).issubset(attempted_tickers)

    local_cfg = _research_cfg(cfg, True)
    adjusted_series, _adjusted_meta, adjusted_errors = _load_series(
        local_cfg,
        logical_tickers,
    )
    adjusted_version = _price_universe_version(adjusted_series)
    coverage = point_in_time_raw_price_rows(
        identity_plan,
        raw_series,
        raw_manifests,
        adjusted_series=adjusted_series,
        policy=policy,
    )
    split_events = raw_split_events(raw_series, policy=policy)
    share_events_path = fundamental_dir / "valuation_readiness_v1" / "share_basis_events.csv"
    share_events_manifest = None
    if share_events_path.exists() and manifest_path(share_events_path).exists():
        share_events, share_events_manifest = _verified_csv(share_events_path)
        crosscheck = crosscheck_share_split_events(
            share_events,
            split_events,
            policy=policy,
        )
        # A current-only raw-price window cannot disprove an older share-count
        # event.  Keep that event visible, but do not turn absent out-of-window
        # bars into a false basis failure.
        prior_dates = pd.to_datetime(
            crosscheck["prior_selection_date"], errors="coerce"
        )
        interval_begin = pd.Timestamp(begin)
        outside_raw_interval = prior_dates.lt(interval_begin)
        crosscheck.loc[
            outside_raw_interval,
            "crosscheck_status",
        ] = "raw_interval_does_not_cover_event"
        unavailable = ~crosscheck["ticker"].isin(raw_series)
        crosscheck.loc[unavailable, "crosscheck_status"] = "raw_ticker_unavailable"
    else:
        crosscheck = pd.DataFrame(columns=[
            "ticker", "prior_selection_date", "selection_date", "share_split_like",
            "expected_split_factor", "matched_raw_event_date",
            "matched_raw_split_factor", "factor_relative_error", "crosscheck_status",
        ])
    failed_basis_tickers = set(
        crosscheck.loc[
            crosscheck["share_split_like"].fillna(False).astype(bool)
            & crosscheck["crosscheck_status"].isin(
                ["expected_raw_split_missing", "raw_split_factor_mismatch"]
            ),
            "ticker",
        ].astype(str)
    )
    if failed_basis_tickers:
        mask = coverage["ticker"].isin(failed_basis_tickers)
        coverage.loc[mask, "price_basis_ready"] = False
        coverage.loc[mask, "price_ready"] = False
        coverage.loc[mask, "price_status"] = "basis_crosscheck_failed"
    by_date = _coverage_by_date(coverage)
    sources = _source_summary(tickers, matching, raw_manifests, coverage)
    failures = _integrity_failures(plan, coverage, crosscheck)
    if adjusted_errors:
        failures = pd.concat([
            failures,
            pd.DataFrame([
                {
                    "scope": "adjusted_reference",
                    "ticker": ticker,
                    "selection_date": "",
                    "reason": error,
                }
                for ticker, error in sorted(adjusted_errors.items())
            ]),
        ], ignore_index=True)
    versions = pd.DataFrame([
        {
            "source_name": f"raw_price:{ticker}",
            "version_id": manifest.get("version_id"),
            "content_sha256": manifest.get("content_sha256"),
            "rows": manifest.get("row_count"),
            "note": (
                f"{manifest.get('source_request_mode')}；未复权名义成交价；"
                "不覆盖收益回测缓存"
            ),
        }
        for ticker, manifest in sorted(raw_manifests.items())
    ] + [
        {
            "source_name": "adjusted_return_price_reference",
            "version_id": adjusted_version.get("version_id"),
            "content_sha256": adjusted_version.get("content_sha256"),
            "rows": adjusted_version.get("row_count"),
            "note": "仅用于同日比值诊断，不参与市值乘法",
        },
        {
            "source_name": "share_basis_events",
            "version_id": (share_events_manifest or {}).get("version_id"),
            "content_sha256": (share_events_manifest or {}).get("content_sha256"),
            "rows": (share_events_manifest or {}).get("row_count"),
            "note": "用于拆股方向交叉核验，不计算估值分数",
        },
        {
            "source_name": "security_ticker_lifecycle",
            "version_id": lifecycle.mapping_id,
            "content_sha256": file_sha256(lifecycle_path),
            "rows": len(lifecycle.rules_by_cik),
            "note": "同一CIK+CUSIP的代码变更生效区间；禁止映射收购方",
        },
    ])
    fetch_plan = pd.DataFrame([
        {
            "ticker": ticker,
            "requested_begin": begin,
            "requested_end": end,
            "first_selection_date": identity_plan.loc[
                identity_plan["price_ticker"].eq(ticker), "selection_date"
            ].min(),
            "last_selection_date": identity_plan.loc[
                identity_plan["price_ticker"].eq(ticker), "selection_date"
            ].max(),
            "status": (
                "ready" if ticker in raw_series
                else "failed" if ticker in failed_checkpoints or ticker in set(
                    matching.loc[matching["status"].eq("failed"), "ticker"]
                )
                else "pending"
            ),
            "note": cache_errors.get(ticker, ""),
        }
        for ticker in tickers
    ])

    frames = {
        "attempts": matching,
        "security_identity": identity_plan[[
            "selection_date",
            "ticker",
            "cik",
            "price_ticker",
            "ticker_alias_applied",
            "ticker_lifecycle_rule_id",
            "ticker_lifecycle_status",
            "security_cusip",
            "ticker_lifecycle_evidence_url",
        ]],
        "fetch_plan": fetch_plan,
        "coverage": coverage,
        "coverage_by_date": by_date,
        "sources": sources,
        "split_events": split_events,
        "share_split_crosscheck": crosscheck,
        "integrity_failures": failures,
        "versions": versions,
    }
    paths = {key: output / f"{key}.csv" for key in frames}
    evidence = {}
    for key, frame in frames.items():
        _atomic_csv(frame, paths[key])
        manifest = write_dataset_manifest(
            paths[key],
            dataset_type="ndx_price_basis_readiness",
            dataset_key=f"{PRICE_BASIS_POLICY_ID}:{key}",
            source="derived_from_versioned_raw_price_and_point_in_time_plan",
            fetched_at_evidence="derived_runtime_utc",
            row_count=len(frame),
            data_start=(
                str(coverage["selection_date"].min())
                if "selection_date" in frame.columns and not frame.empty else None
            ),
            data_end=(
                str(coverage["selection_date"].max())
                if "selection_date" in frame.columns and not frame.empty else None
            ),
            extra={
                "policy_id": PRICE_BASIS_POLICY_ID,
                "research_only": True,
                "audit_complete": audit_complete,
                "valuation_scoring_allowed": False,
                "security_lifecycle_mapping_id": lifecycle.mapping_id,
                "security_lifecycle_sha256": file_sha256(lifecycle_path),
            },
        )
        evidence[key] = {
            "filename": paths[key].name,
            "rows": len(frame),
            "sha256": manifest["content_sha256"],
            "version_id": manifest["version_id"],
        }

    integrity_passed = failures.empty
    consumption_allowed = bool(audit_complete and integrity_passed)
    minimum_coverage = (
        float(by_date["raw_price_ready_coverage"].min()) if not by_date.empty else 0.0
    )
    latest_coverage = (
        float(by_date.sort_values("selection_date").iloc[-1]["raw_price_ready_coverage"])
        if not by_date.empty else 0.0
    )
    summary = {
        "schema_version": 1,
        "generated_at": utc_now_iso(),
        "policy_id": PRICE_BASIS_POLICY_ID,
        "policy": policy.__dict__,
        "research_only": True,
        "plan_path": plan_path.as_posix(),
        "plan_sha256": file_sha256(plan_path),
        "selection_rows": len(plan),
        "selection_dates": int(plan["selection_date"].nunique()),
        "ticker_count": len(logical_tickers),
        "price_ticker_count": len(tickers),
        "ticker_alias_selection_rows": int(
            identity_plan["ticker_alias_applied"].sum()
        ),
        "security_lifecycle": {
            "mapping_id": lifecycle.mapping_id,
            "last_reviewed": lifecycle.last_reviewed,
            "path": lifecycle_path.as_posix(),
            "content_sha256": file_sha256(lifecycle_path),
        },
        "raw_cache_dir": raw_cache.as_posix(),
        "requested_begin": begin,
        "requested_end": end,
        "source_chain": source_chain,
        "ready_ticker_count": len(raw_series),
        "failed_ticker_count": int(sources["fetch_status"].eq("failed").sum()),
        "pending_ticker_count": int(sources["fetch_status"].eq("pending").sum()),
        "raw_price_row_count": int(sum(len(frame) for frame in raw_series.values())),
        "minimum_raw_price_ready_coverage": minimum_coverage,
        "latest_raw_price_ready_coverage": latest_coverage,
        "raw_split_event_count": len(split_events),
        "share_split_crosscheck_count": len(crosscheck),
        "share_split_matched_count": int(
            crosscheck["crosscheck_status"].eq("matched").sum()
            if not crosscheck.empty else 0
        ),
        "integrity_failure_count": len(failures),
        "integrity_passed": integrity_passed,
        "audit_complete": audit_complete,
        "consumption_allowed": consumption_allowed,
        "return_engine_cache_modified": False,
        "valuation_scoring_allowed": False,
        "files": {
            **{key: value["filename"] for key, value in evidence.items()},
            "workbook": "ndx_price_basis_readiness.xlsx",
        },
    }
    atomic_write_text(
        output / "summary.json",
        json.dumps(summary, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
    )
    atomic_write_text(
        output / "manifest.json",
        json.dumps(
            {
                "schema_version": 1,
                "generated_at": summary["generated_at"],
                "policy_id": PRICE_BASIS_POLICY_ID,
                "plan_sha256": summary["plan_sha256"],
                "requested_begin": begin,
                "requested_end": end,
                "source_chain": source_chain,
                "security_lifecycle": summary["security_lifecycle"],
                "raw_cache_dir": summary["raw_cache_dir"],
                "audit_complete": audit_complete,
                "integrity_passed": integrity_passed,
                "consumption_allowed": consumption_allowed,
                "evidence_files": evidence,
                "valuation_scoring_allowed": False,
            },
            ensure_ascii=False,
            indent=2,
            allow_nan=False,
        ) + "\n",
    )
    conclusion = pd.DataFrame([
        {"项目": "研究目的", "内容": "建立SEC申报股数可乘用的点时未复权价格证据"},
        {"项目": "策略代码数", "内容": len(logical_tickers)},
        {"项目": "实际取价代码数", "内容": len(tickers)},
        {
            "项目": "生命周期别名行",
            "内容": int(identity_plan["ticker_alias_applied"].sum()),
        },
        {"项目": "生命周期版本", "内容": lifecycle.mapping_id},
        {"项目": "点时候选行", "内容": len(plan)},
        {"项目": "有效未复权缓存", "内容": len(raw_series)},
        {"项目": "最低原始价就绪率", "内容": minimum_coverage},
        {"项目": "最新原始价就绪率", "内容": latest_coverage},
        {"项目": "原始价拆股事件", "内容": len(split_events)},
        {"项目": "股数拆股匹配", "内容": summary["share_split_matched_count"]},
        {"项目": "完整性失败", "内容": len(failures)},
        {"项目": "可供估值审计消费", "内容": consumption_allowed},
        {"项目": "收益回测缓存", "内容": "未修改；原始价使用独立目录"},
        {"项目": "估值评分/实盘", "内容": "仍关闭；本工具只打通价格口径"},
    ])
    workbook = output / "ndx_price_basis_readiness.xlsx"
    _write_workbook(
        workbook,
        conclusion=conclusion,
        identity=frames["security_identity"],
        fetch_plan=fetch_plan,
        coverage=coverage,
        by_date=by_date,
        sources=sources,
        split_events=split_events,
        crosscheck=crosscheck,
        failures=failures,
        versions=versions,
    )
    print("\n未复权价格口径证据包")
    print(
        f"有效股票: {len(raw_series)}/{len(tickers)} | 点时覆盖最低/最新: "
        f"{minimum_coverage:.2%}/{latest_coverage:.2%}"
    )
    print(
        f"拆股事件: {len(split_events)} | 股数匹配: "
        f"{summary['share_split_matched_count']} | 完整性失败: {len(failures)}"
    )
    print(f"估值审计可消费: {consumption_allowed} | 评分仍关闭")
    print(f"Workbook: {workbook}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
