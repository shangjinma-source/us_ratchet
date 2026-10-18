"""Audit point-in-time valuation inputs without calculating valuation scores."""
from __future__ import annotations

import argparse
import json
import os
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from backtest.ndx_competition import MULTI_SECURITY_ISSUER, _price_universe_version
from backtest.preference_walkforward import _load_series, _research_cfg
from core.config import load_config
from core.helpers import atomic_write_text
from data.ndx_filing_identity import FILING_SHARE_RESOLUTION_POLICY_ID
from data.ndx_fundamentals import fundamental_cache_dir
from data.ndx_price_basis import PRICE_BASIS_POLICY_ID, RAW_PRICE_BASIS
from data.ndx_valuation_readiness import (
    DEBT_CONCEPT_POLICY_ID,
    PROMOTED_DEBT_CONCEPTS,
    VALUATION_READINESS_POLICY_ID,
    ValuationReadinessPolicy,
    build_point_in_time_accounting_readiness,
    build_valuation_factor_readiness,
    point_in_time_price_rows,
    share_basis_events,
)
from data.versioning import (
    file_sha256,
    manifest_path,
    utc_now_iso,
    write_dataset_manifest,
)


def _resolve(path: str | Path) -> Path:
    value = Path(path)
    return value if value.is_absolute() else ROOT / value


def _atomic_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        frame.to_csv(temporary, index=False, encoding="utf-8-sig")
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def _read_manifest(path: Path) -> dict:
    if not path.exists():
        raise FileNotFoundError(path)
    return json.loads(path.read_text(encoding="utf-8"))


def _verified_csv(path: Path) -> tuple[pd.DataFrame, dict]:
    sidecar = manifest_path(path)
    manifest = _read_manifest(sidecar)
    if manifest.get("content_sha256") != file_sha256(path):
        raise ValueError(f"versioned input hash mismatch: {path}")
    return pd.read_csv(path, low_memory=False), manifest


def _verified_price_basis(
    directory: Path,
    *,
    plan_path: Path,
    plan: pd.DataFrame,
) -> tuple[pd.DataFrame, dict, dict]:
    root_manifest = _read_manifest(directory / "manifest.json")
    checks = {
        "policy_id": root_manifest.get("policy_id") == PRICE_BASIS_POLICY_ID,
        "plan_sha256": root_manifest.get("plan_sha256") == file_sha256(plan_path),
        "audit_complete": bool(root_manifest.get("audit_complete")),
        "integrity_passed": bool(root_manifest.get("integrity_passed")),
        "consumption_allowed": bool(root_manifest.get("consumption_allowed")),
    }
    failed = sorted(key for key, passed in checks.items() if not passed)
    if failed:
        raise ValueError(f"price basis manifest failed checks: {failed}")
    coverage_path = directory / "coverage.csv"
    coverage, coverage_manifest = _verified_csv(coverage_path)
    declared = (root_manifest.get("evidence_files") or {}).get("coverage") or {}
    if declared.get("sha256") != coverage_manifest.get("content_sha256"):
        raise ValueError("price basis coverage hash differs from root manifest")
    required = {
        "selection_date",
        "ticker",
        "price_status",
        "price_ready",
        "price_basis_ready",
        "price",
        "price_date",
        "price_age_days",
        "price_future_data_guard_passed",
        "historical_price_basis",
    }
    if not required.issubset(coverage.columns):
        raise ValueError(
            f"price basis coverage lacks {sorted(required - set(coverage.columns))}"
        )
    coverage["selection_date"] = coverage["selection_date"].astype(str).str[:10]
    coverage["ticker"] = coverage["ticker"].astype(str).str.upper()
    expected = plan[["selection_date", "ticker"]].copy()
    expected["selection_date"] = expected["selection_date"].astype(str).str[:10]
    expected["ticker"] = expected["ticker"].astype(str).str.upper()
    observed_keys = set(map(tuple, coverage[["selection_date", "ticker"]].to_numpy()))
    expected_keys = set(map(tuple, expected.to_numpy()))
    if len(coverage) != len(expected) or observed_keys != expected_keys:
        raise ValueError("price basis coverage does not match the valuation plan")
    if coverage.duplicated(["selection_date", "ticker"]).any():
        raise ValueError("price basis coverage contains duplicate keys")
    if not coverage.loc[
        coverage["price_basis_ready"].fillna(False).astype(bool),
        "historical_price_basis",
    ].eq(RAW_PRICE_BASIS).all():
        raise ValueError("ready price basis rows are not declared unadjusted")
    return coverage, coverage_manifest, root_manifest


def _verified_debt_concepts(
    directory: Path,
    *,
    plan_path: Path,
    base_facts_manifest: dict,
) -> tuple[pd.DataFrame, dict, dict]:
    root_manifest = _read_manifest(directory / "manifest.json")
    checks = {
        "policy_id": root_manifest.get("policy_id") == DEBT_CONCEPT_POLICY_ID,
        "plan_sha256": root_manifest.get("plan_sha256") == file_sha256(plan_path),
        "base_facts_sha256": (
            root_manifest.get("base_facts_sha256")
            == base_facts_manifest.get("content_sha256")
        ),
        "integrity_passed": bool(root_manifest.get("integrity_passed")),
        "promotion_allowed": bool(root_manifest.get("promotion_allowed")),
        "promoted_concepts": (
            set(root_manifest.get("promoted_concepts") or [])
            == set(PROMOTED_DEBT_CONCEPTS)
        ),
    }
    failed = sorted(key for key, passed in checks.items() if not passed)
    if failed:
        raise ValueError(f"debt concept manifest failed checks: {failed}")
    facts_path = directory / "expanded_debt_facts.csv"
    facts, facts_manifest = _verified_csv(facts_path)
    declared = (root_manifest.get("evidence_files") or {}).get(
        "expanded_debt_facts"
    ) or {}
    if declared.get("sha256") != facts_manifest.get("content_sha256"):
        raise ValueError("expanded debt fact hash differs from root manifest")
    required = {
        "cik",
        "metric",
        "namespace",
        "concept",
        "unit",
        "value",
        "form",
        "available_date",
        "period_end",
        "accession_number",
    }
    if not required.issubset(facts.columns):
        raise ValueError(
            f"expanded debt facts lack {sorted(required - set(facts.columns))}"
        )
    promoted = facts[facts["concept"].isin(PROMOTED_DEBT_CONCEPTS)].copy()
    if promoted.empty:
        raise ValueError("expanded debt evidence contains no promoted facts")
    return promoted, facts_manifest, root_manifest


def _merge_promoted_debt_facts(
    base: pd.DataFrame,
    promoted: pd.DataFrame,
) -> pd.DataFrame:
    combined = pd.concat([base, promoted], ignore_index=True)
    key = [
        "cik",
        "metric",
        "namespace",
        "concept",
        "unit",
        "value",
        "form",
        "available_date",
        "period_start",
        "period_end",
        "accession_number",
        "frame",
    ]
    existing = [column for column in key if column in combined.columns]
    return combined.drop_duplicates(existing, keep="last").reset_index(drop=True)


def _verified_share_promotions(directory: Path) -> tuple[pd.DataFrame, dict, dict]:
    root = _read_manifest(directory / "manifest.json")
    checks = {
        "policy_id": root.get("policy_id") == FILING_SHARE_RESOLUTION_POLICY_ID,
        "integrity_passed": bool(root.get("integrity_passed")),
        "promotion_allowed": bool(root.get("promotion_allowed")),
        "valuation_scoring_closed": root.get("valuation_scoring_allowed") is False,
    }
    failed = sorted(key for key, passed in checks.items() if not passed)
    if failed:
        raise ValueError(f"share identity manifest failed checks: {failed}")
    evidence = (root.get("evidence_files") or {}).get("promotions") or {}
    path = directory / str(evidence.get("filename") or "")
    promotions, manifest = _verified_csv(path)
    if evidence.get("sha256") != manifest.get("content_sha256"):
        raise ValueError("share identity promotion hash differs from root manifest")
    required = {
        "selection_date",
        "ticker",
        "cik",
        "existing_shares_ready",
        "resolved_shares_value",
        "source_accession_number",
        "source_available_date",
        "shares_period_end",
        "shares_age_days",
        "shares_age_passed",
        "share_concept",
        "share_resolution_method",
        "common_class_members",
        "foreign_filer_asof",
        "multi_security_issuer",
        "promotion_allowed",
        "policy_id",
    }
    if not required.issubset(promotions.columns):
        raise ValueError(
            f"share promotions lack {sorted(required - set(promotions.columns))}"
        )
    bool_columns = (
        "existing_shares_ready",
        "shares_age_passed",
        "foreign_filer_asof",
        "multi_security_issuer",
        "promotion_allowed",
    )
    for column in bool_columns:
        if not pd.api.types.is_bool_dtype(promotions[column]):
            promotions[column] = (
                promotions[column].astype(str).str.lower().eq("true")
            )
    invalid = promotions[
        promotions["existing_shares_ready"]
        | ~promotions["shares_age_passed"]
        | promotions["foreign_filer_asof"]
        | promotions["multi_security_issuer"]
        | ~promotions["promotion_allowed"]
        | ~promotions["policy_id"].eq(FILING_SHARE_RESOLUTION_POLICY_ID)
    ]
    if not invalid.empty:
        raise ValueError("share identity promotions contain blocked rows")
    if promotions.duplicated(["selection_date", "ticker"]).any():
        raise ValueError("share identity promotions contain duplicate keys")
    return promotions, manifest, root


def _apply_share_promotions(
    accounting: pd.DataFrame,
    promotions: pd.DataFrame,
) -> pd.DataFrame:
    keys = ["selection_date", "ticker", "cik"]
    frame = accounting.copy()
    frame["selection_date"] = frame["selection_date"].astype(str).str[:10]
    frame["ticker"] = frame["ticker"].astype(str).str.upper()
    selected = promotions.copy()
    selected["selection_date"] = selected["selection_date"].astype(str).str[:10]
    selected["ticker"] = selected["ticker"].astype(str).str.upper()
    columns = keys + [
        "resolved_shares_value",
        "source_accession_number",
        "source_available_date",
        "shares_period_end",
        "shares_age_days",
        "share_concept",
        "share_resolution_method",
        "common_class_members",
        "policy_id",
    ]
    selected = selected[columns].rename(columns={
        column: f"promotion_{column}" for column in columns if column not in keys
    })
    merged = frame.merge(selected, on=keys, how="left", validate="one_to_one")
    promoted = merged["promotion_resolved_shares_value"].notna()
    if merged.loc[promoted, "shares_ready"].fillna(False).astype(bool).any():
        raise ValueError("share identity attempted to overwrite ready Company Facts")
    merged["shares_identity_promoted"] = promoted
    merged["shares_identity_policy_id"] = merged["promotion_policy_id"].fillna("")
    merged["shares_identity_resolution_method"] = merged[
        "promotion_share_resolution_method"
    ].fillna("")
    merged["shares_identity_accession_number"] = merged[
        "promotion_source_accession_number"
    ].fillna("")
    merged["shares_identity_common_class_members"] = merged[
        "promotion_common_class_members"
    ].fillna("")
    merged.loc[promoted, "shares_status"] = "ready_filing_xbrl_identity_v1"
    merged.loc[promoted, "shares_ready"] = True
    merged.loc[promoted, "shares_value"] = merged.loc[
        promoted, "promotion_resolved_shares_value"
    ]
    merged.loc[promoted, "shares_unit"] = "shares"
    merged.loc[promoted, "shares_namespace"] = "sec-filing-xbrl"
    merged.loc[promoted, "shares_concept"] = merged.loc[
        promoted, "promotion_share_concept"
    ]
    merged.loc[promoted, "shares_period_end"] = merged.loc[
        promoted, "promotion_shares_period_end"
    ]
    merged.loc[promoted, "shares_age_days"] = merged.loc[
        promoted, "promotion_shares_age_days"
    ]
    merged.loc[promoted, "shares_source_available_date"] = merged.loc[
        promoted, "promotion_source_available_date"
    ]
    merged.loc[promoted, "shares_same_date_candidate_count"] = 1
    merged.loc[promoted, "shares_same_date_value_spread"] = 0.0
    merged.loc[promoted, "shares_future_data_guard_passed"] = True
    promotion_columns = [
        column for column in merged.columns if column.startswith("promotion_")
    ]
    return merged.drop(columns=promotion_columns)


def _coverage_by_date(rows: pd.DataFrame) -> pd.DataFrame:
    fields = (
        "price_ready",
        "shares_ready",
        "debt_ready",
        "debt_strict_complete",
        "fcf_yield_accounting_ready",
        "ev_sales_accounting_ready",
        "fcf_yield_ready",
        "ev_sales_ready",
    )
    output = []
    for selection_date, group in rows.groupby("selection_date", sort=True):
        item = {
            "selection_date": selection_date,
            "universe_count": len(group),
            "usd_statement_count": int(group["statement_unit"].eq("USD").sum()),
            "foreign_filer_count": int(group["foreign_filer_asof"].sum()),
            "multi_security_issuer_count": int(group["multi_security_issuer"].sum()),
        }
        for field in fields:
            count = int(group[field].fillna(False).astype(bool).sum())
            item[f"{field}_count"] = count
            item[f"{field}_coverage"] = count / len(group) if len(group) else None
        output.append(item)
    return pd.DataFrame(output)


def _blocker_summary(rows: pd.DataFrame) -> pd.DataFrame:
    output = []
    for factor, column in (
        ("fcf_yield", "fcf_yield_blockers"),
        ("ev_sales", "ev_sales_blockers"),
    ):
        exploded = (
            rows[["selection_date", "ticker", column]]
            .assign(blocker=lambda frame: frame[column].fillna("").str.split(";"))
            .explode("blocker")
        )
        exploded = exploded[exploded["blocker"].astype(str).str.len() > 0]
        grouped = exploded.groupby("blocker").agg(
            affected_rows=("ticker", "size"),
            affected_tickers=("ticker", "nunique"),
            first_selection_date=("selection_date", "min"),
            last_selection_date=("selection_date", "max"),
        ).reset_index()
        grouped.insert(0, "factor", factor)
        grouped["affected_row_rate"] = grouped["affected_rows"] / len(rows)
        output.append(grouped)
    return pd.concat(output, ignore_index=True).sort_values(
        ["factor", "affected_rows", "blocker"],
        ascending=[True, False, True],
        kind="stable",
    )


def _method_summary(rows: pd.DataFrame) -> pd.DataFrame:
    share = rows.groupby(
        ["shares_status", "shares_namespace", "shares_concept"],
        dropna=False,
    ).size().reset_index(name="rows")
    share.insert(0, "metric", "shares_outstanding")
    share = share.rename(columns={
        "shares_status": "status",
        "shares_namespace": "namespace",
        "shares_concept": "method_or_concept",
    })
    share["completeness"] = None
    debt = rows.groupby(
        ["debt_status", "debt_method", "debt_completeness"],
        dropna=False,
    ).size().reset_index(name="rows")
    debt.insert(0, "metric", "debt")
    debt = debt.rename(columns={
        "debt_status": "status",
        "debt_method": "method_or_concept",
        "debt_completeness": "completeness",
    })
    debt["namespace"] = None
    return pd.concat(
        [
            share[["metric", "status", "namespace", "method_or_concept", "completeness", "rows"]],
            debt[["metric", "status", "namespace", "method_or_concept", "completeness", "rows"]],
        ],
        ignore_index=True,
    ).sort_values(["metric", "rows"], ascending=[True, False], kind="stable")


def _integrity_failures(rows: pd.DataFrame) -> pd.DataFrame:
    output = []
    duplicate = rows.duplicated(["selection_date", "ticker"], keep=False)
    for row in rows[duplicate].itertuples(index=False):
        output.append({
            "failure_type": "duplicate_selection_ticker",
            "selection_date": row.selection_date,
            "ticker": row.ticker,
            "detail": "valuation readiness identity is duplicated",
        })
    for field in (
        "shares_future_data_guard_passed",
        "debt_future_data_guard_passed",
        "price_future_data_guard_passed",
    ):
        failed = ~rows[field].fillna(False).astype(bool)
        for row in rows[failed].itertuples(index=False):
            output.append({
                "failure_type": field,
                "selection_date": row.selection_date,
                "ticker": row.ticker,
                "detail": "point-in-time guard failed",
            })
    unexpected = rows[
        rows["valuation_scoring_allowed"].fillna(False).astype(bool)
    ]
    for row in unexpected.itertuples(index=False):
        output.append({
            "failure_type": "unexpected_factor_enablement",
            "selection_date": row.selection_date,
            "ticker": row.ticker,
            "detail": "valuation scoring was enabled by a readiness audit",
        })
    return pd.DataFrame(
        output,
        columns=["failure_type", "selection_date", "ticker", "detail"],
    )


READINESS_COLUMNS_CN = {
    "selection_date": "重选日期",
    "ticker": "股票",
    "cik": "SEC_CIK",
    "statement_unit": "报表币种",
    "foreign_filer_asof": "当时是否境外申报人",
    "multi_security_issuer": "是否多证券发行人",
    "price_status": "价格状态",
    "price": "严格前一交易日价格",
    "price_date": "价格日期",
    "price_age_days": "价格滞后天数",
    "price_basis_ready": "价格股数基准是否就绪",
    "shares_status": "股数状态",
    "shares_value": "申报流通股数",
    "shares_concept": "股数概念",
    "shares_period_end": "股数基准日",
    "shares_age_days": "股数账龄天数",
    "shares_same_date_value_spread": "同日股数差异率",
    "shares_identity_promoted": "filing股数已安全提升",
    "shares_identity_policy_id": "filing股数规则编号",
    "shares_identity_resolution_method": "filing股数解析方法",
    "shares_identity_accession_number": "filing股数申报编号",
    "shares_identity_common_class_members": "filing普通股类成员",
    "debt_status": "债务状态",
    "debt_value": "已组装债务",
    "debt_unit": "债务币种",
    "debt_method": "债务组装方法",
    "debt_completeness": "债务完整性",
    "debt_period_end": "债务基准日",
    "debt_age_days": "债务账龄天数",
    "debt_strict_complete": "总债务严格完整",
    "free_cash_flow_value": "自由现金流金额",
    "free_cash_flow_unit": "自由现金流币种",
    "fcf_yield_accounting_ready": "FCF收益率会计口径就绪",
    "ev_sales_accounting_ready": "EV销售额会计口径就绪",
    "fcf_yield_ready": "FCF收益率最终就绪",
    "ev_sales_ready": "EV销售额最终就绪",
    "fcf_yield_blockers": "FCF收益率阻断原因",
    "ev_sales_blockers": "EV销售额阻断原因",
    "valuation_scoring_allowed": "允许估值评分",
}


def _localized_readiness(rows: pd.DataFrame) -> pd.DataFrame:
    columns = [name for name in READINESS_COLUMNS_CN if name in rows.columns]
    return rows[columns].rename(columns=READINESS_COLUMNS_CN)


def _rename(frame: pd.DataFrame, columns: dict[str, str]) -> pd.DataFrame:
    selected = [name for name in columns if name in frame.columns]
    return frame[selected].rename(columns=columns)


def _format_workbook(path: Path) -> None:
    book = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="DCE6F1")
    for sheet in book.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.zoomScale = 90
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column_index, cells in enumerate(sheet.columns, start=1):
            values = [str(cell.value or "") for cell in list(cells)[:200]]
            width = min(52, max(10, max((len(value) for value in values), default=0) + 2))
            sheet.column_dimensions[get_column_letter(column_index)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    temporary = path.with_name(path.stem + ".formatted.xlsx")
    book.save(temporary)
    book.close()
    os.replace(temporary, path)


def _write_workbook(
    path: Path,
    *,
    conclusion: pd.DataFrame,
    rows: pd.DataFrame,
    coverage: pd.DataFrame,
    blockers: pd.DataFrame,
    methods: pd.DataFrame,
    events: pd.DataFrame,
    components: pd.DataFrame,
    versions: pd.DataFrame,
    failures: pd.DataFrame,
) -> None:
    coverage_cn = _rename(coverage, {
        "selection_date": "重选日期",
        "universe_count": "候选数",
        "usd_statement_count": "美元报表数",
        "foreign_filer_count": "境外申报人数",
        "multi_security_issuer_count": "多证券发行人数",
        "price_ready_count": "价格可用数",
        "price_ready_coverage": "价格覆盖率",
        "shares_ready_count": "股数就绪数",
        "shares_ready_coverage": "股数覆盖率",
        "debt_ready_count": "债务可组装数",
        "debt_ready_coverage": "债务可组装率",
        "debt_strict_complete_count": "总债务严格完整数",
        "debt_strict_complete_coverage": "总债务严格完整率",
        "fcf_yield_accounting_ready_count": "FCF收益率会计就绪数",
        "fcf_yield_accounting_ready_coverage": "FCF收益率会计就绪率",
        "ev_sales_accounting_ready_count": "EV销售额会计就绪数",
        "ev_sales_accounting_ready_coverage": "EV销售额会计就绪率",
        "fcf_yield_ready_count": "FCF收益率最终就绪数",
        "fcf_yield_ready_coverage": "FCF收益率最终就绪率",
        "ev_sales_ready_count": "EV销售额最终就绪数",
        "ev_sales_ready_coverage": "EV销售额最终就绪率",
    })
    blockers_cn = _rename(blockers, {
        "factor": "估值因子",
        "blocker": "阻断原因",
        "affected_rows": "影响行数",
        "affected_tickers": "影响股票数",
        "first_selection_date": "最早重选日",
        "last_selection_date": "最晚重选日",
        "affected_row_rate": "影响行比例",
    })
    methods_cn = _rename(methods, {
        "metric": "指标",
        "status": "状态",
        "namespace": "命名空间",
        "method_or_concept": "方法或概念",
        "completeness": "完整性",
        "rows": "行数",
    })
    events_cn = _rename(events, {
        "ticker": "股票",
        "prior_selection_date": "前次重选日",
        "selection_date": "本次重选日",
        "prior_shares": "前次股数",
        "shares": "本次股数",
        "raw_ratio": "原始倍数",
        "ratio_magnitude": "变化倍数",
        "nearest_split_ratio": "最接近拆股倍数",
        "split_ratio_relative_error": "拆股倍数误差",
        "split_like": "是否类似拆股",
    })
    components_cn = _rename(components, {
        "selection_date": "重选日期",
        "ticker": "股票",
        "cik": "SEC_CIK",
        "metric": "指标",
        "component_role": "来源角色",
        "selected": "是否采用",
        "namespace": "命名空间",
        "concept": "概念",
        "unit": "单位",
        "value": "数值",
        "form": "申报表",
        "available_date": "可用日期",
        "period_end": "期间结束日",
        "accession_number": "申报编号",
    })
    versions_cn = _rename(versions, {
        "source_name": "数据源",
        "version_id": "版本ID",
        "content_sha256": "内容哈希",
        "rows": "行数",
        "note": "说明",
    })
    failures_cn = _rename(failures, {
        "failure_type": "失败类型",
        "selection_date": "重选日期",
        "ticker": "股票",
        "detail": "详情",
    })
    temporary = path.with_name(path.stem + ".tmp.xlsx")
    try:
        with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
            conclusion.to_excel(writer, sheet_name="结论", index=False)
            _localized_readiness(rows).to_excel(
                writer, sheet_name="逐期可计算性", index=False
            )
            coverage_cn.to_excel(writer, sheet_name="日期覆盖率", index=False)
            blockers_cn.to_excel(writer, sheet_name="阻断原因", index=False)
            methods_cn.to_excel(writer, sheet_name="会计方法", index=False)
            events_cn.to_excel(writer, sheet_name="股数跳变", index=False)
            components_cn[components_cn["指标"].eq("debt")].to_excel(
                writer, sheet_name="债务来源", index=False
            )
            components_cn[
                components_cn["指标"].eq("shares_outstanding")
            ].to_excel(writer, sheet_name="股数来源", index=False)
            versions_cn.to_excel(writer, sheet_name="数据版本", index=False)
            failures_cn.to_excel(writer, sheet_name="完整性失败", index=False)
        os.replace(temporary, path)
        _format_workbook(path)
    finally:
        if temporary.exists():
            temporary.unlink()


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Audit point-in-time NDX valuation inputs; never calculate scores "
            "or modify the live strategy."
        )
    )
    parser.add_argument("--config", default="config/config.yaml")
    parser.add_argument("--input-dir", default="")
    parser.add_argument("--normalization-dir", default="")
    parser.add_argument("--output-dir", default="")
    parser.add_argument("--price-basis-dir", default="")
    parser.add_argument("--debt-concept-dir", default="")
    parser.add_argument("--share-identity-dir", default="")
    parser.add_argument("--maximum-shares-age-days", type=int, default=200)
    parser.add_argument("--maximum-debt-age-days", type=int, default=200)
    args = parser.parse_args()

    cfg = load_config(_resolve(args.config))
    source = (
        _resolve(args.input_dir)
        if args.input_dir
        else fundamental_cache_dir(cfg)
    )
    normalization = (
        _resolve(args.normalization_dir)
        if args.normalization_dir
        else source / "normalization_v1"
    )
    output = (
        _resolve(args.output_dir)
        if args.output_dir
        else source / "valuation_readiness_v4"
    )
    output.mkdir(parents=True, exist_ok=True)

    facts, facts_manifest = _verified_csv(source / "fact_availability.csv")
    factor_inputs, factors_manifest = _verified_csv(
        normalization / "quality_factor_inputs.csv"
    )
    normalization_manifest_path = normalization / "manifest.json"
    normalization_manifest = _read_manifest(normalization_manifest_path)
    plan_path = source / "plan.csv"
    if file_sha256(plan_path) != normalization_manifest.get("plan_sha256"):
        raise ValueError("normalization plan hash changed")
    plan = pd.read_csv(plan_path, low_memory=False)
    debt_concept_directory = (
        _resolve(args.debt_concept_dir)
        if args.debt_concept_dir
        else source / "debt_concepts_v2"
    )
    debt_concept_manifest = None
    debt_facts_manifest = None
    if (debt_concept_directory / "manifest.json").exists():
        promoted_debt, debt_facts_manifest, debt_concept_manifest = (
            _verified_debt_concepts(
                debt_concept_directory,
                plan_path=plan_path,
                base_facts_manifest=facts_manifest,
            )
        )
        facts = _merge_promoted_debt_facts(facts, promoted_debt)
        debt_concepts_ready = True
    else:
        debt_concepts_ready = False
    policy = ValuationReadinessPolicy(
        maximum_shares_age_days=int(args.maximum_shares_age_days),
        maximum_debt_age_days=int(args.maximum_debt_age_days),
    )
    policy.validate()
    statement_units = factor_inputs[
        ["selection_date", "ticker", "cik", "revenue_unit"]
    ].rename(columns={"revenue_unit": "statement_unit"})

    accounting, components = build_point_in_time_accounting_readiness(
        plan,
        facts,
        statement_units,
        policy=policy,
    )
    share_identity_directory = (
        _resolve(args.share_identity_dir)
        if args.share_identity_dir
        else source / "share_identity_v1"
    )
    share_promotions, share_promotions_manifest, share_identity_manifest = (
        _verified_share_promotions(share_identity_directory)
    )
    accounting = _apply_share_promotions(accounting, share_promotions)
    share_identity_ready = True
    price_basis_directory = (
        _resolve(args.price_basis_dir)
        if args.price_basis_dir
        else source / "price_basis_v2"
    )
    price_basis_manifest = None
    price_basis_coverage_manifest = None
    if (price_basis_directory / "manifest.json").exists():
        prices, price_basis_coverage_manifest, price_basis_manifest = (
            _verified_price_basis(
                price_basis_directory,
                plan_path=plan_path,
                plan=plan,
            )
        )
        price_basis_ready = True
        price_version = {
            "version_id": price_basis_coverage_manifest.get("version_id"),
            "content_sha256": price_basis_coverage_manifest.get("content_sha256"),
            "ticker_count": int(prices["ticker"].nunique()),
            "row_count": len(prices),
            "data_start": str(prices["selection_date"].min()),
            "data_end": str(prices["selection_date"].max()),
            "price_basis": RAW_PRICE_BASIS,
        }
    else:
        local_cfg = _research_cfg(cfg, True)
        tickers = sorted(plan["ticker"].astype(str).str.upper().unique())
        series, _price_meta, price_errors = _load_series(local_cfg, tickers)
        if price_errors:
            raise ValueError(f"local historical price load failed: {price_errors}")
        price_version = _price_universe_version(series)
        prices = point_in_time_price_rows(
            plan,
            series,
            maximum_age_days=policy.maximum_price_age_days,
        )
        price_basis_ready = False
    rows = build_valuation_factor_readiness(
        accounting,
        factor_inputs,
        prices,
        multi_security_tickers=set(MULTI_SECURITY_ISSUER),
        price_basis_ready=price_basis_ready,
    )
    if price_basis_ready:
        rows["historical_price_basis"] = RAW_PRICE_BASIS
        rows["historical_price_basis_note"] = (
            "独立版本化未复权名义成交价；严格取重选日前价格，不覆盖收益回测缓存。"
        )
    else:
        rows["historical_price_basis"] = "adjusted_or_unknown_local_cache"
        rows["historical_price_basis_note"] = (
            "本地OHLC未绑定逐日未复权价格或复权因子，禁止与历史申报股数相乘。"
        )
    events = share_basis_events(accounting, policy=policy)
    coverage = _coverage_by_date(rows)
    blockers = _blocker_summary(rows)
    methods = _method_summary(rows)
    failures = _integrity_failures(rows)
    versions = pd.DataFrame([
        {
            "source_name": "fact_availability",
            "version_id": facts_manifest.get("version_id"),
            "content_sha256": facts_manifest.get("content_sha256"),
            "rows": facts_manifest.get("row_count"),
            "note": "SEC Company Facts，filed+1 后可用",
        },
        {
            "source_name": "quality_factor_inputs",
            "version_id": factors_manifest.get("version_id"),
            "content_sha256": factors_manifest.get("content_sha256"),
            "rows": len(factor_inputs),
            "note": "点时LTM与存量值标准化输入",
        },
        {
            "source_name": "normalization_manifest",
            "version_id": normalization_manifest.get("normalization_policy_id"),
            "content_sha256": file_sha256(normalization_manifest_path),
            "rows": int(normalization_manifest.get("evidence_files", {}).get(
                "normalized_values", {}
            ).get("rows") or 0),
            "note": "既有质量标准化证据；本工具不改写",
        },
        {
            "source_name": "filing_share_identity_promotions",
            "version_id": share_promotions_manifest.get("version_id"),
            "content_sha256": share_promotions_manifest.get("content_sha256"),
            "rows": len(share_promotions),
            "note": (
                "仅提升美国申报、单一交易证券、200天内且旧值为空的XBRL股数"
            ),
        },
        {
            "source_name": (
                "unadjusted_price_basis" if price_basis_ready else "price_universe"
            ),
            "version_id": price_version.get("version_id"),
            "content_sha256": price_version.get("content_sha256"),
            "rows": price_version.get("row_count"),
            "note": (
                "严格取重选日前未复权名义成交价；可与申报股数同基准相乘"
                if price_basis_ready
                else "严格取重选日前价格；价格/股数基准仍阻断"
            ),
        },
    ])
    if debt_concepts_ready:
        versions = pd.concat([
            versions,
            pd.DataFrame([{
                "source_name": "expanded_debt_concepts",
                "version_id": debt_facts_manifest.get("version_id"),
                "content_sha256": debt_facts_manifest.get("content_sha256"),
                "rows": debt_facts_manifest.get("row_count"),
                "note": (
                    "仅提升DebtLongtermAndShorttermCombinedAmount直接总额；"
                    "DebtCurrent与CommercialPaper保持诊断"
                ),
            }]),
        ], ignore_index=True)

    paths = {
        "readiness": output / "valuation_readiness.csv",
        "components": output / "accounting_components.csv",
        "coverage": output / "coverage_by_date.csv",
        "blockers": output / "blocker_summary.csv",
        "methods": output / "accounting_method_summary.csv",
        "events": output / "share_basis_events.csv",
        "versions": output / "versions.csv",
        "failures": output / "integrity_failures.csv",
    }
    frames = {
        "readiness": rows,
        "components": components,
        "coverage": coverage,
        "blockers": blockers,
        "methods": methods,
        "events": events,
        "versions": versions,
        "failures": failures,
    }
    evidence = {}
    for key, path in paths.items():
        frame = frames[key]
        _atomic_csv(frame, path)
        manifest = write_dataset_manifest(
            path,
            dataset_type="ndx_valuation_readiness",
            dataset_key=f"{VALUATION_READINESS_POLICY_ID}:{key}",
            source="derived_from_versioned_sec_and_local_price_evidence",
            fetched_at_evidence="derived_runtime_utc",
            row_count=len(frame),
            data_start=(
                str(rows["selection_date"].min())
                if "selection_date" in frame and not frame.empty else None
            ),
            data_end=(
                str(rows["selection_date"].max())
                if "selection_date" in frame and not frame.empty else None
            ),
            extra={
                "research_only": True,
                "policy_id": VALUATION_READINESS_POLICY_ID,
                "valuation_scoring_allowed": False,
                "price_basis_ready": price_basis_ready,
                "debt_concepts_ready": debt_concepts_ready,
                "share_identity_ready": share_identity_ready,
            },
        )
        evidence[key] = {
            "filename": path.name,
            "rows": len(frame),
            "sha256": manifest["content_sha256"],
            "version_id": manifest["version_id"],
        }

    min_coverage = {
        column: float(coverage[column].min())
        for column in (
            "price_ready_coverage",
            "shares_ready_coverage",
            "debt_ready_coverage",
            "debt_strict_complete_coverage",
            "fcf_yield_accounting_ready_coverage",
            "ev_sales_accounting_ready_coverage",
            "fcf_yield_ready_coverage",
            "ev_sales_ready_coverage",
        )
    }
    latest = coverage.sort_values("selection_date").iloc[-1]
    summary = {
        "schema_version": 1,
        "generated_at": utc_now_iso(),
        "research_only": True,
        "policy_id": VALUATION_READINESS_POLICY_ID,
        "policy": policy.__dict__,
        "selection_dates": int(rows["selection_date"].nunique()),
        "universe_rows": len(rows),
        "ticker_count": int(rows["ticker"].nunique()),
        "minimum_coverage": min_coverage,
        "latest_selection_date": str(latest["selection_date"]),
        "latest_coverage": {
            column: float(latest[column])
            for column in min_coverage
        },
        "share_basis_event_rows": len(events),
        "split_like_share_event_rows": int(
            events["split_like"].sum() if not events.empty else 0
        ),
        "integrity_failure_rows": len(failures),
        "integrity_passed": failures.empty,
        "price_version": price_version,
        "price_basis_ready": price_basis_ready,
        "price_basis_blocker": (
            None
            if price_basis_ready
            else "unadjusted_price_or_adjustment_factor_missing"
        ),
        "price_basis_manifest": (
            {
                "policy_id": price_basis_manifest.get("policy_id"),
                "plan_sha256": price_basis_manifest.get("plan_sha256"),
                "coverage_version_id": price_basis_coverage_manifest.get("version_id"),
                "coverage_sha256": price_basis_coverage_manifest.get("content_sha256"),
            }
            if price_basis_manifest and price_basis_coverage_manifest
            else None
        ),
        "share_identity_ready": share_identity_ready,
        "share_identity_promoted_rows": int(
            rows["shares_identity_promoted"].fillna(False).astype(bool).sum()
        ),
        "share_identity_manifest": {
            "policy_id": share_identity_manifest.get("policy_id"),
            "promotions_version_id": share_promotions_manifest.get("version_id"),
            "promotions_sha256": share_promotions_manifest.get("content_sha256"),
            "promotion_allowed": share_identity_manifest.get("promotion_allowed"),
        },
        "debt_concepts_ready": debt_concepts_ready,
        "debt_concept_manifest": (
            {
                "policy_id": debt_concept_manifest.get("policy_id"),
                "plan_sha256": debt_concept_manifest.get("plan_sha256"),
                "facts_version_id": debt_facts_manifest.get("version_id"),
                "facts_sha256": debt_facts_manifest.get("content_sha256"),
                "promoted_concepts": debt_concept_manifest.get(
                    "promoted_concepts"
                ),
            }
            if debt_concept_manifest and debt_facts_manifest
            else None
        ),
        "valuation_scoring_allowed": False,
        "next_required_evidence": ([
            "point_in_time_fx_for_non_usd_statements",
            "explicit_adr_ratio_and_multi_security_share_allocation",
            (
                "remaining_debt_completeness_and_zero_debt_confirmation"
                if debt_concepts_ready
                else "expanded_debt_concept_policy_with_no_double_counting"
            ),
        ] if price_basis_ready else [
            "versioned_unadjusted_close_or_daily_adjustment_factor",
            "point_in_time_fx_for_non_usd_statements",
            "explicit_adr_ratio_and_multi_security_share_allocation",
            (
                "remaining_debt_completeness_and_zero_debt_confirmation"
                if debt_concepts_ready
                else "expanded_debt_concept_policy_with_no_double_counting"
            ),
        ]),
        "files": {
            **{key: value["filename"] for key, value in evidence.items()},
            "workbook": "ndx_valuation_readiness.xlsx",
        },
    }
    atomic_write_text(
        output / "summary.json",
        json.dumps(summary, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
    )
    atomic_write_text(
        output / "manifest.json",
        json.dumps(
            {
                "schema_version": 1,
                "generated_at": summary["generated_at"],
                "policy_id": VALUATION_READINESS_POLICY_ID,
                "source_versions": versions.to_dict(orient="records"),
                "evidence_files": evidence,
                "integrity_passed": failures.empty,
                "price_basis_ready": price_basis_ready,
                "debt_concepts_ready": debt_concepts_ready,
                "share_identity_ready": share_identity_ready,
                "share_identity_promotions_sha256": (
                    share_promotions_manifest.get("content_sha256")
                ),
                "valuation_scoring_allowed": False,
            },
            ensure_ascii=False,
            indent=2,
            allow_nan=False,
        ) + "\n",
    )
    conclusion = pd.DataFrame([
        {"项目": "研究目的", "内容": "建立估值因子可计算性守门，不计算估值分数"},
        {"项目": "点时候选行", "内容": len(rows)},
        {"项目": "历史股票数", "内容": rows["ticker"].nunique()},
        {"项目": "最低股数就绪率", "内容": min_coverage["shares_ready_coverage"]},
        {
            "项目": "filing股数安全提升行",
            "内容": summary["share_identity_promoted_rows"],
        },
        {"项目": "最低未复权价格就绪率", "内容": min_coverage["price_ready_coverage"]},
        {"项目": "最低债务可组装率", "内容": min_coverage["debt_ready_coverage"]},
        {"项目": "最低总债务严格完整率", "内容": min_coverage["debt_strict_complete_coverage"]},
        {"项目": "债务概念v2准入", "内容": debt_concepts_ready},
        {"项目": "最低FCF收益率会计就绪率", "内容": min_coverage["fcf_yield_accounting_ready_coverage"]},
        {"项目": "最低EV销售额会计就绪率", "内容": min_coverage["ev_sales_accounting_ready_coverage"]},
        {
            "项目": "最低FCF收益率最终就绪率",
            "内容": min_coverage["fcf_yield_ready_coverage"],
        },
        {
            "项目": "最低EV销售额最终就绪率",
            "内容": min_coverage["ev_sales_ready_coverage"],
        },
        {"项目": "完整性失败", "内容": len(failures)},
        {"项目": "实盘/评分状态", "内容": "关闭；不改配置、不改今日建议"},
        {
            "项目": "下一证据",
            "内容": (
                "点时FX、ADR比例、多股类分摊、剩余债务完整性/零债务确认"
                if price_basis_ready and debt_concepts_ready
                else "点时FX、ADR比例、多股类分摊、扩展债务词表"
                if price_basis_ready
                else "未复权价格/复权因子、点时FX、ADR比例、多股类分摊、扩展债务词表"
            ),
        },
    ])
    workbook = output / "ndx_valuation_readiness.xlsx"
    _write_workbook(
        workbook,
        conclusion=conclusion,
        rows=rows,
        coverage=coverage,
        blockers=blockers,
        methods=methods,
        events=events,
        components=components,
        versions=versions,
        failures=failures,
    )
    print("NDX point-in-time valuation readiness (research only)")
    print(
        f"Rows: {len(rows)} | shares min {min_coverage['shares_ready_coverage']:.2%} | "
        f"debt min {min_coverage['debt_ready_coverage']:.2%} | "
        f"strict debt min {min_coverage['debt_strict_complete_coverage']:.2%}"
    )
    print(
        "Valuation readiness min coverage: "
        f"FCF yield {min_coverage['fcf_yield_ready_coverage']:.2%} | "
        f"EV/Sales {min_coverage['ev_sales_ready_coverage']:.2%}; "
        "scoring remains disabled."
    )
    print(f"Workbook: {workbook}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
