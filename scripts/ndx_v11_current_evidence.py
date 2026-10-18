"""Build current v11 fundamental evidence without ranking or backtesting.

PowerShell / PyCharm terminal:
    python -B scripts\ndx_v11_current_evidence.py --observation-date 2026-07-13

The command consumes the isolated Company Facts bundle created by
``ndx_v11_forward_watch.py``.  It writes a separate research directory,
normalizes point-in-time accounting facts, audits unranked quality inputs and
valuation readiness, and keeps candidate admission disabled.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from backtest.ndx_forward_watch import NdxForwardWatchError  # noqa: E402
from core.helpers import atomic_write_text  # noqa: E402
from data.ndx_fundamental_normalization import (  # noqa: E402
    NORMALIZATION_POLICY_ID,
    build_quality_factor_inputs,
    normalize_point_in_time_facts,
)
from data.ndx_fundamentals import (  # noqa: E402
    NdxFundamentalError,
    raw_bundle_sha256,
)
from data.ndx_v11_current_evidence import (  # noqa: E402
    CURRENT_EVIDENCE_POLICY_ID,
    build_quality_input_readiness,
    build_security_evidence_gate,
    build_valuation_gate_summary,
)
from data.ndx_valuation_readiness import (  # noqa: E402
    ValuationReadinessPolicy,
    build_point_in_time_accounting_readiness,
    build_valuation_factor_readiness,
    point_in_time_price_rows,
)
from data.versioning import (  # noqa: E402
    file_sha256,
    manifest_path,
    utc_now_iso,
    write_dataset_manifest,
)


DEFAULT_INPUT_DIR = ROOT / "data" / "cache" / "research" / "ndx_v11_forward_watch_v1"
DEFAULT_OUTPUT_DIR = ROOT / "data" / "cache" / "research" / "ndx_v11_current_fundamental_evidence_v1"
DEFAULT_PRICE_DIR = ROOT / "data" / "cache"

COLUMN_LABELS = {
    "selection_date": "观察日期",
    "observation_date": "观察日期",
    "ticker": "股票代码",
    "name": "证券名称",
    "cik": "SEC CIK",
    "issuer_id": "发行人编号",
    "security_id": "证券编号",
    "risk_chain": "风险链编号",
    "risk_chain_label": "风险链名称",
    "taxonomy_status": "产业链分类状态",
    "taxonomy_confidence": "产业链置信度",
    "taxonomy_source": "产业链证据来源",
    "metric": "指标",
    "status": "状态",
    "usable": "可用",
    "confidence": "置信度",
    "value": "数值",
    "unit": "单位",
    "normalization_method": "标准化方法",
    "period_start": "期间起始",
    "period_end": "期间期末",
    "age_days": "账龄天数",
    "component_count": "来源组件数",
    "concepts": "XBRL概念",
    "source_available_date_max": "来源最晚可用日",
    "comparison_value": "同比比较值",
    "comparison_unit": "同比比较单位",
    "comparison_method": "同比比较方法",
    "comparison_period_end": "同比比较期末",
    "comparison_gap_days": "同比期间差天数",
    "comparison_concepts": "同比概念",
    "comparison_component_count": "同比组件数",
    "comparison_source_available_date_max": "同比来源最晚可用日",
    "arithmetic_passed": "算术守门通过",
    "future_data_guard_passed": "未来数据守门通过",
    "factor": "因子",
    "ready_count": "就绪数",
    "coverage": "覆盖率",
    "quality_factor_input_ready": "质量因子输入就绪",
    "quality_factor_blockers": "质量因子阻断原因",
    "quality_score_generated": "已生成质量分数",
    "quality_score_allowed": "允许质量评分",
    "core_factor_count": "核心因子数",
    "growth_factor_count": "成长因子数",
    "profitability_factor_count": "盈利因子数",
    "balance_factor_count": "资产负债表因子数",
    "quality_factor_input_ready_count": "质量因子输入就绪数",
    "quality_factor_input_coverage": "质量因子输入覆盖率",
    "minimum_core_factors": "核心因子最低数",
    "minimum_profitability_factors": "盈利因子最低数",
    "minimum_balance_factors": "资产负债表因子最低数",
    "price_status": "价格状态",
    "price_ready": "价格就绪",
    "price": "严格前一交易日价格",
    "price_date": "价格日期",
    "price_age_days": "价格滞后天数",
    "price_future_data_guard_passed": "价格未来数据守门通过",
    "statement_unit": "报表币种",
    "foreign_filer_asof": "点时境外申报人",
    "shares_status": "股数状态",
    "shares_ready": "股数就绪",
    "shares_value": "申报股数",
    "shares_unit": "股数单位",
    "shares_namespace": "股数命名空间",
    "shares_concept": "股数概念",
    "shares_period_end": "股数基准日",
    "shares_age_days": "股数账龄天数",
    "shares_same_date_candidate_count": "同日股数候选数",
    "shares_same_date_value_spread": "同日股数差异率",
    "shares_source_available_date": "股数来源可用日",
    "debt_status": "债务状态",
    "debt_ready": "债务就绪",
    "debt_strict_complete": "债务严格完整",
    "debt_value": "债务数值",
    "debt_unit": "债务单位",
    "debt_method": "债务组装方法",
    "debt_completeness": "债务完整性",
    "debt_period_end": "债务基准日",
    "debt_age_days": "债务账龄天数",
    "debt_source_available_date_max": "债务来源最晚可用日",
    "debt_crosscheck_relative_error": "债务交叉误差",
    "debt_crosscheck_status": "债务交叉状态",
    "free_cash_flow_value": "自由现金流金额",
    "free_cash_flow_unit": "自由现金流单位",
    "free_cash_flow_components_ready": "自由现金流组件就绪",
    "fcf_yield_accounting_ready": "FCF收益率会计证据就绪",
    "ev_sales_accounting_ready": "EV销售额会计证据就绪",
    "fcf_yield_ready": "FCF收益率最终就绪",
    "ev_sales_ready": "EV销售额最终就绪",
    "valuation_scoring_allowed": "允许估值评分",
    "valuation_score_generated": "已生成估值分数",
    "valuation_score_allowed": "允许估值评分",
    "valuation_accounting_evidence_ready": "估值会计证据就绪",
    "fcf_yield_blockers": "FCF收益率阻断原因",
    "ev_sales_blockers": "EV销售额阻断原因",
    "valuation_blockers": "估值总阻断原因",
    "evidence_gate_status": "证据准入状态",
    "evidence_gate_blockers": "证据准入阻断原因",
    "candidate_admission_allowed": "允许候选准入",
    "old_holding_sale_allowed": "允许卖出旧持仓",
    "return_data_read": "读取收益数据",
    "universe_count": "证券数",
    "normalized_count": "标准化数",
    "normalized_coverage": "标准化覆盖率",
    "high_confidence_count": "高置信度数",
    "stale_count": "过期数",
    "missing_count": "缺失数",
    "policy_pending_count": "策略待定数",
    "age_days_p50": "账龄P50",
    "age_days_p90": "账龄P90",
    "failure_type": "失败类型",
    "detail": "详情",
    "blocker": "阻断项",
    "affected_rows": "受影响行数",
    "affected_tickers": "受影响股票数",
    "affected_row_rate": "受影响行比例",
    "source_role": "来源角色",
    "path": "路径",
    "version_id": "版本编号",
    "content_sha256": "内容哈希",
    "rows": "行数",
    "note": "说明",
}


def _resolve(value: str | Path) -> Path:
    path = Path(value)
    return path if path.is_absolute() else ROOT / path


def _read_json(path: Path) -> dict:
    if not path.exists():
        raise FileNotFoundError(path)
    return json.loads(path.read_text(encoding="utf-8"))


def _verified_csv(path: Path) -> tuple[pd.DataFrame, dict]:
    sidecar = manifest_path(path)
    if not path.exists() or not sidecar.exists():
        raise FileNotFoundError(f"versioned input missing: {path}")
    manifest = _read_json(sidecar)
    observed_hash = file_sha256(path)
    if manifest.get("content_sha256") != observed_hash:
        raise ValueError(f"versioned input hash mismatch: {path}")
    frame = pd.read_csv(path, low_memory=False)
    declared_rows = manifest.get("row_count")
    if declared_rows is not None and int(declared_rows) != len(frame):
        raise ValueError(f"versioned input row count mismatch: {path}")
    return frame, manifest


def _verify_raw_index(raw_index: pd.DataFrame) -> str:
    required = {"cik", "path", "content_sha256"}
    if not required.issubset(raw_index.columns):
        raise ValueError(f"raw index columns missing: {sorted(required - set(raw_index.columns))}")
    for row in raw_index.itertuples(index=False):
        path = Path(str(row.path))
        if not path.exists():
            raise FileNotFoundError(path)
        if file_sha256(path) != str(row.content_sha256):
            raise ValueError(f"raw Company Facts hash mismatch: {path}")
    return raw_bundle_sha256(raw_index)


def _load_price_series(price_dir: Path, tickers: list[str]) -> dict[str, pd.DataFrame]:
    series: dict[str, pd.DataFrame] = {}
    for ticker in sorted(set(tickers)):
        path = price_dir / f"{ticker}.csv"
        if not path.exists():
            series[ticker] = pd.DataFrame(columns=["Close"])
            continue
        frame = pd.read_csv(path, usecols=lambda column: column in {"Date", "Close"})
        if not {"Date", "Close"}.issubset(frame.columns):
            raise ValueError(f"price cache columns missing for {ticker}: {path}")
        frame["Date"] = pd.to_datetime(frame["Date"], errors="coerce")
        frame["Close"] = pd.to_numeric(frame["Close"], errors="coerce")
        frame = frame.dropna(subset=["Date", "Close"]).drop_duplicates("Date")
        frame = frame[frame["Close"] > 0].sort_values("Date")
        frame = frame.set_index("Date")
        series[ticker] = frame[["Close"]]
    return series


def _normalization_failures(values: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for field, failure_type in (
        ("arithmetic_passed", "normalization_arithmetic"),
        ("future_data_guard_passed", "normalization_future_data_guard"),
    ):
        failed = ~values[field].fillna(False).astype(bool)
        for row in values[failed].itertuples(index=False):
            rows.append({
                "failure_type": failure_type,
                "selection_date": row.selection_date,
                "ticker": row.ticker,
                "cik": row.cik,
                "metric": row.metric,
                "detail": getattr(row, field),
            })
    return pd.DataFrame(rows, columns=[
        "failure_type", "selection_date", "ticker", "cik", "metric", "detail",
    ])


def _valuation_blocker_summary(rows: pd.DataFrame) -> pd.DataFrame:
    output = []
    for factor, column in (
        ("fcf_yield", "fcf_yield_blockers"),
        ("ev_sales", "ev_sales_blockers"),
    ):
        exploded = rows[["selection_date", "ticker", column]].copy()
        exploded["blocker"] = exploded[column].fillna("").astype(str).str.split(";")
        exploded = exploded.explode("blocker")
        exploded = exploded[exploded["blocker"].str.len() > 0]
        if exploded.empty:
            continue
        grouped = exploded.groupby("blocker").agg(
            affected_rows=("ticker", "size"),
            affected_tickers=("ticker", "nunique"),
            first_selection_date=("selection_date", "min"),
            last_selection_date=("selection_date", "max"),
        ).reset_index()
        grouped.insert(0, "factor", factor)
        grouped["affected_row_rate"] = grouped["affected_rows"] / len(rows)
        output.append(grouped)
    if not output:
        return pd.DataFrame(columns=[
            "factor", "blocker", "affected_rows", "affected_tickers",
            "first_selection_date", "last_selection_date", "affected_row_rate",
        ])
    return pd.concat(output, ignore_index=True).sort_values(
        ["factor", "affected_rows", "blocker"],
        ascending=[True, False, True],
        kind="stable",
    ).reset_index(drop=True)


def _write_frame(frame: pd.DataFrame, path: Path, *, key: str, observation_date: str) -> dict:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    frame.to_csv(temporary, index=False, encoding="utf-8-sig", lineterminator="\n")
    temporary.replace(path)
    return write_dataset_manifest(
        path,
        dataset_type="ndx_v11_current_fundamental_evidence",
        dataset_key=key,
        source="derived_from_sec_companyfacts_forward_watch",
        source_locator="https://data.sec.gov/api/xbrl/companyfacts/",
        fetched_at_evidence="derived_runtime_utc",
        row_count=len(frame),
        data_start=observation_date,
        data_end=observation_date,
        extra={
            "research_only": True,
            "observation_date": observation_date,
            "return_data_read": False,
            "backtest_results_used": False,
            "market_price_values_used": True,
            "price_return_series_used": False,
            "quality_score_generated": False,
            "valuation_score_generated": False,
            "candidate_admission_allowed": False,
            "live_config_mutation_allowed": False,
            "today_advice_impact": "none",
        },
    )


def _zh(frame: pd.DataFrame) -> pd.DataFrame:
    renamed = {}
    for column in frame.columns:
        renamed[column] = COLUMN_LABELS.get(column, f"字段：{column}")
    return frame.rename(columns=renamed)


def _write_workbook(path: Path, frames: dict[str, pd.DataFrame]) -> None:
    temporary = path.with_name(path.stem + ".tmp.xlsx")
    with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
        for sheet, frame in frames.items():
            _zh(frame).to_excel(writer, sheet_name=sheet[:31], index=False)
    workbook = load_workbook(temporary)
    fill = PatternFill("solid", fgColor="DCE6F1")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            width = max(len(str(cell.value or "")) for cell in column[:300]) + 2
            sheet.column_dimensions[letter].width = min(max(width, 10), 48)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(temporary)
    temporary.replace(path)


def build_current_evidence(
    *,
    input_dir: str | Path = DEFAULT_INPUT_DIR,
    output_dir: str | Path = DEFAULT_OUTPUT_DIR,
    price_dir: str | Path = DEFAULT_PRICE_DIR,
    observation_date: str,
) -> dict:
    input_path = _resolve(input_dir)
    output_path = _resolve(output_dir)
    price_path = _resolve(price_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    observed = str(pd.Timestamp(observation_date).date())

    forward_manifest_path = input_path / "manifest.json"
    forward_manifest = _read_json(forward_manifest_path)
    if forward_manifest.get("observation_date") != observed:
        raise NdxForwardWatchError(
            f"forward-watch observation date mismatch: {forward_manifest.get('observation_date')} != {observed}"
        )
    if forward_manifest.get("return_data_read") is True:
        raise NdxForwardWatchError("forward-watch input unexpectedly read return data")
    observations, observations_manifest = _verified_csv(
        input_path / "forward_observations.csv"
    )
    plan, plan_manifest = _verified_csv(input_path / "companyfacts_plan.csv")
    records, records_manifest = _verified_csv(input_path / "fact_availability.csv")
    raw_index, raw_index_manifest = _verified_csv(input_path / "raw_index.csv")
    download_status, download_manifest = _verified_csv(input_path / "download_status.csv")
    if download_status["status"].astype(str).isin({"failed", "pending_limit", "missing"}).any():
        raise NdxForwardWatchError("Company Facts download status is incomplete")
    raw_bundle_hash = _verify_raw_index(raw_index)

    if len(plan) != len(observations):
        raise NdxForwardWatchError("Company Facts plan and forward observations differ")
    if plan["ticker"].astype(str).str.upper().eq("QQQ").any():
        raise NdxForwardWatchError("QQQ cannot enter current evidence")
    plan = plan.rename(columns={"observation_date": "selection_date"}).copy()
    plan["selection_date"] = plan["selection_date"].astype(str).str[:10]
    plan["ticker"] = plan["ticker"].astype(str).str.upper()
    plan["cik"] = pd.to_numeric(plan["cik"], errors="coerce")
    if plan[["selection_date", "cik"]].isna().any().any():
        raise NdxFundamentalError("current Company Facts plan has invalid date or CIK")
    plan["cik"] = plan["cik"].astype(int)
    if plan["selection_date"].ne(observed).any():
        raise NdxForwardWatchError("Company Facts plan contains a different observation date")

    values, components, normalization_summary = normalize_point_in_time_facts(
        plan[["selection_date", "ticker", "cik"]],
        records,
    )
    factor_inputs, factor_coverage = build_quality_factor_inputs(values)
    quality_readiness, quality_coverage = build_quality_input_readiness(factor_inputs)
    statement_units = factor_inputs[[
        "selection_date", "ticker", "cik", "revenue_unit"
    ]].rename(columns={"revenue_unit": "statement_unit"})
    accounting, accounting_components = build_point_in_time_accounting_readiness(
        plan[["selection_date", "ticker", "cik"]],
        records,
        statement_units,
        policy=ValuationReadinessPolicy(),
    )
    series = _load_price_series(price_path, sorted(plan["ticker"].unique()))
    prices = point_in_time_price_rows(
        plan[["selection_date", "ticker"]],
        series,
        maximum_age_days=7,
    )
    issuer_counts = plan.groupby("cik")["ticker"].nunique()
    multi_security_tickers = set(
        plan.loc[plan["cik"].isin(issuer_counts[issuer_counts.gt(1)].index), "ticker"]
    )
    valuation_raw = build_valuation_factor_readiness(
        accounting,
        factor_inputs,
        prices,
        multi_security_tickers=multi_security_tickers,
        price_basis_ready=False,
    )
    valuation_readiness, valuation_coverage = build_valuation_gate_summary(valuation_raw)
    valuation_blockers = _valuation_blocker_summary(valuation_readiness)
    normalization_failures = _normalization_failures(values)
    security_gate = build_security_evidence_gate(
        observations,
        quality_readiness,
        valuation_readiness,
    )

    source_rows = [
        {
            "source_role": "forward_watch_manifest",
            "path": forward_manifest_path.as_posix(),
            "version_id": forward_manifest.get("status"),
            "content_sha256": file_sha256(forward_manifest_path),
            "rows": 1,
            "note": "规则注册和当前观察边界",
        },
        {
            "source_role": "forward_observations",
            "path": (input_path / "forward_observations.csv").as_posix(),
            "version_id": observations_manifest.get("version_id"),
            "content_sha256": observations_manifest.get("content_sha256"),
            "rows": len(observations),
            "note": "当前 Nasdaq 官方成分、身份和 taxonomy 状态",
        },
        {
            "source_role": "companyfacts_plan",
            "path": (input_path / "companyfacts_plan.csv").as_posix(),
            "version_id": plan_manifest.get("version_id"),
            "content_sha256": plan_manifest.get("content_sha256"),
            "rows": len(plan),
            "note": "SEC CIK 计划",
        },
        {
            "source_role": "fact_availability",
            "path": (input_path / "fact_availability.csv").as_posix(),
            "version_id": records_manifest.get("version_id"),
            "content_sha256": records_manifest.get("content_sha256"),
            "rows": len(records),
            "note": "SEC Company Facts，available_date <= observation_date",
        },
        {
            "source_role": "raw_companyfacts_bundle",
            "path": (input_path / "raw_index.csv").as_posix(),
            "version_id": raw_index_manifest.get("version_id"),
            "content_sha256": raw_bundle_hash,
            "rows": len(raw_index),
            "note": "逐个原始 JSON 内容哈希复核",
        },
        {
            "source_role": "download_status",
            "path": (input_path / "download_status.csv").as_posix(),
            "version_id": download_manifest.get("version_id"),
            "content_sha256": download_manifest.get("content_sha256"),
            "rows": len(download_status),
            "note": "全部为 cached/fetched，未使用失败或限额行",
        },
    ]
    sources = pd.DataFrame(source_rows)
    conclusion = pd.DataFrame([
        {"项目": "当前证据状态", "结论": "已完成点时标准化；质量/估值评分和候选准入继续关闭"},
        {"项目": "质量证据", "结论": "仅输出未排名因子输入与完整性审计，不使用旧 SIC 或收益结果"},
        {"项目": "估值证据", "结论": "输出股数、债务、价格和 FCF/EV-Sales 就绪审计；不生成估值分数"},
        {"项目": "未来数据守门", "结论": "标准化组件要求 available_date <= observation_date 且 period_end < observation_date"},
        {"项目": "实盘影响", "结论": "不修改 config.yaml、state、今日建议或买入计划"},
    ])

    frames = {
        "规范化摘要": normalization_summary,
        "质量因子输入": factor_inputs,
        "因子覆盖审计": factor_coverage,
        "质量就绪审计": quality_readiness,
        "质量覆盖审计": quality_coverage,
        "会计估值就绪": accounting,
        "估值就绪审计": valuation_readiness,
        "估值覆盖审计": valuation_coverage,
        "估值阻断原因": valuation_blockers,
        "证券证据准入": security_gate,
        "标准化完整性失败": normalization_failures,
        "原始会计组件": accounting_components,
        "来源版本": sources,
        "结论": conclusion,
    }
    evidence = {}
    for sheet, frame in frames.items():
        filename = {
            "规范化摘要": "normalization_summary.csv",
            "质量因子输入": "quality_factor_inputs.csv",
            "因子覆盖审计": "factor_coverage.csv",
            "质量就绪审计": "quality_input_readiness.csv",
            "质量覆盖审计": "quality_input_coverage.csv",
            "会计估值就绪": "accounting_readiness.csv",
            "估值就绪审计": "valuation_readiness.csv",
            "估值覆盖审计": "valuation_coverage.csv",
            "估值阻断原因": "valuation_blockers.csv",
            "证券证据准入": "security_evidence_gate.csv",
            "标准化完整性失败": "normalization_integrity_failures.csv",
            "原始会计组件": "accounting_components.csv",
            "来源版本": "source_versions.csv",
            "结论": "conclusion.csv",
        }[sheet]
        manifest = _write_frame(
            frame,
            output_path / filename,
            key=filename.removesuffix(".csv"),
            observation_date=observed,
        )
        evidence[filename] = {
            "rows": len(frame),
            "sha256": manifest["content_sha256"],
            "version_id": manifest["version_id"],
        }

    workbook = output_path / "ndx_v11_current_fundamental_evidence.xlsx"
    _write_workbook(workbook, frames)
    workbook_manifest = write_dataset_manifest(
        workbook,
        dataset_type="ndx_v11_current_fundamental_evidence_workbook",
        dataset_key=CURRENT_EVIDENCE_POLICY_ID,
        source="derived_from_sec_companyfacts_forward_watch",
        fetched_at_evidence="derived_runtime_utc",
        row_count=len(frames),
        data_start=observed,
        data_end=observed,
        extra={
            "research_only": True,
            "return_data_read": False,
            "backtest_results_used": False,
            "market_price_values_used": True,
            "price_return_series_used": False,
            "quality_score_generated": False,
            "valuation_score_generated": False,
            "candidate_admission_allowed": False,
        },
    )

    current_taxonomy_accepted = int(observations["taxonomy_status"].eq("accepted").sum())
    quality_ready_count = int(quality_readiness["quality_factor_input_ready"].sum())
    fcf_ready_count = int(valuation_readiness["fcf_yield_ready"].fillna(False).astype(bool).sum())
    ev_ready_count = int(valuation_readiness["ev_sales_ready"].fillna(False).astype(bool).sum())
    blockers = [
        "quality_score_generation_disabled_until_current_risk_audit",
        "valuation_score_generation_disabled_until_price_fx_identity_audit",
        "fresh_oos_backtest_not_run",
    ]
    if current_taxonomy_accepted < len(observations):
        blockers.insert(0, "taxonomy_current_evidence_missing_or_unclassified")
    if not bool(valuation_readiness["fcf_yield_ready"].fillna(False).astype(bool).any()):
        blockers.append("fcf_yield_no_current_candidate_ready")
    if not bool(valuation_readiness["ev_sales_ready"].fillna(False).astype(bool).any()):
        blockers.append("ev_sales_no_current_candidate_ready")
    summary = {
        "schema_version": 1,
        "policy_id": CURRENT_EVIDENCE_POLICY_ID,
        "generated_at": utc_now_iso(),
        "observation_date": observed,
        "research_only": True,
        "current_security_count": len(observations),
        "unique_cik_count": int(plan["cik"].nunique()),
        "taxonomy_accepted_count": current_taxonomy_accepted,
        "taxonomy_unclassified_count": int((~observations["taxonomy_status"].eq("accepted")).sum()),
        "normalized_value_rows": len(values),
        "normalized_usable_rows": int(values["usable"].sum()),
        "normalization_component_rows": len(components),
        "normalization_integrity_failure_rows": len(normalization_failures),
        "quality_factor_input_ready_count": quality_ready_count,
        "quality_factor_input_coverage": float(
            quality_ready_count / len(quality_readiness)
        ) if len(quality_readiness) else 0.0,
        "quality_score_generated": False,
        "valuation_fcf_yield_ready_count": fcf_ready_count,
        "valuation_ev_sales_ready_count": ev_ready_count,
        "valuation_score_generated": False,
        "return_data_read": False,
        "backtest_results_used": False,
        "market_price_values_used": True,
        "price_return_series_used": False,
        "candidate_admission_allowed": False,
        "old_holding_sale_allowed": False,
        "blockers": list(dict.fromkeys(blockers)),
        "raw_bundle_sha256": raw_bundle_hash,
        "normalization_policy_id": NORMALIZATION_POLICY_ID,
        "evidence_files": evidence,
        "workbook": {
            "filename": workbook.name,
            "rows": len(frames),
            "sha256": workbook_manifest["content_sha256"],
            "version_id": workbook_manifest["version_id"],
        },
        "source_versions": source_rows,
    }
    summary_path = output_path / "summary.json"
    atomic_write_text(
        summary_path,
        json.dumps(summary, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
    )
    summary_manifest = write_dataset_manifest(
        summary_path,
        dataset_type="ndx_v11_current_fundamental_evidence_summary",
        dataset_key=CURRENT_EVIDENCE_POLICY_ID,
        source="derived_from_sec_companyfacts_forward_watch",
        fetched_at_evidence="derived_runtime_utc",
        row_count=1,
        data_start=observed,
        data_end=observed,
        extra={
            "research_only": True,
            "return_data_read": False,
            "backtest_results_used": False,
            "candidate_admission_allowed": False,
        },
    )
    summary["summary_manifest_sha256"] = summary_manifest["content_sha256"]
    manifest_path_main = output_path / "manifest.json"
    atomic_write_text(
        manifest_path_main,
        json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True, allow_nan=False) + "\n",
    )
    write_dataset_manifest(
        manifest_path_main,
        dataset_type="ndx_v11_current_fundamental_evidence_manifest",
        dataset_key=CURRENT_EVIDENCE_POLICY_ID,
        source="derived_from_sec_companyfacts_forward_watch",
        fetched_at_evidence="derived_runtime_utc",
        row_count=1,
        data_start=observed,
        data_end=observed,
        extra={
            "research_only": True,
            "return_data_read": False,
            "backtest_results_used": False,
            "candidate_admission_allowed": False,
        },
    )
    return summary


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Build current v11 fundamental evidence without ranking or backtesting"
    )
    parser.add_argument("--input-dir", default=str(DEFAULT_INPUT_DIR))
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR))
    parser.add_argument("--price-dir", default=str(DEFAULT_PRICE_DIR))
    parser.add_argument("--observation-date", required=True)
    args = parser.parse_args()
    try:
        summary = build_current_evidence(
            input_dir=args.input_dir,
            output_dir=args.output_dir,
            price_dir=args.price_dir,
            observation_date=args.observation_date,
        )
        print(
            "NDX v11当前基本面证据完成 | "
            f"securities={summary['current_security_count']} | "
            f"normalized={summary['normalized_usable_rows']}/{summary['normalized_value_rows']} | "
            f"quality_input_ready={summary['quality_factor_input_ready_count']} | "
            f"fcf_ready={summary['valuation_fcf_yield_ready_count']} | "
            f"ev_sales_ready={summary['valuation_ev_sales_ready_count']} | "
            "quality_score=False | valuation_score=False | return_data_read=False"
        )
        print(f"输出: {_resolve(args.output_dir)}")
        print(f"Manifest: {_resolve(args.output_dir) / 'manifest.json'}")
        return 0
    except (NdxForwardWatchError, NdxFundamentalError, OSError, ValueError) as exc:
        print(f"NDX v11当前基本面证据失败闭环: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
