"""Download SEC XBRL instances and audit dimensional share identity evidence."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
import time
from pathlib import Path

import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from core.config import load_config
from core.helpers import atomic_write_text
from data.ndx_filing_identity import (
    FILING_IDENTITY_POLICY_ID,
    build_filing_identity_plan,
    file_content_sha256,
    filing_document_url,
    parse_filing_share_facts,
    select_xbrl_instance,
)
from data.ndx_fundamentals import fundamental_cache_dir
from data.ndx_history import _sec_headers, _sec_user_agent
from data.versioning import file_sha256, utc_now_iso, write_dataset_manifest


ATTEMPT_COLUMNS = [
    "policy_id",
    "job_id",
    "cik",
    "accession_number",
    "status",
    "index_url",
    "instance_filename",
    "selection_method",
    "instance_url",
    "content_bytes",
    "content_sha256",
    "share_fact_count",
    "error",
    "attempted_at",
    "duration_seconds",
]
TERMINAL_STATUSES = {"ready", "no_share_facts", "no_explicit_xbrl_instance"}


def _resolve(path: str | Path) -> Path:
    value = Path(path)
    return value if value.is_absolute() else ROOT / value


def _atomic_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        frame.to_csv(temporary, index=False, encoding="utf-8-sig")
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def _atomic_bytes(content: bytes, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        temporary.write_bytes(content)
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def _load_attempts(path: Path) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame(columns=ATTEMPT_COLUMNS)
    frame = pd.read_csv(path, dtype=str, keep_default_na=False)
    for column in ATTEMPT_COLUMNS:
        if column not in frame.columns:
            frame[column] = ""
    return frame[ATTEMPT_COLUMNS].drop_duplicates("job_id", keep="last")


def _upsert_attempt(frame: pd.DataFrame, row: dict) -> pd.DataFrame:
    remaining = frame[~frame["job_id"].eq(str(row["job_id"]))].copy()
    result = pd.concat(
        [remaining, pd.DataFrame([row], columns=ATTEMPT_COLUMNS)],
        ignore_index=True,
    )
    return result.sort_values(["cik", "accession_number"], kind="stable").reset_index(
        drop=True
    )


def _request_bytes(
    session: requests.Session,
    url: str,
    *,
    timeout: float,
    retries: int,
) -> bytes:
    last_error: Exception | None = None
    for attempt in range(max(1, int(retries))):
        try:
            response = session.get(url, timeout=timeout)
            response.raise_for_status()
            return response.content
        except requests.RequestException as exc:
            last_error = exc
            if attempt + 1 < max(1, int(retries)):
                time.sleep(min(8.0, 1.0 * (2 ** attempt)))
    assert last_error is not None
    raise last_error


def _instance_path(
    raw_dir: Path,
    *,
    cik: int,
    accession_number: str,
    filename: str,
) -> Path:
    safe_accession = accession_number.replace("-", "")
    return raw_dir / str(int(cik)) / safe_accession / filename


def _cached_facts(raw_dir: Path, attempts: pd.DataFrame) -> tuple[pd.DataFrame, list[dict]]:
    facts = []
    versions = []
    for row in attempts.itertuples(index=False):
        if row.status not in {"ready", "no_share_facts"}:
            continue
        path = _instance_path(
            raw_dir,
            cik=int(row.cik),
            accession_number=str(row.accession_number),
            filename=str(row.instance_filename),
        )
        if not path.exists() or file_content_sha256(path) != row.content_sha256:
            continue
        parsed = parse_filing_share_facts(
            path.read_bytes(),
            cik=int(row.cik),
            accession_number=str(row.accession_number),
            source_url=str(row.instance_url),
        )
        if not parsed.empty:
            facts.append(parsed)
        versions.append({
            "source_name": f"sec_xbrl:{row.accession_number}",
            "version_id": f"sec_xbrl:{row.accession_number}:{row.content_sha256[:16]}",
            "content_sha256": row.content_sha256,
            "rows": len(parsed),
            "note": row.instance_url,
        })
    if facts:
        return pd.concat(facts, ignore_index=True), versions
    return pd.DataFrame(columns=[
        "cik",
        "accession_number",
        "concept_namespace",
        "concept",
        "context_ref",
        "period_end",
        "entity_identifier",
        "entity_identifier_matches_cik",
        "unit_ref",
        "unit_measure",
        "raw_value",
        "value",
        "decimals",
        "scale",
        "dimensions_count",
        "dimensions_json",
        "source_url",
    ]), versions


def _style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    fill = PatternFill("solid", fgColor="DCE6F1")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            width = max(len(str(cell.value or "")) for cell in column[:300]) + 2
            sheet.column_dimensions[letter].width = min(max(width, 10), 48)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)


def _write_workbook(path: Path, frames: dict[str, pd.DataFrame]) -> None:
    translations = {
        "selection_date": "重选日",
        "ticker": "股票代码",
        "tickers": "涉及股票代码",
        "cik": "SEC CIK",
        "accession_number": "申报编号",
        "form": "申报表类型",
        "filed_date": "申报日",
        "available_date": "可用日",
        "first_selection_date": "首个使用重选日",
        "last_selection_date": "末个使用重选日",
        "selection_rows": "重选行数",
        "identity_reason": "身份审计原因",
        "identity_reasons": "身份审计原因集合",
        "filing_future_data_guard_passed": "未来数据守门通过",
        "index_url": "SEC目录地址",
        "job_id": "作业编号",
        "status": "状态",
        "instance_filename": "XBRL实例文件",
        "selection_method": "实例选择方法",
        "instance_url": "XBRL实例地址",
        "content_bytes": "内容字节数",
        "content_sha256": "内容哈希",
        "share_fact_count": "流通股候选事实数",
        "error": "错误",
        "attempted_at": "执行时间",
        "duration_seconds": "耗时秒数",
        "concept_namespace": "概念命名空间",
        "concept": "流通股概念",
        "context_ref": "上下文编号",
        "period_end": "事实日期",
        "entity_identifier": "事实实体标识",
        "entity_identifier_matches_cik": "实体CIK匹配",
        "unit_ref": "单位编号",
        "unit_measure": "计量单位",
        "raw_value": "原始值",
        "value": "解析值",
        "decimals": "精度声明",
        "scale": "缩放声明",
        "dimensions_count": "维度数",
        "dimensions_json": "股类维度明细",
        "source_url": "来源地址",
        "filing_evidence_status": "申报证据状态",
        "identity_candidate_evidence_ready": "候选身份证据就绪",
        "identity_promoted": "身份规则已准入",
        "source_name": "来源名称",
        "version_id": "版本号",
        "rows": "行数",
        "note": "说明",
    }
    sheet_names = {
        "conclusion": "结论",
        "jobs": "SEC申报作业",
        "attempts": "下载检查点",
        "facts": "流通股候选事实",
        "selection_coverage": "逐重选身份覆盖",
        "unresolved": "未解决身份",
        "versions": "数据版本",
    }
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for key, frame in frames.items():
            frame.rename(columns=translations).to_excel(
                writer,
                sheet_name=sheet_names[key],
                index=False,
            )
    _style_workbook(path)


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Extract point-in-time dimensional share facts from exact SEC filings; "
            "never infer or promote a class/ADR allocation."
        )
    )
    parser.add_argument("--config", default="config/config.yaml")
    parser.add_argument("--valuation-dir", default="")
    parser.add_argument("--sic-dir", default="")
    parser.add_argument("--output-dir", default="")
    parser.add_argument("--max-filings", type=int, default=0)
    parser.add_argument("--request-delay", type=float, default=0.15)
    parser.add_argument("--timeout", type=float, default=45.0)
    parser.add_argument("--retries", type=int, default=3)
    parser.add_argument("--plan-only", action="store_true")
    parser.add_argument("--refresh", action="store_true")
    parser.add_argument("--retry-failed", action="store_true")
    args = parser.parse_args()

    cfg = load_config(_resolve(args.config))
    fundamental_dir = fundamental_cache_dir(cfg)
    valuation_dir = (
        _resolve(args.valuation_dir)
        if args.valuation_dir
        else fundamental_dir / "valuation_readiness_v3"
    )
    sic_dir = (
        _resolve(args.sic_dir)
        if args.sic_dir
        else fundamental_dir / "sic_history_v1"
    )
    output = (
        _resolve(args.output_dir)
        if args.output_dir
        else fundamental_dir / "filing_identity_v1"
    )
    output.mkdir(parents=True, exist_ok=True)
    raw_indexes = output / "raw_indexes"
    raw_filings = output / "raw_filings"
    readiness_path = valuation_dir / "valuation_readiness.csv"
    sic_path = sic_dir / "selection_sic.csv"
    readiness = pd.read_csv(readiness_path, low_memory=False)
    sic_selection = pd.read_csv(sic_path, low_memory=False)
    jobs, mapped = build_filing_identity_plan(readiness, sic_selection)
    attempts_path = output / "attempts.csv"
    attempts = _load_attempts(attempts_path)
    prior = attempts.set_index("job_id").to_dict(orient="index") if not attempts.empty else {}
    pending = []
    for row in jobs.itertuples(index=False):
        checkpoint = prior.get(str(row.job_id)) or {}
        status = str(checkpoint.get("status") or "")
        if args.refresh:
            pending.append(row)
        elif status in TERMINAL_STATUSES:
            continue
        elif status == "failed" and not args.retry_failed:
            continue
        else:
            pending.append(row)
    selected = (
        pending[: int(args.max_filings)]
        if int(args.max_filings) > 0
        else pending
    )
    print("NDX SEC多股类/ADR身份取证计划")
    print(
        f"阻断重选行: {len(mapped)} | CIK: {mapped['cik'].nunique()} | "
        f"申报作业: {len(jobs)} | 已检查: {len(jobs) - len(pending)} | "
        f"待执行: {len(pending)}"
    )
    print("仅抽取候选事实；不会自动相加股类、换算ADR或开放估值评分。")
    print(f"输出: {output}")
    if args.plan_only:
        return 0

    explicit_agent = os.environ.get("SEC_USER_AGENT")
    if not explicit_agent:
        raise ValueError(
            "SEC_USER_AGENT is required, for example: "
            "$env:SEC_USER_AGENT='us_ratchet/1.0 your_email@example.com'"
        )
    session = requests.Session()
    session.headers.update(_sec_headers(_sec_user_agent(explicit_agent)))
    durations = []
    for index, job in enumerate(selected, start=1):
        started = time.perf_counter()
        status = "failed"
        instance_filename = ""
        selection_method = ""
        instance_url = ""
        content = b""
        fact_count = 0
        error = ""
        try:
            index_content = _request_bytes(
                session,
                str(job.index_url),
                timeout=float(args.timeout),
                retries=int(args.retries),
            )
            index_path = (
                raw_indexes
                / str(int(job.cik))
                / f"{job.accession_number.replace('-', '')}.index.json"
            )
            _atomic_bytes(index_content, index_path)
            index_payload = json.loads(index_content.decode("utf-8"))
            try:
                selected_instance = select_xbrl_instance(index_payload)
            except Exception as exc:
                status = "no_explicit_xbrl_instance"
                error = f"{type(exc).__name__}: {exc}"
            else:
                instance_filename = str(selected_instance["name"])
                selection_method = str(selected_instance["selection_method"])
                instance_url = filing_document_url(
                    job.cik,
                    job.accession_number,
                    instance_filename,
                )
                time.sleep(max(0.0, float(args.request_delay)))
                content = _request_bytes(
                    session,
                    instance_url,
                    timeout=float(args.timeout),
                    retries=int(args.retries),
                )
                parsed = parse_filing_share_facts(
                    content,
                    cik=job.cik,
                    accession_number=job.accession_number,
                    source_url=instance_url,
                )
                fact_count = len(parsed)
                instance_path = _instance_path(
                    raw_filings,
                    cik=int(job.cik),
                    accession_number=str(job.accession_number),
                    filename=instance_filename,
                )
                _atomic_bytes(content, instance_path)
                write_dataset_manifest(
                    instance_path,
                    dataset_type="sec_filing_xbrl_instance",
                    dataset_key=str(job.accession_number),
                    source="sec_edgar_archives",
                    source_locator=instance_url,
                    row_count=fact_count,
                    data_start=(
                        str(parsed["period_end"].min()) if not parsed.empty else None
                    ),
                    data_end=(
                        str(parsed["period_end"].max()) if not parsed.empty else None
                    ),
                    extra={
                        "policy_id": FILING_IDENTITY_POLICY_ID,
                        "cik": int(job.cik),
                        "accession_number": str(job.accession_number),
                        "selection_method": selection_method,
                        "identity_promotion_allowed": False,
                    },
                )
                status = "ready" if fact_count else "no_share_facts"
        except Exception as exc:
            error = f"{type(exc).__name__}: {exc}"
            status = "failed"
        duration = time.perf_counter() - started
        row = {
            "policy_id": FILING_IDENTITY_POLICY_ID,
            "job_id": str(job.job_id),
            "cik": int(job.cik),
            "accession_number": str(job.accession_number),
            "status": status,
            "index_url": str(job.index_url),
            "instance_filename": instance_filename,
            "selection_method": selection_method,
            "instance_url": instance_url,
            "content_bytes": len(content),
            "content_sha256": hashlib.sha256(content).hexdigest() if content else "",
            "share_fact_count": fact_count,
            "error": error,
            "attempted_at": utc_now_iso(),
            "duration_seconds": round(duration, 3),
        }
        attempts = _upsert_attempt(attempts, row)
        _atomic_csv(attempts, attempts_path)
        durations.append(duration)
        remaining = len(selected) - index
        eta = (sum(durations) / len(durations)) * remaining if durations else 0.0
        if index == 1 or index % 10 == 0 or index == len(selected):
            print(
                f"SEC身份: {index}/{len(selected)} | {job.accession_number} | "
                f"{status} | facts={fact_count} | ETA {eta / 60:.1f}m",
                flush=True,
            )
        time.sleep(max(0.0, float(args.request_delay)))

    attempts = _load_attempts(attempts_path)
    facts, raw_versions = _cached_facts(raw_filings, attempts)
    attempt_view = attempts[[
        "job_id", "status", "instance_filename", "share_fact_count", "error"
    ]].copy()
    coverage = mapped.merge(
        jobs[["cik", "accession_number", "job_id"]],
        on=["cik", "accession_number"],
        how="left",
        validate="many_to_one",
    ).merge(attempt_view, on="job_id", how="left", validate="many_to_one")
    coverage["filing_evidence_status"] = coverage["status"].fillna("pending")
    coverage["share_fact_count"] = pd.to_numeric(
        coverage["share_fact_count"], errors="coerce"
    ).fillna(0).astype(int)
    coverage["identity_candidate_evidence_ready"] = coverage["share_fact_count"] > 0
    coverage["identity_promoted"] = False
    unresolved = coverage[
        ~coverage["identity_candidate_evidence_ready"]
    ].copy()
    attempts_by_job = attempts.set_index("job_id") if not attempts.empty else pd.DataFrame()
    statuses = {
        str(job.job_id): (
            str(attempts_by_job.at[str(job.job_id), "status"])
            if not attempts_by_job.empty and str(job.job_id) in attempts_by_job.index
            else "pending"
        )
        for job in jobs.itertuples(index=False)
    }
    failed_jobs = sum(status == "failed" for status in statuses.values())
    pending_jobs = sum(status == "pending" for status in statuses.values())
    audit_complete = failed_jobs == 0 and pending_jobs == 0
    versions = pd.DataFrame(raw_versions + [
        {
            "source_name": "valuation_readiness_v3",
            "version_id": "input_hash",
            "content_sha256": file_sha256(readiness_path),
            "rows": len(readiness),
            "note": readiness_path.as_posix(),
        },
        {
            "source_name": "sic_selection_v1",
            "version_id": "input_hash",
            "content_sha256": file_sha256(sic_path),
            "rows": len(sic_selection),
            "note": sic_path.as_posix(),
        },
    ])
    frames = {
        "jobs": jobs,
        "attempts": attempts,
        "facts": facts,
        "selection_coverage": coverage,
        "unresolved": unresolved,
        "versions": versions,
    }
    evidence = {}
    for key, frame in frames.items():
        path = output / f"{key}.csv"
        _atomic_csv(frame, path)
        manifest = write_dataset_manifest(
            path,
            dataset_type="ndx_filing_identity_readiness",
            dataset_key=f"{FILING_IDENTITY_POLICY_ID}:{key}",
            source="derived_from_exact_point_in_time_sec_filings",
            fetched_at_evidence="derived_runtime_utc",
            row_count=len(frame),
            data_start=str(mapped["selection_date"].min()),
            data_end=str(mapped["selection_date"].max()),
            extra={
                "policy_id": FILING_IDENTITY_POLICY_ID,
                "research_only": True,
                "audit_complete": audit_complete,
                "identity_promotion_allowed": False,
                "valuation_scoring_allowed": False,
            },
        )
        evidence[key] = {
            "filename": path.name,
            "rows": len(frame),
            "sha256": manifest["content_sha256"],
            "version_id": manifest["version_id"],
        }
    summary = {
        "schema_version": 1,
        "generated_at": utc_now_iso(),
        "policy_id": FILING_IDENTITY_POLICY_ID,
        "research_only": True,
        "selection_rows": len(mapped),
        "cik_count": int(mapped["cik"].nunique()),
        "filing_jobs": len(jobs),
        "terminal_jobs": sum(status in TERMINAL_STATUSES for status in statuses.values()),
        "failed_jobs": failed_jobs,
        "pending_jobs": pending_jobs,
        "share_fact_rows": len(facts),
        "candidate_evidence_selection_rows": int(
            coverage["identity_candidate_evidence_ready"].sum()
        ),
        "unresolved_selection_rows": len(unresolved),
        "audit_complete": audit_complete,
        "identity_promotion_allowed": False,
        "valuation_scoring_allowed": False,
        "next_step": (
            "review class dimensions and build explicit issuer/ADR allocation rules"
        ),
        "files": {
            **{key: value["filename"] for key, value in evidence.items()},
            "workbook": "ndx_filing_identity_readiness.xlsx",
        },
    }
    atomic_write_text(
        output / "summary.json",
        json.dumps(summary, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
    )
    atomic_write_text(
        output / "manifest.json",
        json.dumps(
            {
                "schema_version": 1,
                "generated_at": summary["generated_at"],
                "policy_id": FILING_IDENTITY_POLICY_ID,
                "valuation_readiness_sha256": file_sha256(readiness_path),
                "sic_selection_sha256": file_sha256(sic_path),
                "audit_complete": audit_complete,
                "evidence_files": evidence,
                "identity_promotion_allowed": False,
                "valuation_scoring_allowed": False,
            },
            ensure_ascii=False,
            indent=2,
            allow_nan=False,
        ) + "\n",
    )
    conclusion = pd.DataFrame([
        {"项目": "研究目的", "内容": "抽取实际历史申报中的股类维度与流通股候选事实"},
        {"项目": "阻断重选行", "内容": len(mapped)},
        {"项目": "涉及CIK", "内容": mapped["cik"].nunique()},
        {"项目": "SEC申报作业", "内容": len(jobs)},
        {"项目": "候选事实行", "内容": len(facts)},
        {
            "项目": "有候选证据重选行",
            "内容": int(coverage["identity_candidate_evidence_ready"].sum()),
        },
        {"项目": "未解决重选行", "内容": len(unresolved)},
        {"项目": "下载/解析完整", "内容": audit_complete},
        {"项目": "身份规则准入", "内容": "关闭；候选事实不得自动相加或换算"},
        {"项目": "估值评分/实盘", "内容": "关闭；不改配置、不改今日建议"},
    ])
    workbook = output / "ndx_filing_identity_readiness.xlsx"
    _write_workbook(
        workbook,
        {
            "conclusion": conclusion,
            "jobs": jobs,
            "attempts": attempts,
            "facts": facts,
            "selection_coverage": coverage,
            "unresolved": unresolved,
            "versions": versions,
        },
    )
    print("\nSEC多股类/ADR身份证据包")
    print(
        f"作业完成: {summary['terminal_jobs']}/{len(jobs)} | 失败: {failed_jobs} | "
        f"候选事实: {len(facts)} | 未解决重选行: {len(unresolved)}"
    )
    print("身份准入与估值评分仍关闭。")
    print(f"Workbook: {workbook}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
