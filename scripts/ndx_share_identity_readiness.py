"""Resolve and crosscheck filing-level share facts without opening valuation scoring."""
from __future__ import annotations

import argparse
import json
import os
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from core.config import load_config
from core.helpers import atomic_write_text
from data.ndx_filing_identity import (
    FILING_IDENTITY_POLICY_ID,
    FILING_SHARE_RESOLUTION_POLICY_ID,
    FilingShareResolutionPolicy,
    resolve_filing_share_rows,
)
from data.ndx_fundamentals import fundamental_cache_dir
from data.ndx_valuation_readiness import VALUATION_READINESS_V3_POLICY_ID
from data.versioning import file_sha256, manifest_path, utc_now_iso, write_dataset_manifest


def _resolve(path: str | Path) -> Path:
    value = Path(path)
    return value if value.is_absolute() else ROOT / value


def _atomic_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        frame.to_csv(temporary, index=False, encoding="utf-8-sig")
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def _verified_csv(directory: Path, root: dict, key: str) -> tuple[pd.DataFrame, dict]:
    evidence = (root.get("evidence_files") or {}).get(key) or {}
    path = directory / str(evidence.get("filename") or "")
    sidecar = manifest_path(path)
    if not path.exists() or not sidecar.exists():
        raise ValueError(f"versioned evidence is missing: {key}")
    manifest = json.loads(sidecar.read_text(encoding="utf-8"))
    expected = str(evidence.get("sha256") or "")
    if file_sha256(path) != expected or manifest.get("content_sha256") != expected:
        raise ValueError(f"versioned evidence hash changed: {key}")
    return pd.read_csv(path, low_memory=False), manifest


def _style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    fill = PatternFill("solid", fgColor="DCE6F1")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            width = max(len(str(cell.value or "")) for cell in column[:300]) + 2
            sheet.column_dimensions[letter].width = min(max(width, 10), 48)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Resolve exact SEC filing share facts; promote only safe domestic rows."
    )
    parser.add_argument("--config", default="config/config.yaml")
    parser.add_argument("--filing-identity-dir", default="")
    parser.add_argument("--valuation-dir", default="")
    parser.add_argument("--output-dir", default="")
    parser.add_argument("--maximum-share-age-days", type=int, default=200)
    parser.add_argument("--existing-value-tolerance", type=float, default=0.05)
    args = parser.parse_args()

    cfg = load_config(_resolve(args.config))
    fundamental_dir = fundamental_cache_dir(cfg)
    filing_dir = (
        _resolve(args.filing_identity_dir)
        if args.filing_identity_dir
        else fundamental_dir / "filing_identity_v1"
    )
    valuation_dir = (
        _resolve(args.valuation_dir)
        if args.valuation_dir
        else fundamental_dir / "valuation_readiness_v3"
    )
    output = (
        _resolve(args.output_dir)
        if args.output_dir
        else fundamental_dir / "share_identity_v1"
    )
    output.mkdir(parents=True, exist_ok=True)

    filing_root = json.loads((filing_dir / "manifest.json").read_text(encoding="utf-8"))
    if filing_root.get("policy_id") != FILING_IDENTITY_POLICY_ID:
        raise ValueError("filing identity policy changed")
    if filing_root.get("audit_complete") is not True:
        raise ValueError("filing identity audit is incomplete")
    facts, facts_manifest = _verified_csv(filing_dir, filing_root, "facts")
    jobs, jobs_manifest = _verified_csv(filing_dir, filing_root, "jobs")
    identity, identity_manifest = _verified_csv(
        filing_dir,
        filing_root,
        "selection_coverage",
    )
    valuation_root = json.loads(
        (valuation_dir / "manifest.json").read_text(encoding="utf-8")
    )
    if valuation_root.get("policy_id") != VALUATION_READINESS_V3_POLICY_ID:
        raise ValueError("valuation readiness policy changed")
    if valuation_root.get("integrity_passed") is not True:
        raise ValueError("valuation readiness integrity gate failed")
    readiness, readiness_manifest = _verified_csv(
        valuation_dir,
        valuation_root,
        "readiness",
    )
    policy = FilingShareResolutionPolicy(
        maximum_domestic_share_age_days=int(args.maximum_share_age_days),
        existing_value_tolerance=float(args.existing_value_tolerance),
    )
    rows, comparisons, promotions, failures = resolve_filing_share_rows(
        identity,
        facts,
        jobs,
        readiness,
        policy=policy,
    )
    blockers = (
        rows.groupby("promotion_blocker", dropna=False, as_index=False)
        .agg(
            selection_rows=("ticker", "size"),
            ticker_count=("ticker", "nunique"),
            tickers=("ticker", lambda values: ";".join(sorted(set(values)))),
        )
        .sort_values(["selection_rows", "promotion_blocker"], ascending=[False, True])
    )
    by_date = rows.groupby("selection_date", as_index=False).agg(
        audited_rows=("ticker", "size"),
        candidate_ready_rows=("resolution_status", lambda values: int((values == "ready_candidate").sum())),
        promoted_rows=("promotion_allowed", "sum"),
        existing_crosscheck_rows=("existing_shares_ready", "sum"),
    )
    by_date["promoted_coverage_of_audited_rows"] = (
        by_date["promoted_rows"] / by_date["audited_rows"]
    )
    versions = pd.DataFrame([
        {
            "source_name": "filing_share_facts",
            "version_id": facts_manifest.get("version_id"),
            "content_sha256": facts_manifest.get("content_sha256"),
            "rows": len(facts),
            "note": "SEC XBRL逐事实候选；仅消费严格存量概念",
        },
        {
            "source_name": "filing_jobs",
            "version_id": jobs_manifest.get("version_id"),
            "content_sha256": jobs_manifest.get("content_sha256"),
            "rows": len(jobs),
            "note": "实际历史重选使用的filing可用日",
        },
        {
            "source_name": "filing_identity_selection",
            "version_id": identity_manifest.get("version_id"),
            "content_sha256": identity_manifest.get("content_sha256"),
            "rows": len(identity),
            "note": "598个身份阻断重选行",
        },
        {
            "source_name": "valuation_readiness_v3",
            "version_id": readiness_manifest.get("version_id"),
            "content_sha256": readiness_manifest.get("content_sha256"),
            "rows": len(readiness),
            "note": "与既有Company Facts股数交叉核验",
        },
    ])
    frames = {
        "resolved_rows": rows,
        "existing_crosschecks": comparisons,
        "promotions": promotions,
        "blockers": blockers,
        "coverage_by_date": by_date,
        "integrity_failures": failures,
        "versions": versions,
    }
    promotion_allowed = bool(failures.empty and not promotions.empty)
    evidence = {}
    for key, frame in frames.items():
        path = output / f"{key}.csv"
        _atomic_csv(frame, path)
        manifest = write_dataset_manifest(
            path,
            dataset_type="ndx_share_identity_readiness",
            dataset_key=f"{FILING_SHARE_RESOLUTION_POLICY_ID}:{key}",
            source="derived_from_exact_sec_xbrl_share_facts",
            fetched_at_evidence="derived_runtime_utc",
            row_count=len(frame),
            data_start=str(rows["selection_date"].min()),
            data_end=str(rows["selection_date"].max()),
            extra={
                "policy_id": FILING_SHARE_RESOLUTION_POLICY_ID,
                "research_only": True,
                "promotion_allowed": promotion_allowed,
                "valuation_scoring_allowed": False,
            },
        )
        evidence[key] = {
            "filename": path.name,
            "rows": len(frame),
            "sha256": manifest["content_sha256"],
            "version_id": manifest["version_id"],
        }
    maximum_difference = float(
        pd.to_numeric(comparisons["shares_relative_difference"], errors="coerce").max()
    )
    summary = {
        "schema_version": 1,
        "generated_at": utc_now_iso(),
        "policy_id": FILING_SHARE_RESOLUTION_POLICY_ID,
        "policy": policy.__dict__,
        "research_only": True,
        "audited_selection_rows": len(rows),
        "candidate_ready_rows": int(rows["resolution_status"].eq("ready_candidate").sum()),
        "existing_crosscheck_rows": len(comparisons),
        "maximum_existing_relative_difference": maximum_difference,
        "promoted_selection_rows": len(promotions),
        "promoted_ticker_count": int(promotions["ticker"].nunique()),
        "promoted_tickers": sorted(promotions["ticker"].unique()),
        "foreign_blocked_rows": int(rows["promotion_blocker"].eq("foreign_filer_fx_or_adr_pending").sum()),
        "multi_security_blocked_rows": int(rows["promotion_blocker"].eq("multi_security_market_cap_allocation_pending").sum()),
        "integrity_failure_count": len(failures),
        "integrity_passed": failures.empty,
        "promotion_allowed": promotion_allowed,
        "valuation_scoring_allowed": False,
        "files": {
            **{key: value["filename"] for key, value in evidence.items()},
            "workbook": "ndx_share_identity_readiness.xlsx",
        },
    }
    atomic_write_text(
        output / "summary.json",
        json.dumps(summary, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
    )
    atomic_write_text(
        output / "manifest.json",
        json.dumps(
            {
                "schema_version": 1,
                "generated_at": summary["generated_at"],
                "policy_id": FILING_SHARE_RESOLUTION_POLICY_ID,
                "upstream_policy_id": FILING_IDENTITY_POLICY_ID,
                "source_facts_sha256": facts_manifest.get("content_sha256"),
                "source_jobs_sha256": jobs_manifest.get("content_sha256"),
                "source_identity_sha256": identity_manifest.get("content_sha256"),
                "source_readiness_sha256": readiness_manifest.get("content_sha256"),
                "integrity_passed": failures.empty,
                "promotion_allowed": promotion_allowed,
                "valuation_scoring_allowed": False,
                "evidence_files": evidence,
            },
            ensure_ascii=False,
            indent=2,
            allow_nan=False,
        ) + "\n",
    )
    conclusion = pd.DataFrame([
        {"项目": "研究目的", "内容": "从精确SEC filing解析并交叉核验点时发行人股数"},
        {"项目": "审计重选行", "内容": len(rows)},
        {"项目": "候选股数就绪行", "内容": summary["candidate_ready_rows"]},
        {"项目": "旧股数交叉核验行", "内容": len(comparisons)},
        {"项目": "旧股数最大相对差异", "内容": maximum_difference},
        {"项目": "安全提升行", "内容": len(promotions)},
        {"项目": "安全提升股票", "内容": ";".join(summary["promoted_tickers"])},
        {"项目": "境外/ADR继续阻断行", "内容": summary["foreign_blocked_rows"]},
        {"项目": "多证券发行人继续阻断行", "内容": summary["multi_security_blocked_rows"]},
        {"项目": "完整性失败", "内容": len(failures)},
        {"项目": "允许估值层消费提升", "内容": promotion_allowed},
        {"项目": "估值评分/实盘", "内容": "仍关闭；不改配置、不改今日建议"},
    ])
    translations = {
        "selection_date": "重选日", "ticker": "股票代码", "cik": "SEC CIK",
        "existing_shares_ready": "旧股数就绪", "existing_shares_value": "旧股数值",
        "resolved_shares_value": "filing解析股数", "shares_relative_difference": "股数相对差异",
        "existing_crosscheck_passed": "旧股数交叉核验通过",
        "source_accession_number": "来源申报编号", "source_form": "来源表格",
        "source_available_date": "来源可用日", "shares_period_end": "股数事实日",
        "shares_age_days": "股数年龄天数", "shares_age_passed": "股数新鲜度通过",
        "share_concept": "股数概念", "share_resolution_method": "股数解析方法",
        "common_class_members": "普通股类成员", "foreign_filer_asof": "境外申报人",
        "multi_security_issuer": "多证券发行人", "resolution_status": "解析状态",
        "promotion_allowed": "允许提升", "promotion_blocker": "提升阻断",
        "policy_id": "规则编号", "selection_rows": "重选行数",
        "ticker_count": "股票数", "tickers": "股票代码集合",
        "audited_rows": "审计行数", "candidate_ready_rows": "候选就绪行数",
        "promoted_rows": "提升行数", "existing_crosscheck_rows": "旧值核验行数",
        "promoted_coverage_of_audited_rows": "审计行提升率",
        "scope": "范围", "reason": "原因", "source_name": "来源名称",
        "version_id": "版本号", "content_sha256": "内容哈希", "rows": "行数",
        "note": "说明",
    }
    workbook = output / "ndx_share_identity_readiness.xlsx"
    with pd.ExcelWriter(workbook, engine="openpyxl") as writer:
        conclusion.to_excel(writer, sheet_name="结论", index=False)
        rows.rename(columns=translations).to_excel(writer, sheet_name="逐期股数解析", index=False)
        comparisons.rename(columns=translations).to_excel(writer, sheet_name="旧股数交叉核验", index=False)
        promotions.rename(columns=translations).to_excel(writer, sheet_name="安全提升", index=False)
        blockers.rename(columns=translations).to_excel(writer, sheet_name="提升阻断", index=False)
        by_date.rename(columns=translations).to_excel(writer, sheet_name="逐期覆盖", index=False)
        failures.rename(columns=translations).to_excel(writer, sheet_name="完整性失败", index=False)
        versions.rename(columns=translations).to_excel(writer, sheet_name="数据版本", index=False)
    _style_workbook(workbook)
    print("NDX filing股数身份准入（research only）")
    print(
        f"审计: {len(rows)} | 旧值核验: {len(comparisons)} | "
        f"最大差异: {maximum_difference:.2%} | 完整性失败: {len(failures)}"
    )
    print(
        f"安全提升: {len(promotions)} 行/{summary['promoted_ticker_count']} 只 | "
        f"境外阻断: {summary['foreign_blocked_rows']} | "
        f"多证券阻断: {summary['multi_security_blocked_rows']}"
    )
    print(f"估值层可消费提升: {promotion_allowed} | 评分仍关闭")
    print(f"Workbook: {workbook}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
