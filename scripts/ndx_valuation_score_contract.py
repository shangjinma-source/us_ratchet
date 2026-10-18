"""Freeze and audit a neutral-missing FCF Yield score contract."""
from __future__ import annotations

import argparse
import json
import os
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from core.config import load_config
from core.helpers import atomic_write_text
from data.ndx_filing_identity import FILING_SHARE_RESOLUTION_POLICY_ID
from data.ndx_fundamentals import fundamental_cache_dir
from data.ndx_valuation_diagnostics import VALUATION_DIAGNOSTIC_POLICY_ID
from data.ndx_valuation_score import (
    VALUATION_OVERPAY_GUARD_POLICY_ID,
    VALUATION_SCORE_POLICY_ID,
    ValuationScorePolicy,
    build_neutral_missing_scores,
    build_overpay_guard_scores,
)
from data.versioning import file_sha256, manifest_path, utc_now_iso, write_dataset_manifest


def _resolve(path: str | Path) -> Path:
    value = Path(path)
    return value if value.is_absolute() else ROOT / value


def _atomic_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        frame.to_csv(temporary, index=False, encoding="utf-8-sig")
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def _verified_values(directory: Path) -> tuple[pd.DataFrame, dict, dict]:
    root = json.loads((directory / "manifest.json").read_text(encoding="utf-8"))
    if root.get("policy_id") != VALUATION_DIAGNOSTIC_POLICY_ID:
        raise ValueError("valuation diagnostic policy identity changed")
    if root.get("integrity_passed") is not True:
        raise ValueError("valuation diagnostic integrity gate failed")
    if root.get("training_spec_allowed") is not False:
        raise ValueError("diagnostic input unexpectedly enables training")
    evidence = (root.get("evidence_files") or {}).get("values") or {}
    path = directory / str(evidence.get("filename") or "")
    sidecar = manifest_path(path)
    manifest = json.loads(sidecar.read_text(encoding="utf-8"))
    if (
        not path.exists()
        or file_sha256(path) != evidence.get("sha256")
        or manifest.get("content_sha256") != evidence.get("sha256")
    ):
        raise ValueError("valuation diagnostic values hash mismatch")
    return pd.read_csv(path, low_memory=False), manifest, root


def _identity_gate(directory: Path) -> dict:
    path = directory / "manifest.json"
    if not path.exists():
        return {
            "audit_complete": False,
            "identity_promotion_allowed": False,
            "reason": "filing_identity_manifest_missing",
        }
    payload = json.loads(path.read_text(encoding="utf-8"))
    policy_matches = payload.get("policy_id") == FILING_SHARE_RESOLUTION_POLICY_ID
    integrity_passed = bool(payload.get("integrity_passed"))
    promotion_allowed = bool(payload.get("promotion_allowed"))
    return {
        "audit_complete": bool(policy_matches and integrity_passed),
        "identity_promotion_allowed": bool(policy_matches and promotion_allowed),
        "reason": (
            "share_identity_promotions_ready_unresolved_stay_neutral"
            if policy_matches and integrity_passed and promotion_allowed
            else "share_identity_not_promoted"
        ),
        "manifest_sha256": file_sha256(path),
    }


def _style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    fill = PatternFill("solid", fgColor="DCE6F1")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            width = max(len(str(cell.value or "")) for cell in column[:300]) + 2
            sheet.column_dimensions[letter].width = min(max(width, 10), 48)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Build neutral-missing valuation score previews; never backtest."
    )
    parser.add_argument("--config", default="config/config.yaml")
    parser.add_argument("--input-dir", default="")
    parser.add_argument("--identity-dir", default="")
    parser.add_argument("--output-dir", default="")
    parser.add_argument("--winsor-lower", type=float, default=0.05)
    parser.add_argument("--winsor-upper", type=float, default=0.95)
    parser.add_argument("--minimum-coverage", type=float, default=0.70)
    parser.add_argument("--minimum-observations", type=int, default=20)
    parser.add_argument(
        "--score-mode", choices=("full_rank", "overpay_guard"), default="full_rank"
    )
    parser.add_argument("--overpay-penalty-quantile", type=float, default=0.30)
    args = parser.parse_args()

    cfg = load_config(_resolve(args.config))
    fundamental_dir = fundamental_cache_dir(cfg)
    source = (
        _resolve(args.input_dir)
        if args.input_dir
        else fundamental_dir / "valuation_diagnostics_v3"
    )
    identity_dir = (
        _resolve(args.identity_dir)
        if args.identity_dir
        else fundamental_dir / "share_identity_v1"
    )
    default_output_name = (
        "valuation_score_overpay_guard_v1"
        if args.score_mode == "overpay_guard"
        else "valuation_score_contract_v2"
    )
    output = (
        _resolve(args.output_dir)
        if args.output_dir
        else fundamental_dir / default_output_name
    )
    output.mkdir(parents=True, exist_ok=True)
    values, values_manifest, values_root = _verified_values(source)
    policy = ValuationScorePolicy(
        winsor_lower=float(args.winsor_lower),
        winsor_upper=float(args.winsor_upper),
        minimum_factor_coverage=float(args.minimum_coverage),
        minimum_observations=int(args.minimum_observations),
    )
    if args.score_mode == "overpay_guard":
        score_policy_id = VALUATION_OVERPAY_GUARD_POLICY_ID
        scores, coverage = build_overpay_guard_scores(
            values,
            policy=policy,
            penalty_quantile=float(args.overpay_penalty_quantile),
        )
    else:
        score_policy_id = VALUATION_SCORE_POLICY_ID
        scores, coverage = build_neutral_missing_scores(values, policy=policy)
    identity_gate = _identity_gate(identity_dir)
    training_allowed = bool(
        identity_gate["audit_complete"]
        and identity_gate["identity_promotion_allowed"]
    )
    scores["training_spec_generation_allowed"] = training_allowed
    failures = []
    missing = scores["valuation_missing_neutral_applied"].astype(bool)
    if not scores.loc[missing, "valuation_score"].eq(
        policy.missing_neutral_score
    ).all():
        failures.append({"scope": "missing_neutral", "reason": "score_not_exactly_neutral"})
    if scores["return_data_read"].astype(bool).any():
        failures.append({"scope": "return_guard", "reason": "return_data_was_read"})
    if scores["ev_sales_included_in_score"].astype(bool).any():
        failures.append({"scope": "factor_scope", "reason": "ev_sales_was_included"})
    failures_frame = pd.DataFrame(failures, columns=["scope", "reason"])
    versions = pd.DataFrame([
        {
            "source_name": "valuation_diagnostics_v3",
            "version_id": values_manifest.get("version_id"),
            "content_sha256": values_manifest.get("content_sha256"),
            "rows": len(values),
            "note": "仅消费完整性通过的值域诊断，不读取收益",
        },
        {
            "source_name": "filing_identity_gate",
            "version_id": identity_gate.get("reason"),
            "content_sha256": identity_gate.get("manifest_sha256"),
            "rows": 0,
            "note": "身份审计与显式分摊规则未准入前禁止生成训练规格",
        },
    ])
    frames = {
        "scores": scores,
        "coverage": coverage,
        "integrity_failures": failures_frame,
        "versions": versions,
    }
    evidence = {}
    for key, frame in frames.items():
        path = output / f"{key}.csv"
        _atomic_csv(frame, path)
        manifest = write_dataset_manifest(
            path,
            dataset_type="ndx_valuation_score_contract",
            dataset_key=f"{score_policy_id}:{key}",
            source="derived_from_versioned_valuation_diagnostics",
            fetched_at_evidence="derived_runtime_utc",
            row_count=len(frame),
            data_start=str(scores["selection_date"].min()),
            data_end=str(scores["selection_date"].max()),
            extra={
                "policy_id": score_policy_id,
                "return_data_read": False,
                "training_spec_generation_allowed": training_allowed,
                "live_scoring_allowed": False,
            },
        )
        evidence[key] = {
            "filename": path.name,
            "rows": len(frame),
            "sha256": manifest["content_sha256"],
            "version_id": manifest["version_id"],
        }
    active_dates = int(coverage["factor_active"].astype(bool).sum())
    summary = {
        "schema_version": 1,
        "generated_at": utc_now_iso(),
        "policy_id": score_policy_id,
        "score_mode": args.score_mode,
        "overpay_penalty_quantile": (
            float(args.overpay_penalty_quantile)
            if args.score_mode == "overpay_guard" else None
        ),
        "policy": policy.__dict__,
        "research_only": True,
        "selection_rows": len(scores),
        "selection_dates": int(scores["selection_date"].nunique()),
        "active_factor_dates": active_dates,
        "minimum_fcf_yield_coverage": float(coverage["fcf_yield_coverage"].min()),
        "latest_fcf_yield_coverage": float(coverage.iloc[-1]["fcf_yield_coverage"]),
        "missing_neutral_rows": int(missing.sum()),
        "ev_sales_included_in_score": False,
        "return_data_read": False,
        "integrity_failure_count": len(failures_frame),
        "integrity_passed": failures_frame.empty,
        "identity_gate": identity_gate,
        "training_spec_generation_allowed": training_allowed,
        "live_scoring_allowed": False,
        "files": {
            **{key: value["filename"] for key, value in evidence.items()},
            "workbook": "ndx_valuation_score_contract.xlsx",
        },
    }
    atomic_write_text(
        output / "summary.json",
        json.dumps(summary, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
    )
    atomic_write_text(
        output / "manifest.json",
        json.dumps(
            {
                "schema_version": 1,
                "generated_at": summary["generated_at"],
                "policy_id": score_policy_id,
                "source_policy_id": values_root.get("policy_id"),
                "source_values_sha256": values_manifest.get("content_sha256"),
                "identity_gate": identity_gate,
                "integrity_passed": failures_frame.empty,
                "return_data_read": False,
                "training_spec_generation_allowed": training_allowed,
                "live_scoring_allowed": False,
                "evidence_files": evidence,
            },
            ensure_ascii=False,
            indent=2,
            allow_nan=False,
        ) + "\n",
    )
    conclusion = pd.DataFrame([
        {
            "项目": "研究目的",
            "内容": (
                "冻结FCF Yield单侧过贵惩罚契约"
                if args.score_mode == "overpay_guard"
                else "冻结FCF Yield缺失值中性评分契约"
            ),
        },
        {"项目": "逐重选行", "内容": len(scores)},
        {"项目": "重选日", "内容": scores["selection_date"].nunique()},
        {"项目": "因子启用日", "内容": active_dates},
        {"项目": "最低FCF Yield覆盖率", "内容": summary["minimum_fcf_yield_coverage"]},
        {"项目": "缺失中性行", "内容": summary["missing_neutral_rows"]},
        {"项目": "缺失中性分", "内容": policy.missing_neutral_score},
        {
            "项目": "低估值奖励",
            "内容": "关闭；只惩罚低FCF Yield尾部"
            if args.score_mode == "overpay_guard" else "开启全截面排序",
        },
        {"项目": "EV/Sales", "内容": "覆盖不足，仅诊断，不进入分数"},
        {"项目": "读取收益数据", "内容": False},
        {"项目": "生成训练规格", "内容": training_allowed},
        {"项目": "阻断原因", "内容": identity_gate["reason"]},
        {"项目": "实盘评分", "内容": "关闭；不改配置、不改今日建议"},
    ])
    translations = {
        "selection_date": "重选日",
        "ticker": "股票代码",
        "cik": "SEC CIK",
        "fcf_yield_value": "FCF收益率原值",
        "fcf_yield_observed": "FCF收益率已观测",
        "fcf_yield_missing": "FCF收益率缺失",
        "fcf_yield_winsorized": "FCF收益率缩尾值",
        "fcf_yield_score": "FCF收益率分数",
        "valuation_score": "估值分数",
        "valuation_factor_active": "估值因子启用",
        "valuation_factor_coverage": "估值因子覆盖率",
        "valuation_missing_neutral_score": "缺失中性分",
        "valuation_missing_neutral_applied": "已应用缺失中性",
        "valuation_policy_id": "估值规则编号",
        "fcf_yield_full_rank_diagnostic": "FCF收益率全排名仅诊断",
        "fcf_yield_overpay_guard_score": "FCF收益率过贵守门分",
        "valuation_overpay_penalty_quantile": "过贵惩罚分位",
        "valuation_cheap_reward_allowed": "允许低估值奖励",
        "ev_sales_included_in_score": "EV销售额纳入分数",
        "return_data_read": "读取收益数据",
        "training_spec_generation_allowed": "允许生成训练规格",
        "live_scoring_allowed": "允许实盘评分",
        "universe_count": "候选数",
        "fcf_yield_observed_count": "FCF收益率观测数",
        "fcf_yield_coverage": "FCF收益率覆盖率",
        "factor_active": "因子启用",
        "winsor_lower_value": "缩尾下界",
        "winsor_upper_value": "缩尾上界",
        "missing_neutral_count": "缺失中性数",
        "missing_neutral_score": "缺失中性分",
        "observed_score_median": "观测分数中位数",
        "neutral_matches_observed_median": "中性分等于观测中位数",
        "ev_sales_ready_count_diagnostic_only": "EV销售额就绪数仅诊断",
        "ev_sales_coverage_diagnostic_only": "EV销售额覆盖率仅诊断",
        "overpay_penalty_quantile": "过贵惩罚分位",
        "cheap_reward_allowed": "允许低估值奖励",
        "observed_guard_score_median": "观测守门分中位数",
        "scope": "范围",
        "reason": "原因",
        "source_name": "来源名称",
        "version_id": "版本号",
        "content_sha256": "内容哈希",
        "rows": "行数",
        "note": "说明",
    }
    workbook = output / "ndx_valuation_score_contract.xlsx"
    with pd.ExcelWriter(workbook, engine="openpyxl") as writer:
        conclusion.to_excel(writer, sheet_name="结论", index=False)
        scores.rename(columns=translations).to_excel(
            writer, sheet_name="逐股估值预览分", index=False
        )
        coverage.rename(columns=translations).to_excel(
            writer, sheet_name="逐期覆盖审计", index=False
        )
        failures_frame.rename(columns=translations).to_excel(
            writer, sheet_name="完整性失败", index=False
        )
        versions.rename(columns=translations).to_excel(
            writer, sheet_name="数据版本", index=False
        )
    _style_workbook(workbook)
    print("NDX估值缺失中性契约（research only）")
    print(
        f"行: {len(scores)} | 因子启用日: {active_dates}/{len(coverage)} | "
        f"最低/最新覆盖: {summary['minimum_fcf_yield_coverage']:.2%}/"
        f"{summary['latest_fcf_yield_coverage']:.2%}"
    )
    print(
        f"缺失中性行: {summary['missing_neutral_rows']} | 完整性失败: "
        f"{len(failures_frame)} | 生成训练规格: {training_allowed}"
    )
    print(f"Workbook: {workbook}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
