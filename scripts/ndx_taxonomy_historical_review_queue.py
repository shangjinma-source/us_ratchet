"""Generate a hash-bound, non-promoting historical taxonomy review queue."""
from __future__ import annotations

import argparse
import json
import os
import sys
from pathlib import Path

import pandas as pd
import yaml

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from core.helpers import atomic_write_text  # noqa: E402
from data.ndx_taxonomy import read_versioned_csv  # noqa: E402
from data.ndx_taxonomy_review_queue import build_historical_review_queue  # noqa: E402
from data.versioning import file_sha256, utc_now_iso, write_dataset_manifest  # noqa: E402


COLUMN_LABELS = {
    "evidence_id": "证据编号", "issuer_id": "发行人编号", "tickers": "股票代码",
    "form": "年度申报类型", "accession_number": "SEC申报编号",
    "evidence_available_date": "证据可用日", "extraction_status": "章节提取状态",
    "raw_manifest_verified": "原始文件哈希通过", "raw_sha256": "原始文件SHA256",
    "business_excerpt_sha256": "业务摘录SHA256", "section_text_chars": "业务章节字数",
    "historical_evidence_excerpt": "历史业务证据摘录",
    "historical_evidence_reviewable": "历史证据可人工审核",
    "current_hint_status": "当前分类提示状态", "current_hint_risk_chain": "当前风险链提示（不可回填）",
    "current_hint_confidence": "当前提示置信度", "current_hint_evidence_phrases": "当前提示证据短语",
    "current_hint_decision_ids": "当前提示决策编号", "review_status": "历史审核状态",
    "risk_chain_assignment_performed": "已执行风险链分类",
    "historical_backfill_from_current_allowed": "允许当前标签回填历史",
    "return_data_read": "读取收益数据",
}


def _resolve(value: str | Path) -> Path:
    path = Path(value)
    return path if path.is_absolute() else ROOT / path


def _atomic_csv(frame: pd.DataFrame, path: Path) -> None:
    temporary = path.with_name(path.name + ".tmp")
    frame.to_csv(temporary, index=False, encoding="utf-8-sig", lineterminator="\n")
    os.replace(temporary, path)


def _localized(frame: pd.DataFrame) -> pd.DataFrame:
    return frame.rename(columns={column: COLUMN_LABELS.get(column, column) for column in frame.columns})


def _style_workbook(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(path)
    fill = PatternFill("solid", fgColor="DCE6F1")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            width = max(len(str(cell.value or "")) for cell in list(column)[:200]) + 2
            sheet.column_dimensions[letter].width = min(max(width, 10), 48)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)


def build_queue(output_dir: Path, evidence_path: Path, documents_path: Path, registry_path: Path) -> dict:
    output_dir.mkdir(parents=True, exist_ok=True)
    evidence, evidence_version = read_versioned_csv(evidence_path, name="historical_reviewed_evidence", used_for="历史SEC证据待审核状态")
    documents, documents_version = read_versioned_csv(documents_path, name="historical_review_documents", used_for="历史SEC原文与摘录哈希")
    registry = yaml.safe_load(registry_path.read_text(encoding="utf-8")) or {}
    queue = build_historical_review_queue(evidence, documents, registry)
    queue_path = output_dir / "historical_review_queue.csv"
    _atomic_csv(queue, queue_path)
    queue_manifest = write_dataset_manifest(
        queue_path,
        dataset_type="ndx_taxonomy_historical_review_queue",
        dataset_key="ndx_taxonomy_historical_review_queue_v1",
        source="derived_from_manifest_verified_historical_sec_evidence",
        row_count=len(queue),
        extra={"research_only": True, "risk_chain_assignment_performed": False, "historical_backfill_allowed": False, "return_data_read": False},
    )
    workbook = output_dir / "ndx_taxonomy_historical_review_queue.xlsx"
    with pd.ExcelWriter(workbook, engine="openpyxl") as writer:
        pd.DataFrame([
            {"项目": "用途", "内容": "逐条复核历史年度报告；当前风险链仅作路由提示，不能直接回填历史。"},
            {"项目": "晋级", "内容": "本产物不生成 accepted evidence、不修改决策 registry。"},
            {"项目": "数据", "内容": "不读取价格、收益或回测结果。"},
        ]).to_excel(writer, sheet_name="使用说明", index=False)
        _localized(queue).to_excel(writer, sheet_name="历史证据审核队列", index=False)
    _style_workbook(workbook)
    workbook_manifest = write_dataset_manifest(
        workbook,
        dataset_type="ndx_taxonomy_historical_review_queue_workbook",
        dataset_key="ndx_taxonomy_historical_review_queue_v1",
        source="derived_from_manifest_verified_historical_sec_evidence",
        row_count=len(queue),
        extra={"research_only": True, "risk_chain_assignment_performed": False, "historical_backfill_allowed": False, "return_data_read": False},
    )
    summary = {
        "schema_version": 1,
        "generated_at": utc_now_iso(),
        "queue_rows": len(queue),
        "historical_evidence_reviewable_rows": int(queue["historical_evidence_reviewable"].sum()),
        "current_hint_rows": int(queue["current_hint_status"].eq("current_hint_available_for_manual_routing_only").sum()),
        "source_extraction_repair_rows": int(queue["review_status"].eq("source_extraction_repair_required").sum()),
        "risk_chain_assignment_performed": False,
        "historical_backfill_allowed": False,
        "return_data_read": False,
        "candidate_admission_allowed": False,
        "inputs": [evidence_version.to_dict(), documents_version.to_dict(), {"path": registry_path.as_posix(), "sha256": file_sha256(registry_path), "used_for": "仅作人工审核路由提示"}],
        "artifacts": {"queue": {"path": queue_path.as_posix(), "sha256": queue_manifest["content_sha256"]}, "workbook": {"path": workbook.as_posix(), "sha256": workbook_manifest["content_sha256"]}},
    }
    manifest_path = output_dir / "manifest.json"
    atomic_write_text(manifest_path, json.dumps(summary, ensure_ascii=False, indent=2, allow_nan=False) + "\n")
    write_dataset_manifest(
        manifest_path,
        dataset_type="ndx_taxonomy_historical_review_queue_manifest",
        dataset_key="ndx_taxonomy_historical_review_queue_v1",
        source="derived_from_manifest_verified_historical_sec_evidence",
        row_count=1,
        extra={"research_only": True, "risk_chain_assignment_performed": False, "historical_backfill_allowed": False, "return_data_read": False},
    )
    return summary


def main() -> int:
    parser = argparse.ArgumentParser(description="生成历史 taxonomy SEC 原文人工审核队列")
    parser.add_argument("--evidence", default="data/cache/candidate_pools/ndx_taxonomy_v2/review_pack/reviewed/evidence_reviewed.csv")
    parser.add_argument("--documents", default="data/cache/candidate_pools/ndx_taxonomy_v2/review_pack/sec_review_documents.csv")
    parser.add_argument("--current-registry", default="config/ndx_v11_current_taxonomy_review_decisions.yaml")
    parser.add_argument("--output-dir", default="data/cache/candidate_pools/ndx_taxonomy_v2/historical_review_queue")
    args = parser.parse_args()
    try:
        summary = build_queue(_resolve(args.output_dir), _resolve(args.evidence), _resolve(args.documents), _resolve(args.current_registry))
    except (OSError, ValueError, TypeError, yaml.YAMLError) as exc:
        print(f"历史 taxonomy 审核队列失败闭环: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1
    print(json.dumps(summary, ensure_ascii=True, separators=(",", ":")))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
