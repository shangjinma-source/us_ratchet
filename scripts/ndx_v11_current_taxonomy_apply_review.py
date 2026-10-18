"""Apply authored, point-in-time risk-chain decisions to the current NDX pack.

The input YAML is a compact decision seed written by the research process.
This command resolves each decision against the current SEC evidence pack,
re-reads the manifest-verified filing, checks the cited business phrase in the
bounded source section, and writes a hash-locked review bundle.  It preserves
every security_id, applies one issuer decision to all share classes, and keeps
everything without reliable evidence explicitly unclassified.

It never reads prices, returns or backtest output and never changes config,
state, today's advice or a buy plan.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
from pathlib import Path
from typing import Mapping

import pandas as pd
import yaml

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from core.helpers import atomic_write_text  # noqa: E402
from data.ndx_taxonomy import load_taxonomy_definition, read_versioned_csv  # noqa: E402
from data.ndx_taxonomy_review import (  # noqa: E402
    REVIEW_TEMPLATE_COLUMNS,
    build_filing_plan_gaps,
    extract_business_section,
)
from data.versioning import (  # noqa: E402
    file_sha256,
    manifest_path,
    utc_now_iso,
    write_dataset_manifest,
)


SCHEMA_VERSION = 1
DEFAULT_REVIEW_DIR = (
    "data/cache/candidate_pools/ndx_taxonomy_v2/"
    "current_review_20260713/review_pack"
)
DEFAULT_DECISION_FILE = "config/ndx_v11_current_taxonomy_review_decisions.yaml"
DEFAULT_OUTPUT_DIR = (
    "data/cache/candidate_pools/ndx_taxonomy_v2/"
    "current_review_20260713/reviewed"
)

DECISION_AUDIT_COLUMNS = [
    "decision_id",
    "issuer_id",
    "tickers",
    "security_id",
    "ticker",
    "evidence_id",
    "job_id",
    "risk_chain",
    "confidence",
    "status",
    "effective_from",
    "effective_to",
    "form",
    "accession_number",
    "evidence_available_date",
    "selection_date",
    "extraction_status",
    "allow_fallback",
    "raw_manifest_verified",
    "raw_sha256",
    "section_text_sha256",
    "source_phrase",
    "source_phrase_found",
    "decision_status",
    "decision_note",
]

PENDING_COLUMNS = [
    "security_id",
    "ticker",
    "issuer_id",
    "job_id",
    "extraction_status",
    "unclassified_reason",
    "source_available_date",
    "selection_date",
    "raw_manifest_verified",
]

CN_COLUMNS = {
    "decision_id": "审核决定编号",
    "issuer_id": "发行人编号",
    "tickers": "发行人证券代码",
    "security_id": "证券编号",
    "ticker": "股票代码",
    "evidence_id": "证据编号",
    "job_id": "取证作业编号",
    "risk_chain": "风险链编号",
    "effective_from": "分类生效日",
    "effective_to": "分类失效日",
    "form": "申报表类型",
    "accession_number": "SEC申报编号",
    "evidence_available_date": "证据可用日",
    "selection_date": "选择日",
    "extraction_status": "业务段提取状态",
    "allow_fallback": "允许后备段",
    "raw_manifest_verified": "原始文件Manifest已验证",
    "raw_sha256": "原始文件哈希",
    "section_text_sha256": "业务段哈希",
    "source_phrase": "原文核验短语",
    "source_phrase_found": "原文短语已找到",
    "decision_status": "决定审计状态",
    "decision_note": "决定备注",
    "status": "状态",
    "confidence": "置信度",
    "evidence_source_type": "证据来源类型",
    "evidence_source": "证据来源",
    "evidence_locator": "证据地址",
    "evidence_available_date": "证据可用日",
    "accession_number": "SEC申报编号",
    "evidence_excerpt": "证据摘录",
    "notes": "备注",
    "unclassified_reason": "未分类原因",
    "count": "证券数量",
    "issuer_count": "发行人数量",
    "coverage_pct": "证券覆盖率",
    "source_name": "来源名称",
    "path": "文件路径",
    "version_id": "版本编号",
    "content_sha256": "内容哈希",
    "manifest_sha256": "Manifest哈希",
    "manifest_verified": "Manifest已验证",
    "row_count": "行数",
    "used_for": "用途",
}


class CurrentTaxonomyReviewError(ValueError):
    """Raised when current evidence cannot be applied without guessing."""


def _resolve(value: str | Path) -> Path:
    target = Path(value)
    return target if target.is_absolute() else ROOT / target


def _atomic_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        frame.to_csv(temporary, index=False, encoding="utf-8-sig", lineterminator="\n")
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def _normalise_spaces(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\u00a0", " ")).strip().casefold()


def _valid_tickers(decision: Mapping[str, object]) -> list[str]:
    values = decision.get("tickers")
    if values is None and decision.get("ticker"):
        values = [decision.get("ticker")]
    if not isinstance(values, list) or not values:
        raise CurrentTaxonomyReviewError("each decision requires a non-empty tickers list")
    tickers = [str(value or "").strip().upper() for value in values]
    if any(not value for value in tickers):
        raise CurrentTaxonomyReviewError("decision contains an empty ticker")
    if "QQQ" in tickers:
        raise CurrentTaxonomyReviewError("QQQ cannot enter current taxonomy")
    if len(set(tickers)) != len(tickers):
        raise CurrentTaxonomyReviewError(f"decision has duplicate tickers: {tickers}")
    return tickers


def _read_seed(path: Path, taxonomy: Mapping[str, object]) -> dict:
    try:
        payload = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    except (OSError, ValueError, TypeError, yaml.YAMLError) as exc:
        raise CurrentTaxonomyReviewError(f"invalid decision seed: {path}") from exc
    if int(payload.get("schema_version") or 0) != SCHEMA_VERSION:
        raise CurrentTaxonomyReviewError("decision seed schema_version must be 1")
    if str(payload.get("taxonomy_id") or "") != str(taxonomy.get("taxonomy_id") or ""):
        raise CurrentTaxonomyReviewError("decision seed taxonomy_id mismatch")
    decisions = payload.get("decisions")
    if not isinstance(decisions, list) or not decisions:
        raise CurrentTaxonomyReviewError("decision seed has no decisions")
    categories = set((taxonomy.get("risk_chains") or {}).keys())
    seen_tickers: set[str] = set()
    for position, decision in enumerate(decisions, start=1):
        if not isinstance(decision, dict):
            raise CurrentTaxonomyReviewError(f"decision {position} is not a mapping")
        decision_id = str(decision.get("decision_id") or f"current:{position:03d}").strip()
        if not decision_id:
            raise CurrentTaxonomyReviewError(f"decision {position} has empty decision_id")
        tickers = _valid_tickers(decision)
        overlap = seen_tickers.intersection(tickers)
        if overlap:
            raise CurrentTaxonomyReviewError(f"ticker appears in multiple decisions: {sorted(overlap)}")
        seen_tickers.update(tickers)
        chain = str(decision.get("risk_chain") or "").strip().lower()
        confidence = str(decision.get("confidence") or "").strip().lower()
        status = str(decision.get("status") or "accepted").strip().lower()
        phrase = str(decision.get("evidence_phrase") or "").strip()
        if chain not in categories or chain == "unclassified":
            raise CurrentTaxonomyReviewError(f"unsupported accepted risk chain in {decision_id}: {chain}")
        if confidence not in {"high", "medium"} or status != "accepted":
            raise CurrentTaxonomyReviewError(f"accepted decision requires high/medium confidence: {decision_id}")
        if not phrase:
            raise CurrentTaxonomyReviewError(f"decision requires an exact source phrase: {decision_id}")
        effective_to = str(decision.get("effective_to") or "").strip()
        if effective_to and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", effective_to):
            raise CurrentTaxonomyReviewError(f"invalid effective_to in {decision_id}")
    return payload


def _verify_raw_manifest(path: Path) -> tuple[bool, str, str]:
    sidecar = manifest_path(path)
    if not path.exists() or not sidecar.exists():
        return False, "", "manifest_missing"
    try:
        payload = json.loads(sidecar.read_text(encoding="utf-8"))
        digest = file_sha256(path)
        if str(payload.get("content_sha256") or "") != digest:
            return False, digest, "content_hash_mismatch"
        return True, digest, "verified"
    except (OSError, ValueError, TypeError, json.JSONDecodeError) as exc:
        return False, "", f"manifest_invalid:{type(exc).__name__}"


def _evidence_excerpt(section_text: str, phrase: str) -> str:
    text = str(section_text or "").strip()
    target = str(phrase or "").strip()
    if not text or not target:
        return text[:6000]
    position = _normalise_spaces(text).find(_normalise_spaces(target))
    if position < 0:
        return text[:6000]
    # Normalised offsets are not byte offsets.  Use a bounded source excerpt
    # from the beginning when whitespace differs; exact phrase presence is
    # still checked against the normalised full section.
    raw_position = text.casefold().find(target.casefold())
    if raw_position < 0:
        return text[:6000]
    start = max(0, raw_position - 220)
    return text[start:start + 6000]


def _write_artifact(
    path: Path,
    frame: pd.DataFrame,
    key: str,
    *,
    registry_sha256: str,
    extra: Mapping[str, object] | None = None,
) -> dict:
    _atomic_csv(frame, path)
    manifest_extra = {
        "research_only": True,
        "current_snapshot_only": True,
        "decision_registry_sha256": registry_sha256,
        "return_data_read": False,
        "price_data_used": False,
        "backtest_results_used": False,
        "config_state_today_advice_changed": False,
        "historical_backfill_allowed": False,
    }
    if extra:
        manifest_extra.update(dict(extra))
    payload = write_dataset_manifest(
        path,
        dataset_type="ndx_v11_current_taxonomy_reviewed_artifact",
        dataset_key=f"reviewed:{key}",
        source="derived_from_manifest_verified_current_sec_evidence",
        fetched_at_evidence="derived_runtime_utc",
        row_count=len(frame),
        extra=manifest_extra,
    )
    return {
        "filename": path.name,
        "manifest_filename": manifest_path(path).name,
        "rows": len(frame),
        "sha256": payload["content_sha256"],
        "version_id": payload["version_id"],
    }


def _write_json_artifact(
    path: Path,
    payload: Mapping[str, object],
    key: str,
    *,
    registry_sha256: str,
) -> dict:
    atomic_write_text(path, json.dumps(dict(payload), ensure_ascii=False, indent=2) + "\n")
    manifest = write_dataset_manifest(
        path,
        dataset_type="ndx_v11_current_taxonomy_reviewed_artifact",
        dataset_key=f"reviewed:{key}",
        source="derived_from_manifest_verified_current_sec_evidence",
        fetched_at_evidence="derived_runtime_utc",
        row_count=1,
        extra={
            "research_only": True,
            "decision_registry_sha256": registry_sha256,
            "return_data_read": False,
            "price_data_used": False,
            "backtest_results_used": False,
        },
    )
    return {
        "filename": path.name,
        "manifest_filename": manifest_path(path).name,
        "rows": 1,
        "sha256": manifest["content_sha256"],
        "version_id": manifest["version_id"],
    }


def _localized(frame: pd.DataFrame) -> pd.DataFrame:
    output = frame.copy()
    output = output.rename(columns={
        column: CN_COLUMNS.get(str(column), f"技术字段({column})")
        for column in output.columns
    })
    for column in output.columns:
        if pd.api.types.is_bool_dtype(output[column]):
            output[column] = output[column].map({True: "是", False: "否"})
    return output


def _write_workbook(
    path: Path,
    reviewed: pd.DataFrame,
    audit: pd.DataFrame,
    pending: pd.DataFrame,
    coverage: pd.DataFrame,
    sources: pd.DataFrame,
    summary: Mapping[str, object],
) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame([
            {"项目": "用途", "内容": "当前NDX点时SEC业务证据审核结果；风险链只作为风险分类证据，不直接提高股票分数。"},
            {"项目": "接受规则", "内容": "必须通过原始文件Manifest、业务段哈希、证据可用日、精确原文短语和风险链词表核验。"},
            {"项目": "发行人多股类", "内容": "同一issuer_id共享一条发行人决定，同时保留每个security_id和股票代码。"},
            {"项目": "未分类规则", "内容": "缺少可靠业务证据、文件缺口、非标准段落或短语核验失败时保持unclassified。"},
            {"项目": "研究边界", "内容": "未读取收益/价格/回测；未修改config、state、今日建议或买入计划。"},
        ]).to_excel(writer, sheet_name="使用说明", index=False)
        _localized(reviewed).to_excel(writer, sheet_name="已审核证据", index=False)
        _localized(audit).to_excel(writer, sheet_name="审核审计", index=False)
        _localized(pending).to_excel(writer, sheet_name="未分类", index=False)
        _localized(coverage).to_excel(writer, sheet_name="覆盖率", index=False)
        _localized(sources).to_excel(writer, sheet_name="来源与版本", index=False)
        pd.DataFrame([dict(summary)]).T.reset_index().set_axis(
            ["项目", "数值"], axis=1
        ).to_excel(writer, sheet_name="审计摘要", index=False)

    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(path)
    fill = PatternFill("solid", fgColor="DCE6F1")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        if sheet.max_row and sheet.max_column:
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            width = max(len(str(cell.value or "")) for cell in list(column)[:200]) + 2
            sheet.column_dimensions[letter].width = min(max(width, 10), 48)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)


def _load_inputs(review_dir: Path) -> tuple[dict, ...]:
    plan, plan_version = read_versioned_csv(
        review_dir.parent / "current_taxonomy_plan.csv",
        name="ndx_v11_current_taxonomy_plan",
        used_for="current point-in-time SEC filing locator",
    )
    attempts, attempts_version = read_versioned_csv(
        review_dir.parent / "fetch_attempts.csv",
        name="ndx_v11_current_taxonomy_fetch_attempts",
        used_for="current SEC filing fetch checkpoint",
    )
    template, template_version = read_versioned_csv(
        review_dir / "evidence_review_template.csv",
        name="ndx_v11_current_taxonomy_review_template",
        used_for="current security-level evidence rows",
    )
    documents, documents_version = read_versioned_csv(
        review_dir / "sec_review_documents.csv",
        name="ndx_v11_current_taxonomy_review_documents",
        used_for="manifest-verified SEC business-section evidence",
    )
    gaps = build_filing_plan_gaps(plan, attempts)
    return (
        {"frame": plan, "version": plan_version.to_dict()},
        {"frame": attempts, "version": attempts_version.to_dict()},
        {"frame": template, "version": template_version.to_dict()},
        {"frame": documents, "version": documents_version.to_dict()},
        {"frame": gaps, "version": {}},
    )


def _apply_seed(
    plan: pd.DataFrame,
    template: pd.DataFrame,
    documents: pd.DataFrame,
    seed: Mapping[str, object],
    taxonomy: Mapping[str, object],
    *,
    allow_gaps: bool,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, dict]:
    reviewed = template.copy()
    for column in REVIEW_TEMPLATE_COLUMNS:
        if column not in reviewed.columns:
            reviewed[column] = ""
    reviewed = reviewed[REVIEW_TEMPLATE_COLUMNS].copy()
    if reviewed["ticker"].astype(str).str.upper().eq("QQQ").any():
        raise CurrentTaxonomyReviewError("QQQ found in current evidence template")
    if reviewed["security_id"].astype(str).duplicated().any():
        raise CurrentTaxonomyReviewError("duplicate security_id in current evidence template")
    if reviewed.empty:
        raise CurrentTaxonomyReviewError("current evidence template is empty")

    plan = plan.copy()
    plan["job_id"] = plan["job_id"].astype(str)
    plan["security_id"] = plan["security_id"].astype(str)
    plan["ticker"] = plan["ticker"].astype(str).str.upper().str.strip()
    plan_by_security = plan.set_index("security_id", drop=False)
    plan_by_job = {str(key): group for key, group in plan.groupby("job_id", sort=False)}
    documents = documents.copy()
    documents["job_id"] = documents["job_id"].astype(str)
    document_by_job = documents.set_index("job_id", drop=False)

    reasons: dict[str, str] = {}
    for row in reviewed.itertuples(index=False):
        security_id = str(row.security_id)
        plan_row = plan_by_security.loc[security_id] if security_id in plan_by_security.index else None
        if plan_row is None:
            reasons[security_id] = "security_missing_from_current_plan"
        else:
            job_id = str(plan_row.job_id)
            if job_id not in document_by_job.index:
                reasons[security_id] = "sec_filing_not_fetched_or_review_document_missing"
            else:
                status = str(document_by_job.loc[job_id].extraction_status or "")
                if status != "business_section_extracted":
                    reasons[security_id] = f"extraction_status={status}"
                else:
                    reasons[security_id] = "no_authored_decision"

    audit_rows: list[dict] = []
    locked_decisions: list[dict] = []
    seen_tickers: set[str] = set()
    categories = set((taxonomy.get("risk_chains") or {}).keys())
    # The vocabulary is kept in the taxonomy file; the seed repeats no labels.
    for position, raw_decision in enumerate(seed["decisions"], start=1):
        decision = dict(raw_decision)
        decision_id = str(decision.get("decision_id") or f"current:{position:03d}").strip()
        tickers = _valid_tickers(decision)
        seen_tickers.update(tickers)
        target_rows = reviewed[reviewed["ticker"].astype(str).str.upper().isin(tickers)].copy()
        if len(target_rows) != len(tickers):
            raise CurrentTaxonomyReviewError(f"decision tickers are absent from current target: {decision_id}")
        target_issuer_ids = set(target_rows["issuer_id"].astype(str))
        if len(target_issuer_ids) != 1:
            raise CurrentTaxonomyReviewError(f"decision tickers do not share one issuer_id: {decision_id}")
        issuer_id = next(iter(target_issuer_ids))
        explicit_issuer = str(decision.get("issuer_id") or "").strip()
        if explicit_issuer and explicit_issuer != issuer_id:
            raise CurrentTaxonomyReviewError(f"issuer_id mismatch in {decision_id}")
        issuer_rows = reviewed[reviewed["issuer_id"].astype(str).eq(issuer_id)]
        issuer_tickers = set(issuer_rows["ticker"].astype(str).str.upper())
        if issuer_tickers != set(tickers):
            raise CurrentTaxonomyReviewError(
                f"decision must cover every current security of issuer {issuer_id}: "
                f"seed={sorted(tickers)} current={sorted(issuer_tickers)}"
            )
        job_ids = set()
        for security_id in issuer_rows["security_id"].astype(str):
            if security_id not in plan_by_security.index:
                raise CurrentTaxonomyReviewError(f"security absent from plan in {decision_id}: {security_id}")
            job_ids.add(str(plan_by_security.loc[security_id].job_id))
        if len(job_ids) != 1:
            raise CurrentTaxonomyReviewError(f"issuer decision maps to multiple SEC jobs: {decision_id}")
        job_id = next(iter(job_ids))
        if job_id not in document_by_job.index:
            raise CurrentTaxonomyReviewError(f"SEC evidence missing for {decision_id}; fetch first")
        document = document_by_job.loc[job_id]
        plan_rows = plan_by_job[job_id]
        plan_row = plan_rows.iloc[0]
        raw_path = Path(str(document.raw_path))
        manifest_ok, raw_sha256, manifest_reason = _verify_raw_manifest(raw_path)
        if not manifest_ok:
            raise CurrentTaxonomyReviewError(f"raw SEC manifest failed for {decision_id}: {manifest_reason}")
        if raw_sha256 != str(document.raw_sha256 or ""):
            raise CurrentTaxonomyReviewError(f"raw SEC hash changed for {decision_id}")
        extraction = extract_business_section(raw_path.read_bytes(), str(plan_row.form))
        extraction_status = str(extraction.get("extraction_status") or "")
        allow_fallback = bool(decision.get("allow_fallback"))
        accepted_statuses = {"business_section_extracted"}
        if allow_fallback:
            accepted_statuses.add("business_text_fallback_candidate")
        if extraction_status not in accepted_statuses:
            raise CurrentTaxonomyReviewError(
                f"decision requires a reliable business section in {decision_id}: {extraction_status}"
            )
        actual_section_sha256 = hashlib.sha256(
            str(extraction.get("section_text") or "").encode("utf-8")
        ).hexdigest()
        expected_section_sha256 = str(document.section_text_sha256 or "")
        if expected_section_sha256 and actual_section_sha256 != expected_section_sha256:
            raise CurrentTaxonomyReviewError(f"business section hash changed for {decision_id}")
        phrase = str(decision.get("evidence_phrase") or "").strip()
        phrase_found = _normalise_spaces(phrase) in _normalise_spaces(extraction.get("section_text"))
        if not phrase_found:
            raise CurrentTaxonomyReviewError(f"source phrase not found in {decision_id}: {phrase}")
        source_date = str(plan_row.source_available_date or "").strip()
        selection_date = str(plan_row.observation_from or "").strip()
        if not source_date or not selection_date or source_date > selection_date:
            raise CurrentTaxonomyReviewError(f"future-data guard failed for {decision_id}")
        effective_to = str(decision.get("effective_to") or "").strip()
        if effective_to and source_date > effective_to:
            raise CurrentTaxonomyReviewError(f"effective interval reversed for {decision_id}")
        risk_chain = str(decision.get("risk_chain") or "").strip().lower()
        confidence = str(decision.get("confidence") or "").strip().lower()
        if risk_chain not in categories or risk_chain == "unclassified":
            # The seed is checked against the taxonomy in main; this branch is
            # kept here so direct unit callers also fail closed.
            raise CurrentTaxonomyReviewError(f"unsupported risk chain in {decision_id}")
        snippet = _evidence_excerpt(str(extraction.get("section_text") or ""), phrase)
        evidence_id_base = f"sec:{job_id}"
        evidence_source = f"SEC {plan_row.form} filing business description (reviewed)"
        for row_index, row in issuer_rows.iterrows():
            security_id = str(row.security_id)
            evidence_id = f"{evidence_id_base}:{security_id}"
            reviewed.at[row_index, "risk_chain"] = risk_chain
            reviewed.at[row_index, "effective_from"] = source_date
            reviewed.at[row_index, "effective_to"] = effective_to
            reviewed.at[row_index, "evidence_source_type"] = "sec_filing_business_description"
            reviewed.at[row_index, "evidence_source"] = evidence_source
            reviewed.at[row_index, "evidence_locator"] = str(document.document_url)
            reviewed.at[row_index, "evidence_available_date"] = source_date
            reviewed.at[row_index, "accession_number"] = str(plan_row.accession_number)
            reviewed.at[row_index, "evidence_excerpt"] = snippet
            reviewed.at[row_index, "confidence"] = confidence
            reviewed.at[row_index, "status"] = "accepted"
            reviewed.at[row_index, "notes"] = (
                f"decision_id={decision_id}; issuer_id={issuer_id}; security_id={security_id}; "
                f"source_phrase={phrase}; allow_fallback={allow_fallback}; "
                f"selection_date={selection_date}; {str(decision.get('notes') or '').strip()}"
            ).strip()
            reasons[security_id] = "accepted"
            audit_rows.append({
                "decision_id": decision_id,
                "issuer_id": issuer_id,
                "tickers": ";".join(sorted(issuer_tickers)),
                "security_id": security_id,
                "ticker": str(row.ticker),
                "evidence_id": evidence_id,
                "job_id": job_id,
                "risk_chain": risk_chain,
                "confidence": confidence,
                "status": "accepted",
                "effective_from": source_date,
                "effective_to": effective_to,
                "form": str(plan_row.form),
                "accession_number": str(plan_row.accession_number),
                "evidence_available_date": source_date,
                "selection_date": selection_date,
                "extraction_status": extraction_status,
                "allow_fallback": allow_fallback,
                "raw_manifest_verified": manifest_ok,
                "raw_sha256": raw_sha256,
                "section_text_sha256": actual_section_sha256,
                "source_phrase": phrase,
                "source_phrase_found": phrase_found,
                "decision_status": "applied",
                "decision_note": str(decision.get("notes") or ""),
            })
        locked_decisions.append({
            "decision_id": decision_id,
            "issuer_id": issuer_id,
            "tickers": sorted(issuer_tickers),
            "security_ids": sorted(issuer_rows["security_id"].astype(str)),
            "job_id": job_id,
            "evidence_id": evidence_id_base,
            "risk_chain": risk_chain,
            "confidence": confidence,
            "status": "accepted",
            "effective_from": source_date,
            "effective_to": effective_to,
            "form": str(plan_row.form),
            "accession_number": str(plan_row.accession_number),
            "evidence_available_date": source_date,
            "selection_date": selection_date,
            "document_url": str(document.document_url),
            "raw_sha256": raw_sha256,
            "section_text_sha256": actual_section_sha256,
            "evidence_excerpt_sha256": hashlib.sha256(snippet.encode("utf-8")).hexdigest(),
            "source_phrase": phrase,
            "allow_fallback": allow_fallback,
            "notes": str(decision.get("notes") or ""),
        })

    for row_index, row in reviewed.iterrows():
        if str(row.status).lower() == "accepted":
            continue
        security_id = str(row.security_id)
        reviewed.at[row_index, "risk_chain"] = "unclassified"
        reviewed.at[row_index, "status"] = "unclassified"
        reviewed.at[row_index, "confidence"] = ""
        reviewed.at[row_index, "effective_from"] = ""
        reviewed.at[row_index, "effective_to"] = ""
        reviewed.at[row_index, "notes"] = (
            f"unclassified_reason={reasons.get(security_id, 'no_authored_decision')}; "
            "current taxonomy evidence is not a score or historical backfill"
        )

    pending_rows = []
    for _, row in reviewed[reviewed["status"].astype(str).ne("accepted")].iterrows():
        security_id = str(row.security_id)
        plan_row = plan_by_security.loc[security_id] if security_id in plan_by_security.index else None
        job_id = str(plan_row.job_id) if plan_row is not None else ""
        document = document_by_job.loc[job_id] if job_id in document_by_job.index else None
        pending_rows.append({
            "security_id": security_id,
            "ticker": str(row.ticker),
            "issuer_id": str(row.issuer_id),
            "job_id": job_id,
            "extraction_status": str(document.extraction_status) if document is not None else "",
            "unclassified_reason": reasons.get(security_id, "no_authored_decision"),
            "source_available_date": str(plan_row.source_available_date) if plan_row is not None else "",
            "selection_date": str(plan_row.observation_from) if plan_row is not None else "",
            "raw_manifest_verified": str(document.raw_manifest_verified) if document is not None else False,
        })
    pending = pd.DataFrame(pending_rows, columns=PENDING_COLUMNS)
    audit = pd.DataFrame(audit_rows, columns=DECISION_AUDIT_COLUMNS)
    coverage = (
        reviewed.groupby("risk_chain", dropna=False)
        .agg(count=("security_id", "count"), issuer_count=("issuer_id", "nunique"))
        .reset_index()
    )
    coverage["coverage_pct"] = (coverage["count"] / len(reviewed) * 100).round(2)
    coverage = coverage.sort_values(["risk_chain"], kind="stable").reset_index(drop=True)
    return reviewed, audit, pending, {
        "locked_decisions": locked_decisions,
        "coverage": coverage,
        "seen_tickers": sorted(seen_tickers),
        "gap_count": 0,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Apply current NDX SEC taxonomy decisions with hash and phrase locks")
    parser.add_argument("--review-dir", default=DEFAULT_REVIEW_DIR)
    parser.add_argument("--decision-file", default=DEFAULT_DECISION_FILE)
    parser.add_argument("--taxonomy", default="config/candidate_taxonomy_v2.yaml")
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--allow-gaps", action="store_true")
    args = parser.parse_args()
    try:
        review_dir = _resolve(args.review_dir)
        decision_path = _resolve(args.decision_file)
        output_dir = _resolve(args.output_dir)
        taxonomy = load_taxonomy_definition(_resolve(args.taxonomy))
        seed = _read_seed(decision_path, taxonomy)
        inputs = _load_inputs(review_dir)
        plan, attempts, template, documents, gaps = [item["frame"] for item in inputs]
        if not gaps.empty and not args.allow_gaps:
            raise CurrentTaxonomyReviewError(
                f"current SEC evidence has {len(gaps)} filing gaps; fetch and rebuild review pack first"
            )
        reviewed, audit, pending, derived = _apply_seed(
            plan,
            template,
            documents,
            seed,
            taxonomy,
            allow_gaps=bool(args.allow_gaps),
        )
        registry_sha256 = file_sha256(decision_path)
        output_dir.mkdir(parents=True, exist_ok=True)
        source_rows = [item["version"] for item in inputs[:4]]
        source_rows.append({
            "source_name": "taxonomy_definition",
            "path": _resolve(args.taxonomy).as_posix(),
            "version_id": str(taxonomy.get("version_id") or taxonomy.get("schema_version") or ""),
            "content_sha256": file_sha256(_resolve(args.taxonomy)),
            "manifest_sha256": file_sha256(manifest_path(_resolve(args.taxonomy))) if manifest_path(_resolve(args.taxonomy)).exists() else "",
            "manifest_verified": True,
            "row_count": len((taxonomy.get("risk_chains") or {})),
            "used_for": "controlled risk-chain vocabulary only",
        })
        sources = pd.DataFrame(source_rows)
        summary = {
            "schema_version": SCHEMA_VERSION,
            "generated_at": utc_now_iso(),
            "taxonomy_id": str(taxonomy.get("taxonomy_id") or ""),
            "taxonomy_version": str(taxonomy.get("version_id") or ""),
            "observation_date": str(plan["observation_from"].astype(str).str[:10].iloc[0]) if not plan.empty else "",
            "current_security_count": int(len(reviewed)),
            "current_issuer_count": int(reviewed["issuer_id"].nunique()),
            "accepted_security_count": int(reviewed["status"].eq("accepted").sum()),
            "accepted_issuer_count": int(reviewed.loc[reviewed["status"].eq("accepted"), "issuer_id"].nunique()),
            "unclassified_security_count": int(reviewed["status"].eq("unclassified").sum()),
            "accepted_decision_count": int(audit["decision_id"].nunique()) if not audit.empty else 0,
            "fallback_accepted_security_count": int(audit["allow_fallback"].eq(True).sum()) if not audit.empty else 0,
            "filing_plan_gap_count": int(len(gaps)),
            "risk_chain_assigned": bool(reviewed["status"].eq("accepted").any()),
            "classification_automatic": False,
            "return_data_read": False,
            "price_data_used": False,
            "backtest_results_used": False,
            "config_state_today_advice_changed": False,
            "future_data_guard": "evidence_available_date <= effective_from <= selection_date",
            "historical_backfill_allowed": False,
            "decision_registry_id": str(seed.get("review_registry_id") or ""),
            "decision_registry_path": decision_path.as_posix(),
            "decision_registry_sha256": registry_sha256,
            "input_versions": [item["version"] for item in inputs[:4]],
        }
        artifacts = {}
        artifacts["evidence_reviewed"] = _write_artifact(
            output_dir / "evidence_reviewed.csv", reviewed, "evidence_reviewed",
            registry_sha256=registry_sha256,
        )
        artifacts["decision_audit"] = _write_artifact(
            output_dir / "review_decision_audit.csv", audit, "review_decision_audit",
            registry_sha256=registry_sha256,
        )
        artifacts["pending"] = _write_artifact(
            output_dir / "unclassified.csv", pending, "unclassified",
            registry_sha256=registry_sha256,
        )
        artifacts["coverage"] = _write_artifact(
            output_dir / "coverage.csv", derived["coverage"], "coverage",
            registry_sha256=registry_sha256,
        )
        artifacts["source_versions"] = _write_artifact(
            output_dir / "source_versions.csv", sources, "source_versions",
            registry_sha256=registry_sha256,
        )
        snapshot = {
            "schema_version": SCHEMA_VERSION,
            "review_registry_id": str(seed.get("review_registry_id") or ""),
            "taxonomy_id": str(taxonomy.get("taxonomy_id") or ""),
            "research_only": True,
            "source_seed_sha256": registry_sha256,
            "decision_policy": seed.get("decision_policy") or {},
            "decisions": derived["locked_decisions"],
        }
        snapshot_path = output_dir / "review_decisions_snapshot.yaml"
        atomic_write_text(snapshot_path, yaml.safe_dump(snapshot, allow_unicode=True, sort_keys=False))
        snapshot_manifest = write_dataset_manifest(
            snapshot_path,
            dataset_type="ndx_v11_current_taxonomy_decision_snapshot",
            dataset_key="current_taxonomy_decision_snapshot",
            source="derived_from_authored_current_taxonomy_decision_seed",
            fetched_at_evidence="derived_runtime_utc",
            row_count=len(derived["locked_decisions"]),
            extra={
                "research_only": True,
                "decision_registry_sha256": registry_sha256,
                "return_data_read": False,
            },
        )
        artifacts["decision_snapshot"] = {
            "filename": snapshot_path.name,
            "manifest_filename": manifest_path(snapshot_path).name,
            "rows": len(derived["locked_decisions"]),
            "sha256": snapshot_manifest["content_sha256"],
            "version_id": snapshot_manifest["version_id"],
        }
        workbook_path = output_dir / "ndx_v11_current_taxonomy_reviewed.xlsx"
        _write_workbook(
            workbook_path, reviewed, audit, pending, derived["coverage"], sources, summary
        )
        workbook_manifest = write_dataset_manifest(
            workbook_path,
            dataset_type="ndx_v11_current_taxonomy_reviewed_workbook",
            dataset_key="current_taxonomy_reviewed",
            source="derived_from_manifest_verified_current_sec_evidence",
            fetched_at_evidence="derived_runtime_utc",
            row_count=len(reviewed),
            extra={
                "research_only": True,
                "decision_registry_sha256": registry_sha256,
                "return_data_read": False,
            },
        )
        artifacts["workbook"] = {
            "filename": workbook_path.name,
            "manifest_filename": manifest_path(workbook_path).name,
            "rows": len(reviewed),
            "sha256": workbook_manifest["content_sha256"],
            "version_id": workbook_manifest["version_id"],
        }
        summary["artifacts"] = artifacts
        summary_artifact = _write_json_artifact(
            output_dir / "summary.json", summary, "summary", registry_sha256=registry_sha256
        )
        artifacts["summary"] = summary_artifact
        root_manifest = {
            "schema_version": SCHEMA_VERSION,
            "manifest_type": "ndx_v11_current_taxonomy_reviewed",
            "generated_at": summary["generated_at"],
            "research_only": True,
            "review_only": True,
            "risk_chain_assigned": summary["risk_chain_assigned"],
            "classification_automatic": False,
            "return_data_read": False,
            "price_data_used": False,
            "backtest_results_used": False,
            "config_state_today_advice_changed": False,
            "historical_backfill_allowed": False,
            "future_data_guard": {"rule": summary["future_data_guard"], "passed": True},
            "decision_registry_sha256": registry_sha256,
            "summary": summary,
            "artifacts": artifacts,
            "prohibited_uses": [
                "direct_score_boost",
                "historical_backfill_from_current_label",
                "automatic_config_change",
                "automatic_today_advice_change",
            ],
        }
        manifest_path_out = output_dir / "manifest.json"
        atomic_write_text(manifest_path_out, json.dumps(root_manifest, ensure_ascii=False, indent=2) + "\n")
        write_dataset_manifest(
            manifest_path_out,
            dataset_type="ndx_v11_current_taxonomy_reviewed_manifest",
            dataset_key="current_taxonomy_reviewed",
            source="derived_from_manifest_verified_current_sec_evidence",
            fetched_at_evidence="derived_runtime_utc",
            row_count=1,
            extra={
                "research_only": True,
                "decision_registry_sha256": registry_sha256,
                "return_data_read": False,
            },
        )
        print(
            "NDX v11 当前 taxonomy 审核导入完成 | "
            f"accepted={summary['accepted_security_count']} | "
            f"unclassified={summary['unclassified_security_count']} | "
            f"issuers={summary['current_issuer_count']} | "
            f"gaps={summary['filing_plan_gap_count']} | "
            "return_data_read=False"
        )
        print(f"输出: {output_dir}")
        print(f"Manifest: {manifest_path_out}")
        return 0
    except Exception as exc:
        print(
            f"NDX v11 current taxonomy review apply failed closed: {type(exc).__name__}: {exc}",
            file=sys.stderr,
        )
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
