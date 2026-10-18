"""Build a point-in-time SEC evidence pack for the current NDX universe.

This workflow is deliberately review-only.  It selects the latest eligible
10-K/20-F/40-F filed on or before the observation date, fetches the SEC-hosted
primary document with resumable checkpoints, and exposes source text for
manual taxonomy review.  It never assigns a risk chain, reads returns, scores
securities, or changes live configuration/state/today advice.

PowerShell / PyCharm terminal:
    python -B scripts\\ndx_v11_current_taxonomy.py --plan-only --observation-date 2026-07-13
    python -B scripts\\ndx_v11_current_taxonomy.py --smoke --observation-date 2026-07-13
    $env:SEC_USER_AGENT="us_ratchet/1.0 your-real-email@example.com"
    python -B scripts\\ndx_v11_current_taxonomy.py --fetch-sec --resume --max-jobs 0 --observation-date 2026-07-13
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
import time
from pathlib import Path
from typing import Mapping

import pandas as pd
import requests

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from core.helpers import atomic_write_text  # noqa: E402
from data.ndx_current_taxonomy import (  # noqa: E402
    build_current_taxonomy_plan,
    filing_document_url,
    filing_html_index_url,
    normalized_cik,
    select_latest_business_filing,
    submissions_url,
)
from data.ndx_history import _sec_headers, _sec_user_agent  # noqa: E402
from data.ndx_taxonomy import (  # noqa: E402
    NdxTaxonomyError,
    empty_fetch_attempts,
    read_versioned_csv,
    select_business_document,
    select_business_document_from_index_html,
)
from data.ndx_taxonomy_review import (  # noqa: E402
    REVIEW_TEMPLATE_COLUMNS,
    build_filing_plan_gaps,
    build_review_documents,
    review_summary,
)
from data.versioning import (  # noqa: E402
    file_sha256,
    manifest_path,
    utc_now_iso,
    write_dataset_manifest,
)
from scripts.ndx_taxonomy_review_pack import write_bundle as write_review_bundle  # noqa: E402


OBSERVATION_DATE = "2026-07-13"
DEFAULT_INPUT = "data/cache/research/ndx_v11_forward_watch_v1/forward_observations.csv"
DEFAULT_OUTPUT_ROOT = "data/cache/candidate_pools/ndx_taxonomy_v2"
PLAN_SCHEMA_VERSION = 1

SUBMISSION_ATTEMPT_COLUMNS = [
    "cik", "status", "raw_path", "content_sha256", "attempted_at", "error",
]

PLAN_EXTRA_COLUMNS = [
    "submission_cache_status", "submission_raw_path", "submission_raw_sha256",
    "submission_manifest_verified",
]

CHINESE_COLUMNS = {
    "observation_date": "观察日期",
    "ticker": "股票代码",
    "name": "证券名称",
    "issuer_id": "发行人编号",
    "security_id": "证券编号",
    "identity_status": "身份状态",
    "cik": "SEC CIK",
    "ndx_as_of_date": "NDX快照日期",
    "ndx_membership_status": "NDX成员状态",
    "risk_chain": "风险链编号（当前未用于分类）",
    "risk_chain_label": "风险链名称（当前未用于分类）",
    "taxonomy_confidence": "Taxonomy置信度（当前未用于分类）",
    "taxonomy_status": "Taxonomy状态（当前未用于分类）",
    "taxonomy_source": "Taxonomy来源（当前未用于分类）",
    "taxonomy_effective_from": "Taxonomy生效日（当前未用于分类）",
    "taxonomy_effective_to": "Taxonomy失效日（当前未用于分类）",
    "taxonomy_evidence_available_date": "Taxonomy证据可用日（当前未用于分类）",
    "job_id": "取证作业编号",
    "cusip": "CUSIP",
    "observation_from": "选择日期",
    "evidence_required_from": "证据需要起日",
    "form": "申报表类型",
    "filing_date": "申报日期",
    "acceptance_datetime": "SEC接收时间",
    "source_available_date": "来源可用日",
    "accession_number": "SEC申报编号",
    "primary_document": "SEC主文档",
    "sec_index_url": "SEC申报目录",
    "document_url": "SEC主文档地址",
    "plan_status": "计划状态",
    "selection_method": "选择方法",
    "reason": "原因",
    "future_data_guard": "未来数据守门",
    "classification_allowed": "允许直接分类",
    "submission_cache_status": "SEC提交缓存状态",
    "submission_raw_path": "SEC提交原始文件",
    "submission_raw_sha256": "SEC提交原始文件哈希",
    "submission_manifest_verified": "SEC提交Manifest已核验",
    "status": "状态",
    "raw_path": "原始文件路径",
    "content_bytes": "内容字节数",
    "content_sha256": "内容哈希",
    "attempted_at": "尝试时间",
    "error": "错误信息",
    "raw_document_count": "SEC原始主文档数",
    "business_section_extracted_count": "业务段落提取数",
    "filing_plan_gap_count": "申报计划缺口数",
    "pending_manual_review": "待人工复核数",
}


def _resolve(value: str | Path) -> Path:
    path = Path(value)
    return path if path.is_absolute() else ROOT / path


def _atomic_bytes(path: Path, content: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        temporary.write_bytes(content)
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def _atomic_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        frame.to_csv(temporary, index=False, encoding="utf-8-sig", lineterminator="\n")
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def _write_csv_manifest(
    frame: pd.DataFrame,
    path: Path,
    *,
    dataset_key: str,
    observation_date: str,
    used_for: str,
) -> dict:
    _atomic_csv(frame, path)
    return write_dataset_manifest(
        path,
        dataset_type="ndx_v11_current_taxonomy_artifact",
        dataset_key=dataset_key,
        source="derived_from_pit_ndx_observation_and_sec_evidence",
        fetched_at_evidence="derived_runtime_utc",
        row_count=len(frame),
        data_start=observation_date,
        data_end=observation_date,
        extra={
            "research_only": True,
            "observation_date": observation_date,
            "used_for": used_for,
            "risk_chain_assigned": False,
            "classification_automatic": False,
            "return_data_read": False,
            "price_data_used": False,
            "backtest_results_used": False,
            "live_config_mutation_allowed": False,
            "today_advice_impact": "none",
        },
    )


def _write_json_manifest(
    payload: Mapping[str, object],
    path: Path,
    *,
    dataset_key: str,
    observation_date: str,
    dataset_type: str = "ndx_v11_current_taxonomy_artifact",
) -> dict:
    atomic_write_text(
        path,
        json.dumps(dict(payload), ensure_ascii=False, indent=2, allow_nan=False) + "\n",
    )
    return write_dataset_manifest(
        path,
        dataset_type=dataset_type,
        dataset_key=dataset_key,
        source="derived_from_pit_ndx_observation_and_sec_evidence",
        fetched_at_evidence="derived_runtime_utc",
        row_count=1,
        data_start=observation_date,
        data_end=observation_date,
        extra={
            "research_only": True,
            "observation_date": observation_date,
            "risk_chain_assigned": False,
            "return_data_read": False,
            "price_data_used": False,
            "backtest_results_used": False,
            "live_config_mutation_allowed": False,
            "today_advice_impact": "none",
        },
    )


def _artifact_info(path: Path, rows: int | None = None) -> dict:
    sidecar = manifest_path(path)
    payload = json.loads(sidecar.read_text(encoding="utf-8"))
    return {
        "path": path.as_posix(),
        "manifest_path": sidecar.as_posix(),
        "content_sha256": file_sha256(path),
        "manifest_sha256": file_sha256(sidecar),
        "version_id": payload.get("version_id"),
        "row_count": int(rows if rows is not None else payload.get("row_count") or 0),
    }


def _verify_manifest(path: Path) -> tuple[bool, str, str]:
    sidecar = manifest_path(path)
    if not path.exists() or not sidecar.exists():
        return False, "", "manifest_missing"
    try:
        payload = json.loads(sidecar.read_text(encoding="utf-8"))
        digest = file_sha256(path)
        if str(payload.get("content_sha256") or "") != digest:
            return False, digest, "content_hash_mismatch"
        return True, digest, "verified"
    except (OSError, ValueError, TypeError, json.JSONDecodeError) as exc:
        return False, "", f"manifest_invalid:{type(exc).__name__}"


def _load_checkpoint(path: Path, columns: list[str]) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame(columns=columns)
    frame = pd.read_csv(path, dtype=str, keep_default_na=False)
    for column in columns:
        if column not in frame.columns:
            frame[column] = ""
    return frame[columns].drop_duplicates(columns[0], keep="last").reset_index(drop=True)


def _upsert(frame: pd.DataFrame, row: dict, key: str) -> pd.DataFrame:
    remaining = frame[~frame[key].astype(str).eq(str(row[key]))].copy()
    return pd.concat([remaining, pd.DataFrame([row])], ignore_index=True).sort_values(
        key, kind="stable"
    ).reset_index(drop=True)


def _load_current_observations(input_path: Path, observation_date: str) -> tuple[pd.DataFrame, dict]:
    observations, version = read_versioned_csv(
        input_path,
        name="ndx_v11_forward_observations",
        used_for="当前NDX证券范围和点时发行人身份；不用于收益或分类推断",
    )
    required = {"observation_date", "ticker", "issuer_id", "security_id", "cik"}
    missing = sorted(required - set(observations.columns))
    if missing:
        raise NdxTaxonomyError(f"current observations missing columns: {missing}")
    observations = observations.copy()
    observations["observation_date"] = observations["observation_date"].astype(str).str[:10]
    if not observations["observation_date"].eq(observation_date).all():
        raise NdxTaxonomyError("current observations contain a date other than the requested observation")
    observations["ticker"] = observations["ticker"].astype(str).str.upper().str.strip()
    if observations["ticker"].eq("QQQ").any():
        raise NdxTaxonomyError("QQQ cannot enter current taxonomy evidence")
    observations["cik"] = observations["cik"].map(normalized_cik)
    if observations["cik"].eq("").any():
        raise NdxTaxonomyError("current observations contain an unresolved CIK")
    if observations["security_id"].astype(str).duplicated().any():
        raise NdxTaxonomyError("current observations contain duplicate security_id")
    for cik, group in observations.groupby("cik", sort=False):
        if group["issuer_id"].astype(str).nunique() != 1:
            raise NdxTaxonomyError(
                f"one CIK maps to multiple issuer_id values; refusing to guess: {cik}"
            )
    return observations, version.to_dict()


def _load_cached_submissions(
    output: Path,
    ciks: list[str],
) -> tuple[dict[str, Mapping[str, object]], dict[str, dict]]:
    payloads: dict[str, Mapping[str, object]] = {}
    metadata: dict[str, dict] = {}
    for cik in sorted(set(ciks)):
        raw_path = output / "raw_submissions" / f"CIK{cik}.json"
        base = {
            "cik": cik,
            "raw_path": raw_path.as_posix(),
            "raw_sha256": "",
            "manifest_verified": False,
        }
        if not raw_path.exists():
            metadata[cik] = {**base, "cache_status": "not_cached", "error": "raw SEC submissions JSON is not cached"}
            continue
        verified, digest, reason = _verify_manifest(raw_path)
        if not verified:
            metadata[cik] = {**base, "raw_sha256": digest, "cache_status": "integrity_failed", "error": reason}
            continue
        try:
            payload = json.loads(raw_path.read_text(encoding="utf-8"))
            if normalized_cik(payload.get("cik")) != cik:
                raise NdxTaxonomyError("SEC submissions payload CIK does not match filename")
            if not isinstance(payload.get("filings"), Mapping):
                raise NdxTaxonomyError("SEC submissions payload lacks filings object")
        except (OSError, ValueError, TypeError, json.JSONDecodeError, NdxTaxonomyError) as exc:
            metadata[cik] = {
                **base,
                "raw_sha256": digest,
                "cache_status": "invalid_payload",
                "error": f"{type(exc).__name__}: {exc}",
            }
            continue
        payloads[cik] = payload
        metadata[cik] = {
            **base,
            "raw_sha256": digest,
            "manifest_verified": True,
            "cache_status": "cached_valid",
            "error": "",
        }
    return payloads, metadata


def _fetch_submissions(
    output: Path,
    ciks: list[str],
    *,
    observation_date: str,
    resume: bool,
    force: bool,
    max_jobs: int,
    timeout: float,
    request_delay: float,
) -> tuple[dict[str, Mapping[str, object]], dict[str, dict], pd.DataFrame]:
    output.mkdir(parents=True, exist_ok=True)
    attempts_path = output / "submission_attempts.csv"
    attempts = _load_checkpoint(attempts_path, SUBMISSION_ATTEMPT_COLUMNS)
    _, cached_metadata = _load_cached_submissions(output, ciks)
    pending = []
    for cik in sorted(set(ciks)):
        cached = cached_metadata[cik]["cache_status"] == "cached_valid"
        if cached and resume and not force:
            continue
        pending.append(cik)
    if max_jobs > 0:
        pending = pending[:max_jobs]
    if pending:
        explicit_agent = os.environ.get("SEC_USER_AGENT")
        if not explicit_agent:
            raise NdxTaxonomyError(
                "SEC_USER_AGENT is required for --fetch-sec, e.g. "
                "$env:SEC_USER_AGENT='us_ratchet/1.0 your_email@example.com'"
            )
        session = requests.Session()
        session.headers.update(_sec_headers(_sec_user_agent(explicit_agent)))
    else:
        session = None
    for position, cik in enumerate(pending, start=1):
        status = "failed"
        content = b""
        error = ""
        raw_path = output / "raw_submissions" / f"CIK{cik}.json"
        try:
            response = session.get(submissions_url(cik), timeout=timeout)
            response.raise_for_status()
            content = bytes(response.content)
            payload = json.loads(content.decode("utf-8"))
            if normalized_cik(payload.get("cik")) != cik:
                raise NdxTaxonomyError("SEC submissions payload CIK does not match request")
            if not isinstance(payload.get("filings"), Mapping):
                raise NdxTaxonomyError("SEC submissions payload lacks filings object")
            _atomic_bytes(raw_path, content)
            write_dataset_manifest(
                raw_path,
                dataset_type="ndx_v11_current_taxonomy_sec_submissions_raw",
                dataset_key=f"CIK{cik}",
                source="sec_submissions_api",
                source_locator=submissions_url(cik),
                fetched_at_evidence="runtime_fetch_utc",
                row_count=1,
                data_start=observation_date,
                data_end=observation_date,
                extra={
                    "research_only": True,
                    "observation_date": observation_date,
                    "selection_rule": "annual filing filed_date <= observation_date",
                    "risk_chain_assigned": False,
                    "return_data_read": False,
                    "price_data_used": False,
                    "backtest_results_used": False,
                },
            )
            status = "fetched"
        except Exception as exc:  # pragma: no cover - network branch
            error = f"{type(exc).__name__}: {exc}"
        attempt = {
            "cik": cik,
            "status": status,
            "raw_path": raw_path.as_posix() if status == "fetched" else "",
            "content_sha256": hashlib.sha256(content).hexdigest() if content else "",
            "attempted_at": utc_now_iso(),
            "error": error,
        }
        attempts = _upsert(attempts, attempt, "cik")
        _write_csv_manifest(
            attempts,
            attempts_path,
            dataset_key="submission_attempts",
            observation_date=observation_date,
            used_for="SEC submissions 原始文件断点和完整性审计",
        )
        print(f"current taxonomy SEC submissions {position}/{len(pending)} | CIK:{cik} | {status}", flush=True)
        if request_delay > 0:
            time.sleep(request_delay)
    if not attempts_path.exists():
        _write_csv_manifest(
            attempts,
            attempts_path,
            dataset_key="submission_attempts",
            observation_date=observation_date,
            used_for="SEC submissions 原始文件断点和完整性审计",
        )
    payloads, metadata = _load_cached_submissions(output, ciks)
    return payloads, metadata, attempts


def _build_plan(
    observations: pd.DataFrame,
    submissions_by_cik: Mapping[str, Mapping[str, object]],
    metadata: Mapping[str, Mapping[str, object]],
    *,
    observation_date: str,
) -> pd.DataFrame:
    plan = build_current_taxonomy_plan(
        observations,
        submissions_by_cik,
        observation_date=observation_date,
    )
    meta_frame = pd.DataFrame([
        {
            "cik": cik,
            "submission_cache_status": str(value.get("cache_status") or ""),
            "submission_raw_path": str(value.get("raw_path") or ""),
            "submission_raw_sha256": str(value.get("raw_sha256") or ""),
            "submission_manifest_verified": bool(value.get("manifest_verified")),
            "submission_error": str(value.get("error") or ""),
        }
        for cik, value in metadata.items()
    ])
    plan = plan.merge(meta_frame, on="cik", how="left", validate="many_to_one")
    invalid = plan["submission_cache_status"].isin({"integrity_failed", "invalid_payload"})
    plan.loc[invalid, "plan_status"] = "submission_cache_integrity_failed"
    plan.loc[invalid, "reason"] = plan.loc[invalid, "submission_error"]
    plan.loc[invalid, ["form", "filing_date", "acceptance_datetime", "source_available_date", "accession_number", "primary_document", "document_url"]] = ""
    plan["submission_manifest_verified"] = plan["submission_manifest_verified"].fillna(False).astype(bool)
    plan = plan.drop(columns=["submission_error"])
    return plan


def _load_fetch_attempts(output: Path) -> pd.DataFrame:
    path = output / "fetch_attempts.csv"
    if not path.exists():
        return empty_fetch_attempts()
    frame = _load_checkpoint(path, list(empty_fetch_attempts().columns))
    return frame


def _align_fetch_attempts(
    output: Path,
    plan: pd.DataFrame,
    *,
    observation_date: str,
) -> pd.DataFrame:
    """Keep only current-plan jobs and archive superseded locator attempts."""
    attempts = _load_fetch_attempts(output)
    current_jobs = set(plan["job_id"].astype(str))
    stale = attempts[~attempts["job_id"].astype(str).isin(current_jobs)].copy()
    if not stale.empty:
        stale["status"] = "stale_plan_job"
        _write_csv_manifest(
            stale,
            output / "stale_fetch_attempts.csv",
            dataset_key="stale_fetch_attempts",
            observation_date=observation_date,
            used_for="申报定位规则变化后被当前计划替代的SEC作业审计",
        )
    active = attempts[attempts["job_id"].astype(str).isin(current_jobs)].copy()
    _write_csv_manifest(
        active,
        output / "fetch_attempts.csv",
        dataset_key="fetch_attempts",
        observation_date=observation_date,
        used_for="当前SEC主文档断点、内容哈希和未来数据守门审计",
    )
    return active


def _write_raw_document_manifest(path: Path, row, document_url: str, observation_date: str) -> dict:
    return write_dataset_manifest(
        path,
        dataset_type="ndx_v11_current_taxonomy_sec_filing_raw",
        dataset_key=f"{row.job_id}:{path.name}",
        source="sec_filing_business_description",
        source_locator=document_url,
        fetched_at_evidence="sec_public_archive_http_200",
        row_count=1,
        data_start=str(row.source_available_date),
        data_end=str(row.source_available_date),
        extra={
            "research_only": True,
            "classification_automatic": False,
            "observation_date": observation_date,
            "issuer_id": str(row.issuer_id),
            "cik": str(row.cik),
            "accession_number": str(row.accession_number),
            "form": str(row.form),
            "evidence_available_date": str(row.source_available_date),
            "selection_date": observation_date,
            "future_data_guard": "evidence_available_date <= selection_date",
            "return_data_read": False,
            "price_data_used": False,
            "backtest_results_used": False,
        },
    )


def _fetch_filing_documents(
    plan: pd.DataFrame,
    *,
    output: Path,
    observation_date: str,
    resume: bool,
    force: bool,
    max_jobs: int,
    timeout: float,
    request_delay: float,
) -> pd.DataFrame:
    attempts_path = output / "fetch_attempts.csv"
    attempts = _align_fetch_attempts(
        output,
        plan,
        observation_date=observation_date,
    )
    unique_plan = plan[plan["plan_status"].astype(str).eq("planned")].drop_duplicates(
        "job_id", keep="first"
    )
    pending = []
    for row in unique_plan.itertuples(index=False):
        prior = attempts[attempts["job_id"].astype(str).eq(str(row.job_id))]
        if not prior.empty and not force and resume and str(prior.iloc[-1]["status"]) == "fetched":
            raw_path = Path(str(prior.iloc[-1]["raw_path"]))
            verified, digest, _ = _verify_manifest(raw_path)
            if verified and digest == str(prior.iloc[-1]["content_sha256"]):
                continue
        pending.append(row)
    if max_jobs > 0:
        pending = pending[:max_jobs]
    if pending:
        explicit_agent = os.environ.get("SEC_USER_AGENT")
        if not explicit_agent:
            raise NdxTaxonomyError(
                "SEC_USER_AGENT is required for --fetch-sec, e.g. "
                "$env:SEC_USER_AGENT='us_ratchet/1.0 your_email@example.com'"
            )
        session = requests.Session()
        session.headers.update(_sec_headers(_sec_user_agent(explicit_agent)))
    else:
        session = None
    for position, row in enumerate(pending, start=1):
        status = "failed"
        document_url = ""
        raw_path = ""
        content = b""
        error = ""
        try:
            index_url = str(row.sec_index_url)
            index_response = session.get(index_url, timeout=timeout)
            index_response.raise_for_status()
            if not index_url.endswith("/index.json"):
                raise NdxTaxonomyError("current filing plan index URL is not SEC index.json")
            base = index_url.rsplit("/index.json", 1)[0]
            try:
                item = select_business_document(index_response.json(), str(row.form))
            except (NdxTaxonomyError, ValueError, TypeError) as json_error:
                html_url = filing_html_index_url(row.cik, row.accession_number)
                html_response = session.get(html_url, timeout=timeout)
                html_response.raise_for_status()
                try:
                    item = select_business_document_from_index_html(
                        html_response.content, str(row.form)
                    )
                except NdxTaxonomyError as html_error:
                    raise NdxTaxonomyError(
                        f"SEC index JSON failed: {json_error}; HTML fallback failed: {html_error}"
                    ) from html_error
            name = Path(str(item.get("name") or "")).name
            if not name or name != str(item.get("name") or "") or not name.lower().endswith((".htm", ".html")):
                raise NdxTaxonomyError("SEC index selected an unsafe or non-HTML document")
            document_url = filing_document_url(row.cik, row.accession_number, name)
            if not document_url:
                raise NdxTaxonomyError("SEC primary document URL could not be built")
            document_response = session.get(document_url, timeout=timeout)
            document_response.raise_for_status()
            content = bytes(document_response.content)
            if len(content) < 256:
                raise NdxTaxonomyError("SEC primary business document is unexpectedly small")
            raw_file = (
                output / "raw_sec" / str(row.issuer_id).replace(":", "_")
                / str(row.accession_number).replace("-", "") / name
            )
            _atomic_bytes(raw_file, content)
            _write_raw_document_manifest(raw_file, row, document_url, observation_date)
            raw_path = raw_file.as_posix()
            status = "fetched"
        except Exception as exc:  # pragma: no cover - network branch
            error = f"{type(exc).__name__}: {exc}"
        attempt = {
            "job_id": str(row.job_id),
            "issuer_id": str(row.issuer_id),
            "cik": str(row.cik),
            "accession_number": str(row.accession_number),
            "status": status,
            "index_url": str(row.sec_index_url),
            "document_url": document_url,
            "raw_path": raw_path,
            "content_bytes": len(content),
            "content_sha256": hashlib.sha256(content).hexdigest() if content else "",
            "attempted_at": utc_now_iso(),
            "error": error,
        }
        attempts = _upsert(attempts, attempt, "job_id")
        _write_csv_manifest(
            attempts,
            attempts_path,
            dataset_key="fetch_attempts",
            observation_date=observation_date,
            used_for="SEC主文档断点、内容哈希和未来数据守门审计",
        )
        print(f"current taxonomy SEC filing {position}/{len(pending)} | {row.issuer_id} | {status}", flush=True)
        if request_delay > 0:
            time.sleep(request_delay)
    return attempts


def _localize(frame: pd.DataFrame) -> pd.DataFrame:
    output = frame.copy()
    output = output.rename(columns={
        column: CHINESE_COLUMNS.get(str(column), f"技术字段（{column}）")
        for column in output.columns
    })
    for column in output.columns:
        if pd.api.types.is_bool_dtype(output[column]):
            output[column] = output[column].map({True: "是", False: "否"})
    return output


def _style_workbook(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(path)
    fill = PatternFill("solid", fgColor="DCE6F1")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            width = max(len(str(cell.value or "")) for cell in list(column)[:200]) + 2
            sheet.column_dimensions[letter].width = min(max(width, 10), 48)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)


def _write_plan_workbook(
    path: Path,
    target: pd.DataFrame,
    plan: pd.DataFrame,
    submission_attempts: pd.DataFrame,
    fetch_attempts: pd.DataFrame,
    summary: Mapping[str, object],
    *,
    observation_date: str,
) -> dict:
    overview = pd.DataFrame([
        {"项目": "工作目的", "内容": "建立当前NDX发行人年度SEC业务描述的点时证据包；风险链、有效期、置信度和状态保留人工复核，不自动分类。"},
        {"项目": "当前范围", "内容": f"观察日期 {observation_date}；目标证券 {len(target)}；发行人 {target['issuer_id'].nunique()}；不包含QQQ。"},
        {"项目": "选择规则", "内容": "只选择SEC submissions中截至观察日已提交的10-K、20-F或40-F；优先非/A版本，同日不猜测。"},
        {"项目": "未来数据守门", "内容": "evidence_available_date <= selection_date；未通过时保持未分类，不能回填历史。"},
        {"项目": "投资数据", "内容": "未读取收益、回测或交易成本；不产生评分，不修改config.yaml、state或今日建议。"},
    ])
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        overview.to_excel(writer, sheet_name="使用说明", index=False)
        _localize(target).to_excel(writer, sheet_name="当前目标", index=False)
        _localize(plan).to_excel(writer, sheet_name="当前取证计划", index=False)
        _localize(submission_attempts).to_excel(writer, sheet_name="SEC提交缓存", index=False)
        _localize(fetch_attempts).to_excel(writer, sheet_name="SEC取证检查点", index=False)
        _localize(pd.DataFrame([dict(summary)])).to_excel(writer, sheet_name="审计摘要", index=False)
    _style_workbook(path)
    return write_dataset_manifest(
        path,
        dataset_type="ndx_v11_current_taxonomy_workbook",
        dataset_key="current_taxonomy_plan",
        source="derived_from_pit_ndx_observation_and_sec_evidence",
        fetched_at_evidence="derived_runtime_utc",
        row_count=len(plan),
        data_start=observation_date,
        data_end=observation_date,
        extra={
            "research_only": True,
            "risk_chain_assigned": False,
            "return_data_read": False,
            "price_data_used": False,
            "backtest_results_used": False,
        },
    )


def _expand_review_template(template: pd.DataFrame, plan: pd.DataFrame) -> pd.DataFrame:
    if template.empty:
        return pd.DataFrame(columns=REVIEW_TEMPLATE_COLUMNS)
    plan_groups = {str(key): group for key, group in plan.groupby("job_id", sort=False)}
    rows = []
    for base in template.to_dict(orient="records"):
        evidence_id = str(base.get("evidence_id") or "")
        job_id = evidence_id.removeprefix("sec:").split(":", 1)[0]
        group = plan_groups.get(job_id)
        if group is None or group.empty:
            rows.append(base)
            continue
        for item in group.itertuples(index=False):
            row = dict(base)
            row["evidence_id"] = f"{evidence_id}:{item.security_id}"
            row["security_id"] = str(item.security_id)
            row["ticker"] = str(item.ticker)
            row["issuer_id"] = str(item.issuer_id)
            row["notes"] = (
                f"issuer_id证据可覆盖同发行人多股类；当前security_id需保留。"
                f" job_id={job_id}; current_observation={item.observation_from}; "
                "风险链、有效期、confidence和accepted状态必须人工依据原文填写。"
            )
            rows.append(row)
    return pd.DataFrame(rows, columns=REVIEW_TEMPLATE_COLUMNS)


def _aggregate_raw_hashes(output: Path) -> dict:
    files = [
        path for root_name in ("raw_submissions", "raw_sec")
        for path in (output / root_name).rglob("*")
        if path.is_file() and not path.name.endswith(".manifest.json")
    ]
    digest = hashlib.sha256(
        "\n".join(f"{path.as_posix()}:{file_sha256(path)}" for path in sorted(files)).encode("utf-8")
    ).hexdigest()
    return {"file_count": len(files), "aggregate_sha256": digest}


def _build_review_pack(
    output: Path,
    *,
    observation_date: str,
    input_version: dict,
    target: pd.DataFrame,
) -> tuple[dict, dict, dict]:
    plan_path = output / "current_taxonomy_plan.csv"
    attempts_path = output / "fetch_attempts.csv"
    plan, plan_version = read_versioned_csv(
        plan_path,
        name="ndx_v11_current_taxonomy_plan",
        used_for="当前发行人SEC年度业务文件取证计划",
    )
    _align_fetch_attempts(
        output,
        plan,
        observation_date=observation_date,
    )
    attempts, attempts_version = read_versioned_csv(
        attempts_path,
        name="ndx_v11_current_taxonomy_fetch_attempts",
        used_for="当前SEC主文档断点和哈希审计",
    )
    documents, template = build_review_documents(
        plan,
        attempts,
        raw_root=output / "raw_sec",
    )
    template = _expand_review_template(template, plan)
    gaps = build_filing_plan_gaps(plan, attempts)
    summary = review_summary(documents, gaps)
    summary.update({
        "taxonomy_scope": "current_ndx_snapshot_only",
        "observation_date": observation_date,
        "current_ticker_count": int(target["ticker"].nunique()),
        "current_security_count": int(target["security_id"].nunique()),
        "current_issuer_count": int(target["issuer_id"].nunique()),
        "current_cik_count": int(target["cik"].nunique()),
        "current_target_rows": int(len(target)),
        "current_plan_rows": int(len(plan)),
        "current_job_count": int(plan["job_id"].nunique()),
        "planned_job_count": int(plan["plan_status"].eq("planned").sum()),
        "submission_not_cached_count": int(plan["plan_status"].eq("submission_not_cached").sum()),
        "submission_integrity_failed_count": int(plan["plan_status"].eq("submission_cache_integrity_failed").sum()),
        "risk_chain_assigned": False,
        "classification_automatic": False,
        "return_data_read": False,
        "price_data_used": False,
        "backtest_results_used": False,
        "future_data_guard_rule": "evidence_available_date <= selection_date",
        "historical_backfill_allowed": False,
        "manual_template_rows": int(len(template)),
        "input_versions": [input_version, plan_version.to_dict(), attempts_version.to_dict()],
        "raw_hashes": _aggregate_raw_hashes(output),
    })
    versions = pd.DataFrame([
        input_version,
        plan_version.to_dict(),
        attempts_version.to_dict(),
        {
            "source_name": "current_taxonomy_raw_documents_aggregate",
            "path": f"{(output / 'raw_sec').as_posix()}/**/*.htm",
            "version_id": f"aggregate:current_sec_raw:{summary['raw_hashes']['aggregate_sha256'][:16]}",
            "content_sha256": summary["raw_hashes"]["aggregate_sha256"],
            "manifest_sha256": "",
            "manifest_verified": True,
            "row_count": summary["raw_hashes"]["file_count"],
            "used_for": "当前SEC业务段落人工复核候选，不自动分类",
        },
    ])
    review_output = output / "review_pack"
    review_manifest = write_review_bundle(
        review_output,
        documents,
        template,
        gaps,
        versions,
        summary,
    )
    return review_manifest, summary, {
        "documents": documents,
        "template": template,
        "gaps": gaps,
        "plan": plan,
        "attempts": attempts,
        "plan_version": plan_version.to_dict(),
        "attempts_version": attempts_version.to_dict(),
    }


def _write_root_manifest(
    output: Path,
    *,
    observation_date: str,
    mode: str,
    summary: dict,
    artifacts: dict,
) -> dict:
    manifest = {
        "schema_version": PLAN_SCHEMA_VERSION,
        "manifest_type": "ndx_v11_current_taxonomy",
        "generated_at": utc_now_iso(),
        "mode": mode,
        "observation_date": observation_date,
        "research_only": True,
        "review_only": True,
        "risk_chain_assigned": False,
        "classification_automatic": False,
        "return_data_read": False,
        "price_data_used": False,
        "backtest_results_used": False,
        "historical_backfill_allowed": False,
        "future_data_guard": {
            "rule": "evidence_available_date <= selection_date",
            "passed": True,
        },
        "summary": summary,
        "artifacts": artifacts,
        "prohibited_uses": [
            "automatic_risk_chain_assignment",
            "current_label_historical_backfill",
            "return_or_price_score_boost",
            "automatic_config_change",
            "automatic_today_advice_change",
        ],
    }
    path = output / "manifest.json"
    _write_json_manifest(
        manifest,
        path,
        dataset_key="current_taxonomy_manifest",
        observation_date=observation_date,
        dataset_type="ndx_v11_current_taxonomy_manifest",
    )
    return manifest


def _prepare_output(args) -> Path:
    if args.output_dir:
        output = _resolve(args.output_dir)
    else:
        output = _resolve(DEFAULT_OUTPUT_ROOT) / f"current_review_{args.observation_date.replace('-', '')}"
        if args.smoke:
            output = output / "smoke"
    output.mkdir(parents=True, exist_ok=True)
    return output


def main() -> int:
    parser = argparse.ArgumentParser(description="Build current NDX point-in-time SEC taxonomy evidence")
    modes = parser.add_mutually_exclusive_group(required=True)
    modes.add_argument("--plan-only", action="store_true")
    modes.add_argument("--smoke", action="store_true")
    modes.add_argument("--fetch-sec", action="store_true")
    modes.add_argument("--review-pack", action="store_true")
    parser.add_argument("--observation-date", default=OBSERVATION_DATE)
    parser.add_argument("--input", default=DEFAULT_INPUT)
    parser.add_argument("--output-dir", default="")
    parser.add_argument("--resume", action="store_true")
    parser.add_argument("--force", action="store_true")
    parser.add_argument("--max-jobs", type=int, default=0, help="new issuer jobs per run; zero means all")
    parser.add_argument("--timeout", type=float, default=90.0)
    parser.add_argument("--request-delay", type=float, default=0.20)
    args = parser.parse_args()
    try:
        if args.max_jobs < 0:
            raise NdxTaxonomyError("max-jobs must be >= 0")
        output = _prepare_output(args)
        input_path = _resolve(args.input)
        observations, input_version = _load_current_observations(input_path, args.observation_date)
        target = observations.copy()
        ciks = sorted(set(observations["cik"].astype(str)))

        if args.review_pack:
            review_manifest, review_summary_payload, review_data = _build_review_pack(
                output,
                observation_date=args.observation_date,
                input_version=input_version,
                target=target,
            )
            fetch_attempts = review_data["attempts"]
            plan = review_data["plan"]
            submission_attempts = _load_checkpoint(output / "submission_attempts.csv", SUBMISSION_ATTEMPT_COLUMNS)
            fetch_attempts = review_data["attempts"]
        else:
            if args.fetch_sec:
                submissions, metadata, submission_attempts = _fetch_submissions(
                    output,
                    ciks,
                    observation_date=args.observation_date,
                    resume=bool(args.resume),
                    force=bool(args.force),
                    max_jobs=int(args.max_jobs),
                    timeout=float(args.timeout),
                    request_delay=float(args.request_delay),
                )
            else:
                submissions, metadata = _load_cached_submissions(output, ciks)
                submission_attempts = _load_checkpoint(output / "submission_attempts.csv", SUBMISSION_ATTEMPT_COLUMNS)
            plan = _build_plan(
                observations,
                submissions,
                metadata,
                observation_date=args.observation_date,
            )
            plan_manifest = _write_csv_manifest(
                plan,
                output / "current_taxonomy_plan.csv",
                dataset_key="current_taxonomy_plan",
                observation_date=args.observation_date,
                used_for="当前NDX证券的点时SEC年度申报定位；不直接分类",
            )
            target_manifest = _write_csv_manifest(
                target,
                output / "current_target_universe.csv",
                dataset_key="current_target_universe",
                observation_date=args.observation_date,
                used_for="当前NDX证券、发行人和security_id范围",
            )
            if not (output / "fetch_attempts.csv").exists():
                _write_csv_manifest(
                    empty_fetch_attempts(),
                    output / "fetch_attempts.csv",
                    dataset_key="fetch_attempts",
                    observation_date=args.observation_date,
                    used_for="SEC主文档取证断点",
                )
            if args.fetch_sec:
                _fetch_filing_documents(
                    plan,
                    output=output,
                    observation_date=args.observation_date,
                    resume=bool(args.resume),
                    force=bool(args.force),
                    max_jobs=int(args.max_jobs),
                    timeout=float(args.timeout),
                    request_delay=float(args.request_delay),
                )
            fetch_attempts = _load_fetch_attempts(output)
            review_manifest, review_summary_payload, review_data = _build_review_pack(
                output,
                observation_date=args.observation_date,
                input_version=input_version,
                target=target,
            )

        # The workbook is separate from the generic review workbook so the
        # current target mapping and plan remain visible in Chinese headers.
        summary = dict(review_summary_payload)
        summary["submission_attempt_count"] = int(len(submission_attempts))
        summary["fetch_attempt_count"] = int(len(fetch_attempts))
        summary["fetched_document_count"] = int(fetch_attempts["status"].astype(str).eq("fetched").sum()) if not fetch_attempts.empty else 0
        summary["status"] = (
            "CURRENT_SEC_EVIDENCE_FETCHED_PENDING_MANUAL_REVIEW"
            if summary["fetched_document_count"]
            else "CURRENT_SEC_EVIDENCE_PLAN_ONLY_PENDING_FETCH"
        )
        workbook_path = output / "ndx_v11_current_taxonomy_plan.xlsx"
        workbook_manifest = _write_plan_workbook(
            workbook_path,
            target,
            plan,
            submission_attempts,
            fetch_attempts,
            summary,
            observation_date=args.observation_date,
        )
        summary_manifest = _write_json_manifest(
            summary,
            output / "summary.json",
            dataset_key="current_taxonomy_summary",
            observation_date=args.observation_date,
        )
        artifacts = {
            "input_observations": {
                "path": input_version["path"],
                "content_sha256": input_version["content_sha256"],
                "manifest_sha256": input_version["manifest_sha256"],
                "version_id": input_version["version_id"],
                "row_count": input_version["row_count"],
            },
            "current_target_universe": _artifact_info(output / "current_target_universe.csv", len(target)),
            "current_taxonomy_plan": _artifact_info(output / "current_taxonomy_plan.csv", len(plan)),
            "submission_attempts": _artifact_info(output / "submission_attempts.csv", len(submission_attempts)) if (output / "submission_attempts.csv").exists() else None,
            "fetch_attempts": _artifact_info(output / "fetch_attempts.csv", len(fetch_attempts)),
            "plan_workbook": {
                "path": workbook_path.as_posix(),
                "manifest_path": manifest_path(workbook_path).as_posix(),
                "content_sha256": workbook_manifest["content_sha256"],
                "manifest_sha256": file_sha256(manifest_path(workbook_path)),
                "version_id": workbook_manifest["version_id"],
                "row_count": len(plan),
            },
            "summary": {
                "path": (output / "summary.json").as_posix(),
                "manifest_path": manifest_path(output / "summary.json").as_posix(),
                "content_sha256": summary_manifest["content_sha256"],
                "manifest_sha256": file_sha256(manifest_path(output / "summary.json")),
                "version_id": summary_manifest["version_id"],
                "row_count": 1,
            },
            "review_pack_manifest": {
                "path": (output / "review_pack" / "manifest.json").as_posix(),
                "manifest_sha256": file_sha256(output / "review_pack" / "manifest.json"),
                "content_sha256": file_sha256(output / "review_pack" / "manifest.json"),
                "row_count": int(review_summary_payload.get("raw_document_count") or 0),
            },
        }
        artifacts = {key: value for key, value in artifacts.items() if value is not None}
        root_manifest = _write_root_manifest(
            output,
            observation_date=args.observation_date,
            mode=("smoke" if args.smoke else "fetch-sec" if args.fetch_sec else "review-pack" if args.review_pack else "plan-only"),
            summary=summary,
            artifacts=artifacts,
        )
        if args.smoke:
            if len(target) != 101 or target["cik"].nunique() != 100:
                raise NdxTaxonomyError(
                    f"smoke expected current NDX 101 securities/100 CIK, observed {len(target)}/{target['cik'].nunique()}"
                )
            if bool(target["ticker"].eq("QQQ").any()):
                raise NdxTaxonomyError("smoke found QQQ in current target")
            if bool(summary.get("risk_chain_assigned")) or bool(summary.get("return_data_read")):
                raise NdxTaxonomyError("smoke found prohibited classification or return-data state")
        print(
            "NDX 当前 taxonomy v2 SEC点时证据包完成 | "
            f"securities={len(target)} | issuer={target['issuer_id'].nunique()} | "
            f"planned_jobs={summary.get('planned_job_count', 0)} | "
            f"fetched={summary.get('fetched_document_count', 0)} | "
            f"business_sections={summary.get('business_section_extracted_count', 0)} | "
            f"gaps={summary.get('filing_plan_gap_count', 0)} | "
            "risk_chain_assigned=False | return_data_read=False"
        )
        print(f"输出: {output}")
        print(f"Manifest: {output / 'manifest.json'}")
        return 0
    except (NdxTaxonomyError, OSError, ValueError, requests.RequestException) as exc:
        print(f"NDX current taxonomy failed closed: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
