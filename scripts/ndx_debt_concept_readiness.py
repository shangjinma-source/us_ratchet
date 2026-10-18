"""Audit expanded SEC debt concepts before allowing valuation use."""
from __future__ import annotations

import argparse
import json
import os
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from core.config import load_config
from core.helpers import atomic_write_text
from data.ndx_fundamentals import (
    concept_inventory,
    fundamental_cache_dir,
    load_companyfacts_records,
    raw_bundle_sha256,
)
from data.ndx_valuation_readiness import (
    DEBT_CONCEPT_POLICY_ID,
    EXPANDED_DEBT_METRIC_CONCEPTS,
    PROMOTED_DEBT_CONCEPTS,
    ValuationReadinessPolicy,
    build_point_in_time_accounting_readiness,
)
from data.versioning import file_sha256, manifest_path, utc_now_iso, write_dataset_manifest


def _resolve(path: str | Path) -> Path:
    value = Path(path)
    return value if value.is_absolute() else ROOT / value


def _atomic_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        frame.to_csv(temporary, index=False, encoding="utf-8-sig")
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def _read_json(path: Path) -> dict:
    if not path.exists():
        raise FileNotFoundError(path)
    return json.loads(path.read_text(encoding="utf-8"))


def _verified_csv(path: Path) -> tuple[pd.DataFrame, dict]:
    sidecar = manifest_path(path)
    manifest = _read_json(sidecar)
    if manifest.get("content_sha256") != file_sha256(path):
        raise ValueError(f"versioned input hash mismatch: {path}")
    return pd.read_csv(path, low_memory=False), manifest


def _verify_raw_bundle(
    expected_index: pd.DataFrame,
    observed_index: pd.DataFrame,
) -> None:
    expected = expected_index[["cik", "content_sha256"]].copy()
    observed = observed_index[["cik", "content_sha256"]].copy()
    expected["cik"] = pd.to_numeric(expected["cik"], errors="raise").astype(int)
    observed["cik"] = pd.to_numeric(observed["cik"], errors="raise").astype(int)
    expected = expected[expected["cik"].isin(observed["cik"])].copy()
    merged = expected.merge(
        observed,
        on="cik",
        how="outer",
        suffixes=("_expected", "_observed"),
        indicator=True,
        validate="one_to_one",
    )
    mismatch = merged[
        ~merged["_merge"].eq("both")
        | ~merged["content_sha256_expected"].eq(merged["content_sha256_observed"])
    ]
    if not mismatch.empty or len(expected) != len(observed):
        raise ValueError(
            "raw Company Facts bundle differs from readiness index: "
            f"{mismatch['cik'].astype(str).tolist()[:10]}"
        )


def _merge_promoted_debt_facts(
    base: pd.DataFrame,
    expanded: pd.DataFrame,
) -> pd.DataFrame:
    promoted = expanded[expanded["concept"].isin(PROMOTED_DEBT_CONCEPTS)].copy()
    combined = pd.concat([base, promoted], ignore_index=True)
    key = [
        "cik",
        "metric",
        "namespace",
        "concept",
        "unit",
        "value",
        "form",
        "available_date",
        "period_start",
        "period_end",
        "accession_number",
        "frame",
    ]
    existing = [column for column in key if column in combined.columns]
    return combined.drop_duplicates(existing, keep="last").reset_index(drop=True)


def _comparison(
    base: pd.DataFrame,
    expanded: pd.DataFrame,
) -> pd.DataFrame:
    keys = ["selection_date", "ticker", "cik"]
    fields = [
        "debt_status",
        "debt_ready",
        "debt_strict_complete",
        "debt_value",
        "debt_unit",
        "debt_method",
        "debt_completeness",
        "debt_period_end",
        "debt_age_days",
        "debt_crosscheck_relative_error",
        "debt_crosscheck_status",
        "debt_concepts",
        "debt_component_count",
        "debt_future_data_guard_passed",
    ]
    rows = base[keys + fields].merge(
        expanded[keys + fields],
        on=keys,
        how="outer",
        suffixes=("_base", "_expanded"),
        validate="one_to_one",
        indicator=True,
    )
    if not rows["_merge"].eq("both").all():
        raise ValueError("base and expanded debt readiness keys differ")
    base_value = pd.to_numeric(rows["debt_value_base"], errors="coerce")
    expanded_value = pd.to_numeric(rows["debt_value_expanded"], errors="coerce")
    scale = pd.concat([base_value.abs(), expanded_value.abs()], axis=1).max(axis=1)
    rows["debt_value_relative_change"] = (
        (expanded_value - base_value).abs() / scale.where(scale > 0, 1.0)
    )
    rows["strict_complete_gain"] = (
        rows["debt_strict_complete_expanded"].fillna(False).astype(bool)
        & ~rows["debt_strict_complete_base"].fillna(False).astype(bool)
    )
    rows["strict_complete_regression"] = (
        rows["debt_strict_complete_base"].fillna(False).astype(bool)
        & ~rows["debt_strict_complete_expanded"].fillna(False).astype(bool)
    )
    rows["resolver_changed"] = (
        rows["debt_method_base"].fillna("").astype(str)
        != rows["debt_method_expanded"].fillna("").astype(str)
    ) | rows["strict_complete_gain"] | rows["strict_complete_regression"]
    return rows.drop(columns="_merge").sort_values(
        ["selection_date", "ticker"], kind="stable"
    ).reset_index(drop=True)


def _coverage_comparison(rows: pd.DataFrame) -> pd.DataFrame:
    output = []
    for selection_date, group in rows.groupby("selection_date", sort=True):
        item = {"selection_date": selection_date, "universe_count": len(group)}
        for side in ("base", "expanded"):
            ready = group[f"debt_ready_{side}"].fillna(False).astype(bool)
            strict = group[f"debt_strict_complete_{side}"].fillna(False).astype(bool)
            item[f"debt_ready_count_{side}"] = int(ready.sum())
            item[f"debt_ready_coverage_{side}"] = float(ready.mean())
            item[f"debt_strict_count_{side}"] = int(strict.sum())
            item[f"debt_strict_coverage_{side}"] = float(strict.mean())
        item["debt_ready_count_gain"] = (
            item["debt_ready_count_expanded"] - item["debt_ready_count_base"]
        )
        item["debt_strict_count_gain"] = (
            item["debt_strict_count_expanded"] - item["debt_strict_count_base"]
        )
        output.append(item)
    return pd.DataFrame(output)


def _concept_summary(
    records: pd.DataFrame,
    components: pd.DataFrame,
) -> pd.DataFrame:
    inventory = concept_inventory(records)
    summary = inventory.groupby(
        ["namespace", "concept", "unit"], dropna=False
    ).agg(
        fact_rows=("record_count", "sum"),
        cik_count=("cik", "nunique"),
        first_available_date=("first_available_date", "min"),
        last_available_date=("last_available_date", "max"),
        first_period_end=("first_period_end", "min"),
        last_period_end=("last_period_end", "max"),
    ).reset_index()
    selected = components[
        components["metric"].eq("debt") & components["selected"].fillna(False).astype(bool)
    ].groupby(["namespace", "concept", "unit"]).agg(
        selected_rows=("ticker", "size"),
        selected_tickers=("ticker", "nunique"),
    ).reset_index()
    summary = summary.merge(
        selected,
        on=["namespace", "concept", "unit"],
        how="left",
        validate="one_to_one",
    )
    summary[["selected_rows", "selected_tickers"]] = summary[
        ["selected_rows", "selected_tickers"]
    ].fillna(0).astype(int)
    roles = {
        "DebtLongtermAndShorttermCombinedAmount": (
            "promoted_reported_total_no_component_addition"
        ),
        "DebtCurrent": "diagnostic_only_possible_finance_lease_scope",
        "CommercialPaper": "diagnostic_only_possible_short_term_subset",
    }
    summary["policy_role"] = summary["concept"].map(roles).fillna("diagnostic_only")
    return summary.sort_values(
        ["selected_rows", "fact_rows", "concept"],
        ascending=[False, False, True],
        kind="stable",
    ).reset_index(drop=True)


def _integrity_failures(
    comparison: pd.DataFrame,
    expanded: pd.DataFrame,
    components: pd.DataFrame,
    *,
    tolerance: float,
) -> pd.DataFrame:
    output = []
    regressions = comparison[comparison["strict_complete_regression"]]
    for row in regressions.itertuples(index=False):
        output.append({
            "selection_date": row.selection_date,
            "ticker": row.ticker,
            "scope": "strict_complete_regression",
            "reason": f"{row.debt_method_base}->{row.debt_method_expanded}",
        })
    future = expanded[
        expanded["debt_ready"].fillna(False).astype(bool)
        & ~expanded["debt_future_data_guard_passed"].fillna(False).astype(bool)
    ]
    for row in future.itertuples(index=False):
        output.append({
            "selection_date": row.selection_date,
            "ticker": row.ticker,
            "scope": "future_data_guard",
            "reason": "ready debt uses unavailable evidence",
        })
    selected_promoted = components[
        components["selected"].fillna(False).astype(bool)
        & components["concept"].isin(PROMOTED_DEBT_CONCEPTS)
    ]
    selected_counts = selected_promoted.groupby(
        ["selection_date", "ticker"]
    ).size()
    bad_counts = selected_counts[selected_counts.ne(1)]
    for (selection_date, ticker), count in bad_counts.items():
        output.append({
            "selection_date": selection_date,
            "ticker": ticker,
            "scope": "promoted_component_count",
            "reason": f"selected_count={count}",
        })
    return pd.DataFrame(
        output,
        columns=["selection_date", "ticker", "scope", "reason"],
    )


def _crosscheck_warnings(expanded: pd.DataFrame) -> pd.DataFrame:
    rows = expanded[
        expanded["debt_crosscheck_status"].eq(
            "component_mismatch_reported_total_preferred"
        )
    ].copy()
    columns = [
        "selection_date",
        "ticker",
        "cik",
        "debt_value",
        "debt_unit",
        "debt_method",
        "debt_crosscheck_relative_error",
        "debt_crosscheck_status",
        "debt_concepts",
    ]
    return rows[columns].sort_values(
        ["selection_date", "ticker"], kind="stable"
    ).reset_index(drop=True)


def _style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    fill = PatternFill("solid", fgColor="DCE6F1")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            width = max(len(str(cell.value or "")) for cell in column[:300]) + 2
            sheet.column_dimensions[letter].width = min(max(width, 10), 48)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)


def _write_workbook(
    path: Path,
    frames: dict[str, pd.DataFrame],
) -> None:
    translations = {
        "selection_date": "重选日",
        "ticker": "股票代码",
        "cik": "SEC CIK",
        "concept": "SEC概念",
        "namespace": "命名空间",
        "unit": "单位",
        "fact_rows": "事实行数",
        "cik_count": "覆盖CIK数",
        "selected_rows": "被选用行数",
        "selected_tickers": "被选用股票数",
        "policy_role": "策略角色",
        "universe_count": "候选数",
        "debt_ready_count_base": "原策略债务就绪数",
        "debt_ready_coverage_base": "原策略债务就绪率",
        "debt_strict_count_base": "原策略严格完整数",
        "debt_strict_coverage_base": "原策略严格完整率",
        "debt_ready_count_expanded": "扩展策略债务就绪数",
        "debt_ready_coverage_expanded": "扩展策略债务就绪率",
        "debt_strict_count_expanded": "扩展策略严格完整数",
        "debt_strict_coverage_expanded": "扩展策略严格完整率",
        "debt_ready_count_gain": "债务就绪净增",
        "debt_strict_count_gain": "严格完整净增",
        "scope": "范围",
        "reason": "原因",
        "note": "说明",
        "version_id": "版本号",
        "content_sha256": "内容哈希",
        "rows": "行数",
    }

    def zh(frame: pd.DataFrame) -> pd.DataFrame:
        return frame.rename(columns=translations)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet, frame in frames.items():
            zh(frame).to_excel(writer, sheet_name=sheet, index=False)
    _style_workbook(path)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Audit expanded SEC debt concepts without valuation scoring."
    )
    parser.add_argument("--config", default="config/config.yaml")
    parser.add_argument("--input-dir", default="")
    parser.add_argument("--output-dir", default="")
    parser.add_argument("--maximum-debt-age-days", type=int, default=200)
    args = parser.parse_args()

    cfg = load_config(_resolve(args.config))
    source = (
        _resolve(args.input_dir)
        if args.input_dir
        else fundamental_cache_dir(cfg)
    )
    output = (
        _resolve(args.output_dir)
        if args.output_dir
        else source / "debt_concepts_v2"
    )
    output.mkdir(parents=True, exist_ok=True)
    plan_path = source / "plan.csv"
    plan = pd.read_csv(plan_path, low_memory=False)
    base_facts, base_facts_manifest = _verified_csv(source / "fact_availability.csv")
    expected_raw_index, raw_index_manifest = _verified_csv(source / "raw_index.csv")
    factor_inputs, factor_manifest = _verified_csv(
        source / "normalization_v1" / "quality_factor_inputs.csv"
    )
    ciks = sorted(pd.to_numeric(plan["cik"], errors="raise").astype(int).unique())
    expanded_records, observed_raw_index = load_companyfacts_records(
        source / "raw_companyfacts",
        ciks,
        metric_concepts=EXPANDED_DEBT_METRIC_CONCEPTS,
    )
    expected_used_index = expected_raw_index[
        pd.to_numeric(expected_raw_index["cik"], errors="coerce").isin(ciks)
    ].copy()
    _verify_raw_bundle(expected_used_index, observed_raw_index)
    if raw_bundle_sha256(expected_used_index) != raw_bundle_sha256(observed_raw_index):
        raise ValueError("raw Company Facts bundle hash changed")
    statement_units = factor_inputs[
        ["selection_date", "ticker", "cik", "revenue_unit"]
    ].rename(columns={"revenue_unit": "statement_unit"})
    policy = ValuationReadinessPolicy(
        maximum_debt_age_days=int(args.maximum_debt_age_days)
    )
    policy.validate()
    base_accounting, _base_components = build_point_in_time_accounting_readiness(
        plan,
        base_facts,
        statement_units,
        policy=policy,
    )
    merged_facts = _merge_promoted_debt_facts(base_facts, expanded_records)
    expanded_accounting, expanded_components = (
        build_point_in_time_accounting_readiness(
            plan,
            merged_facts,
            statement_units,
            policy=policy,
        )
    )
    comparison = _comparison(base_accounting, expanded_accounting)
    coverage = _coverage_comparison(comparison)
    changes = comparison[comparison["resolver_changed"]].reset_index(drop=True)
    concepts = _concept_summary(expanded_records, expanded_components)
    failures = _integrity_failures(
        comparison,
        expanded_accounting,
        expanded_components,
        tolerance=policy.debt_crosscheck_tolerance,
    )
    crosscheck_warnings = _crosscheck_warnings(expanded_accounting)
    strict_gain_rows = int(comparison["strict_complete_gain"].sum())
    strict_regression_rows = int(comparison["strict_complete_regression"].sum())
    promotion_allowed = bool(
        failures.empty and strict_gain_rows > 0 and strict_regression_rows == 0
    )
    versions = pd.DataFrame([
        {
            "source_name": "base_fact_availability",
            "version_id": base_facts_manifest.get("version_id"),
            "content_sha256": base_facts_manifest.get("content_sha256"),
            "rows": len(base_facts),
            "note": "既有质量/股数/债务证据，不改写",
        },
        {
            "source_name": "raw_companyfacts_index",
            "version_id": raw_index_manifest.get("version_id"),
            "content_sha256": raw_index_manifest.get("content_sha256"),
            "rows": len(expected_raw_index),
            "note": "159份SEC原始Company Facts逐文件哈希",
        },
        {
            "source_name": "quality_factor_inputs",
            "version_id": factor_manifest.get("version_id"),
            "content_sha256": factor_manifest.get("content_sha256"),
            "rows": len(factor_inputs),
            "note": "只读取报表单位，不读取收益",
        },
    ])
    frames = {
        "expanded_debt_facts": expanded_records,
        "concept_summary": concepts,
        "coverage_comparison": coverage,
        "row_changes": changes,
        "expanded_accounting": expanded_accounting,
        "expanded_components": expanded_components,
        "crosscheck_warnings": crosscheck_warnings,
        "integrity_failures": failures,
        "versions": versions,
    }
    paths = {key: output / f"{key}.csv" for key in frames}
    evidence = {}
    for key, frame in frames.items():
        _atomic_csv(frame, paths[key])
        manifest = write_dataset_manifest(
            paths[key],
            dataset_type="ndx_debt_concept_readiness",
            dataset_key=f"{DEBT_CONCEPT_POLICY_ID}:{key}",
            source="derived_from_versioned_sec_companyfacts",
            fetched_at_evidence="derived_runtime_utc",
            row_count=len(frame),
            data_start=(
                str(plan["selection_date"].min())
                if "selection_date" in frame.columns and not frame.empty else None
            ),
            data_end=(
                str(plan["selection_date"].max())
                if "selection_date" in frame.columns and not frame.empty else None
            ),
            extra={
                "policy_id": DEBT_CONCEPT_POLICY_ID,
                "research_only": True,
                "promotion_allowed": promotion_allowed,
                "valuation_scoring_allowed": False,
            },
        )
        evidence[key] = {
            "filename": paths[key].name,
            "rows": len(frame),
            "sha256": manifest["content_sha256"],
            "version_id": manifest["version_id"],
        }
    base_min = float(coverage["debt_strict_coverage_base"].min())
    expanded_min = float(coverage["debt_strict_coverage_expanded"].min())
    base_latest = float(
        coverage.sort_values("selection_date").iloc[-1]["debt_strict_coverage_base"]
    )
    expanded_latest = float(
        coverage.sort_values("selection_date").iloc[-1][
            "debt_strict_coverage_expanded"
        ]
    )
    summary = {
        "schema_version": 1,
        "generated_at": utc_now_iso(),
        "policy_id": DEBT_CONCEPT_POLICY_ID,
        "research_only": True,
        "plan_sha256": file_sha256(plan_path),
        "base_facts_sha256": base_facts_manifest.get("content_sha256"),
        "raw_bundle_sha256": raw_bundle_sha256(observed_raw_index),
        "selection_rows": len(plan),
        "selection_dates": int(plan["selection_date"].nunique()),
        "ticker_count": int(plan["ticker"].nunique()),
        "expanded_fact_rows": len(expanded_records),
        "promoted_concepts": sorted(PROMOTED_DEBT_CONCEPTS),
        "diagnostic_only_concepts": sorted(
            set(expanded_records["concept"].unique()) - set(PROMOTED_DEBT_CONCEPTS)
        ),
        "strict_complete_gain_rows": strict_gain_rows,
        "strict_complete_regression_rows": strict_regression_rows,
        "base_minimum_strict_coverage": base_min,
        "expanded_minimum_strict_coverage": expanded_min,
        "base_latest_strict_coverage": base_latest,
        "expanded_latest_strict_coverage": expanded_latest,
        "integrity_failure_rows": len(failures),
        "component_mismatch_warning_rows": len(crosscheck_warnings),
        "promotion_allowed": promotion_allowed,
        "valuation_scoring_allowed": False,
        "files": {
            **{key: value["filename"] for key, value in evidence.items()},
            "workbook": "ndx_debt_concept_readiness.xlsx",
        },
    }
    atomic_write_text(
        output / "summary.json",
        json.dumps(summary, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
    )
    atomic_write_text(
        output / "manifest.json",
        json.dumps(
            {
                "schema_version": 1,
                "generated_at": summary["generated_at"],
                "policy_id": DEBT_CONCEPT_POLICY_ID,
                "plan_sha256": summary["plan_sha256"],
                "base_facts_sha256": summary["base_facts_sha256"],
                "raw_bundle_sha256": summary["raw_bundle_sha256"],
                "promoted_concepts": summary["promoted_concepts"],
                "evidence_files": evidence,
                "integrity_passed": failures.empty,
                "promotion_allowed": promotion_allowed,
                "valuation_scoring_allowed": False,
            },
            ensure_ascii=False,
            indent=2,
            allow_nan=False,
        ) + "\n",
    )
    conclusion = pd.DataFrame([
        {"项目": "研究目的", "内容": "扩展SEC债务概念并验证防重复聚合"},
        {"项目": "点时候选行", "内容": len(plan)},
        {"项目": "新增债务事实", "内容": len(expanded_records)},
        {"项目": "严格完整净增行", "内容": strict_gain_rows},
        {"项目": "严格完整回退行", "内容": strict_regression_rows},
        {"项目": "最低严格完整率（原）", "内容": base_min},
        {"项目": "最低严格完整率（扩展）", "内容": expanded_min},
        {"项目": "最新严格完整率（原）", "内容": base_latest},
        {"项目": "最新严格完整率（扩展）", "内容": expanded_latest},
        {"项目": "完整性失败", "内容": len(failures)},
        {"项目": "分量不可相加警告", "内容": len(crosscheck_warnings)},
        {"项目": "允许进入估值准备度", "内容": promotion_allowed},
        {"项目": "估值评分/实盘", "内容": "关闭；不读取收益、不改今日建议"},
    ])
    workbook = output / "ndx_debt_concept_readiness.xlsx"
    _write_workbook(
        workbook,
        {
            "结论": conclusion,
            "概念覆盖": concepts,
            "日期覆盖对比": coverage,
            "逐行变化": changes,
            "扩展会计结果": expanded_accounting,
            "债务来源组件": expanded_components,
            "分量核验警告": crosscheck_warnings,
            "完整性失败": failures,
            "数据版本": versions,
        },
    )
    print("NDX SEC债务概念扩展审计（research only）")
    print(
        f"新增事实: {len(expanded_records)} | 严格完整净增/回退: "
        f"{strict_gain_rows}/{strict_regression_rows}"
    )
    print(
        f"最低严格完整率: {base_min:.2%} -> {expanded_min:.2%} | "
        f"最新: {base_latest:.2%} -> {expanded_latest:.2%}"
    )
    print(
        f"完整性失败: {len(failures)} | 分量警告: {len(crosscheck_warnings)} | "
        f"允许进入估值准备度: {promotion_allowed} | "
        "评分仍关闭"
    )
    print(f"Workbook: {workbook}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
