"""Build a Chinese, review-only evidence pack from fetched SEC filings."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from core.helpers import atomic_write_text
from data.ndx_taxonomy import NdxTaxonomyError, read_versioned_csv
from data.ndx_taxonomy_review import (
    REVIEW_DOCUMENT_COLUMNS,
    REVIEW_GAP_COLUMNS,
    REVIEW_SCHEMA_VERSION,
    REVIEW_TEMPLATE_COLUMNS,
    build_filing_plan_gaps,
    build_review_documents,
    review_summary,
)
from data.versioning import file_sha256, manifest_path, utc_now_iso, write_dataset_manifest


TRANSLATIONS = {
    "job_id": "取证作业编号",
    "gap_id": "补证缺口编号",
    "issuer_id": "发行人编号",
    "security_id": "证券编号",
    "security_ids": "证券编号集合",
    "ticker": "股票代码",
    "tickers": "股票代码集合",
    "cik": "SEC CIK",
    "cusip": "CUSIP",
    "form": "申报表类型",
    "accession_number": "申报编号",
    "selection_date_start": "选择起日",
    "selection_date_end": "选择止日",
    "observation_from": "历史观察起日",
    "evidence_required_from": "证据需要起日",
    "source_available_date": "来源可用日",
    "document_url": "SEC主文档地址",
    "raw_path": "原始文件路径",
    "raw_manifest_path": "原始文件Manifest",
    "raw_sha256": "原始文件哈希",
    "raw_manifest_sha256": "原始文件Manifest哈希",
    "raw_manifest_verified": "原始文件Manifest校验",
    "document_title": "文档标题",
    "document_text_chars": "文档文本字数",
    "document_text_sha256": "文档文本哈希",
    "extraction_status": "段落提取状态",
    "section_label": "标准业务段落",
    "section_text_chars": "业务段落字数",
    "section_text_sha256": "业务段落哈希",
    "business_excerpt": "业务段落候选摘录",
    "business_excerpt_sha256": "摘录哈希",
    "business_excerpt_truncated": "摘录是否截断",
    "review_status": "复核状态",
    "review_note": "复核说明",
    "evidence_id": "证据编号",
    "risk_chain": "风险链编号（人工填写）",
    "effective_from": "分类生效起日（人工填写）",
    "effective_to": "分类生效止日（人工填写）",
    "evidence_source_type": "证据来源类型",
    "evidence_source": "证据来源名称",
    "evidence_locator": "证据地址",
    "evidence_available_date": "证据可用日",
    "evidence_excerpt": "证据摘录",
    "confidence": "置信度（人工填写）",
    "status": "状态（人工填写）",
    "notes": "备注",
    "plan_status": "计划状态",
    "issue_type": "缺口类型",
    "sec_submissions_url": "SEC发行人申报索引",
    "source_policy": "来源政策",
    "classification_allowed": "允许直接分类",
    "future_data_guard": "未来数据守门",
    "next_action": "下一动作",
    "source_name": "来源名称",
    "path": "文件路径",
    "version_id": "版本编号",
    "content_sha256": "内容哈希",
    "manifest_sha256": "Manifest哈希",
    "manifest_verified": "Manifest校验通过",
    "row_count": "行数",
    "used_for": "用途",
}


def _resolve(path: str | Path) -> Path:
    target = Path(path)
    return target if target.is_absolute() else ROOT / target


def _atomic_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        frame.to_csv(temporary, index=False, encoding="utf-8-sig", lineterminator="\n")
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def _write_artifact(path: Path, frame: pd.DataFrame, key: str) -> dict:
    _atomic_csv(frame, path)
    payload = write_dataset_manifest(
        path,
        dataset_type="ndx_taxonomy_v2_review_artifact",
        dataset_key=f"review_pack:{key}",
        source="derived_from_manifest_verified_sec_taxonomy_inputs",
        fetched_at_evidence="derived_runtime_utc",
        row_count=len(frame),
        extra={
            "review_only": True,
            "risk_chain_assigned": False,
            "return_data_read": False,
            "historical_backfill_allowed": False,
        },
    )
    return {
        "filename": path.name,
        "manifest_filename": manifest_path(path).name,
        "rows": len(frame),
        "sha256": payload["content_sha256"],
        "version_id": payload["version_id"],
    }


def _localized(frame: pd.DataFrame) -> pd.DataFrame:
    output = frame.copy()
    output = output.rename(columns={
        column: TRANSLATIONS.get(str(column), f"技术字段（{column}）")
        for column in output.columns
    })
    for column in output.columns:
        if output[column].dtype == object or pd.api.types.is_bool_dtype(output[column]):
            localized = output[column].map({True: "是", False: "否"})
            output[column] = localized.where(localized.notna(), output[column])
    return output


def _write_workbook(
    path: Path,
    documents: pd.DataFrame,
    excerpts: pd.DataFrame,
    template: pd.DataFrame,
    gaps: pd.DataFrame,
    versions: pd.DataFrame,
    summary: dict,
) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame([
            {"项目": "工作目的", "内容": "把已抓取 SEC 原文整理为人工复核证据包；不自动赋予风险链。"},
            {"项目": "原始文件", "内容": f"{summary.get('raw_document_count', 0)} 个；每个文件须通过内容哈希和 Manifest 校验。"},
            {"项目": "提取规则", "内容": "仅对 10-K/40-F Item 1 Business、20-F Item 4 Information on the Company 提取候选段落。"},
            {"项目": "人工门", "内容": "风险链、有效期、置信度和 accepted 状态必须人工核验填写；空值保持 pending/unclassified。"},
            {"项目": "未来数据守门", "内容": "evidence_available_date <= selection_date；不满足时不得用于历史分类。"},
            {"项目": "未定位发行人", "内容": f"{summary.get('filing_plan_gap_count', 0)} 行进入补证计划，不猜 accession。"},
            {"项目": "收益与实盘", "内容": "未读取收益/价格/回测；不修改 config.yaml、state 或今日建议。"},
        ]).to_excel(writer, sheet_name="使用说明", index=False)
        _localized(documents).to_excel(writer, sheet_name="SEC文档索引", index=False)
        _localized(excerpts).to_excel(writer, sheet_name="业务段落摘录", index=False)
        _localized(template).to_excel(writer, sheet_name="证据导入模板", index=False)
        _localized(gaps).to_excel(writer, sheet_name="待补证计划", index=False)
        _localized(versions).to_excel(writer, sheet_name="版本与审计", index=False)

    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(path)
    fill = PatternFill("solid", fgColor="DCE6F1")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        if sheet.max_row >= 1 and sheet.max_column >= 1:
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            width = max(len(str(cell.value or "")) for cell in list(column)[:200]) + 2
            sheet.column_dimensions[letter].width = min(max(width, 10), 48)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)


def _source_versions(plan_version, attempts_version, documents: pd.DataFrame) -> pd.DataFrame:
    rows = [plan_version.to_dict(), attempts_version.to_dict()]
    if not documents.empty:
        raw_hash = hashlib.sha256(
            "\n".join(sorted(documents["raw_sha256"].astype(str))).encode("utf-8")
        ).hexdigest()
        rows.append({
            "source_name": "sec_raw_documents_aggregate",
            "path": "raw_sec/**/*.htm",
            "version_id": f"aggregate:sec_raw_documents:{raw_hash[:16]}",
            "content_sha256": raw_hash,
            "manifest_sha256": "",
            "manifest_verified": bool(documents["raw_manifest_verified"].all()),
            "row_count": len(documents),
            "used_for": "人工业务描述复核候选，不直接分类",
        })
    return pd.DataFrame(rows)


def write_bundle(
    output: Path,
    documents: pd.DataFrame,
    template: pd.DataFrame,
    gaps: pd.DataFrame,
    versions: pd.DataFrame,
    summary: dict,
) -> dict:
    output.mkdir(parents=True, exist_ok=True)
    excerpts_columns = [
        "job_id", "issuer_id", "tickers", "form", "accession_number",
        "selection_date_start", "source_available_date", "document_url",
        "extraction_status", "section_label", "business_excerpt",
        "business_excerpt_sha256", "business_excerpt_truncated", "raw_sha256",
        "review_status", "review_note",
    ]
    excerpts = documents[excerpts_columns].copy()
    artifacts = {
        "documents": _write_artifact(output / "sec_review_documents.csv", documents, "documents"),
        "excerpts": _write_artifact(output / "business_excerpts.csv", excerpts, "excerpts"),
        "evidence_template": _write_artifact(
            output / "evidence_review_template.csv", template, "evidence_template"
        ),
        "filing_plan_gaps": _write_artifact(
            output / "filing_plan_gaps.csv", gaps, "filing_plan_gaps"
        ),
        "source_versions": _write_artifact(
            output / "review_source_versions.csv", versions, "source_versions"
        ),
    }
    workbook_path = output / "ndx_taxonomy_v2_review_pack.xlsx"
    _write_workbook(path=workbook_path, documents=documents, excerpts=excerpts,
                    template=template, gaps=gaps, versions=versions, summary=summary)
    workbook_manifest = write_dataset_manifest(
        workbook_path,
        dataset_type="ndx_taxonomy_v2_review_workbook",
        dataset_key="ndx_taxonomy_v2_review_pack",
        source="derived_from_manifest_verified_sec_taxonomy_inputs",
        fetched_at_evidence="derived_runtime_utc",
        row_count=len(documents),
        extra={
            "review_only": True,
            "risk_chain_assigned": False,
            "return_data_read": False,
            "historical_backfill_allowed": False,
        },
    )
    artifacts["workbook"] = {
        "filename": workbook_path.name,
        "manifest_filename": manifest_path(workbook_path).name,
        "rows": len(documents),
        "sha256": workbook_manifest["content_sha256"],
        "version_id": workbook_manifest["version_id"],
    }
    summary = dict(summary)
    summary["generated_at"] = utc_now_iso()
    summary["artifacts"] = artifacts
    summary_path = output / "summary.json"
    atomic_write_text(summary_path, json.dumps(summary, ensure_ascii=False, indent=2) + "\n")
    summary_manifest = write_dataset_manifest(
        summary_path,
        dataset_type="ndx_taxonomy_v2_review_summary",
        dataset_key="ndx_taxonomy_v2_review_pack",
        source="derived_from_manifest_verified_sec_taxonomy_inputs",
        fetched_at_evidence="derived_runtime_utc",
        row_count=1,
        extra={"review_only": True, "risk_chain_assigned": False, "return_data_read": False},
    )
    artifacts["summary"] = {
        "filename": summary_path.name,
        "manifest_filename": manifest_path(summary_path).name,
        "rows": 1,
        "sha256": summary_manifest["content_sha256"],
        "version_id": summary_manifest["version_id"],
    }
    manifest = {
        "schema_version": REVIEW_SCHEMA_VERSION,
        "manifest_type": "ndx_taxonomy_v2_review_pack",
        "generated_at": summary["generated_at"],
        "research_only": True,
        "review_only": True,
        "risk_chain_assigned": False,
        "return_data_read": False,
        "price_data_used": False,
        "backtest_results_used": False,
        "historical_backfill_allowed": False,
        "future_data_guard": summary.get("future_data_guard"),
        "summary": summary,
        "artifacts": artifacts,
        "input_versions": summary.get("input_versions", []),
        "prohibited_uses": [
            "automatic_risk_chain_assignment",
            "current_label_historical_backfill",
            "direct_score_boost",
            "automatic_config_change",
            "automatic_today_advice_change",
        ],
    }
    manifest_path_out = output / "manifest.json"
    atomic_write_text(manifest_path_out, json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    write_dataset_manifest(
        manifest_path_out,
        dataset_type="ndx_taxonomy_v2_review_manifest",
        dataset_key="ndx_taxonomy_v2_review_pack",
        source="derived_from_manifest_verified_sec_taxonomy_inputs",
        fetched_at_evidence="derived_runtime_utc",
        row_count=1,
        extra={"review_only": True, "risk_chain_assigned": False},
    )
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser(description="生成 NDX taxonomy v2 中文 SEC 证据复核包")
    parser.add_argument(
        "--input-dir",
        default="data/cache/candidate_pools/ndx_taxonomy_v2",
        help="已完成 taxonomy v2 审计的输入目录",
    )
    parser.add_argument(
        "--output-dir",
        default="data/cache/candidate_pools/ndx_taxonomy_v2/review_pack",
        help="中文证据复核包输出目录",
    )
    args = parser.parse_args()
    try:
        input_dir = _resolve(args.input_dir)
        output_dir = _resolve(args.output_dir)
        plan, plan_version = read_versioned_csv(
            input_dir / "evidence_plan.csv",
            name="ndx_taxonomy_v2_evidence_plan",
            used_for="SEC证据复核计划",
        )
        attempts, attempts_version = read_versioned_csv(
            input_dir / "fetch_attempts.csv",
            name="ndx_taxonomy_v2_fetch_attempts",
            used_for="SEC原始文件断点检查",
        )
        documents, template = build_review_documents(
            plan,
            attempts,
            raw_root=input_dir / "raw_sec",
        )
        gaps = build_filing_plan_gaps(plan, attempts)
        versions = _source_versions(plan_version, attempts_version, documents)
        summary = review_summary(documents, gaps)
        summary["input_versions"] = [plan_version.to_dict(), attempts_version.to_dict()]
        manifest = write_bundle(output_dir, documents, template, gaps, versions, summary)
        print(
            "NDX taxonomy v2 中文证据复核包完成 | "
            f"documents={len(documents)} | "
            f"business_sections={summary['business_section_extracted_count']} | "
            f"gaps={len(gaps)} | "
            f"pending_manual_review={len(template)}"
        )
        print(f"输出: {output_dir}")
        print(f"Manifest: {output_dir / 'manifest.json'}")
        return 0
    except Exception as exc:
        print(f"NDX taxonomy review pack failed closed: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
