"""Build and audit the versioned NDX risk-chain taxonomy v2.

Default execution is offline and research-only.  ``--plan-only`` creates a
point-in-time SEC review plan without network calls.  ``--fetch-sec`` is an
explicit, resumable raw-document acquisition step; it never assigns a
risk-chain.  A human-reviewed evidence CSV is required before any row can be
marked ``accepted``.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
import time
from pathlib import Path

import pandas as pd
import requests

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from core.config import load_config
from core.helpers import atomic_write_text
from data.ndx_history import _sec_headers, _sec_user_agent
from data.ndx_security_identity import load_security_lifecycle
from data.ndx_taxonomy import (
    EVIDENCE_COLUMNS,
    NdxTaxonomyError,
    build_coverage_report,
    build_evidence_plan,
    build_security_universe,
    build_taxonomy_records,
    empty_evidence_frame,
    empty_fetch_attempts,
    load_evidence_csv,
    load_taxonomy_definition,
    normalize_and_expand_evidence,
    read_versioned_csv,
    select_business_document,
    select_business_document_from_index_html,
    taxonomy_digest,
)
from data.versioning import file_sha256, manifest_path, utc_now_iso, write_dataset_manifest


OUTPUT_SCHEMA_VERSION = 1
FETCH_TERMINAL_STATUSES = {"fetched", "failed", "ambiguous_document"}

ARTIFACT_TRANSLATIONS = {
    "scope": "范围",
    "source_name": "来源名称",
    "path": "文件路径",
    "version_id": "版本编号",
    "content_sha256": "内容哈希",
    "manifest_sha256": "Manifest哈希",
    "manifest_verified": "Manifest校验通过",
    "used_for": "用途",
    "security_id": "证券编号",
    "issuer_id": "发行人编号",
    "ticker": "股票代码",
    "cik": "SEC CIK",
    "cusip": "CUSIP",
    "observation_from": "历史观察起日",
    "observation_to": "历史观察止日",
    "available_from": "最早可用日",
    "available_to": "最晚可用日",
    "observed_reports": "观察季度数",
    "ticker_count": "代码数量",
    "ticker_set": "代码集合",
    "identity_status": "身份状态",
    "identity_issue": "身份问题",
    "selection_date": "选择日",
    "report_date": "报告期",
    "filed_date": "申报日",
    "available_date": "成分可用日",
    "accession_number": "申报编号",
    "price_ticker": "取价代码",
    "name": "证券名称",
    "ticker_lifecycle_rule_id": "代码生命周期规则",
    "ticker_lifecycle_status": "代码生命周期状态",
    "job_id": "取证作业编号",
    "evidence_required_from": "需要证据起日",
    "form": "申报表类型",
    "source_available_date": "来源可用日",
    "source_url": "来源地址",
    "sec_index_url": "SEC申报目录",
    "source_role": "来源角色",
    "classification_allowed": "允许直接分类",
    "point_in_time_for_observation_start": "起始观察点时成立",
    "plan_status": "计划状态",
    "review_note": "复核说明",
    "evidence_id": "证据编号",
    "risk_chain": "风险链编号",
    "risk_chain_label": "风险链名称",
    "effective_from": "分类生效起日",
    "effective_to": "分类生效止日",
    "selection_date_start": "分类观察起日",
    "selection_date_end": "分类观察止日",
    "selection_dates": "覆盖选择日",
    "observations_count": "覆盖观察数",
    "evidence_source": "证据来源名称",
    "evidence_source_type": "证据来源类型",
    "evidence_locator": "证据地址",
    "evidence_available_date": "证据可用日",
    "evidence_ids": "证据编号集合",
    "evidence_excerpt": "证据摘录",
    "confidence": "置信度",
    "status": "状态",
    "reason": "原因",
    "notes": "备注",
    "validation_status": "证据校验状态",
    "validation_reason": "证据校验原因",
    "scope": "证据作用范围",
    "issue_type": "问题类型",
    "detail": "问题详情",
    "risk_chains": "冲突风险链",
    "future_data_guard_passed": "未来数据守门通过",
    "guard_rule": "守门规则",
    "classified_observations": "已分类观察数",
    "conflict_observations": "冲突观察数",
    "unclassified_observations": "未分类观察数",
    "classification_coverage": "分类覆盖率",
    "historical_ticker_count": "历史Ticker数量",
    "security_count": "证券身份数量",
    "issuer_count": "发行人数量",
    "classified_ticker_count": "已分类Ticker数量",
    "classified_security_count": "已分类证券数量",
    "unclassified_security_count": "未分类证券数量",
    "conflict_security_count": "冲突证券数量",
    "observation_count": "历史观察行数",
    "classified_observation_count": "已分类观察行数",
    "return_data_read": "是否读取收益数据",
    "recognized_ticker": "已识别代码",
    "rule_id": "生命周期规则编号",
    "evidence_urls": "生命周期证据地址",
    "attempted_at": "尝试时间",
    "index_url": "申报目录地址",
    "document_url": "业务文件地址",
    "raw_path": "原始文件路径",
    "content_bytes": "内容字节数",
    "error": "错误信息",
}


def _resolve(path: str | Path) -> Path:
    value = Path(path)
    return value if value.is_absolute() else ROOT / value


def _atomic_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        frame.to_csv(temporary, index=False, encoding="utf-8-sig", lineterminator="\n")
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def _atomic_bytes(content: bytes, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        temporary.write_bytes(content)
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def _version_for_local_file(path: Path, name: str, used_for: str) -> dict:
    digest = file_sha256(path)
    sidecar = manifest_path(path)
    sidecar_digest = file_sha256(sidecar) if sidecar.exists() else None
    return {
        "source_name": name,
        "path": path.as_posix(),
        "version_id": f"local:{name}:{digest[:16]}",
        "content_sha256": digest,
        "manifest_sha256": sidecar_digest,
        "manifest_verified": bool(sidecar.exists()),
        "row_count": None,
        "used_for": used_for,
    }


def _input_paths(cfg, args) -> dict[str, Path]:
    cache = Path(cfg.cache_dir)
    fundamental = cache / "candidate_pools" / "ndx_fundamentals"
    return {
        "history": _resolve(args.history) if args.history else cache / "candidate_pools" / "ndx_history" / "qqq_nport_constituents.csv",
        "crosswalk": _resolve(args.crosswalk) if args.crosswalk else fundamental / "cik_crosswalk.csv",
        "sic_selection": _resolve(args.sic_selection) if args.sic_selection else fundamental / "sic_history_v1" / "selection_sic.csv",
        "lifecycle": _resolve(args.lifecycle) if args.lifecycle else ROOT / "config" / "ndx_security_lifecycle.yaml",
        "taxonomy": _resolve(args.taxonomy) if args.taxonomy else ROOT / "config" / "candidate_taxonomy_v2.yaml",
    }


def _load_inputs(cfg, args) -> tuple[dict, object, pd.DataFrame | None, list[dict]]:
    paths = _input_paths(cfg, args)
    history, history_version = read_versioned_csv(
        paths["history"], name="ndx_point_in_time_constituents", used_for="历史证券观察和点时身份"
    )
    crosswalk, crosswalk_version = read_versioned_csv(
        paths["crosswalk"], name="ndx_sec_cik_crosswalk", used_for="ticker-CIK-CUSIP发行人身份"
    )
    sic_selection = None
    versions = [history_version.to_dict(), crosswalk_version.to_dict()]
    if paths["sic_selection"].exists():
        sic_selection, sic_version = read_versioned_csv(
            paths["sic_selection"], name="sic_selection_v1", used_for="只定位可复核SEC申报，不直接分类"
        )
        versions.append(sic_version.to_dict())
    taxonomy = load_taxonomy_definition(paths["taxonomy"])
    lifecycle = load_security_lifecycle(paths["lifecycle"])
    versions.extend([
        _version_for_local_file(paths["taxonomy"], "candidate_taxonomy_v2", "受控风险链定义"),
        _version_for_local_file(paths["lifecycle"], "ndx_security_lifecycle_v1", "代码生命周期与证券身份辅助审计"),
    ])
    universe = build_security_universe(history, crosswalk, lifecycle=lifecycle)
    return {
        "paths": paths,
        "taxonomy": taxonomy,
        "universe": universe,
        "versions": versions,
    }, universe, sic_selection, versions


def _load_fetch_attempts(path: Path) -> pd.DataFrame:
    if not path.exists():
        return empty_fetch_attempts()
    frame = pd.read_csv(path, dtype=str, keep_default_na=False)
    for column in empty_fetch_attempts().columns:
        if column not in frame.columns:
            frame[column] = ""
    return frame[empty_fetch_attempts().columns].drop_duplicates("job_id", keep="last")


def _upsert_attempt(frame: pd.DataFrame, row: dict) -> pd.DataFrame:
    remaining = frame[~frame["job_id"].eq(str(row["job_id"]))].copy()
    return pd.concat([remaining, pd.DataFrame([row])], ignore_index=True).sort_values(
        ["issuer_id", "accession_number", "job_id"], kind="stable"
    ).reset_index(drop=True)


def _write_raw_document_manifest(path: Path, row, document_url: str) -> dict:
    return write_dataset_manifest(
        path,
        dataset_type="ndx_taxonomy_v2_sec_raw_filing",
        dataset_key=f"{row.issuer_id}:{row.accession_number}:{path.name}",
        source="sec_filing_business_description",
        source_locator=document_url or str(row.sec_index_url),
        fetched_at_evidence="sec_public_archive_http_200",
        row_count=1,
        extra={
            "research_only": True,
            "classification_automatic": False,
            "issuer_id": str(row.issuer_id),
            "cik": str(row.cik),
            "accession_number": str(row.accession_number),
            "form": str(row.form),
            "security_id": str(row.security_id),
            "ticker": str(row.ticker),
            "evidence_available_date": str(row.source_available_date),
            "sec_index_url": str(row.sec_index_url),
        },
    )


def _fetch_sec_documents(
    plan: pd.DataFrame,
    *,
    output: Path,
    max_jobs: int,
    refresh: bool,
    timeout: float,
    request_delay: float,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    agent = os.environ.get("SEC_USER_AGENT")
    if not agent:
        raise NdxTaxonomyError(
            "SEC_USER_AGENT is required for --fetch-sec, e.g. "
            "$env:SEC_USER_AGENT='us_ratchet/1.0 your_email@example.com'"
        )
    headers = _sec_headers(_sec_user_agent(agent))
    attempts_path = output / "fetch_attempts.csv"
    attempts = _load_fetch_attempts(attempts_path)
    unique_plan = plan[
        plan["sec_index_url"].astype(str).str.startswith("https://")
        & plan["accession_number"].astype(str).str.strip().ne("")
    ].drop_duplicates("job_id", keep="first")
    pending = []
    for row in unique_plan.itertuples(index=False):
        prior = attempts[attempts["job_id"].eq(str(row.job_id))]
        if not refresh and not prior.empty and str(prior.iloc[-1]["status"]) == "fetched":
            raw_path = Path(str(prior.iloc[-1]["raw_path"]))
            expected = str(prior.iloc[-1]["content_sha256"])
            if raw_path.exists() and file_sha256(raw_path) == expected:
                if not manifest_path(raw_path).exists():
                    _write_raw_document_manifest(
                        raw_path,
                        row,
                        str(prior.iloc[-1]["document_url"] or prior.iloc[-1]["index_url"]),
                    )
                continue
        pending.append(row)
    if max_jobs > 0:
        pending = pending[:max_jobs]
    session = requests.Session()
    session.headers.update(headers)
    fetched_rows = []
    for position, row in enumerate(pending, start=1):
        status = "failed"
        document_url = ""
        raw_path = ""
        content = b""
        error = ""
        try:
            response = session.get(str(row.sec_index_url), timeout=timeout)
            response.raise_for_status()
            base = str(row.sec_index_url).rsplit("/index.json", 1)[0]
            try:
                item = select_business_document(response.json(), str(row.form))
            except NdxTaxonomyError as json_error:
                index_html_url = f"{base}/{str(row.accession_number)}-index.html"
                index_html_response = session.get(index_html_url, timeout=timeout)
                index_html_response.raise_for_status()
                try:
                    item = select_business_document_from_index_html(
                        index_html_response.content, str(row.form)
                    )
                except NdxTaxonomyError as html_error:
                    raise NdxTaxonomyError(
                        f"{json_error}; SEC index HTML fallback failed: {html_error}"
                    ) from html_error
            name = Path(str(item.get("name") or "")).name
            if not name:
                raise NdxTaxonomyError("SEC index selected a blank document name")
            document_url = f"{base}/{name}"
            document_response = session.get(document_url, timeout=timeout)
            document_response.raise_for_status()
            content = bytes(document_response.content)
            if len(content) < 256:
                raise NdxTaxonomyError("SEC business document is unexpectedly small")
            raw_file = output / "raw_sec" / str(row.issuer_id).replace(":", "_") / str(row.accession_number).replace("-", "") / name
            _atomic_bytes(content, raw_file)
            _write_raw_document_manifest(raw_file, row, document_url)
            raw_path = raw_file.as_posix()
            status = "fetched"
        except Exception as exc:
            error = f"{type(exc).__name__}: {exc}"
        attempt = {
            "job_id": str(row.job_id),
            "issuer_id": str(row.issuer_id),
            "cik": str(row.cik),
            "accession_number": str(row.accession_number),
            "status": status,
            "index_url": str(row.sec_index_url),
            "document_url": document_url,
            "raw_path": raw_path,
            "content_bytes": len(content),
            "content_sha256": hashlib.sha256(content).hexdigest() if content else "",
            "attempted_at": utc_now_iso(),
            "error": error,
        }
        attempts = _upsert_attempt(attempts, attempt)
        _atomic_csv(attempts, attempts_path)
        if status == "fetched":
            fetched_rows.append({**attempt, "security_id": str(row.security_id), "ticker": str(row.ticker), "form": str(row.form), "source_available_date": str(row.source_available_date)})
        print(
            f"taxonomy SEC取证 {position}/{len(pending)} | {row.issuer_id} | {status}",
            flush=True,
        )
        if request_delay > 0:
            time.sleep(request_delay)
    return attempts, pd.DataFrame(fetched_rows)


def _align_fetch_attempts(
    attempts: pd.DataFrame,
    plan: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Separate current annual-report jobs from superseded locator attempts."""
    if attempts.empty:
        return attempts.copy(), attempts.copy()
    current_jobs = set(plan["job_id"].astype(str))
    active = attempts[attempts["job_id"].astype(str).isin(current_jobs)].copy()
    stale = attempts[~attempts["job_id"].astype(str).isin(current_jobs)].copy()
    if not stale.empty:
        stale["status"] = "stale_plan_job_preserved_not_reviewable"
    return active.reset_index(drop=True), stale.reset_index(drop=True)


def _localized(frame: pd.DataFrame) -> pd.DataFrame:
    out = frame.copy()
    if out.empty:
        return out.rename(columns={column: ARTIFACT_TRANSLATIONS.get(column, f"技术字段（{column}）") for column in out.columns})
    mapping = {column: ARTIFACT_TRANSLATIONS.get(str(column), f"技术字段（{column}）") for column in out.columns}
    out = out.rename(columns=mapping)
    for column in out.columns:
        if out[column].dtype == object or pd.api.types.is_bool_dtype(out[column]):
            localized = out[column].map({True: "是", False: "否"})
            out[column] = localized.where(localized.notna(), out[column])
    return out


def _style_workbook(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(path)
    fill = PatternFill("solid", fgColor="DCE6F1")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        if sheet.max_row >= 1 and sheet.max_column >= 1:
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            width = max(len(str(cell.value or "")) for cell in list(column)[:200]) + 2
            sheet.column_dimensions[letter].width = min(max(width, 10), 48)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)


def _write_workbook(path: Path, frames: dict[str, pd.DataFrame], summary: dict) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame([
            {"项目": "研究目的", "内容": "建立点时风险链证据，不直接改变选股分数或实盘配置"},
            {"项目": "历史覆盖", "内容": f"{summary.get('historical_ticker_count', 0)} 个历史 ticker；证券身份不按名称猜测"},
            {"项目": "未来数据守门", "内容": "evidence_available_date <= selection_date；失败则未分类"},
            {"项目": "QQQ角色", "内容": "仅作为历史 N-PORT 来源代理；QQQ 本身不进入证券候选"},
            {"项目": "收益数据", "内容": "未读取；分类不使用收益、价格或回测结果"},
            {"项目": "实盘影响", "内容": "research_only；不修改 config.yaml、state 或今日建议"},
        ]).to_excel(writer, sheet_name="使用说明", index=False)
        sheet_names = {
            "coverage_summary": "覆盖率汇总",
            "coverage_detail": "证券覆盖率",
            "identities": "身份映射",
            "lifecycle_aliases": "代码生命周期",
            "evidence_plan": "证据计划",
            "evidence_input": "证据导入",
            "evidence_issues": "证据问题",
            "records": "产业链分类",
            "conflicts": "冲突审计",
            "unclassified": "未分类",
            "future_audit": "未来数据审计",
            "identity_issues": "身份问题",
            "versions": "数据版本",
            "fetch_attempts": "取证检查点",
            "stale_fetch_attempts": "旧取证审计",
            "evidence_template": "证据导入模板",
        }
        dictionary = []
        for key, frame in frames.items():
            localized = _localized(frame)
            localized.to_excel(writer, sheet_name=sheet_names[key], index=False)
            dictionary.extend({
                "所属工作表": sheet_names[key],
                "程序字段": str(column),
                "中文表头": ARTIFACT_TRANSLATIONS.get(str(column), f"技术字段（{column}）"),
            } for column in frame.columns)
        pd.DataFrame(dictionary).drop_duplicates().to_excel(
            writer, sheet_name="字段字典", index=False
        )
    _style_workbook(path)


def _write_artifact(path: Path, frame: pd.DataFrame, key: str, taxonomy_version: str) -> dict:
    _atomic_csv(frame, path)
    manifest = write_dataset_manifest(
        path,
        dataset_type="ndx_taxonomy_v2_artifact",
        dataset_key=f"{taxonomy_version}:{key}",
        source="derived_from_versioned_pit_identity_and_reviewed_evidence",
        fetched_at_evidence="derived_runtime_utc",
        row_count=len(frame),
        extra={
            "taxonomy_version": taxonomy_version,
            "research_only": True,
            "return_data_read": False,
            "future_data_guard": "evidence_available_date <= selection_date",
            "historical_backfill_allowed": False,
        },
    )
    return {
        "filename": path.name,
        "manifest_filename": manifest_path(path).name,
        "rows": len(frame),
        "sha256": manifest["content_sha256"],
        "version_id": manifest["version_id"],
    }


def build_audit_bundle(
    cfg,
    args,
    *,
    fetched_attempts: pd.DataFrame | None = None,
    stale_fetch_attempts: pd.DataFrame | None = None,
    fetched_rows: pd.DataFrame | None = None,
) -> tuple[dict[str, pd.DataFrame], dict, list[dict], object, pd.DataFrame]:
    inputs, universe, sic_selection, versions = _load_inputs(cfg, args)
    evidence = load_evidence_csv(_resolve(args.evidence_csv) if args.evidence_csv else None)
    if args.evidence_csv:
        evidence_path = _resolve(args.evidence_csv)
        versions = [*versions, _version_for_local_file(evidence_path, "reviewed_taxonomy_evidence", "人工审核分类证据")]
    expanded, evidence_issues = normalize_and_expand_evidence(
        evidence, universe, inputs["taxonomy"]
    )
    evidence_plan = build_evidence_plan(universe, sic_selection=sic_selection)
    records, conflicts, future_audit = build_taxonomy_records(
        universe, inputs["taxonomy"], expanded
    )
    coverage_summary, coverage_detail = build_coverage_report(universe, records)
    unclassified = records[records["status"].ne("accepted")].copy()
    source_versions = pd.DataFrame(versions)
    if not future_audit.empty:
        future_guard_passed = bool(future_audit["future_data_guard_passed"].all())
    else:
        future_guard_passed = True
    summary_row = coverage_summary.iloc[0].to_dict() if not coverage_summary.empty else {}
    summary = {
        "schema_version": OUTPUT_SCHEMA_VERSION,
        "generated_at": utc_now_iso(),
        "taxonomy_id": inputs["taxonomy"].get("taxonomy_id"),
        "taxonomy_version": inputs["taxonomy"].get("version_id"),
        "taxonomy_digest": taxonomy_digest(inputs["taxonomy"], universe, expanded),
        "research_only": True,
        "today_advice_impact": "none",
        "return_data_read": False,
        "price_data_used": False,
        "backtest_results_used": False,
        "historical_backfill_allowed": False,
        "future_data_guard_rule": "evidence_available_date <= selection_date",
        "future_data_guard_passed": future_guard_passed,
        "historical_ticker_count": int(universe.observations["ticker"].nunique()),
        "security_count": int(universe.identities["security_id"].nunique()),
        "issuer_count": int(universe.identities["issuer_id"].nunique()),
        "identity_issue_count": len(universe.issues),
        "evidence_input_rows": len(evidence),
        "evidence_expanded_rows": len(expanded),
        "evidence_valid_rows": int(expanded["validation_status"].eq("valid").sum()) if not expanded.empty else 0,
        "evidence_pending_rows": int(expanded["validation_status"].eq("pending").sum()) if not expanded.empty else 0,
        "evidence_rejected_rows": len(evidence_issues),
        "classified_security_count": int(summary_row.get("classified_security_count") or 0),
        "unclassified_security_count": int(summary_row.get("unclassified_security_count") or 0),
        "conflict_security_count": int(summary_row.get("conflict_security_count") or 0),
        "classified_observation_count": int(summary_row.get("classified_observation_count") or 0),
        "unclassified_observation_count": int(
            summary_row.get("observation_count") or 0
        ) - int(summary_row.get("classified_observation_count") or 0),
        "classification_status": (
            "accepted_evidence_applied" if int(summary_row.get("classified_security_count") or 0) else "all_unclassified_pending_first_party_review"
        ),
        "input_versions": versions,
        "source_exclusion": "candidate_taxonomy_v1 and all returns/backtests were not used for classification",
    }
    frames = {
        "coverage_summary": coverage_summary,
        "coverage_detail": coverage_detail,
        "identities": universe.identities,
        "lifecycle_aliases": universe.lifecycle_aliases,
        "evidence_plan": evidence_plan,
        "evidence_input": expanded,
        "evidence_issues": evidence_issues,
        "records": records,
        "conflicts": conflicts,
        "unclassified": unclassified,
        "future_audit": future_audit,
        "identity_issues": universe.issues,
        "versions": source_versions,
        "fetch_attempts": fetched_attempts if fetched_attempts is not None else empty_fetch_attempts(),
        "stale_fetch_attempts": (
            stale_fetch_attempts
            if stale_fetch_attempts is not None
            else empty_fetch_attempts()
        ),
        "evidence_template": empty_evidence_frame(),
    }
    if fetched_attempts is not None:
        summary["fetched_document_count"] = int(
            fetched_attempts["status"].astype(str).eq("fetched").sum()
        ) if not fetched_attempts.empty else 0
    elif fetched_rows is not None and not fetched_rows.empty:
        summary["fetched_document_count"] = len(fetched_rows)
    else:
        checkpoint_dir = _resolve(args.output_dir) if args.output_dir else Path(cfg.cache_dir) / "candidate_pools" / "ndx_taxonomy_v2"
        checkpoint_path = checkpoint_dir / "fetch_attempts.csv"
        if checkpoint_path.exists():
            checkpoint = _load_fetch_attempts(checkpoint_path)
            summary["fetched_document_count"] = int(
                checkpoint["status"].astype(str).eq("fetched").sum()
            )
        else:
            summary["fetched_document_count"] = 0
    summary["stale_fetch_attempt_count"] = (
        len(stale_fetch_attempts) if stale_fetch_attempts is not None else 0
    )
    return frames, summary, versions, universe, evidence_plan


def write_bundle(
    output: Path,
    frames: dict[str, pd.DataFrame],
    summary: dict,
    *,
    mode: str,
) -> dict:
    output.mkdir(parents=True, exist_ok=True)
    taxonomy_version = str(summary.get("taxonomy_version") or "taxonomy_v2")
    file_map = {
        "coverage_summary": "coverage_summary.csv",
        "coverage_detail": "coverage_by_security.csv",
        "identities": "security_identity.csv",
        "lifecycle_aliases": "ticker_lifecycle_aliases.csv",
        "evidence_plan": "evidence_plan.csv",
        "evidence_input": "evidence_import.csv",
        "evidence_issues": "evidence_issues.csv",
        "records": "taxonomy_records.csv",
        "conflicts": "conflicts.csv",
        "unclassified": "unclassified.csv",
        "future_audit": "future_data_audit.csv",
        "identity_issues": "identity_issues.csv",
        "versions": "source_versions.csv",
        "fetch_attempts": "fetch_attempts.csv",
        "stale_fetch_attempts": "stale_fetch_attempts.csv",
        "evidence_template": "evidence_template.csv",
    }
    artifacts = {}
    for key, filename in file_map.items():
        artifacts[key] = _write_artifact(output / filename, frames[key], key, taxonomy_version)
    workbook_path = output / "ndx_taxonomy_v2_report.xlsx"
    _write_workbook(workbook_path, frames, summary)
    workbook_manifest = write_dataset_manifest(
        workbook_path,
        dataset_type="ndx_taxonomy_v2_artifact",
        dataset_key=f"{taxonomy_version}:workbook",
        source="derived_from_versioned_pit_identity_and_reviewed_evidence",
        fetched_at_evidence="derived_runtime_utc",
        row_count=len(frames["records"]),
        extra={
            "taxonomy_version": taxonomy_version,
            "research_only": True,
            "return_data_read": False,
            "future_data_guard": summary.get("future_data_guard_rule"),
        },
    )
    artifacts["workbook"] = {
        "filename": workbook_path.name,
        "manifest_filename": manifest_path(workbook_path).name,
        "rows": len(frames["records"]),
        "sha256": workbook_manifest["content_sha256"],
        "version_id": workbook_manifest["version_id"],
    }
    summary_path = output / "summary.json"
    atomic_write_text(
        summary_path,
        json.dumps(summary, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
    )
    summary_manifest = write_dataset_manifest(
        summary_path,
        dataset_type="ndx_taxonomy_v2_artifact",
        dataset_key=f"{taxonomy_version}:summary",
        source="derived_from_versioned_pit_identity_and_reviewed_evidence",
        fetched_at_evidence="derived_runtime_utc",
        row_count=1,
        extra={
            "taxonomy_version": taxonomy_version,
            "research_only": True,
            "return_data_read": False,
            "future_data_guard": summary.get("future_data_guard_rule"),
        },
    )
    artifacts["summary"] = {
        "filename": summary_path.name,
        "manifest_filename": manifest_path(summary_path).name,
        "rows": 1,
        "sha256": summary_manifest["content_sha256"],
        "version_id": summary_manifest["version_id"],
    }
    manifest = {
        "schema_version": OUTPUT_SCHEMA_VERSION,
        "manifest_type": "ndx_taxonomy_v2_bundle",
        "generated_at": summary.get("generated_at"),
        "mode": mode,
        "taxonomy_id": summary.get("taxonomy_id"),
        "taxonomy_version": taxonomy_version,
        "taxonomy_digest": summary.get("taxonomy_digest"),
        "research_only": True,
        "today_advice_impact": "none",
        "return_data_read": False,
        "price_data_used": False,
        "backtest_results_used": False,
        "historical_backfill_allowed": False,
        "future_data_guard": {
            "rule": summary.get("future_data_guard_rule"),
            "passed": bool(summary.get("future_data_guard_passed")),
        },
        "summary": summary,
        "artifacts": artifacts,
        "input_versions": summary.get("input_versions") or [],
        "prohibited_uses": [
            "direct_score_boost",
            "automatic_config_change",
            "automatic_today_advice_change",
            "historical_backfill_from_current_label",
        ],
    }
    manifest_path_out = output / "manifest.json"
    atomic_write_text(
        manifest_path_out,
        json.dumps(manifest, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
    )
    write_dataset_manifest(
        manifest_path_out,
        dataset_type="ndx_taxonomy_v2_bundle_manifest",
        dataset_key=f"{taxonomy_version}:bundle_manifest",
        source="derived_from_versioned_pit_identity_and_reviewed_evidence",
        fetched_at_evidence="derived_runtime_utc",
        row_count=1,
        extra={
            "taxonomy_version": taxonomy_version,
            "research_only": True,
            "manifest_self_sidecar": True,
        },
    )
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser(description="构建和审计点时 NDX 产业链 taxonomy v2")
    parser.add_argument("--config", default="config/config.yaml")
    parser.add_argument("--taxonomy", default="")
    parser.add_argument("--history", default="")
    parser.add_argument("--crosswalk", default="")
    parser.add_argument("--sic-selection", default="")
    parser.add_argument("--lifecycle", default="")
    parser.add_argument("--evidence-csv", default="")
    parser.add_argument("--output-dir", default="")
    parser.add_argument("--plan-only", action="store_true", help="只生成本地计划和审计，不联网")
    parser.add_argument("--smoke", action="store_true", help="离线烟雾模式，等价于 plan-only 并检查 163 ticker 覆盖")
    parser.add_argument("--fetch-sec", action="store_true", help="显式抓取 SEC 申报业务文件；不自动分类")
    parser.add_argument("--resume", action="store_true", help="复用已成功的 SEC 取证检查点")
    parser.add_argument("--refresh", action="store_true", help="重新抓取已成功的 SEC 取证作业")
    parser.add_argument("--max-jobs", type=int, default=0, help="本次最多抓取多少个 SEC 作业，0 表示全部")
    parser.add_argument("--timeout", type=float, default=45.0)
    parser.add_argument("--request-delay", type=float, default=0.15)
    args = parser.parse_args()
    if args.fetch_sec and (args.plan_only or args.smoke):
        parser.error("--fetch-sec cannot be combined with --plan-only/--smoke")
    if args.refresh and not args.fetch_sec:
        parser.error("--refresh requires --fetch-sec")

    cfg = load_config(_resolve(args.config))
    default_output = Path(cfg.cache_dir) / "candidate_pools" / "ndx_taxonomy_v2"
    output = _resolve(args.output_dir) if args.output_dir else default_output
    if args.smoke and not args.output_dir:
        output = default_output / "smoke"

    fetched_attempts = None
    stale_fetch_attempts = None
    fetched_rows = None
    try:
        inputs, universe, sic_selection, _ = _load_inputs(cfg, args)
        evidence_plan = build_evidence_plan(universe, sic_selection=sic_selection)
        output.mkdir(parents=True, exist_ok=True)
        if args.fetch_sec:
            fetched_attempts, fetched_rows = _fetch_sec_documents(
                evidence_plan,
                output=output,
                max_jobs=max(0, int(args.max_jobs)),
                refresh=bool(args.refresh),
                timeout=float(args.timeout),
                request_delay=max(0.0, float(args.request_delay)),
            )
            fetched_total = int(
                fetched_attempts["status"].astype(str).eq("fetched").sum()
            ) if not fetched_attempts.empty else 0
            print(
                "SEC原始业务文件抓取完成: "
                f"本次新增={len(fetched_rows)} | 累计成功={fetched_total} | "
                f"checkpoint={output / 'fetch_attempts.csv'}"
            )
        else:
            # Offline audits must preserve an existing SEC checkpoint.  The
            # checkpoint is provenance for the raw documents, not a disposable
            # empty frame regenerated by the report writer.
            checkpoint_path = output / "fetch_attempts.csv"
            if checkpoint_path.exists():
                fetched_attempts = _load_fetch_attempts(checkpoint_path)
        if fetched_attempts is not None:
            fetched_attempts, stale_fetch_attempts = _align_fetch_attempts(
                fetched_attempts,
                evidence_plan,
            )
            archived_stale = _load_fetch_attempts(
                output / "stale_fetch_attempts.csv"
            )
            if not archived_stale.empty:
                stale_fetch_attempts = pd.concat(
                    [archived_stale, stale_fetch_attempts], ignore_index=True
                ).drop_duplicates("job_id", keep="last")
                stale_fetch_attempts["status"] = (
                    "stale_plan_job_preserved_not_reviewable"
                )
        frames, summary, _, _, _ = build_audit_bundle(
            cfg,
            args,
            fetched_attempts=fetched_attempts,
            stale_fetch_attempts=stale_fetch_attempts,
            fetched_rows=fetched_rows,
        )
        mode = "smoke" if args.smoke else "plan_only" if args.plan_only or not args.fetch_sec else "offline_audit"
        manifest = write_bundle(output, frames, summary, mode=mode)
        if args.smoke:
            if summary["historical_ticker_count"] != 163:
                raise NdxTaxonomyError(
                    f"smoke expected 163 historical tickers, observed {summary['historical_ticker_count']}"
                )
            if not summary["future_data_guard_passed"]:
                raise NdxTaxonomyError("smoke future-data guard failed")
        print(
            "NDX taxonomy v2 审计完成 | "
            f"ticker={summary['historical_ticker_count']} | security={summary['security_count']} | "
            f"issuer={summary['issuer_count']} | classified={summary['classified_security_count']} | "
            f"unclassified={summary['unclassified_security_count']} | conflicts={summary['conflict_security_count']}"
        )
        print(f"输出: {output}")
        print(f"Manifest: {output / 'manifest.json'}")
        return 0
    except Exception as exc:
        print(f"NDX taxonomy v2 failed closed: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
