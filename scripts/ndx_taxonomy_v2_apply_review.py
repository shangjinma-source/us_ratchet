"""Apply the authored, hash-locked SEC taxonomy review registry.

The registry is deliberately separate from the raw review template.  This
command validates every decision against the manifest-verified SEC document,
its extracted business excerpt, the taxonomy vocabulary and the point-in-time
date rule before producing the evidence CSV consumed by ``ndx_taxonomy_v2``.
It never reads prices, returns or backtest output.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
from pathlib import Path

import pandas as pd
import yaml

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from core.helpers import atomic_write_text
from data.ndx_taxonomy import load_taxonomy_definition, read_versioned_csv
from data.ndx_taxonomy_review import REVIEW_TEMPLATE_COLUMNS
from data.versioning import file_sha256, manifest_path, utc_now_iso, write_dataset_manifest
from scripts.ndx_taxonomy_review_pack import TRANSLATIONS


DECISION_AUDIT_COLUMNS = [
    "decision_id",
    "ticker",
    "evidence_id",
    "issuer_id",
    "risk_chain",
    "confidence",
    "status",
    "effective_from",
    "effective_to",
    "extraction_status",
    "raw_manifest_verified",
    "raw_sha256_match",
    "excerpt_sha256_match",
    "decision_status",
    "decision_note",
]


def _resolve(path: str | Path) -> Path:
    target = Path(path)
    return target if target.is_absolute() else ROOT / target


def _atomic_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        frame.to_csv(temporary, index=False, encoding="utf-8-sig", lineterminator="\n")
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def _write_artifact(
    path: Path,
    frame: pd.DataFrame,
    key: str,
    *,
    decision_registry_sha256: str,
    extra: dict | None = None,
) -> dict:
    _atomic_csv(frame, path)
    manifest_extra = {
        "research_only": True,
        "return_data_read": False,
        "price_data_used": False,
        "backtest_results_used": False,
        "historical_backfill_allowed": False,
        "decision_registry_sha256": decision_registry_sha256,
    }
    if extra:
        manifest_extra.update(extra)
    payload = write_dataset_manifest(
        path,
        dataset_type="ndx_taxonomy_v2_reviewed_artifact",
        dataset_key=f"reviewed:{key}",
        source="derived_from_hash_locked_sec_review_decisions",
        fetched_at_evidence="derived_runtime_utc",
        row_count=len(frame),
        extra=manifest_extra,
    )
    return {
        "filename": path.name,
        "manifest_filename": manifest_path(path).name,
        "rows": len(frame),
        "sha256": payload["content_sha256"],
        "version_id": payload["version_id"],
    }


def _localized(frame: pd.DataFrame, *, mode: str = "default") -> pd.DataFrame:
    output = frame.copy()
    translations = dict(TRANSLATIONS)
    if mode == "reviewed":
        translations.update({
            "risk_chain": "风险链编号（Codex审核）",
            "effective_from": "分类生效起日（Codex审核）",
            "effective_to": "分类生效止日（Codex审核）",
            "confidence": "置信度（Codex审核）",
            "status": "状态（Codex审核）",
        })
    elif mode == "pending":
        translations.update({
            "risk_chain": "风险链编号（待审核）",
            "effective_from": "分类生效起日（待审核）",
            "effective_to": "分类生效止日（待审核）",
            "confidence": "置信度（待审核）",
            "status": "状态（待审核）",
        })
    output = output.rename(columns={
        column: translations.get(str(column), f"技术字段（{column}）")
        for column in output.columns
    })
    for column in output.columns:
        if output[column].dtype == object or pd.api.types.is_bool_dtype(output[column]):
            localized = output[column].map({True: "是", False: "否"})
            output[column] = localized.where(localized.notna(), output[column])
    return output


def _read_registry(path: Path) -> dict:
    try:
        payload = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    except (OSError, ValueError, TypeError, yaml.YAMLError) as exc:
        raise ValueError(f"invalid decision registry: {path}") from exc
    if int(payload.get("schema_version") or 0) != 1:
        raise ValueError("decision registry schema_version must be 1")
    if str(payload.get("taxonomy_id") or "") != "us_growth_risk_chain_taxonomy_v2":
        raise ValueError("decision registry taxonomy_id mismatch")
    decisions = payload.get("decisions")
    if not isinstance(decisions, list) or not decisions:
        raise ValueError("decision registry has no decisions")
    return payload


def _resolved_effective_from(
    decision: dict, template_row: pd.Series, registry: dict
) -> str:
    value = str(decision.get("effective_from") or "").strip()
    policy = registry.get("decision_policy") or {}
    if not value and str(policy.get("effective_from_rule") or "") == "evidence_available_date":
        value = "evidence_available_date"
    if value == "evidence_available_date":
        value = str(template_row.get("evidence_available_date") or "").strip()
    return value


def _validate_and_apply(
    template: pd.DataFrame,
    documents: pd.DataFrame,
    registry: dict,
    taxonomy: dict,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    reviewed = template.copy()
    for column in REVIEW_TEMPLATE_COLUMNS:
        if column not in reviewed.columns:
            reviewed[column] = ""
    reviewed = reviewed[REVIEW_TEMPLATE_COLUMNS].copy()
    evidence_by_id = reviewed.set_index("evidence_id", drop=False)
    documents = documents.copy()
    documents["evidence_id"] = documents["job_id"].map(lambda value: f"sec:{value}")
    document_by_evidence = documents.set_index("evidence_id", drop=False)
    categories = set((taxonomy.get("risk_chains") or {}).keys())
    seen = set()
    audit_rows = []

    for position, decision in enumerate(registry["decisions"], start=1):
        if not isinstance(decision, dict):
            raise ValueError(f"decision {position} is not a mapping")
        decision_id = f"decision:{position:03d}"
        evidence_id = str(decision.get("evidence_id") or "").strip()
        ticker = str(decision.get("ticker") or "").strip().upper()
        if not evidence_id or evidence_id in seen:
            raise ValueError(f"duplicate or empty evidence_id in {decision_id}: {evidence_id}")
        seen.add(evidence_id)
        if evidence_id not in evidence_by_id.index or evidence_id not in document_by_evidence.index:
            raise ValueError(f"decision references absent review evidence: {evidence_id}")
        template_row = evidence_by_id.loc[evidence_id]
        document_row = document_by_evidence.loc[evidence_id]
        document_tickers = {
            item.strip().upper()
            for item in str(document_row.get("tickers") or "").split(";")
            if item.strip()
        }
        if ticker not in document_tickers:
            raise ValueError(f"ticker/evidence mismatch for {decision_id}: {ticker} vs {document_tickers}")
        expected_issuer = str(decision.get("issuer_id") or "").strip()
        if expected_issuer != str(document_row.get("issuer_id") or ""):
            raise ValueError(f"issuer/evidence mismatch for {decision_id}: {expected_issuer}")
        if str(document_row.get("extraction_status") or "") != "business_section_extracted":
            raise ValueError(f"decision requires a standardized business section: {evidence_id}")
        raw_match = str(decision.get("expected_raw_sha256") or "").strip() == str(document_row.get("raw_sha256") or "")
        excerpt_match = str(decision.get("expected_excerpt_sha256") or "").strip() == str(document_row.get("business_excerpt_sha256") or "")
        if not raw_match or not excerpt_match:
            raise ValueError(
                f"hash-locked SEC evidence changed for {decision_id}: "
                f"raw_match={raw_match}, excerpt_match={excerpt_match}"
            )
        if str(document_row.get("raw_manifest_verified") or "").lower() not in {"true", "1"}:
            raise ValueError(f"SEC raw manifest is not verified for {decision_id}")
        risk_chain = str(decision.get("risk_chain") or "").strip().lower()
        confidence = str(decision.get("confidence") or "").strip().lower()
        status = str(decision.get("status") or "").strip().lower()
        if risk_chain not in categories or risk_chain == "unclassified":
            raise ValueError(f"unsupported accepted risk chain for {decision_id}: {risk_chain}")
        if confidence not in {"high", "medium"} or status != "accepted":
            raise ValueError(f"accepted decision requires high/medium confidence: {decision_id}")
        effective_from = _resolved_effective_from(decision, template_row, registry)
        available_date = str(template_row.get("evidence_available_date") or "").strip()
        effective_to = str(decision.get("effective_to") or "").strip()
        if not available_date or not effective_from or available_date > effective_from:
            raise ValueError(f"future-data guard failed for {decision_id}")
        if effective_to and effective_from > effective_to:
            raise ValueError(f"effective interval reversed for {decision_id}")
        index = reviewed.index[reviewed["evidence_id"].eq(evidence_id)]
        if len(index) != 1:
            raise ValueError(f"evidence_id is not unique in review template: {evidence_id}")
        row_index = index[0]
        reviewed.at[row_index, "risk_chain"] = risk_chain
        reviewed.at[row_index, "effective_from"] = effective_from
        reviewed.at[row_index, "effective_to"] = effective_to
        reviewed.at[row_index, "confidence"] = confidence
        reviewed.at[row_index, "status"] = status
        reviewed.at[row_index, "notes"] = (
            f"{str(template_row.get('notes') or '').strip()} | "
            f"review_registry={registry.get('review_registry_id')}; "
            f"{str(decision.get('notes') or '').strip()}"
        ).strip(" |")
        audit_rows.append({
            "decision_id": decision_id,
            "ticker": ticker,
            "evidence_id": evidence_id,
            "issuer_id": expected_issuer,
            "risk_chain": risk_chain,
            "confidence": confidence,
            "status": status,
            "effective_from": effective_from,
            "effective_to": effective_to,
            "extraction_status": document_row.get("extraction_status"),
            "raw_manifest_verified": True,
            "raw_sha256_match": raw_match,
            "excerpt_sha256_match": excerpt_match,
            "decision_status": "applied",
            "decision_note": decision.get("notes", ""),
        })

    return reviewed, pd.DataFrame(audit_rows, columns=DECISION_AUDIT_COLUMNS)


def _write_workbook(
    path: Path,
    reviewed: pd.DataFrame,
    audit: pd.DataFrame,
    pending: pd.DataFrame,
    summary: dict,
) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame([
            {"项目": "用途", "内容": "把 Codex 已核验的 SEC 一手业务证据转换为 taxonomy v2 导入文件；不是股票评分文件。"},
            {"项目": "接受规则", "内容": "只有业务段落、原始文件 Manifest、原文哈希和摘录哈希全部通过，才写入 accepted。"},
            {"项目": "有效期", "内容": "首个 effective_from = evidence_available_date；不把当前标签回填到更早历史。"},
            {"项目": "置信度", "内容": "high=业务描述直接且链条单一；medium=业务描述直接但有明显旁支或平台混合。"},
            {"项目": "未分类", "内容": f"{len(pending)} 行继续 pending_review；没有用股票名称、收益或回测补齐。"},
            {"项目": "收益与实盘", "内容": "未读取收益/价格/回测；不修改 config.yaml、state、今日建议或买入计划。"},
        ]).to_excel(writer, sheet_name="审核说明", index=False)
        _localized(reviewed, mode="reviewed").to_excel(writer, sheet_name="已审核证据", index=False)
        _localized(audit, mode="reviewed").to_excel(writer, sheet_name="审核审计", index=False)
        _localized(pending, mode="pending").to_excel(writer, sheet_name="待审核证据", index=False)
        pd.DataFrame([summary]).T.reset_index().set_axis(["项目", "值"], axis=1).to_excel(
            writer, sheet_name="版本与审计", index=False
        )

    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(path)
    fill = PatternFill("solid", fgColor="DCE6F1")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        if sheet.max_row >= 1 and sheet.max_column >= 1:
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            width = max(len(str(cell.value or "")) for cell in list(column)[:200]) + 2
            sheet.column_dimensions[letter].width = min(max(width, 10), 48)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description="应用 hash 锁定的 NDX taxonomy v2 SEC 审核决定")
    parser.add_argument("--review-dir", default="data/cache/candidate_pools/ndx_taxonomy_v2/review_pack")
    parser.add_argument("--decision-file", default="config/ndx_taxonomy_v2_review_decisions.yaml")
    parser.add_argument("--taxonomy", default="config/candidate_taxonomy_v2.yaml")
    parser.add_argument("--output-dir", default="data/cache/candidate_pools/ndx_taxonomy_v2/review_pack/reviewed")
    args = parser.parse_args()
    try:
        review_dir = _resolve(args.review_dir)
        decision_path = _resolve(args.decision_file)
        output_dir = _resolve(args.output_dir)
        taxonomy = load_taxonomy_definition(_resolve(args.taxonomy))
        template, template_version = read_versioned_csv(
            review_dir / "evidence_review_template.csv",
            name="ndx_taxonomy_v2_review_template",
            used_for="hash 锁定证据审核输入",
        )
        documents, documents_version = read_versioned_csv(
            review_dir / "sec_review_documents.csv",
            name="ndx_taxonomy_v2_review_documents",
            used_for="SEC 原文和摘录哈希核验",
        )
        registry = _read_registry(decision_path)
        registry_sha256 = file_sha256(decision_path)
        reviewed, audit = _validate_and_apply(template, documents, registry, taxonomy)
        pending = reviewed[reviewed["status"].astype(str).str.lower().ne("accepted")].copy()
        summary = {
            "schema_version": 1,
            "generated_at": utc_now_iso(),
            "review_registry_id": registry.get("review_registry_id"),
            "decision_registry_path": decision_path.as_posix(),
            "decision_registry_sha256": registry_sha256,
            "taxonomy_id": taxonomy.get("taxonomy_id"),
            "taxonomy_version": taxonomy.get("version_id"),
            "reviewed_document_count": len(documents),
            "template_row_count": len(reviewed),
            "accepted_decision_count": len(audit),
            "accepted_evidence_row_count": int(reviewed["status"].eq("accepted").sum()),
            "pending_review_row_count": len(pending),
            "future_data_guard": "evidence_available_date <= effective_from <= selection_date when resolved",
            "historical_backfill_allowed": False,
            "risk_chain_from_backtest": False,
            "return_data_read": False,
            "price_data_used": False,
            "backtest_results_used": False,
            "config_state_today_advice_changed": False,
            "input_versions": [template_version.to_dict(), documents_version.to_dict()],
        }
        output_dir.mkdir(parents=True, exist_ok=True)
        artifacts = {}
        artifacts["evidence_reviewed"] = _write_artifact(
            output_dir / "evidence_reviewed.csv",
            reviewed,
            "evidence_reviewed",
            decision_registry_sha256=registry_sha256,
            extra={"accepted_decision_count": len(audit)},
        )
        artifacts["decision_audit"] = _write_artifact(
            output_dir / "review_decision_audit.csv",
            audit,
            "review_decision_audit",
            decision_registry_sha256=registry_sha256,
        )
        snapshot_path = output_dir / "review_decisions_snapshot.yaml"
        atomic_write_text(snapshot_path, decision_path.read_text(encoding="utf-8"))
        snapshot_manifest = write_dataset_manifest(
            snapshot_path,
            dataset_type="ndx_taxonomy_v2_review_decision_snapshot",
            dataset_key=str(registry.get("review_registry_id") or "review_decisions"),
            source="copied_from_authored_hash_locked_review_registry",
            fetched_at_evidence="authored_source_file_hash",
            row_count=len(registry.get("decisions") or []),
            extra={"source_sha256": registry_sha256, "research_only": True},
        )
        artifacts["decision_snapshot"] = {
            "filename": snapshot_path.name,
            "manifest_filename": manifest_path(snapshot_path).name,
            "rows": len(registry.get("decisions") or []),
            "sha256": snapshot_manifest["content_sha256"],
            "version_id": snapshot_manifest["version_id"],
        }
        workbook_path = output_dir / "ndx_taxonomy_v2_reviewed.xlsx"
        _write_workbook(workbook_path, reviewed, audit, pending, summary)
        workbook_manifest = write_dataset_manifest(
            workbook_path,
            dataset_type="ndx_taxonomy_v2_reviewed_workbook",
            dataset_key="ndx_taxonomy_v2_reviewed",
            source="derived_from_hash_locked_sec_review_decisions",
            fetched_at_evidence="derived_runtime_utc",
            row_count=len(reviewed),
            extra={
                "decision_registry_sha256": registry_sha256,
                "research_only": True,
                "return_data_read": False,
            },
        )
        artifacts["workbook"] = {
            "filename": workbook_path.name,
            "manifest_filename": manifest_path(workbook_path).name,
            "rows": len(reviewed),
            "sha256": workbook_manifest["content_sha256"],
            "version_id": workbook_manifest["version_id"],
        }
        summary["artifacts"] = artifacts
        summary_path = output_dir / "reviewed_summary.json"
        atomic_write_text(summary_path, json.dumps(summary, ensure_ascii=False, indent=2) + "\n")
        summary_manifest = write_dataset_manifest(
            summary_path,
            dataset_type="ndx_taxonomy_v2_reviewed_summary",
            dataset_key="ndx_taxonomy_v2_reviewed",
            source="derived_from_hash_locked_sec_review_decisions",
            fetched_at_evidence="derived_runtime_utc",
            row_count=1,
            extra={"decision_registry_sha256": registry_sha256, "research_only": True},
        )
        artifacts["summary"] = {
            "filename": summary_path.name,
            "manifest_filename": manifest_path(summary_path).name,
            "rows": 1,
            "sha256": summary_manifest["content_sha256"],
            "version_id": summary_manifest["version_id"],
        }
        manifest = {
            "schema_version": 1,
            "manifest_type": "ndx_taxonomy_v2_reviewed_bundle",
            "generated_at": summary["generated_at"],
            "research_only": True,
            "return_data_read": False,
            "price_data_used": False,
            "backtest_results_used": False,
            "historical_backfill_allowed": False,
            "future_data_guard": summary["future_data_guard"],
            "decision_registry_sha256": registry_sha256,
            "summary": summary,
            "artifacts": artifacts,
            "input_versions": summary["input_versions"],
            "prohibited_uses": [
                "direct_score_boost",
                "automatic_config_change",
                "automatic_today_advice_change",
                "historical_backfill_from_current_label",
                "risk_chain_assignment_from_returns_or_backtests",
            ],
        }
        manifest_path_out = output_dir / "manifest.json"
        atomic_write_text(manifest_path_out, json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
        write_dataset_manifest(
            manifest_path_out,
            dataset_type="ndx_taxonomy_v2_reviewed_manifest",
            dataset_key="ndx_taxonomy_v2_reviewed",
            source="derived_from_hash_locked_sec_review_decisions",
            fetched_at_evidence="derived_runtime_utc",
            row_count=1,
            extra={"decision_registry_sha256": registry_sha256, "research_only": True},
        )
        print(
            "NDX taxonomy v2 SEC审核决定应用完成 | "
            f"accepted={len(audit)} | pending={len(pending)} | "
            f"registry_sha256={registry_sha256}"
        )
        print(f"输出: {output_dir}")
        print(f"导入证据: {output_dir / 'evidence_reviewed.csv'}")
        return 0
    except Exception as exc:
        print(f"NDX taxonomy review apply failed closed: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
